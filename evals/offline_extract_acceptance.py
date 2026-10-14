"""Run parser-to-evaluation acceptance with synthetic PDF and Word files, no LLM calls.

Usage: uv run python evals/offline_extract_acceptance.py --output data/work/offline-extract.json
"""

import argparse
import asyncio
import json
import sys
import uuid
from pathlib import Path

import pymupdf
from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "server"))

from app.providers.local_ocr import LocalOCR  # noqa: E402
from app.schemas.contracts import (  # noqa: E402
    Category,
    ExtractedRequirement,
    Location,
    Source,
)
from app.services.parsing import parse_document  # noqa: E402
from extract_tender import star_coverage, star_units  # noqa: E402

FAILURE_MODES = [
    "Two starred clauses on one PDF page must not count as covered by one unrelated citation.",
    "Two starred clauses in one Word block must be distinguished by clause text.",
    "A starred clause continuing across a line break must remain one source clause.",
    "Repeated PDF text on one page is ambiguous and must not enter coverage.",
    "A citation in one Word block must not cover repeated text in another block.",
    "Repeated identical starred text in one Word block is an ambiguous citation.",
    "An unrelated-only citation on a page with stars must cover zero clauses.",
    "Invalid source quotes and locations must not enter the coverage numerator.",
    "A document with no star markers must report an undefined ratio, not divide by zero.",
    "Coverage of a clause and the extracted item's starred label are separate measures.",
]


def make_pdf() -> bytes:
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "★Clause A required; ★Clause B required.", fontname="china-s")
    page.insert_text((72, 110), "★Cross-line clause requires", fontname="china-s")
    page.insert_text((72, 126), "continuous service for three years.", fontname="china-s")
    page.insert_text((72, 148), "Ordinary clause: submit authorization.", fontname="china-s")
    page.insert_text((72, 186), "★Repeated clause must be performed.", fontname="china-s")
    page.insert_text((72, 224), "★Repeated clause must be performed.", fontname="china-s")
    result = document.tobytes()
    document.close()
    return result


def docx_bytes(document: Document) -> bytes:
    from io import BytesIO

    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def make_chunks(parsed, name: str) -> list[dict]:
    document_id = uuid.uuid5(uuid.NAMESPACE_URL, f"offline-extract:{name}")
    return [
        {
            "id": uuid.uuid5(document_id, f"chunk:{index}"),
            "document_id": document_id,
            "page": None if hasattr(page, "blocks") else page.page,
            "text": page.text,
            "citation_verified": True if hasattr(page, "blocks") else page.citation_verified,
            "blocks": [block.model_dump() for block in page.blocks]
            if hasattr(page, "blocks")
            else None,
        }
        for index, page in enumerate(parsed, start=1)
    ]


def item_for(unit: dict, chunk: dict, *, quote: str | None = None, starred: bool = False):
    return ExtractedRequirement(
        category=Category.substantive,
        starred=starred,
        text=unit["quote"],
        source=Source(
            document_id=chunk["document_id"],
            chunk_id=chunk["id"],
            quote=quote or unit["quote"],
            page=unit["page"],
            location=(
                Location.model_validate(
                    {key: unit["location"].get(key) for key in Location.model_fields}
                )
                if unit["location"]
                else None
            ),
        ),
    )


async def parse(name: str, content: bytes):
    pages, _, warnings = await parse_document(content, Path(name).suffix, LocalOCR("eng", None), 10)
    return make_chunks(pages, name), warnings


async def run() -> dict:
    pdf_chunks, pdf_warnings = await parse("synthetic.pdf", make_pdf())
    pdf_units = star_units(pdf_chunks)
    pdf_page = pdf_chunks[0]
    pdf_items = [
        item_for(pdf_units[0], pdf_page, starred=False),
        item_for(pdf_units[2], pdf_page, starred=True),
        # Same-page and even position-correct text is not a match when the clause differs.
        item_for(
            pdf_units[1], pdf_page, quote="Ordinary clause: submit authorization.", starred=False
        ),
        # This quote repeats twice on the same page, so citation verification rejects it.
        item_for(pdf_units[3], pdf_page, starred=True),
        # A correct quote at an invalid page is rejected before coverage is computed.
        item_for(pdf_units[1], pdf_page, starred=True).model_copy(
            update={
                "source": item_for(pdf_units[1], pdf_page).source.model_copy(update={"page": 2})
            }
        ),
    ]
    pdf_metric = star_coverage(pdf_items, pdf_units, pdf_chunks)
    assert len(pdf_units) == 5, f"expected 5 PDF star clauses, found {len(pdf_units)}"
    assert (
        pdf_units[2]["quote"] == "★Cross-line clause requires\ncontinuous service for three years."
    )
    assert pdf_metric["covered_clauses"] == 2, pdf_metric
    assert pdf_metric["star_labelled_clauses"] == 1, pdf_metric

    doc = Document()
    doc.add_heading("Technical requirements", level=1)
    doc.add_paragraph("★Clause C required; ★Clause D required.")
    doc.add_paragraph("★Repeated text across blocks.")
    doc.add_paragraph("★Repeated text across blocks.")
    doc.add_paragraph("★Same block repeated text. ★Same block repeated text.")
    word_chunks, word_warnings = await parse("synthetic.docx", docx_bytes(doc))
    word_units = star_units(word_chunks)
    word_metric = star_coverage(
        [
            item_for(word_units[0], word_chunks[0], starred=True),
            item_for(word_units[1], word_chunks[0], starred=True),
            item_for(word_units[2], word_chunks[0], starred=True),
        ],
        word_units,
        word_chunks,
    )
    word_cross_block_probe = star_coverage(
        [item_for(word_units[2], word_chunks[0], starred=True)], word_units, word_chunks
    )
    same_block_duplicate_probe = star_coverage(
        [item_for(word_units[4], word_chunks[0], starred=True)], word_units[4:6], word_chunks
    )
    assert len(word_units) == 6, f"expected 6 Word star clauses, found {len(word_units)}"
    assert word_metric["covered_clauses"] == 3, word_metric
    assert word_cross_block_probe["covered_clauses"] == 1, word_cross_block_probe
    assert same_block_duplicate_probe["covered_clauses"] == 0, same_block_duplicate_probe
    assert word_units[0]["position"].startswith("p"), word_units[0]

    unrelated_only_metric = star_coverage(
        [item_for(pdf_units[1], pdf_page, quote="Ordinary clause: submit authorization.")],
        pdf_units,
        pdf_chunks,
    )
    assert unrelated_only_metric["source_star_clauses"] == 5
    assert unrelated_only_metric["covered_clauses"] == 0, unrelated_only_metric

    plain_chunks, _ = await parse("synthetic.pdf", make_plain_pdf())
    plain_metric = star_coverage([], star_units(plain_chunks), plain_chunks)
    assert plain_metric["source_star_clauses"] == 0
    assert plain_metric["clause_coverage"] is None
    assert plain_metric["star_label_coverage"] is None
    return {
        "mode": "offline_parser_acceptance",
        "provider_calls": 0,
        "metric_scope": "Parser-derived ★ clause diagnostic, not human-annotated recall.",
        "failure_modes_checked": FAILURE_MODES,
        "pdf": {"units": pdf_units, "coverage": pdf_metric, "warnings": pdf_warnings},
        "docx": {
            "units": word_units,
            "coverage": word_metric,
            "one_block_citation_probe": word_cross_block_probe,
            "same_block_duplicate_probe": same_block_duplicate_probe,
            "warnings": word_warnings,
        },
        "same_page_unrelated_only_probe": unrelated_only_metric,
        "no_star_document": plain_metric,
        "result": "passed",
    }


def make_plain_pdf() -> bytes:
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "Ordinary requirement: submit authorization.", fontname="china-s")
    content = document.tobytes()
    document.close()
    return content


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    if args.output.exists():
        parser.error("Choose a new output path; existing files are not overwritten")
    artifact = asyncio.run(run())
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(artifact, ensure_ascii=False, indent=2))
    print(f"Offline extraction acceptance passed; artifact written to {args.output}")


if __name__ == "__main__":
    main()
