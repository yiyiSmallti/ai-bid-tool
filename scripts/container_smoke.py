"""Exercise a fresh, local-only Compose deployment with synthetic documents.

Refuses an existing output directory or a remote Docker endpoint. Only its own
unique project is stopped; all test volumes and records are retained.
"""

import argparse
import json
import os
import secrets
import subprocess
import sys
import time
from pathlib import Path
from urllib.parse import quote
from uuid import UUID, uuid4

import httpx
import pymupdf
from app.core.config import Settings
from app.providers.storage import S3Storage
from cryptography.fernet import Fernet


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--host", required=True, help="Explicit socket for an isolated Docker instance"
    )
    parser.add_argument("--root", type=Path, required=True, help="New test-only directory")
    parser.add_argument("--port", type=int, default=48900)
    parser.add_argument("--s3-port", type=int, default=49000)
    parser.add_argument("--compose", default="/opt/homebrew/lib/docker/cli-plugins/docker-compose")
    args = parser.parse_args()
    if not args.host.startswith("unix://") or args.root.exists():
        raise SystemExit("Use an explicit local Docker socket and a new test directory")
    args.root.mkdir(mode=0o700, parents=True)
    project = "bid-test-" + uuid4().hex[:12]
    key = Fernet.generate_key().decode()
    password = secrets.token_urlsafe(24)
    values = {
        "BID_OWNER_PASSWORD": secrets.token_urlsafe(24),
        "BID_DATABASE_PASSWORD": secrets.token_urlsafe(24),
        "BID_ENCRYPTION_KEY": key,
        "BID_CLI_KEY": Fernet.generate_key().decode(),
        "BID_S3_ACCESS_KEY": "synthetic-test-" + uuid4().hex[:12],
        "BID_S3_SECRET_KEY": secrets.token_urlsafe(24),
    }
    values["BID_MIGRATION_DATABASE_URL"] = (
        "postgresql+psycopg://bid_owner:"
        + quote(values["BID_OWNER_PASSWORD"], safe="")
        + "@postgres/bid"
    )
    values["BID_DATABASE_URL"] = (
        "postgresql+psycopg://bid_app:"
        + quote(values["BID_DATABASE_PASSWORD"], safe="")
        + "@postgres/bid"
    )
    env_file = args.root / "test.env"
    with os.fdopen(os.open(env_file, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600), "w") as handle:
        handle.write("\n".join(f"{name}={value}" for name, value in values.items()))
    docker_config = args.root / "docker"
    docker_config.mkdir(mode=0o700)
    (docker_config / "config.json").write_text(
        json.dumps({"cliPluginsExtraDirs": ["/opt/homebrew/lib/docker/cli-plugins"]})
    )
    environment = {
        **os.environ,
        **values,
        "DOCKER_HOST": args.host,
        "DOCKER_CONFIG": str(docker_config),
        "BID_BOOTSTRAP_PASSWORD": password,
        "BID_PASSWORD": password,
    }
    override = args.root / "override.yml"
    s3 = """      BID_STORAGE: s3
      BID_S3_ENDPOINT: http://minio:9000
      BID_S3_BUCKET: bid-test
      BID_S3_ACCESS_KEY: ${BID_S3_ACCESS_KEY}
      BID_S3_SECRET_KEY: ${BID_S3_SECRET_KEY}
"""
    override.write_text(
        f"services:\n  minio:\n    ports: !override\n      - 127.0.0.1:{args.s3_port}:9000\n"
        f"  server:\n    ports: !override\n      - 127.0.0.1:{args.port}:8000\n    environment:\n{s3}"
        f"  worker:\n    environment:\n{s3}"
    )
    repository = Path(__file__).resolve().parents[1]
    compose = [
        args.compose,
        "--env-file",
        str(env_file),
        "--project-name",
        project,
        "-f",
        str(repository / "deploy/docker-compose.yml"),
        "-f",
        str(override),
    ]
    (args.root / "project.json").write_text(json.dumps({"project": project, "host": args.host}))
    log = (args.root / "compose.log").open("w")
    evidence = {"project": project, "synthetic_only": True, "checks": []}

    def run(arguments, capture=False):
        result = subprocess.run(
            arguments,
            env=environment,
            check=True,
            stdout=subprocess.PIPE if capture else log,
            stderr=log,
            text=True,
        )
        return result.stdout or ""

    def cli(arguments):
        process = subprocess.run(
            [
                sys.executable,
                "-m",
                "bid_cli.main",
                "--mode",
                "remote",
                "--server",
                base,
                "--state",
                str(args.root / "session.enc"),
                *arguments,
                "--json",
            ],
            env=environment,
            stdout=subprocess.PIPE,
            stderr=log,
            text=True,
        )
        return process.returncode, json.loads(process.stdout)

    def checked(name):
        evidence["checks"].append(name)
        (args.root / "result.json").write_text(json.dumps(evidence, indent=2))
        print("Passed:", name, flush=True)

    base = f"http://127.0.0.1:{args.port}"
    endpoint = f"http://127.0.0.1:{args.s3_port}"
    try:
        print(
            "Building and starting the isolated test project; log:",
            args.root / "compose.log",
            flush=True,
        )
        run([*compose, "up", "-d", "--build"])
        checked("compose build and startup")
        with httpx.Client(timeout=5) as http:
            for _ in range(60):
                try:
                    if (
                        http.get(base + "/health").status_code == 200
                        and http.get(endpoint + "/minio/health/ready").status_code == 200
                    ):
                        break
                except httpx.HTTPError:
                    pass
                time.sleep(1)
            else:
                raise RuntimeError("Test services did not become ready within one minute")
            checked("API and MinIO ready")
            storage = S3Storage(
                Settings(
                    database_url="postgresql+psycopg://unused@localhost/bid_test",
                    encryption_key=key,
                    storage="s3",
                    s3_endpoint=endpoint,
                    s3_bucket="bid-test",
                    s3_access_key=values["BID_S3_ACCESS_KEY"],
                    s3_secret_key=values["BID_S3_SECRET_KEY"],
                )
            )
            storage.client.create_bucket(Bucket="bid-test")
            assert http.get(endpoint + "/bid-test").status_code == 403
            checked("fresh private test bucket")
            organizations = []
            sessions = []
            for label in ("a", "b"):
                output = run(
                    [
                        *compose,
                        "run",
                        "--rm",
                        "--no-deps",
                        "-e",
                        "BID_BOOTSTRAP_PASSWORD",
                        "migrate",
                        "python",
                        "-m",
                        "app.admin",
                        "bootstrap",
                        "--org-name",
                        "Synthetic container test " + label,
                        "--email",
                        label + "@example.invalid",
                    ],
                    capture=True,
                )
                org = str(UUID(output.strip().split()[-1]))
                organizations.append(org)
                response = http.post(
                    base + "/auth/login",
                    json={"email": label + "@example.invalid", "password": password, "org_id": org},
                )
                assert response.status_code == 200
                sessions.append(
                    {
                        "Authorization": "Bearer " + response.json()["data"]["session"],
                        "X-Org-Id": org,
                    }
                )
            checked("new synthetic organizations and login")
            status, _ = cli(["login", "--email", "a@example.invalid", "--org", organizations[0]])
            assert status == 0
            document_path = args.root / "synthetic.pdf"
            with pymupdf.open() as document:
                page = document.new_page()
                page.insert_text(
                    (40, 60),
                    "Synthetic container fixture, not a real tender. Minimum memory 64 GB.",
                )
                document.save(document_path)
            status, task = cli(
                [
                    "task",
                    "create",
                    "--name",
                    "Synthetic container task",
                    "--tender",
                    str(document_path),
                ]
            )
            assert status == 0
            document_id = task["data"]["document_id"]
            status, parsed = cli(
                ["tender", "parse", "--document", document_id, "--wait", "--timeout", "30"]
            )
            assert (
                status == 0
                and parsed["data"]["pages"] == 1
                and parsed["items"][0]["citation_verified"]
            )
            checked("remote CLI upload and genuine worker parsing over S3")
            status, duplicate = cli(
                ["tender", "upload", "--task", task["data"]["id"], "--file", str(document_path)]
            )
            assert status == 0 and duplicate["data"]["duplicate"]
            checked("duplicate upload is idempotent")
            metadata = {
                "name": "Synthetic container product",
                "vendor": "Synthetic vendor",
                "model": "Exact A",
            }
            input_file = args.root / "product.json"
            input_file.write_text(json.dumps({"data": metadata}))
            status, product = cli(["resource", "product", "add", "--input", str(input_file)])
            assert status == 0 and product["data"]["revision"] == 1
            product_id = product["data"]["product_id"]
            selection_file = args.root / "selection.json"
            selection_file.write_text(
                json.dumps({"product_id": product_id, "lot": "synthetic-lot"})
            )
            selection_args = [
                "task",
                "resource",
                "add",
                "--task",
                task["data"]["id"],
                "--input",
                str(selection_file),
            ]
            status, selected = cli(selection_args)
            assert status == 0 and selected["data"]["revision"] == 1
            input_file.write_text(
                json.dumps({"expected_revision": 1, "data": {**metadata, "model": "Exact B"}})
            )
            update_args = [
                "resource",
                "product",
                "update",
                "--id",
                product_id,
                "--input",
                str(input_file),
            ]
            status, revised = cli(update_args)
            assert status == 0 and revised["data"]["revision"] == 2
            status, conflict = cli(update_args)
            assert status == 2 and conflict["data"]["error"]["code"] == "revision_conflict"
            status, fixed = cli(["task", "resource", "list", "--task", task["data"]["id"]])
            assert status == 0 and fixed["items"][0]["data"]["model"] == "Exact A"
            checked("version update conflicts and immutable task snapshot")
            status, replacement = cli(selection_args)
            assert (
                status == 0
                and replacement["data"]["replaced_snapshot_id"] == selected["data"]["id"]
            )
            status, repeated = cli(selection_args)
            assert status == 0 and repeated["data"]["duplicate"]
            status, history = cli(
                ["task", "resource", "list", "--task", task["data"]["id"], "--history"]
            )
            assert status == 0 and [item["revision"] for item in history["items"]] == [1, 2]
            hidden = http.get(
                base + "/resources/products", headers=sessions[1], params={"product_id": product_id}
            )
            assert hidden.status_code == 404
            checked("selection history, duplicate operation and resource tenant isolation")
            feature_data = {
                "product_id": product_id,
                "name": "Synthetic container feature",
                "description": "Synthetic declaration, no verified screenshot",
                "status": "planned",
            }
            feature_input = args.root / "feature.json"
            feature_input.write_text(json.dumps({"data": feature_data}))
            status, feature = cli(["resource", "feature", "add", "--input", str(feature_input)])
            assert status == 0 and feature["warnings"]
            feature_id = feature["data"]["feature_id"]
            selection_file.write_text(json.dumps({"feature_id": feature_id}))
            feature_selection_args = [
                "task",
                "feature",
                "add",
                "--task",
                task["data"]["id"],
                "--input",
                str(selection_file),
            ]
            status, feature_selected = cli(feature_selection_args)
            assert status == 0 and feature_selected["data"]["revision"] == 1
            feature_input.write_text(
                json.dumps(
                    {"expected_revision": 1, "data": {**feature_data, "status": "developing"}}
                )
            )
            feature_update_args = [
                "resource",
                "feature",
                "update",
                "--id",
                feature_id,
                "--input",
                str(feature_input),
            ]
            status, feature_updated = cli(feature_update_args)
            assert status == 0 and feature_updated["data"]["revision"] == 2
            status, feature_conflict = cli(feature_update_args)
            assert status == 2 and feature_conflict["data"]["error"]["code"] == "revision_conflict"
            status, feature_fixed = cli(["task", "feature", "list", "--task", task["data"]["id"]])
            assert status == 0 and feature_fixed["items"][0]["data"]["status"] == "planned"
            checked("declared feature versions and immutable task snapshot")
            status, feature_replaced = cli(feature_selection_args)
            assert (
                status == 0
                and feature_replaced["data"]["replaced_snapshot_id"]
                == feature_selected["data"]["id"]
            )
            status, feature_repeated = cli(feature_selection_args)
            assert status == 0 and feature_repeated["data"]["duplicate"]
            status, feature_history = cli(
                ["task", "feature", "list", "--task", task["data"]["id"], "--history"]
            )
            assert status == 0 and [item["revision"] for item in feature_history["items"]] == [1, 2]
            assert (
                http.get(
                    base + "/resources/features",
                    headers=sessions[1],
                    params={"feature_id": feature_id},
                ).status_code
                == 404
            )
            checked("feature history, repeat, conflict and tenant isolation")
            certificate_data = {
                "kind": "qualification",
                "name": "Synthetic declaration",
                "number": "SYNTHETIC-ONLY",
                "valid_from": "2026-01-01",
                "valid_until": "2026-12-31",
            }
            certificate_input = args.root / "certificate.json"
            certificate_input.write_text(json.dumps({"data": certificate_data}))
            status, certificate = cli(
                ["resource", "certificate", "add", "--input", str(certificate_input)]
            )
            assert status == 0 and certificate["warnings"]
            certificate_id = certificate["data"]["certificate_id"]
            selection_file.write_text(json.dumps({"certificate_id": certificate_id}))
            certificate_selection_args = [
                "task",
                "certificate",
                "add",
                "--task",
                task["data"]["id"],
                "--input",
                str(selection_file),
            ]
            status, certificate_selected = cli(certificate_selection_args)
            assert status == 0 and certificate_selected["data"]["revision"] == 1
            certificate_input.write_text(
                json.dumps(
                    {
                        "expected_revision": 1,
                        "data": {**certificate_data, "valid_until": "2027-12-31"},
                    }
                )
            )
            certificate_update_args = [
                "resource",
                "certificate",
                "update",
                "--id",
                certificate_id,
                "--input",
                str(certificate_input),
            ]
            status, certificate_updated = cli(certificate_update_args)
            assert status == 0 and certificate_updated["data"]["revision"] == 2
            status, certificate_conflict = cli(certificate_update_args)
            assert (
                status == 2 and certificate_conflict["data"]["error"]["code"] == "revision_conflict"
            )
            status, certificate_fixed = cli(
                ["task", "certificate", "list", "--task", task["data"]["id"]]
            )
            assert (
                status == 0 and certificate_fixed["items"][0]["data"]["valid_until"] == "2026-12-31"
            )
            checked("declared certificate versions and immutable task snapshot")
            status, certificate_replaced = cli(certificate_selection_args)
            assert (
                status == 0
                and certificate_replaced["data"]["replaced_snapshot_id"]
                == certificate_selected["data"]["id"]
            )
            status, certificate_repeated = cli(certificate_selection_args)
            assert status == 0 and certificate_repeated["data"]["duplicate"]
            status, certificate_history = cli(
                ["task", "certificate", "list", "--task", task["data"]["id"], "--history"]
            )
            assert status == 0 and [item["revision"] for item in certificate_history["items"]] == [
                1,
                2,
            ]
            assert (
                http.get(
                    base + "/resources/certificates",
                    headers=sessions[1],
                    params={"certificate_id": certificate_id},
                ).status_code
                == 404
            )
            checked("certificate history, repeat, conflict and tenant isolation")
            status, inspected = cli(
                [
                    "task",
                    "certificate",
                    "list",
                    "--task",
                    task["data"]["id"],
                    "--history",
                    "--as-of",
                    "2027-01-01",
                ]
            )
            assert status == 0 and sorted(
                value["state"] for value in inspected["data"]["validity_by_revision"].values()
            ) == ["expired", "valid"]
            profile_data = {
                "name": "Synthetic organization declaration",
                "registration_details": None,
                "performance_summary": "Synthetic A",
                "standard_wording": None,
            }
            profile_input = args.root / "profile.json"
            profile_input.write_text(json.dumps({"data": profile_data}))
            status, profile = cli(["resource", "profile", "add", "--input", str(profile_input)])
            assert status == 0 and profile["warnings"]
            profile_id = profile["data"]["profile_id"]
            selection_file.write_text(json.dumps({"profile_id": profile_id}))
            profile_selection_args = [
                "task",
                "profile",
                "add",
                "--task",
                task["data"]["id"],
                "--input",
                str(selection_file),
            ]
            status, profile_selected = cli(profile_selection_args)
            assert status == 0 and profile_selected["data"]["revision"] == 1
            profile_input.write_text(
                json.dumps(
                    {
                        "expected_revision": 1,
                        "data": {**profile_data, "performance_summary": "Synthetic B"},
                    }
                )
            )
            profile_update_args = [
                "resource",
                "profile",
                "update",
                "--id",
                profile_id,
                "--input",
                str(profile_input),
            ]
            status, profile_updated = cli(profile_update_args)
            assert status == 0 and profile_updated["data"]["revision"] == 2
            status, profile_conflict = cli(profile_update_args)
            assert status == 2 and profile_conflict["data"]["error"]["code"] == "revision_conflict"
            status, profile_fixed = cli(["task", "profile", "list", "--task", task["data"]["id"]])
            assert (
                status == 0
                and profile_fixed["items"][0]["data"]["performance_summary"] == "Synthetic A"
            )
            checked("declared profile versions and immutable task snapshot")
            status, profile_replaced = cli(profile_selection_args)
            assert (
                status == 0
                and profile_replaced["data"]["replaced_snapshot_id"]
                == profile_selected["data"]["id"]
            )
            status, profile_repeated = cli(profile_selection_args)
            assert status == 0 and profile_repeated["data"]["duplicate"]
            status, profile_history = cli(
                ["task", "profile", "list", "--task", task["data"]["id"], "--history"]
            )
            assert status == 0 and [item["revision"] for item in profile_history["items"]] == [
                1,
                2,
            ]
            assert (
                http.get(
                    base + "/resources/profiles",
                    headers=sessions[1],
                    params={"profile_id": profile_id},
                ).status_code
                == 404
            )
            checked("profile history, repeat, conflict and tenant isolation")
            from io import BytesIO

            from docx import Document

            template_document = Document()
            template_document.add_paragraph("Synthetic container DOCX only")
            initial_bytes = BytesIO()
            template_document.save(initial_bytes)
            template_file = args.root / "synthetic.docx"
            template_file.write_bytes(initial_bytes.getvalue())
            template_input = args.root / "template.json"
            template_data = {"name": "Synthetic template", "project_types": None, "chapters": None}
            template_input.write_text(json.dumps({"data": template_data}))
            status, template = cli(
                [
                    "resource",
                    "template",
                    "add",
                    "--input",
                    str(template_input),
                    "--file",
                    str(template_file),
                ]
            )
            assert status == 0 and template["warnings"]
            template_id = template["data"]["template_id"]
            selection_file.write_text(json.dumps({"template_id": template_id}))
            template_selection_args = [
                "task",
                "template",
                "add",
                "--task",
                task["data"]["id"],
                "--input",
                str(selection_file),
            ]
            status, template_selected = cli(template_selection_args)
            assert status == 0 and template_selected["data"]["revision"] == 1
            template_document.add_paragraph("Synthetic second revision")
            new_bytes = BytesIO()
            template_document.save(new_bytes)
            template_file.write_bytes(new_bytes.getvalue())
            template_input.write_text(
                json.dumps(
                    {"expected_revision": 1, "data": {**template_data, "name": "Synthetic updated"}}
                )
            )
            template_update_args = [
                "resource",
                "template",
                "update",
                "--id",
                template_id,
                "--input",
                str(template_input),
                "--file",
                str(template_file),
            ]
            status, updated = cli(template_update_args)
            assert status == 0 and updated["data"]["revision"] == 2
            status, conflict = cli(template_update_args)
            assert status == 4 and conflict["data"]["error"]["code"] == "revision_conflict"
            status, fixed = cli(["task", "template", "list", "--task", task["data"]["id"]])
            assert (
                status == 0 and fixed["items"][0]["template_revision_id"] == template["data"]["id"]
            )
            checked("DOCX immutable file revisions and fixed task selection")
            status, replaced = cli(template_selection_args)
            assert (
                status == 0
                and replaced["data"]["replaced_snapshot_id"] == template_selected["data"]["id"]
            )
            status, repeat = cli(template_selection_args)
            assert status == 0 and repeat["data"]["duplicate"]
            status, template_history = cli(
                ["task", "template", "list", "--task", task["data"]["id"], "--history"]
            )
            assert status == 0 and [entry["revision"] for entry in template_history["items"]] == [
                1,
                2,
            ]
            assert (
                http.get(
                    base + "/resources/templates",
                    headers=sessions[1],
                    params={"template_id": template_id},
                ).status_code
                == 404
            )
            checked("template history, duplicate, conflict and tenant isolation")
            for identifier, expected in [
                (template["data"]["id"], initial_bytes.getvalue()),
                (updated["data"]["id"], new_bytes.getvalue()),
            ]:
                output = args.root.resolve() / f"{identifier}.docx"
                arguments = [
                    "resource",
                    "template",
                    "download",
                    "--revision",
                    identifier,
                    "--output",
                    str(output),
                ]
                status, downloaded = cli(arguments)
                assert status == 0 and downloaded["ok"] and output.read_bytes() == expected
                status, unchanged = cli(arguments)
                assert status == 2 and output.read_bytes() == expected
            scan_source = args.root / "certificate-synthetic.pdf"
            scan_source.write_bytes(document_path.read_bytes())
            scan_input = args.root / "certificate-file.json"
            scan_metadata = {**certificate_data, "valid_until": "2027-12-31"}
            scan_input.write_text(json.dumps({"expected_revision": 2, "data": scan_metadata}))
            scan_add_args = [
                "resource",
                "certificate",
                "file",
                "add",
                "--id",
                certificate_id,
                "--input",
                str(scan_input),
                "--file",
                str(scan_source),
            ]
            status, scan = cli(scan_add_args)
            assert status == 0 and scan["data"]["revision"] == 3
            scan_revision = scan["data"]["certificate_revision_id"]
            status, conflict = cli(scan_add_args)
            assert status == 4 and conflict["data"]["error"]["code"] == "revision_conflict"
            scan_task_args = ["task", "certificate", "file", "list", "--task", task["data"]["id"]]
            status, old_scan = cli([*scan_task_args])
            assert (
                status == 0
                and old_scan["items"][0]["revision"] == 2
                and old_scan["items"][0]["file"] is None
            )
            selection_file.write_text(json.dumps({"certificate_id": certificate_id}))
            status, scan_selected = cli(certificate_selection_args)
            assert status == 0 and scan_selected["data"]["revision"] == 3
            status, scan_repeated = cli(certificate_selection_args)
            assert status == 0 and scan_repeated["data"]["duplicate"]
            scan_input.write_text(
                json.dumps(
                    {
                        "expected_revision": 3,
                        "data": {**scan_metadata, "name": "Synthetic metadata only"},
                    }
                )
            )
            status, plain = cli(
                [
                    "resource",
                    "certificate",
                    "update",
                    "--id",
                    certificate_id,
                    "--input",
                    str(scan_input),
                ]
            )
            assert status == 0 and plain["data"]["revision"] == 4
            status, current = cli(
                ["resource", "certificate", "file", "list", "--id", certificate_id]
            )
            assert status == 0 and not current["items"] and current["warnings"]
            status, fixed = cli([*scan_task_args])
            assert (
                status == 0
                and fixed["items"][0]["certificate_revision_id"] == scan_revision
                and fixed["items"][0]["file"] == scan["data"]["file"]
            )
            status, file_history = cli(
                ["resource", "certificate", "file", "list", "--id", certificate_id, "--history"]
            )
            assert status == 0 and [v["revision"] for v in file_history["items"]] == [3]
            output = args.root.resolve() / "certificate-download.pdf"
            scan_download_args = [
                "resource",
                "certificate",
                "file",
                "download",
                "--revision",
                scan_revision,
                "--output",
                str(output),
            ]
            status, downloaded = cli(scan_download_args)
            assert (
                status == 0
                and output.read_bytes() == document_path.read_bytes()
                and output.stat().st_mode & 0o777 == 0o600
            )
            status, unchanged = cli(scan_download_args)
            assert status == 2 and output.read_bytes() == document_path.read_bytes()
            status, plain_selected = cli(certificate_selection_args)
            assert status == 0 and plain_selected["data"]["revision"] == 4
            status, scan_history = cli([*scan_task_args, "--history"])
            assert status == 0 and [v["revision"] for v in scan_history["items"]] == [1, 2, 3, 4]
            assert [v["file"] is None for v in scan_history["items"]] == [True, True, False, True]
            assert (
                http.get(
                    base + "/resources/certificates/files",
                    headers=sessions[1],
                    params={"certificate_id": certificate_id},
                ).status_code
                == 404
            )
            assert (
                http.get(
                    base + f"/resources/certificates/revisions/{scan_revision}/file/download-link",
                    headers=sessions[1],
                ).status_code
                == 404
            )
            checked(
                "certificate PDF revisions, fixed selections, no inheritance and tenant isolation"
            )
            checked("genuine remote certificate file CLI and no-overwrite download")
            scan_objects = storage.client.list_objects_v2(Bucket="bid-test")["Contents"]
            scan_keys = [entry["Key"] for entry in scan_objects if "/certificate/" in entry["Key"]]
            assert len(scan_keys) == 1
            raw_stream = storage.client.get_object(Bucket="bid-test", Key=scan_keys[0])["Body"]
            try:
                raw = raw_stream.read()
            finally:
                raw_stream.close()
            assert (
                raw.startswith(storage.cipher.marker)
                and storage.cipher.decrypt(scan_keys[0], raw) == document_path.read_bytes()
            )
            scan_link = http.get(
                base + f"/resources/certificates/revisions/{scan_revision}/file/download-link",
                headers=sessions[0],
            ).json()["data"]
            assert (
                http.get(base + scan_link["url"], headers=sessions[0]).content
                == document_path.read_bytes()
            )
            checked("real encrypted S3 certificate original and authenticated download")
            selection_file.write_text(json.dumps({"certificate_id": certificate_id, "revision": 3}))
            status, source_choice = cli(certificate_selection_args)
            assert status == 0
            source_input = args.root / "evidence-source.json"
            source_input.write_text(
                json.dumps({"task_certificate_id": source_choice["data"]["id"], "page": 1})
            )
            source_add = [
                "evidence",
                "source",
                "add",
                "--task",
                task["data"]["id"],
                "--input",
                str(source_input),
            ]
            status, archived = cli(source_add)
            assert status == 0 and not archived["data"]["duplicate"]
            source = archived["data"]["source"]
            assert (
                source["certificate_revision_id"] == scan_revision
                and source["confirmed_by"] is None
                and source["eligible_for_draft_export"] is False
            )
            status, repeated = cli(source_add)
            assert (
                status == 0
                and repeated["data"]["duplicate"]
                and repeated["data"]["source"]["id"] == source["id"]
            )
            png_output = args.root.resolve() / "source-download.png"
            source_download = [
                "evidence",
                "source",
                "download",
                "--id",
                source["id"],
                "--output",
                str(png_output),
            ]
            status, preview = cli(source_download)
            assert status == 0 and png_output.stat().st_mode & 0o777 == 0o600
            actual = pymupdf.Pixmap(png_output.read_bytes())
            with pymupdf.open(stream=document_path.read_bytes(), filetype="pdf") as pdf:
                expected = pdf[0].get_pixmap(dpi=150, colorspace=pymupdf.csRGB, alpha=False)
                assert (actual.width, actual.height, actual.samples) == (
                    expected.width,
                    expected.height,
                    expected.samples,
                )
            assert cli(source_download)[0] == 2
            checked(
                "genuine remote source CLI, decoded original pixels, repeat and private no-overwrite PNG"
            )
            source_path = f"/evidence-sources/{source['id']}/preview/download"
            for method, path, options in [
                (
                    "POST",
                    f"/tasks/{task['data']['id']}/evidence-sources",
                    {"json": {"task_certificate_id": source_choice["data"]["id"], "page": 1}},
                ),
                ("GET", f"/tasks/{task['data']['id']}/evidence-sources", {}),
                ("GET", source_path + "-link", {}),
                ("GET", source_path + "?signature=invalid", {}),
            ]:
                assert (
                    http.request(method, base + path, headers=sessions[1], **options).status_code
                    == 404
                )
            selection_file.write_text(json.dumps({"certificate_id": certificate_id}))
            assert cli(certificate_selection_args)[0] == 0
            source_list = ["evidence", "source", "list", "--task", task["data"]["id"]]
            assert cli(source_list)[1]["items"] == []
            old = cli([*source_list, "--history"])[1]
            assert len(old["items"]) == 1 and not old["items"][0]["active_selection"]
            checked("all source routes hide foreign objects and replaced source remains history")
            source_objects = storage.client.list_objects_v2(Bucket="bid-test")["Contents"]
            source_keys = [
                entry["Key"] for entry in source_objects if "/evidence-source/" in entry["Key"]
            ]
            assert len(source_keys) == 1
            raw_stream = storage.client.get_object(Bucket="bid-test", Key=source_keys[0])["Body"]
            try:
                raw = raw_stream.read()
            finally:
                raw_stream.close()
            assert (
                raw.startswith(storage.cipher.marker)
                and storage.cipher.decrypt(source_keys[0], raw) == png_output.read_bytes()
            )
            source_link = http.get(base + source_path + "-link", headers=sessions[0]).json()["data"]
            assert (
                http.get(base + source_link["url"], headers=sessions[0]).content
                == png_output.read_bytes()
            )
            checked("real encrypted S3 PNG and authenticated history preview download")
            objects = storage.client.list_objects_v2(Bucket="bid-test")["Contents"]
            template_keys = [entry["Key"] for entry in objects if "/template/" in entry["Key"]]
            assert len(template_keys) == 2
            for key in template_keys:
                raw = storage.client.get_object(Bucket="bid-test", Key=key)["Body"].read()
                assert raw.startswith(storage.cipher.marker) and not raw.startswith(b"PK")
                assert storage.cipher.decrypt(key, raw) in (
                    initial_bytes.getvalue(),
                    new_bytes.getvalue(),
                )
            checked("encrypted S3 DOCX files, authenticated downloads and no overwrite")
            object_key = next(
                entry["Key"]
                for entry in objects
                if entry["Key"].endswith(".pdf") and "/certificate/" not in entry["Key"]
            )
            stored = storage.client.get_object(Bucket="bid-test", Key=object_key)["Body"]
            try:
                raw = stored.read()
            finally:
                stored.close()
            assert raw.startswith(storage.cipher.marker) and b"%PDF" not in raw
            assert storage.cipher.decrypt(object_key, raw) == document_path.read_bytes()
            link = http.get(base + f"/documents/{document_id}/download-link", headers=sessions[0])
            download = http.get(base + link.json()["data"]["url"], headers=sessions[0])
            assert download.status_code == 200 and download.content == document_path.read_bytes()
            checked("encrypted objects and signed authenticated download")
            assert (
                http.get(base + f"/documents/{document_id}", headers=sessions[1]).status_code == 404
            )
            assert (
                http.post(
                    base + f"/tasks/{task['data']['id']}/documents",
                    headers=sessions[0],
                    files={"file": ("invalid.pdf", b"invalid")},
                ).status_code
                == 400
            )
            checked("cross-tenant access and invalid file rejection")
            status, unavailable = cli(
                ["req", "extract", "--document", document_id, "--wait", "--timeout", "30"]
            )
            assert status == 4 and unavailable["data"]["error"]["code"] == "provider_unavailable"
            checked("deferred AI fails explicitly")
            evidence["ok"] = True
    finally:
        run([*compose, "stop"])
        (args.root / "result.json").write_text(json.dumps(evidence, indent=2))
        log.close()
        print("Test services stopped; volumes and local records retained at", args.root, flush=True)


if __name__ == "__main__":
    main()
