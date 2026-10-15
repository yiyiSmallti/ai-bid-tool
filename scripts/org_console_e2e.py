"""Provision the single-user native console fixture through API and job processor.

The script accepts only a PostgreSQL database whose name starts with ``bid_test``.
It creates every business object through the authenticated application API. Parse, extraction and draft jobs are run
by the production Processor. The only vendor is an in-process
``httpx.MockTransport`` returning explicitly synthetic, exact citations.

The generated manifest contains object IDs and fixture metadata needed by
native acceptance consumers. It never contains passwords, sessions, authorization headers, signed
URLs or vendor request bodies. Run artifacts and storage must stay under
``data/work`` in this worktree.
"""

from __future__ import annotations

import argparse
import asyncio
import io
import json
import os
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any
from uuid import uuid4

import httpx
import pymupdf
from app.api.main import create_app
from app.core.config import Settings
from app.providers.llm import OpenAICompatibleExtractor
from docx import Document as WordDocument
from sqlalchemy.engine import make_url

REPO_ROOT = Path(__file__).resolve().parents[1]
WORK_ROOT = (REPO_ROOT / "data" / "work").resolve()
WORD_REQUIREMENTS = 1_200
SYNTHETIC_MODEL = "org-console-synthetic-model"


def env(name: str) -> str:
    value = os.environ.get(name)
    if not value:
        raise SystemExit(f"{name} is required")
    return value


def tcp_port(raw: str) -> int:
    value = int(raw)
    if not 1 <= value <= 65_535:
        raise argparse.ArgumentTypeError("port must be between 1 and 65535")
    return value


def guarded_url(name: str) -> str:
    value = env(name)
    database = make_url(value).database or ""
    if not database.startswith("bid_test"):
        raise SystemExit(f"{name} must select a database whose name starts with bid_test")
    return value


def output_directory(raw: str) -> Path:
    output = Path(raw).resolve()
    if not output.is_relative_to(WORK_ROOT):
        raise SystemExit(f"--output must be under {WORK_ROOT}")
    output.mkdir(parents=True, exist_ok=True)
    return output


def configured_data_directory() -> Path:
    data_dir = Path(env("BID_DATA_DIR")).resolve()
    if not data_dir.is_relative_to(WORK_ROOT):
        raise SystemExit(f"BID_DATA_DIR must be under {WORK_ROOT}")
    data_dir.mkdir(parents=True, exist_ok=True)
    return data_dir


def write_upload_materials(output: Path) -> dict[str, str]:
    pdf_path = output / "upload-synthetic.pdf"
    word_path = output / "upload-synthetic.docx"
    pdf_path.write_bytes(pdf_fixture(2, "A"))
    word_path.write_bytes(word_fixture(4, "A"))
    return {"pdf": str(pdf_path), "word": str(word_path)}


def word_fixture(count: int, label: str) -> bytes:
    """Return distinct paragraphs so every quote has one structural location."""
    document = WordDocument()
    for index in range(1, count + 1):
        document.add_paragraph(
            f"合成条款 {label}-{index:04d} 要求投标人在交付清单中提交唯一标识 "
            f"SYN-{label}-{index:04d} 的说明文件。"
        )
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def pdf_fixture(count: int, label: str) -> bytes:
    """Return a real PDF with a different sentence on every page."""
    with pymupdf.open() as document:
        for index in range(1, count + 1):
            page = document.new_page()
            sentence = (
                f"Synthetic tender {label} page {index} requires unique record "
                f"PDF-{label}-{index:02d} at delivery."
            )
            page.insert_text((54, 72), sentence)
        return document.tobytes()


class SyntheticVendor:
    """Build exact extraction citations from the text actually sent by the adapter."""

    block = re.compile(r'<block id="([^"]+)">\n(.*?)\n</block>', re.DOTALL)
    page = re.compile(r'<page number="(\d+)">\n(.*?)\n</page>', re.DOTALL)

    def __init__(self) -> None:
        self.calls: list[dict[str, Any]] = []

    @staticmethod
    def item(ref: str, quote: str) -> dict[str, Any]:
        marker = re.search(r"(?:SYN-[AB]-|PDF-[AB]-)(\d+)", quote)
        ordinal = int(marker.group(1)) if marker else 1
        category = ("qualification", "technical", "scoring", "substantive")[(ordinal - 1) % 4]
        return {
            "category": category,
            "starred": category == "substantive" or ordinal % 17 == 0,
            "text": f"合成响应要求 {ordinal:04d}",
            "ref": ref,
            "quote": quote,
            "condition": None,
        }

    @staticmethod
    def proposals(sent: dict[str, Any]) -> list[dict[str, Any]]:
        # Commitment proposals need no material, so they never cite model-chosen text.
        return [
            {
                "requirement_id": requirement["requirement_id"],
                "response_kind": "commitment",
                "suggested_disposition": "respond",
                "response_text": "合成模型提议：按招标要求响应，待人工审阅。",
                "deviation": "none",
                "deviation_note": "合成提议，仅用于浏览器验收。",
                "evidence": [],
            }
            for requirement in sent["requirements"]
        ]

    def __call__(self, request: httpx.Request) -> httpx.Response:
        body = json.loads(request.content)
        prompt = body["messages"][-1]["content"]
        if prompt.startswith("{"):
            sent = json.loads(prompt)
            items = self.proposals(sent)
            self.calls.append({"kind": "draft", "items": len(items)})
            return self.reply(items, prompt)
        items = [self.item(ref, text.strip()) for ref, text in self.block.findall(prompt)]
        items.extend(self.item(ref, text.strip()) for ref, text in self.page.findall(prompt))
        self.calls.append(
            {
                "kind": "extract",
                "items": len(items),
                "reasoning_effort": body.get("reasoning_effort"),
            }
        )
        return self.reply(items, prompt)

    @staticmethod
    def reply(items: list[dict[str, Any]], prompt: str) -> httpx.Response:
        return httpx.Response(
            200,
            json={
                "model": SYNTHETIC_MODEL,
                "choices": [
                    {
                        "finish_reason": "stop",
                        "message": {"content": json.dumps({"items": items}, ensure_ascii=False)},
                    }
                ],
                "usage": {"prompt_tokens": max(1, len(prompt) // 4), "completion_tokens": 50},
            },
        )


class SyntheticExtractor(OpenAICompatibleExtractor):
    test_only = True


@dataclass
class FixtureQueue:
    calls: list[str] = field(default_factory=list)
    processor: Any = None
    run_jobs: bool = False
    tasks: set[asyncio.Task[None]] = field(default_factory=set)

    async def enqueue(self, job_id: str) -> int:
        self.calls.append(job_id)
        if self.run_jobs:
            if self.processor is None:
                raise RuntimeError("Fixture processor is not configured")
            task = asyncio.create_task(self.processor(job_id))
            self.tasks.add(task)
            task.add_done_callback(self.tasks.discard)
        return len(self.calls)


def fixture_settings(data_dir: Path) -> Settings:
    return Settings.load().model_copy(
        update={
            "data_dir": data_dir,
            "llm_provider": "openai",
            "llm_model": SYNTHETIC_MODEL,
            "llm_api_key": None,
            "llm_base_url": "https://synthetic.example.invalid/v1",
            "llm_concurrency": 8,
            "llm_max_output_tokens": 64_000,
            "converter_url": "https://synthetic.converter.invalid",
        }
    )


def synthetic_converter(request: httpx.Request) -> httpx.Response:
    with pymupdf.open() as pdf:
        for page_num in (1, 2):
            pdf.new_page().insert_text((72, 72), f"Synthetic export preview page {page_num}")
        pdf_bytes = pdf.tobytes()
    return httpx.Response(200, content=pdf_bytes, headers={"Content-Type": "application/pdf"})


def fixture_app(settings: Settings, vendor: SyntheticVendor, *, run_jobs: bool = False):
    llm = SyntheticExtractor(settings, httpx.MockTransport(vendor))
    llm.reasoning_levels = {
        "low": {
            "name": "low",
            "label": "轻度推理",
            "request_options": {"reasoning_effort": "low"},
            "batch_chars": 16_000,
        },
        "high": {
            "name": "high",
            "label": "深度推理",
            "request_options": {"reasoning_effort": "high"},
            "batch_chars": 8_000,
        },
    }
    llm.default_reasoning = "low"
    llm.model_revision = 1
    app = create_app(settings, llm=llm, queue=FixtureQueue(run_jobs=run_jobs))
    app.state.processor.converter_transport = httpx.MockTransport(synthetic_converter)
    return app


async def response_json(
    client: httpx.AsyncClient,
    method: str,
    path: str,
    *,
    expected: int = 200,
    **kwargs: Any,
) -> dict[str, Any]:
    response = await client.request(method, path, **kwargs)
    if response.status_code != expected:
        raise RuntimeError(
            f"{method} {path} returned {response.status_code}: {response.text[:1000]}"
        )
    payload = response.json()
    if expected < 400 and set(payload) != {
        "ok",
        "command",
        "data",
        "items",
        "warnings",
        "cost",
        "duration_ms",
    }:
        raise RuntimeError(f"{method} {path} did not return the Result seven-key contract")
    return payload


async def human_session(client: httpx.AsyncClient) -> dict[str, str]:
    payload = await response_json(
        client, "POST", "/auth/session", json={"access_key": env("BID_ACCESS_KEY")}
    )
    return {"Authorization": "Bearer " + payload["data"]["session"]}


async def run_job(
    client: httpx.AsyncClient,
    app: Any,
    headers: dict[str, str],
    document_id: str,
    action: str,
    body: dict[str, Any] | None = None,
) -> dict[str, Any]:
    receipt = await response_json(
        client,
        "POST",
        f"/documents/{document_id}/{action}",
        headers=headers,
        json=body or {},
    )
    job_id = receipt["data"]["job_id"]
    await app.state.processor(job_id)
    status = await response_json(client, "GET", f"/jobs/{job_id}", headers=headers)
    if status["data"]["status"] != "succeeded":
        raise RuntimeError(f"{action} job failed: {status['data']}")
    return status["data"]


async def create_task_documents(
    client: httpx.AsyncClient,
    app: Any,
    headers: dict[str, str],
    task_key: str,
    word_count: int,
) -> dict[str, Any]:
    task = await response_json(
        client,
        "POST",
        "/tasks",
        headers=headers,
        json={
            "name": f"合成投标审阅任务 {task_key.upper()}",
            "tender_number": f"E2E-{task_key.upper()}-2026",
            "deadline": "2027-01-15T12:00:00+08:00",
        },
    )
    task_id = task["data"]["id"]
    word = word_fixture(word_count, task_key.upper())
    pdf = pdf_fixture(3, task_key.upper())
    uploaded: dict[str, dict[str, Any]] = {}
    for kind, content, name, media_type in (
        (
            "word",
            word,
            f"synthetic-{task_key}-requirements.docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ),
        ("pdf", pdf, f"synthetic-{task_key}-pages.pdf", "application/pdf"),
    ):
        result = await response_json(
            client,
            "POST",
            f"/tasks/{task_id}/documents",
            headers=headers,
            files={"file": (name, content, media_type)},
        )
        uploaded[kind] = result["data"]
        uploaded[kind]["parse_job"] = await run_job(
            client, app, headers, result["data"]["id"], "parse"
        )

    extraction_preview = await response_json(
        client,
        "POST",
        f"/documents/{uploaded['word']['id']}/extract",
        headers=headers,
        json={"dry_run": True},
    )
    low = await run_job(
        client,
        app,
        headers,
        uploaded["word"]["id"],
        "extract",
        {"reasoning": "low"},
    )
    high = await run_job(
        client,
        app,
        headers,
        uploaded["word"]["id"],
        "extract",
        {"reasoning": "high"},
    )
    pdf_extract = await run_job(
        client,
        app,
        headers,
        uploaded["pdf"]["id"],
        "extract",
        {"reasoning": "low"},
    )
    requirements = await response_json(
        client,
        "GET",
        f"/tasks/{task_id}/requirements",
        headers=headers,
        params={"job": low["id"]},
    )
    pdf_requirements = await response_json(
        client,
        "GET",
        f"/tasks/{task_id}/requirements",
        headers=headers,
        params={"job": pdf_extract["id"]},
    )
    if (
        len(requirements["items"]) != word_count
        or len({row["source"]["quote"] for row in requirements["items"]}) != word_count
        or len(pdf_requirements["items"]) != 3
        or len({row["source"]["quote"] for row in pdf_requirements["items"]}) != 3
    ):
        raise RuntimeError("The processor did not persist the expected exact-citation fixture")
    return {
        "task": task["data"],
        "documents": uploaded,
        "extraction_preview": extraction_preview["data"],
        "extractions": {"low": low, "high": high, "pdf": pdf_extract},
        "requirements": requirements["items"],
        "pdf_requirements": pdf_requirements["items"],
    }


async def create_card(
    client: httpx.AsyncClient,
    headers: dict[str, str],
    task_id: str,
    extraction_job_id: str,
    requirement_id: str,
    content: dict[str, Any],
) -> dict[str, Any]:
    result = await response_json(
        client,
        "POST",
        f"/tasks/{task_id}/cards",
        headers=headers,
        json={
            "extraction_job_id": extraction_job_id,
            "requirement_id": requirement_id,
            "content": content,
        },
    )
    return result["data"]


async def card_action(
    client: httpx.AsyncClient,
    headers: dict[str, str],
    card: dict[str, Any],
    action: str,
    **extra: Any,
) -> dict[str, Any]:
    result = await response_json(
        client,
        "POST",
        f"/cards/{card['id']}/actions",
        headers=headers,
        json={"expected_revision": card["revision"], "action": action, **extra},
    )
    return result["data"]


def commitment_content(index: int, deviation: str = "none") -> dict[str, Any]:
    return {
        "response_kind": "commitment",
        "response_text": f"我方承诺按合成条款 {index:04d} 完成响应。",
        "deviation": deviation,
        "deviation_note": (
            "该合成响应保留一项负偏离供界面醒目展示。"
            if deviation == "negative"
            else "该合成响应与招标要求一致。"
        ),
        "evidence": [],
    }


async def confirm_requirements(
    client: httpx.AsyncClient,
    headers: dict[str, str],
    task_id: str,
    extraction_job_id: str,
    requirements: list[dict[str, Any]],
) -> None:
    """Review synthetic exact citations through the human API before response work."""
    scope_path = f"/tasks/{task_id}/extractions/{extraction_job_id}"
    reviews = []
    params: dict[str, Any] = {"limit": 100}
    # Review every page before writing: confirmations invalidate pagination cursors.
    while True:
        page = await response_json(
            client, "GET", f"{scope_path}/requirement-reviews", headers=headers, params=params
        )
        reviews.extend(page["items"])
        cursor = page["data"]["next_cursor"]
        if cursor is None:
            break
        params["cursor"] = cursor
    sources = {row["id"]: row["source"] for row in requirements}
    if len(reviews) != len(sources) or {row["requirement_id"] for row in reviews} != set(sources):
        raise RuntimeError("Requirement review must cover the exact synthetic extraction")
    for review in reviews:
        source = sources[review["requirement_id"]]
        if (
            not review["citation_valid"]
            or review["source_pin"] is None
            or review["content"]["source"] != source
            or review["source_pin"]["source"] != source
        ):
            raise RuntimeError("Synthetic requirement review needs its verified exact source")
    set_revision = page["data"]["scope"]["revision"]
    for offset in range(0, len(reviews), 100):
        batch = reviews[offset : offset + 100]
        confirmed = await response_json(
            client,
            "POST",
            f"{scope_path}/requirement-confirmations",
            headers=headers,
            json={
                "request_id": str(uuid4()),
                "expected_set_revision": set_revision,
                "reviewed_each": True,
                "reason": "Synthetic human reviewed every selected requirement and its exact source.",
                "items": [
                    {
                        "requirement_id": review["requirement_id"],
                        "expected_revision": review["revision"],
                        "expected_review_hash": review["review_hash"],
                    }
                    for review in batch
                ],
            },
        )
        if (
            confirmed["data"]["changed"] != len(batch)
            or {row["requirement_id"] for row in confirmed["items"]}
            != {row["requirement_id"] for row in batch}
            or any(row["current_state"] != "confirmed" for row in confirmed["items"])
        ):
            raise RuntimeError(
                "The human requirement confirmation batch did not confirm every item"
            )
        set_revision = confirmed["data"]["scope"]["revision"]


async def create_review_fixture(
    client: httpx.AsyncClient,
    app: Any,
    headers: dict[str, str],
    task_data: dict[str, Any],
) -> dict[str, Any]:
    task_id = task_data["task"]["id"]
    extraction = task_data["extractions"]["low"]["id"]
    requirements = task_data["requirements"]
    await confirm_requirements(client, headers, task_id, extraction, requirements)

    product_data = {
        "name": "合成存储设备声明",
        "vendor": "Synthetic Vendor",
        "model": "SYN-MODEL-64",
        "model_version": "2026.10",
        "official_url": "https://vendor.example.invalid/synthetic-model",
    }
    product = await response_json(
        client,
        "POST",
        "/resources/products",
        headers=headers,
        json={"data": product_data},
    )
    selection = await response_json(
        client,
        "POST",
        f"/tasks/{task_id}/products",
        headers=headers,
        json={"product_id": product["data"]["product_id"]},
    )
    evidence_input = {
        "kind": "product",
        "selection_id": selection["data"]["id"],
        "field_path": "model",
        "quote": product_data["model"],
    }
    second_evidence_input = {
        "kind": "product",
        "selection_id": selection["data"]["id"],
        "field_path": "model_version",
        "quote": product_data["model_version"],
    }

    cards: dict[str, dict[str, Any]] = {}
    cards["confirmed_evidence"] = await create_card(
        client,
        headers,
        task_id,
        extraction,
        requirements[0]["id"],
        {
            "response_kind": "evidence",
            "response_text": "合成型号声明可逐字支持本条响应。",
            "deviation": "none",
            "deviation_note": "型号声明与合成要求一致。",
            "evidence": [evidence_input],
        },
    )
    cards["confirmed_evidence"] = await card_action(
        client, headers, cards["confirmed_evidence"], "submit"
    )
    cards["confirmed_evidence"] = await card_action(
        client,
        headers,
        cards["confirmed_evidence"],
        "confirm",
        reviewed_evidence_ids=[item["id"] for item in cards["confirmed_evidence"]["evidence"]],
    )

    cards["pending_evidence"] = await create_card(
        client,
        headers,
        task_id,
        extraction,
        requirements[1]["id"],
        {
            "response_kind": "evidence",
            "response_text": "待人工逐项核对的合成证据响应。",
            "deviation": "none",
            "deviation_note": "证据尚未经过人工确认。",
            "evidence": [evidence_input],
        },
    )
    cards["pending_evidence"] = await card_action(
        client, headers, cards["pending_evidence"], "submit"
    )

    cards["rejected"] = await create_card(
        client,
        headers,
        task_id,
        extraction,
        requirements[4]["id"],
        commitment_content(5),
    )
    cards["rejected"] = await card_action(client, headers, cards["rejected"], "submit")
    cards["rejected"] = await card_action(
        client,
        headers,
        cards["rejected"],
        "reject",
        reason="合成审阅：商务响应文字需要修改。",
    )

    cards["needs_material"] = await create_card(
        client,
        headers,
        task_id,
        extraction,
        requirements[5]["id"],
        commitment_content(6),
    )
    cards["needs_material"] = await card_action(client, headers, cards["needs_material"], "submit")
    cards["needs_material"] = await card_action(
        client,
        headers,
        cards["needs_material"],
        "needs_material",
        reason="合成审阅：需要补充真实材料。",
    )

    cards["pending_reject"] = await create_card(
        client,
        headers,
        task_id,
        extraction,
        requirements[12]["id"],
        commitment_content(13),
    )
    cards["pending_reject"] = await card_action(client, headers, cards["pending_reject"], "submit")

    cards["pending_needs_material"] = await create_card(
        client,
        headers,
        task_id,
        extraction,
        requirements[17]["id"],
        commitment_content(18),
    )
    cards["pending_needs_material"] = await card_action(
        client, headers, cards["pending_needs_material"], "submit"
    )

    cards["confirmed_commitment"] = await create_card(
        client,
        headers,
        task_id,
        extraction,
        requirements[9]["id"],
        commitment_content(10, "negative"),
    )
    cards["confirmed_commitment"] = await card_action(
        client, headers, cards["confirmed_commitment"], "submit"
    )
    cards["confirmed_commitment"] = await card_action(
        client,
        headers,
        cards["confirmed_commitment"],
        "confirm",
        reviewed_evidence_ids=[],
    )

    cards["conflict_target"] = await create_card(
        client,
        headers,
        task_id,
        extraction,
        requirements[13]["id"],
        commitment_content(14),
    )
    cards["evidence_draft"] = await create_card(
        client,
        headers,
        task_id,
        extraction,
        requirements[8]["id"],
        {
            "response_kind": "evidence",
            "response_text": "保留为草稿的合成证据响应。",
            "deviation": "none",
            "deviation_note": "用于核验草稿状态和证据选择。",
            "evidence": [evidence_input],
        },
    )

    multi_evidence_cards = []
    for offset, requirement in enumerate(requirements[820:900]):
        card = await create_card(
            client,
            headers,
            task_id,
            extraction,
            requirement["id"],
            {
                "response_kind": "evidence",
                "response_text": f"大列表合成双材料响应 {offset + 1:03d}。",
                "deviation": "none",
                "deviation_note": "型号与型号版本两个真实声明字段均逐字关联。",
                "evidence": [evidence_input, second_evidence_input],
            },
        )
        if offset % 3 == 0:
            card = await card_action(client, headers, card, "submit")
        multi_evidence_cards.append(
            {
                "id": card["id"],
                "requirement_id": card["requirement_id"],
                "state": card["state"],
                "evidence_count": len(card["evidence"]),
            }
        )
    if len(multi_evidence_cards) != 80 or any(
        card["evidence_count"] != 2 for card in multi_evidence_cards
    ):
        raise RuntimeError("The large-list fixture needs 80 cards with two resolved Evidence rows")

    commercial_items = [
        {
            "requirement_id": row["id"],
            "expected_revision": None,
            "disposition": "comply_only",
            "reason": "合成批次：该商务条款只需遵守。",
        }
        for index, row in enumerate(requirements[20:820], start=20)
        if index % 4 == 0
    ][:200]
    technical_items = [
        {
            "requirement_id": row["id"],
            "expected_revision": None,
            "disposition": "comply_only",
            "reason": "合成批次：该技术条款只需遵守。",
        }
        for index, row in enumerate(requirements[20:820], start=20)
        if index % 4 == 1
    ][:200]
    for items in (commercial_items, technical_items):
        await response_json(
            client,
            "POST",
            f"/tasks/{task_id}/cards/dispositions",
            headers=headers,
            json={"extraction_job_id": extraction, "items": items},
        )

    draft_one_receipt = await response_json(
        client,
        "POST",
        f"/tasks/{task_id}/drafts",
        headers=headers,
        json={"extraction_job_id": extraction},
    )
    draft_one_job = draft_one_receipt["data"]["job_id"]
    await app.state.processor(draft_one_job)
    draft_one_status = await response_json(client, "GET", f"/jobs/{draft_one_job}", headers=headers)
    draft_one_id = draft_one_status["data"]["result"]["draft_id"]

    cards["confirmed_evidence"] = await card_action(
        client,
        headers,
        cards["confirmed_evidence"],
        "reopen",
        reason="合成审阅：使旧初稿进入 stale 状态。",
    )
    stale_draft = await response_json(client, "GET", f"/drafts/{draft_one_id}", headers=headers)
    if stale_draft["data"]["validity"] != "stale":
        raise RuntimeError("Reopening a reviewed response did not stale the existing draft")

    draft_two_receipt = await response_json(
        client,
        "POST",
        f"/tasks/{task_id}/drafts",
        headers=headers,
        json={"extraction_job_id": extraction},
    )
    draft_two_job = draft_two_receipt["data"]["job_id"]
    await app.state.processor(draft_two_job)
    draft_two_status = await response_json(client, "GET", f"/jobs/{draft_two_job}", headers=headers)
    current_draft_id = draft_two_status["data"]["result"]["draft_id"]
    current_draft = await response_json(
        client, "GET", f"/drafts/{current_draft_id}", headers=headers
    )

    generation_preview = await response_json(
        client,
        "POST",
        f"/tasks/{task_id}/cards/generations",
        headers=headers,
        json={
            "extraction_job_id": extraction,
            "requirement_ids": [requirements[18]["id"]],
            "reasoning": "low",
            "dry_run": True,
        },
    )
    preflight = generation_preview["data"]["preflight"]
    if (
        preflight["admission_blocker"] is not None
        or preflight["estimate"]["calls"] <= 0
        or preflight["estimate"]["llm_tokens"] <= 0
    ):
        raise RuntimeError(
            "The generation fixture needs an available provider and resource estimate"
        )

    batch_candidates = [
        row["id"] for index, row in enumerate(requirements[900:1000], start=900) if index % 4 == 1
    ][:3]
    return {
        "product_selection": selection["data"],
        "cards": cards,
        "multi_evidence_cards": multi_evidence_cards,
        "batch_candidates": batch_candidates,
        "disposition_counts": {
            "commercial": len(commercial_items),
            "technical": len(technical_items),
        },
        "drafts": {
            "stale": stale_draft["data"],
            "current": current_draft["data"],
        },
        "generation_preview": generation_preview["data"],
    }


EXPORT_SECTIONS = (
    "substantive",
    "commercial",
    "technical",
    "comply_only",
    "gaps",
    "evidence_appendix",
)
DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def export_template() -> bytes:
    document = WordDocument()
    # Register only style IDs actually present in the package.
    document.add_table(rows=1, cols=1).style = "Table Grid"
    document.tables[0]._element.getparent().remove(document.tables[0]._element)
    document.add_heading("", level=1)
    for section in EXPORT_SECTIONS:
        document.add_paragraph("{{bid." + section + "}}")
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


async def create_export_fixture(
    client: httpx.AsyncClient, app: Any, headers: dict[str, str]
) -> dict[str, Any]:
    """A small task with a confirmed certificate page and a released review copy to preview."""
    data = await create_task_documents(client, app, headers, "c", 12)
    task_id = data["task"]["id"]
    extraction = data["extractions"]["pdf"]["id"]
    requirement = data["pdf_requirements"][0]
    await confirm_requirements(client, headers, task_id, extraction, data["pdf_requirements"])

    certificate_data = {
        "kind": "qualification",
        "name": "合成资质证书",
        "number": "SYN-CERT-2026",
        "valid_from": "2026-01-01",
        "valid_until": "2028-01-01",
    }
    certificate = await response_json(
        client,
        "POST",
        "/resources/certificates",
        headers=headers,
        json={"data": certificate_data},
    )
    certificate_pdf = pdf_fixture(2, "CERT")
    revision = await response_json(
        client,
        "POST",
        f"/resources/certificates/{certificate['data']['certificate_id']}/file-revisions",
        headers=headers,
        data={"metadata": json.dumps({"expected_revision": 1, "data": certificate_data})},
        files={"file": ("synthetic-certificate.pdf", certificate_pdf, "application/pdf")},
    )
    selected = await response_json(
        client,
        "POST",
        f"/tasks/{task_id}/certificates",
        headers=headers,
        json={"certificate_id": certificate["data"]["certificate_id"]},
    )
    source = await response_json(
        client,
        "POST",
        f"/tasks/{task_id}/evidence-sources",
        headers=headers,
        json={"task_certificate_id": selected["data"]["id"], "page": 1},
    )
    source = source["data"]["source"]
    with pymupdf.open(stream=certificate_pdf, filetype="pdf") as pdf:
        page_quote = pdf[0].get_textpage().extractText().strip().splitlines()[0]
    card = await create_card(
        client,
        headers,
        task_id,
        extraction,
        requirement["id"],
        {
            "response_kind": "evidence",
            "response_text": "合成资质证书原件第 1 页载明该记录。",
            "deviation": "none",
            "deviation_note": "证书页原文与招标要求的记录编号一致。",
            "evidence": [
                {
                    "kind": "certificate_pdf_page",
                    "evidence_source_id": source["id"],
                    "quote": page_quote,
                }
            ],
        },
    )
    if card["review_domain"] is None:
        card = (
            await response_json(
                client,
                "POST",
                f"/cards/{card['id']}/classification",
                headers=headers,
                json={
                    "expected_revision": card["revision"],
                    "review_domain": "commercial",
                    "reason": "合成导出：资质证书由商务审核。",
                },
            )
        )["data"]
    card = await card_action(client, headers, card, "submit")
    card = await card_action(
        client,
        headers,
        card,
        "confirm",
        reviewed_evidence_ids=[item["id"] for item in card["evidence"]],
        reviewed_warning_codes=card["warning_codes"],
        reason="合成导出：已核对证书原页。" if card["warning_codes"] else None,
    )

    draft_receipt = await response_json(
        client,
        "POST",
        f"/tasks/{task_id}/drafts",
        headers=headers,
        json={"extraction_job_id": extraction},
    )
    await app.state.processor(draft_receipt["data"]["job_id"])
    draft_status = await response_json(
        client, "GET", f"/jobs/{draft_receipt['data']['job_id']}", headers=headers
    )
    draft_id = draft_status["data"]["result"]["draft_id"]

    template = await response_json(
        client,
        "POST",
        "/resources/templates",
        headers=headers,
        data={"metadata": json.dumps({"data": {"name": "合成导出版式", "chapters": []}})},
        files={"file": ("synthetic-export.docx", export_template(), DOCX_MEDIA_TYPE)},
    )
    selection = await response_json(
        client,
        "POST",
        f"/tasks/{task_id}/templates",
        headers=headers,
        json={"template_id": template["data"]["template_id"]},
    )
    binding_body = {
        "template_revision_id": template["data"]["id"],
        "expected_template_sha256": template["data"]["file"]["sha256"],
        "sections": [
            {
                "section": section,
                "heading_style_id": "Heading1",
                "table_style_id": "TableGrid",
                "columns": [
                    {"key": key, "width_percent": width}
                    for key, width in zip(
                        ("ordinal", "requirement", "response", "compliance"),
                        (6, 36, 44, 14),
                        strict=True,
                    )
                ]
                if index < 3
                else [],
            }
            for index, section in enumerate(EXPORT_SECTIONS)
        ],
        "dry_run": True,
    }
    binding_preview = await response_json(
        client, "POST", "/export-template-bindings", headers=headers, json=binding_body
    )
    binding = await response_json(
        client,
        "POST",
        "/export-template-bindings",
        headers=headers,
        json={
            **binding_body,
            "dry_run": False,
            "expected_static_content_hash": binding_preview["data"]["static_content_hash"],
        },
    )
    run_body = {
        "draft_id": draft_id,
        "task_template_id": selection["data"]["id"],
        "binding_id": binding["data"]["id"],
        "mode": "review_copy",
    }
    run_preview = await response_json(
        client,
        "POST",
        f"/tasks/{task_id}/export-runs",
        headers=headers,
        json={**run_body, "dry_run": True},
    )
    run = await response_json(
        client,
        "POST",
        f"/tasks/{task_id}/export-runs",
        headers=headers,
        json={
            **run_body,
            "expected_input_hash": run_preview["data"]["input_hash"],
            "acknowledged_issue_ids": [
                issue["issue_id"]
                for issue in run_preview["data"]["issues"]
                if issue["severity"] == "acknowledge"
            ],
        },
    )
    await app.state.processor(run["data"]["render_job_id"])
    ready = await response_json(client, "GET", f"/export-runs/{run['data']['id']}", headers=headers)
    released = await response_json(
        client,
        "POST",
        f"/export-runs/{run['data']['id']}/release",
        headers=headers,
        json={
            "expected_input_hash": ready["data"]["input_hash"],
            "expected_candidate_sha256": ready["data"]["candidate_sha256"],
        },
    )
    return {
        "task": data["task"],
        "documents": data["documents"],
        "extraction_id": extraction,
        "card_requirement_id": requirement["id"],
        "certificate": {"revision_id": revision["data"]["id"], "name": certificate_data["name"]},
        "export": released["data"],
    }


def public_task_data(value: dict[str, Any]) -> dict[str, Any]:
    return {
        "task": value["task"],
        "documents": value["documents"],
        "extraction_preview": value["extraction_preview"],
        "extractions": value["extractions"],
        "requirement_count": len(value["requirements"]),
        "requirement_ids": [row["id"] for row in value["requirements"]],
        "pdf_requirement_ids": [row["id"] for row in value["pdf_requirements"]],
    }


async def provision(output: Path) -> Path:
    admin_url = guarded_url("BID_TEST_ADMIN_URL")
    runtime_url = guarded_url("BID_DATABASE_URL")
    if make_url(admin_url).database != make_url(runtime_url).database:
        raise SystemExit(
            "BID_TEST_ADMIN_URL and BID_DATABASE_URL must select the same test database"
        )
    data_dir = await asyncio.to_thread(configured_data_directory)
    upload_materials = await asyncio.to_thread(write_upload_materials, output)

    settings = fixture_settings(data_dir)
    vendor = SyntheticVendor()
    app = fixture_app(settings, vendor)
    async with (
        app.router.lifespan_context(app),
        httpx.AsyncClient(
            transport=httpx.ASGITransport(app=app), base_url="http://org-console.fixture"
        ) as client,
    ):
        headers = await human_session(client)
        task_a = await create_task_documents(client, app, headers, "a", WORD_REQUIREMENTS)
        task_b = await create_task_documents(client, app, headers, "b", 12)
        review = await create_review_fixture(client, app, headers, task_a)
        previews = await create_export_fixture(client, app, headers)
        task_b_id = task_b["task"]["id"]
        for path in (
            f"/tasks/{task_b_id}",
            f"/tasks/{task_b_id}/documents",
            f"/tasks/{task_b_id}/jobs?kind=parse",
        ):
            await response_json(client, "GET", path, headers=headers)

    manifest = {
        "schema": "single-user-console-e2e-v2",
        "fixture": "synthetic-labelled-materials-only",
        "provider": {
            "transport": "httpx.MockTransport",
            "model": SYNTHETIC_MODEL,
            "calls": vendor.calls,
        },
        "tasks": {"a": public_task_data(task_a), "b": public_task_data(task_b)},
        "review": review,
        "previews": previews,
        "upload_materials": upload_materials,
        "expectations": {
            "large_requirement_count": WORD_REQUIREMENTS,
            "page_size": 50,
            "first_load_ms_max": 5_000,
            "interaction_p95_ms_max": 300,
            "g1_job_kind": "parse",
        },
        "secrets": "Access credentials are supplied only through environment variables.",
    }
    manifest_path = output / "fixture-manifest.json"
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n")
    return manifest_path


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", required=True, help="New or existing directory under data/work")
    parser.add_argument(
        "--serve",
        action="store_true",
        help="Serve the synthetic API for native consumer verification after provisioning",
    )
    parser.add_argument("--host", choices=("127.0.0.1",), default="127.0.0.1")
    parser.add_argument("--port", type=tcp_port, default=8000)
    args = parser.parse_args()
    output = output_directory(args.output)
    manifest = asyncio.run(provision(output))
    print(manifest)
    if args.serve:
        import uvicorn

        settings = fixture_settings(configured_data_directory())
        app = fixture_app(settings, SyntheticVendor(), run_jobs=True)
        uvicorn.run(app, host=args.host, port=args.port, access_log=False)


if __name__ == "__main__":
    main()
