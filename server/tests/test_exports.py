"""API/processor export acceptance gates, with synthetic, repeatable artifacts.

Failure modes enumerated before implementation:
- A token, agent or worker must never prepare/release/download.
- Unknown IDs must remain inaccessible on every export route/job.
- Gaps block final_section; review_copy exposes only source/gap metadata and is partial.
- Invalid citations, stale cards/selections and unconfirmed evidence block both modes.
- Missing/duplicate acknowledgments and input/candidate hash conflicts cannot publish.
- Worker completion produces only a candidate; cancellation/old attempts cannot publish.
- DOCX bytes must reproduce, stay encrypted, retain editable text and exact page PNGs.
- Changed selections/reviews revoke already issued download signatures.
- Dry run and export make no provider calls or usage rows; duplicate release audits once.
"""

import asyncio
import hashlib
import json
from io import BytesIO
from pathlib import Path
from uuid import UUID, uuid4
from zipfile import ZipFile

import pytest
from app.models.entities import AuditLog, Job, UsageRecord
from docx import Document
from sqlalchemy import func, select
from test_response_cards import (
    CERTIFICATE_PAGE,
    create_card,
    create_tender,
    phase_one_client,
    require_action,
    select_real_materials,
)

SECTIONS = ("substantive", "commercial", "technical", "comply_only", "gaps", "evidence_appendix")
COLUMNS = (
    "ordinal",
    "requirement",
    "response",
    "compliance",
)


def export_template():
    doc = Document()
    # Register only style IDs actually present in the package.
    doc.add_table(rows=1, cols=1).style = "Table Grid"
    doc.tables[0]._element.getparent().remove(doc.tables[0]._element)
    doc.add_heading("", level=1)
    for section in SECTIONS:
        doc.add_paragraph("{{bid." + section + "}}")
    output = BytesIO()
    doc.save(output)
    return output.getvalue()


async def setup_template(api, header, task):
    uploaded = await api.post(
        "/resources/templates",
        headers=header,
        data={
            "metadata": json.dumps({"data": {"name": "Synthetic export layout", "chapters": []}})
        },
        files={
            "file": (
                "synthetic-export.docx",
                export_template(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert uploaded.status_code == 200, uploaded.text
    template = uploaded.json()["data"]
    selection = await api.post(
        f"/tasks/{task}/templates", headers=header, json={"template_id": template["template_id"]}
    )
    assert selection.status_code == 200, selection.text
    selected = selection.json()["data"]
    body = {
        "template_revision_id": template["id"],
        "expected_template_sha256": template["file"]["sha256"],
        "sections": [
            {
                "section": section,
                "heading_style_id": "Heading1",
                "table_style_id": "TableGrid",
                "columns": [
                    {"key": key, "width_percent": width}
                    for key, width in zip(COLUMNS, [6, 36, 44, 14], strict=True)
                ]
                if index < 3
                else [],
            }
            for index, section in enumerate(SECTIONS)
        ],
        "dry_run": True,
    }
    preview = await api.post("/export-template-bindings", headers=header, json=body)
    assert preview.status_code == 200, preview.text
    body.update(
        dry_run=False, expected_static_content_hash=preview.json()["data"]["static_content_hash"]
    )
    binding = await api.post("/export-template-bindings", headers=header, json=body)
    assert binding.status_code == 200, binding.text
    return selected, binding.json()["data"]


async def draft(api, app, header, task, extraction):
    submitted = await api.post(
        f"/tasks/{task}/drafts", headers=header, json={"extraction_job_id": extraction}
    )
    assert submitted.status_code == 200, submitted.text
    job = submitted.json()["data"]["job_id"]
    await app.state.processor(job)
    result = await api.get(f"/jobs/{job}", headers=header)
    terminal = result.json()
    assert terminal["data"]["status"] == "succeeded", result.text
    assert terminal["data"]["result"]["cost"] == terminal["cost"]
    assert terminal["cost"]["llm_tokens"] == 0
    return terminal["data"]["result"]["draft_id"]


async def prepared(api, app, header, task, body):
    preview = await api.post(
        f"/tasks/{task}/export-runs", headers=header, json={**body, "dry_run": True}
    )
    assert preview.status_code == 200, preview.text
    data = preview.json()["data"]
    submitted = await api.post(
        f"/tasks/{task}/export-runs",
        headers=header,
        json={
            **body,
            "expected_input_hash": data["input_hash"],
            "acknowledged_issue_ids": [
                i["issue_id"] for i in data["issues"] if i["severity"] == "acknowledge"
            ],
        },
    )
    assert submitted.status_code == 200, submitted.text
    run = submitted.json()["data"]
    await app.state.processor(run["render_job_id"])
    ready = await api.get(f"/export-runs/{run['id']}", headers=header)
    assert ready.status_code == 200, ready.text
    status = await api.get(f"/jobs/{run['render_job_id']}", headers=header)
    assert ready.json()["data"]["state"] == "awaiting_release", status.text
    assert ready.json()["data"]["export_id"] is None
    return ready.json()["data"]


@pytest.mark.parametrize("with_page", [False, True])
async def test_review_export_real_chain_and_every_route_gated(
    database, tmp_path, admin_engine, with_page
):
    async with phase_one_client(database, tmp_path) as (api, app, headers, provider):
        header = headers
        task, _, extraction, requirements = await create_tender(
            api, app, header, tmp_path, confirmed=True
        )
        selected, binding = await setup_template(api, header, task)
        confirmed = None
        expected_png = None
        if with_page:
            _, _, _, _, source = await select_real_materials(api, header, task, tmp_path)
            card = await create_card(
                api,
                header,
                task,
                extraction,
                requirements[1],
                {
                    "response_kind": "evidence",
                    "response_text": "Synthetic certificate QMS-2026 is supplied.",
                    "deviation": "none",
                    "deviation_note": "The named synthetic page is attached.",
                    "evidence": [
                        {
                            "kind": "certificate_pdf_page",
                            "evidence_source_id": source["id"],
                            "quote": CERTIFICATE_PAGE,
                        }
                    ],
                },
            )
            submitted = await require_action(api, header, card, "submit")
            confirmed = await require_action(
                api,
                header,
                submitted,
                "confirm",
                reviewed_evidence_ids=[e["id"] for e in submitted["evidence"]],
                reviewed_warning_codes=submitted["warning_codes"],
                reason="Synthetic acceptance review.",
            )
            expected_png = source["preview"]["sha256"]
        draft_id = await draft(api, app, header, task, extraction)
        body = {
            "draft_id": draft_id,
            "task_template_id": selected["id"],
            "binding_id": binding["id"],
            "mode": "final_section",
        }
        calls = provider.calls
        blocked = await api.post(
            f"/tasks/{task}/export-runs", headers=header, json={**body, "dry_run": True}
        )
        assert not blocked.json()["data"]["ready"]
        assert any(i["code"] == "export_gaps_present" for i in blocked.json()["data"]["issues"])
        body["mode"] = "review_copy"
        from dispatch_acceptance import assert_dispatch_rollback

        preview = await api.post(
            f"/tasks/{task}/export-runs", headers=header, json={**body, "dry_run": True}
        )
        assert preview.status_code == 200, preview.text
        preflight = preview.json()["data"]
        await assert_dispatch_rollback(
            api,
            app,
            header,
            f"/tasks/{task}/export-runs",
            {
                **body,
                "expected_input_hash": preflight["input_hash"],
                "acknowledged_issue_ids": [
                    issue["issue_id"]
                    for issue in preflight["issues"]
                    if issue["severity"] == "acknowledge"
                ],
            },
            tmp_path / "export-dispatch-rollback.json",
        )
        run = await prepared(api, app, header, task, body)
        released = await api.post(
            f"/export-runs/{run['id']}/release",
            headers=header,
            json={
                "expected_input_hash": run["input_hash"],
                "expected_candidate_sha256": run["candidate_sha256"],
            },
        )
        assert released.status_code == 200, released.text
        assert released.json()["ok"] is False
        export = released.json()["data"]
        assert export["completion"] == "partial" and export["validity"] == "current"
        repeated = await api.post(
            f"/export-runs/{run['id']}/release",
            headers=header,
            json={
                "expected_input_hash": run["input_hash"],
                "expected_candidate_sha256": run["candidate_sha256"],
            },
        )
        assert repeated.json()["data"]["id"] == export["id"]
        link = await api.get(f"/exports/{export['id']}/download-link", headers=header)
        assert link.status_code == 200, link.text
        downloaded = await api.get(link.json()["data"]["url"], headers=header)
        assert downloaded.status_code == 200, downloaded.text
        assert hashlib.sha256(downloaded.content).hexdigest() == export["file"]["sha256"]
        doc = Document(BytesIO(downloaded.content))
        assert len(doc.tables) >= 3
        with ZipFile(BytesIO(downloaded.content)) as archive:
            assert "不得提交" in "".join(
                archive.read(name).decode() for name in archive.namelist() if name.endswith(".xml")
            )
            media = [
                hashlib.sha256(archive.read(n)).hexdigest()
                for n in archive.namelist()
                if n.startswith("word/media/")
            ]
            assert media == ([expected_png] if with_page else [])
        assert provider.calls == calls
        async with app.state.db.transaction() as session:
            assert (
                await session.scalar(
                    select(func.count())
                    .select_from(UsageRecord)
                    .where(UsageRecord.job_id == UUID(run["render_job_id"]))
                )
                == 0
            )
            assert (
                await session.scalar(
                    select(func.count())
                    .select_from(AuditLog)
                    .where(AuditLog.action == "export.released")
                )
                == 1
            )
        artifact = Path("data/work/export-acceptance") / ("page" if with_page else "gaps")
        await asyncio.to_thread(artifact.mkdir, parents=True, exist_ok=True)
        (artifact / "synthetic-review.docx").write_bytes(downloaded.content)
        (artifact / "receipt.json").write_text(
            json.dumps(
                {"file_sha256": export["file"]["sha256"], "media_sha256": media, "synthetic": True},
                indent=2,
            )
        )
        if confirmed:
            await require_action(
                api, header, confirmed, "reopen", reason="Synthetic invalidation check."
            )
            old_link = await api.get(link.json()["data"]["url"], headers=header)
            assert old_link.status_code == 409
            history = await api.get(f"/exports/{export['id']}", headers=header)
            assert history.json()["ok"] and history.json()["data"]["validity"] == "stale"
            stale_run = await api.get(f"/export-runs/{run['id']}", headers=header)
            assert stale_run.json()["data"]["state"] == "invalidated"
            assert stale_run.json()["data"]["candidate_sha256"] is None
            assert stale_run.json()["data"]["export_id"] == export["id"]
            assert (
                confirmed["requirement_id"] in history.json()["data"]["invalidated_requirement_ids"]
            )


async def test_export_unknown_resource_boundaries(database, tmp_path, admin_engine):
    async with phase_one_client(database, tmp_path) as (api, _, headers, _):
        unknown = str(uuid4())
        for path in (
            f"/export-runs/{unknown}",
            f"/exports/{unknown}",
            f"/exports/{unknown}/download-link",
            f"/tasks/{unknown}/exports",
        ):
            assert (await api.get(path, headers=headers)).status_code == 404


async def complete_inputs(
    api, app, header, database, admin_engine, tmp_path, *, texts=None, profile_wording=None
):
    """Create three confirmed tables plus an independent comply-only decision.
    `texts` replaces the response text of the cards for those requirement indexes;
    `profile_wording` makes requirement 2 cite a selected company profile declaration
    instead of the certificate page."""
    task, _, extraction, requirements = await create_tender(
        api, app, header, tmp_path, confirmed=True
    )
    selected, binding = await setup_template(api, header, task)
    _, _, _, _, page = await select_real_materials(api, header, task, tmp_path)
    proof = [
        {
            "kind": "certificate_pdf_page",
            "evidence_source_id": page["id"],
            "quote": CERTIFICATE_PAGE,
        }
    ]
    if profile_wording is not None:
        profile = await api.post(
            "/resources/profiles",
            headers=header,
            json={"data": {"name": "Synthetic profile", "standard_wording": profile_wording}},
        )
        assert profile.status_code == 200, profile.text
        chosen = await api.post(
            f"/tasks/{task}/profiles",
            headers=header,
            json={"profile_id": profile.json()["data"]["profile_id"]},
        )
        assert chosen.status_code == 200, chosen.text
        proof = [
            {
                "kind": "org_profile",
                "selection_id": chosen.json()["data"]["id"],
                "field_path": "standard_wording",
                "quote": profile_wording,
            }
        ]
    reviewed = {}
    for index in (0, 1, 2, 4):
        card = await create_card(
            api,
            header,
            task,
            extraction,
            requirements[index],
            {
                "response_kind": "evidence" if index == 1 else "commitment",
                "response_text": (texts or {}).get(
                    index, f"Synthetic confirmed response for requirement {index + 1}."
                ),
                "deviation": "negative" if index == 2 else "none",
                "deviation_note": "Offered forty days exceeds required thirty days."
                if index == 2
                else "The specified synthetic response is retained verbatim.",
                "evidence": proof if index == 1 else [],
            },
        )
        if index == 0:
            response = await api.post(
                f"/cards/{card['id']}/classification",
                headers=header,
                json={
                    "expected_revision": card["revision"],
                    "review_domain": "technical",
                    "reason": "The synthetic scored requirement is technical.",
                },
            )
            assert response.status_code == 200, response.text
            card = response.json()["data"]
        reviewed[index] = card
    for index in (0, 1, 2, 4):
        submitted = await require_action(api, header, reviewed[index], "submit")
        reviewed[index] = await require_action(
            api,
            header,
            submitted,
            "confirm",
            reviewed_evidence_ids=[e["id"] for e in submitted["evidence"]],
            reviewed_warning_codes=submitted["warning_codes"],
            reason="Synthetic review of proof obligations.",
        )
    disposition = await api.post(
        f"/tasks/{task}/cards/dispositions",
        headers=header,
        json={
            "extraction_job_id": extraction,
            "items": [
                {
                    "requirement_id": requirements[3]["id"],
                    "expected_revision": None,
                    "disposition": "comply_only",
                    "reason": "The complaint procedure requires compliance only.",
                }
            ],
        },
    )
    assert disposition.status_code == 200, disposition.text
    draft_id = await draft(api, app, header, task, extraction)
    body = {
        "draft_id": draft_id,
        "task_template_id": selected["id"],
        "binding_id": binding["id"],
        "mode": "final_section",
    }
    return (task, body, reviewed, page)


async def test_final_export_verbatim_negative_ack_hash_gates_and_immutable_download(
    database, tmp_path, admin_engine
):
    from app.models.exports import Export
    from app.services.auth import HUMAN_SCOPES, Identity, set_actor_context

    async with phase_one_client(database, tmp_path) as (api, app, headers, provider):
        header = headers
        task, body, _, page = await complete_inputs(
            api, app, header, database, admin_engine, tmp_path
        )
        async with app.state.db.transaction() as session:
            before = await session.scalar(select(func.count()).select_from(AuditLog))
        preview = await api.post(
            f"/tasks/{task}/export-runs", headers=header, json={**body, "dry_run": True}
        )
        assert preview.status_code == 200, preview.text
        fixed = preview.json()["data"]
        assert fixed["ready"] and fixed["gap_count"] == 0
        assert fixed["table_rows"] == {"substantive": 1, "commercial": 1, "technical": 2}
        assert (
            fixed["comply_only_count"] == 1
            and fixed["negative_count"] == 1
            and (fixed["attachment_pages"] == 1)
        )
        async with app.state.db.transaction() as session:
            assert await session.scalar(select(func.count()).select_from(AuditLog)) == before
        missing_ack = await api.post(
            f"/tasks/{task}/export-runs",
            headers=header,
            json={**body, "expected_input_hash": fixed["input_hash"]},
        )
        assert missing_ack.status_code == 400
        assert missing_ack.json()["data"]["error"]["code"] == "export_acknowledgment_mismatch"
        wrong_input = await api.post(
            f"/tasks/{task}/export-runs",
            headers=header,
            json={**body, "expected_input_hash": "0" * 64},
        )
        assert wrong_input.status_code == 409
        run = await prepared(api, app, header, task, body)
        bad_candidate = await api.post(
            f"/export-runs/{run['id']}/release",
            headers=header,
            json={"expected_input_hash": run["input_hash"], "expected_candidate_sha256": "0" * 64},
        )
        assert bad_candidate.status_code == 409
        release_body = {
            "expected_input_hash": run["input_hash"],
            "expected_candidate_sha256": run["candidate_sha256"],
        }
        released = await api.post(
            f"/export-runs/{run['id']}/release", headers=header, json=release_body
        )
        assert released.status_code == 200 and released.json()["ok"], released.text
        export = released.json()["data"]
        assert export["completion"] == "complete"
        link = await api.get(f"/exports/{export['id']}/download-link", headers=header)
        content = (await api.get(link.json()["data"]["url"], headers=header)).content
        with ZipFile(BytesIO(content)) as package:
            document_xml = package.read("word/document.xml").decode()
            assert "负偏离" in document_xml and "不得提交" not in document_xml
            assert "Synthetic confirmed response for requirement 3." in document_xml
            assert [
                hashlib.sha256(package.read(n)).hexdigest()
                for n in package.namelist()
                if n.startswith("word/media/")
            ] == [page["preview"]["sha256"]]
        async with app.state.db.transaction() as session:
            await set_actor_context(session, Identity(set(HUMAN_SCOPES)))
            row = await session.get(Export, UUID(export["id"]))
            encrypted = app.state.storage.path(row.object_key).read_bytes()
            assert encrypted.startswith(b"BIDFILE1\n") and content not in encrypted
            assert (
                await session.scalar(
                    select(func.count())
                    .select_from(UsageRecord)
                    .where(UsageRecord.job_id == UUID(run["render_job_id"]))
                )
                == 0
            )
        artifact = Path("data/work/export-acceptance/final")
        await asyncio.to_thread(artifact.mkdir, parents=True, exist_ok=True)
        (artifact / "synthetic-final.docx").write_bytes(content)
        (artifact / "receipt.json").write_text(
            json.dumps(
                {
                    "file_sha256": hashlib.sha256(content).hexdigest(),
                    "input_hash": run["input_hash"],
                    "synthetic": True,
                },
                indent=2,
            )
        )


@pytest.mark.parametrize("revoke", ["reopen", "cancel"])
async def test_queued_export_rechecks_inputs_attempt_and_authorization(
    database, tmp_path, admin_engine, revoke
):
    async with phase_one_client(database, tmp_path) as (api, app, headers, _):
        header = headers
        task, body, reviewed, _ = await complete_inputs(
            api, app, header, database, admin_engine, tmp_path
        )
        preview = (
            await api.post(
                f"/tasks/{task}/export-runs", headers=header, json={**body, "dry_run": True}
            )
        ).json()["data"]
        submitted = await api.post(
            f"/tasks/{task}/export-runs",
            headers=header,
            json={
                **body,
                "expected_input_hash": preview["input_hash"],
                "acknowledged_issue_ids": [
                    i["issue_id"] for i in preview["issues"] if i["severity"] == "acknowledge"
                ],
            },
        )
        assert submitted.status_code == 200, submitted.text
        run = submitted.json()["data"]
        if revoke == "reopen":
            await require_action(
                api, header, reviewed[1], "reopen", reason="Synthetic queued invalidation."
            )
        elif revoke == "cancel":
            response = await api.post(f"/jobs/{run['render_job_id']}/cancel", headers=header)
            assert response.status_code == 200, response.text
        await app.state.processor(run["render_job_id"])
        shown = await api.get(f"/export-runs/{run['id']}", headers=header)
        assert shown.status_code == 200, shown.text
        assert shown.json()["data"]["state"] in ("invalidated", "cancelled", "failed")
        assert shown.json()["data"]["export_id"] is None
        assert (await api.get(f"/tasks/{task}/exports", headers=header)).json()["items"] == []


async def test_token_cannot_use_export_or_generic_job_read_cancel(database, tmp_path, admin_engine):
    from datetime import UTC, datetime, timedelta

    async with phase_one_client(database, tmp_path) as (api, app, headers, _):
        header = headers
        token_reply = await api.post(
            "/tokens",
            headers=header,
            json={
                "name": "Synthetic export rejection token",
                "scopes": [
                    "task:read",
                    "draft:read",
                    "card:read",
                    "template:read",
                    "job:read",
                    "job:cancel",
                ],
                "expires_at": (datetime.now(UTC) + timedelta(hours=1)).isoformat(),
            },
        )
        assert token_reply.status_code == 200, token_reply.text
        token = token_reply.json()["data"]["token"]
        token_headers = {**header, "Authorization": "Bearer " + token}
        task, body, _, _ = await complete_inputs(api, app, header, database, admin_engine, tmp_path)
        run = await prepared(api, app, header, task, body)
        for method, path, data in [
            ("GET", f"/export-runs/{run['id']}", None),
            ("GET", f"/jobs/{run['render_job_id']}", None),
            ("POST", f"/jobs/{run['render_job_id']}/cancel", None),
            ("POST", f"/tasks/{task}/export-runs", {**body, "dry_run": True}),
            (
                "POST",
                f"/export-runs/{run['id']}/release",
                {
                    "expected_input_hash": run["input_hash"],
                    "expected_candidate_sha256": run["candidate_sha256"],
                },
            ),
        ]:
            result = await api.request(
                method, path, headers=token_headers, **{"json": data} if data is not None else {}
            )
            assert result.status_code == 403, (method, path, result.text)


async def test_human_render_worker_keeps_identity_without_human_export_rights(
    database, tmp_path, admin_engine, monkeypatch
):
    """Candidate rendering must not impersonate a human or trust agent/job JSON grants.

    Failure modes: nonhuman/token/delegated submission, mismatched submitter,
    missing immutable export grant, or worker publication authority.
    """
    from types import SimpleNamespace

    from app.core.errors import ServiceError
    from app.services import exports

    worker_access = exports.worker_access
    checked = []

    async def check_worker(session, job, attempt_id):
        if not checked:
            snapshot = {column.key: getattr(job, column.key) for column in Job.__table__.columns}
            for changed in (
                {"kind": "draft"},
                {"actor_kind": "worker"},
                {"actor_kind": "agent"},
                {"actor_kind": "token"},
                {"actor_token_id": uuid4()},
                {"agent_principal_id": uuid4()},
                {"agent_session_id": uuid4()},
                {"agent_step_id": uuid4()},
                {"actor_scopes": [scope for scope in job.actor_scopes if scope != "export"]},
            ):
                with pytest.raises(ServiceError) as rejected:
                    await worker_access(session, SimpleNamespace(**snapshot | changed), attempt_id)
                assert rejected.value.code == "forbidden"
        actor = await worker_access(session, job, attempt_id)
        assert job.actor_kind == "session" and job.actor_token_id is None
        assert actor.actor_kind == "worker" and actor.token_id is None
        assert actor.principal_id is None and actor.session_id is None and (actor.step_id is None)
        with pytest.raises(ServiceError) as denied:
            actor.require("export")
        assert denied.value.code == "forbidden"
        with pytest.raises(ServiceError) as denied:
            await exports.human_access(session, actor)
        assert denied.value.code == "forbidden"
        checked.append({"actor_kind": actor.actor_kind, "attempt_id": str(attempt_id)})
        return actor

    monkeypatch.setattr(exports, "worker_access", check_worker)
    async with phase_one_client(database, tmp_path) as (api, app, headers, _):
        header = headers
        task, body, _, _ = await complete_inputs(api, app, header, database, admin_engine, tmp_path)
        run = await prepared(api, app, header, task, body)
        assert checked and run["export_id"] is None
        assert (await api.get(f"/tasks/{task}/exports", headers=header)).json()["items"] == []
        released = await api.post(
            f"/export-runs/{run['id']}/release",
            headers=header,
            json={
                "expected_input_hash": run["input_hash"],
                "expected_candidate_sha256": run["candidate_sha256"],
            },
        )
        assert released.status_code == 200, released.text
        (tmp_path / "human-render-authority.json").write_text(
            json.dumps(
                {
                    "run_id": run["id"],
                    "candidate_sha256": run["candidate_sha256"],
                    "worker_checks": checked,
                    "human_release": released.json()["ok"],
                },
                indent=2,
            )
        )


@pytest.mark.parametrize("limit", ["memory", "deadline", "output"])
async def test_export_worker_limits_never_publish_a_partial_file(
    database, tmp_path, admin_engine, limit
):
    from app.providers.base import ProviderFailure

    async with phase_one_client(database, tmp_path) as (api, app, headers, _):
        header = headers
        task, body, _, _ = await complete_inputs(api, app, header, database, admin_engine, tmp_path)
        preview = (
            await api.post(
                f"/tasks/{task}/export-runs", headers=header, json={**body, "dry_run": True}
            )
        ).json()["data"]
        submitted = await api.post(
            f"/tasks/{task}/export-runs",
            headers=header,
            json={
                **body,
                "expected_input_hash": preview["input_hash"],
                "acknowledged_issue_ids": [
                    i["issue_id"] for i in preview["issues"] if i["severity"] == "acknowledge"
                ],
            },
        )
        assert submitted.status_code == 200, submitted.text
        run = submitted.json()["data"]
        settings = app.state.processor.settings
        if limit == "memory":
            settings.export_memory_bytes = 1
        elif limit == "deadline":
            settings.export_deadline_seconds = 0.001
        else:
            settings.export_max_output_bytes = 1
        for attempt in range(3 if limit == "deadline" else 1):
            if limit == "deadline" and attempt < 2:
                with pytest.raises(ProviderFailure):
                    await app.state.processor(run["render_job_id"])
            else:
                await app.state.processor(run["render_job_id"])
        status = await api.get(f"/jobs/{run['render_job_id']}", headers=header)
        assert status.json()["data"]["status"] == "failed", status.text
        assert status.json()["data"]["error"]["code"] in {
            "export_memory_limit",
            "export_render_timeout",
            "export_resource_limit",
        }
        assert (await api.get(f"/tasks/{task}/exports", headers=header)).json()["items"] == []
        shown = (await api.get(f"/export-runs/{run['id']}", headers=header)).json()["data"]
        assert (
            shown["state"] == "failed"
            and shown["candidate_sha256"] is None
            and (shown["export_id"] is None)
        )


LEGACY_COLUMNS = (
    "ordinal",
    "tender_clause",
    "source_location",
    "response",
    "deviation",
    "deviation_note",
    "evidence",
)


async def test_starter_template_provenance_and_outdated_binding(database, tmp_path, admin_engine):
    import re

    from app.models.exports import ExportTemplateBinding
    from app.services.export_template_sample import binding_sections, build

    async with phase_one_client(database, tmp_path) as (api, app, headers, _):
        header = headers
        task, body, _, page = await complete_inputs(
            api, app, header, database, admin_engine, tmp_path
        )
        uploaded = await api.post(
            "/resources/templates",
            headers=header,
            data={"metadata": json.dumps({"data": {"name": "Starter", "chapters": []}})},
            files={"file": ("starter.docx", build(), "application/octet-stream")},
        )
        assert uploaded.status_code == 200, uploaded.text
        template = uploaded.json()["data"]
        selected = (
            await api.post(
                f"/tasks/{task}/templates",
                headers=header,
                json={"template_id": template["template_id"]},
            )
        ).json()["data"]
        request = {
            "template_revision_id": template["id"],
            "expected_template_sha256": template["file"]["sha256"],
            "sections": binding_sections(),
            "dry_run": True,
        }
        preview = await api.post("/export-template-bindings", headers=header, json=request)
        assert preview.status_code == 200, preview.text
        request |= {
            "dry_run": False,
            "expected_static_content_hash": preview.json()["data"]["static_content_hash"],
        }
        binding = (
            await api.post("/export-template-bindings", headers=header, json=request)
        ).json()["data"]
        assert binding["current"] is True
        from app.services.auth import HUMAN_SCOPES, Identity, set_actor_context

        async with app.state.db.transaction() as session:
            await set_actor_context(session, Identity(set(HUMAN_SCOPES)))
            row = await session.get(ExportTemplateBinding, UUID(binding["id"]))
            assert row is not None
            legacy_columns = [
                {"key": key, "width_percent": width}
                for key, width in zip(LEGACY_COLUMNS, (5, 25, 15, 25, 10, 10, 10), strict=True)
            ]
            session.add(
                ExportTemplateBinding(
                    template_revision_id=row.template_revision_id,
                    template_sha256=row.template_sha256,
                    binding_hash="e" * 64,
                    sections=[
                        {**section, "columns": legacy_columns if section["columns"] else []}
                        for section in row.sections
                    ],
                    static_content_hash=row.static_content_hash,
                    adapter_version=row.adapter_version,
                    reviewed_at=row.reviewed_at,
                )
            )
        listed = (
            await api.get(
                "/export-template-bindings",
                headers=header,
                params={"template_revision_id": template["id"]},
            )
        ).json()["items"]
        assert sorted(item["current"] for item in listed) == [False, True]
        outdated = next(item for item in listed if not item["current"])
        stale = {**body, "task_template_id": selected["id"], "binding_id": outdated["id"]}
        refused = await api.post(
            f"/tasks/{task}/export-runs", headers=header, json={**stale, "dry_run": True}
        )
        assert refused.status_code == 400, refused.text
        assert refused.json()["data"]["error"]["code"] == "export_blocked"
        codes = {issue["code"] for issue in refused.json()["data"]["issues"]}
        assert "export_binding_outdated" in codes and (not refused.json()["data"]["ready"])
        current = {**body, "task_template_id": selected["id"], "binding_id": binding["id"]}
        run = await prepared(api, app, header, task, current)
        released = await api.post(
            f"/export-runs/{run['id']}/release",
            headers=header,
            json={
                "expected_input_hash": run["input_hash"],
                "expected_candidate_sha256": run["candidate_sha256"],
            },
        )
        assert released.status_code == 200, released.text
        export = released.json()["data"]
        link = await api.get(f"/exports/{export['id']}/download-link", headers=header)
        content = (await api.get(link.json()["data"]["url"], headers=header)).content
        provenance = await api.get(f"/exports/{export['id']}/provenance", headers=header)
        assert provenance.status_code == 200, provenance.text
        record = provenance.json()["data"]
    document = Document(BytesIO(content))
    section = document.sections[0]
    assert (round(section.page_width.cm, 1), round(section.page_height.cm, 1)) == (21.0, 29.7)
    tables = document.tables
    for index, name in enumerate(("substantive", "commercial", "technical")):
        rows = [r for r in record["items"] if r["section"] == name]
        printed = [row.cells[0].text for row in tables[index].rows[1:]]
        if rows:
            assert printed == [str(r["number"]) for r in rows]
        assert [cell.text for cell in tables[index].rows[0].cells] == [
            "序号",
            "招标文件要求",
            "投标文件响应内容",
            "响应情况",
        ]
    assert [a["label"] for a in record["attachments"]] == ["E001"]
    assert record["attachments"][0]["png_sha256"] == page["preview"]["sha256"]
    assert record["file"]["sha256"] == hashlib.sha256(content).hexdigest()
    text = "\n".join(p.text for p in document.paragraphs) + "\n".join(
        cell.text for table in tables for row in table.rows for cell in row.cells
    )
    assert "（见附件 E001）" in text and "附件 E001" in text
    assert not re.search("[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}", text)
    assert not re.search("\\b[0-9a-f]{64}\\b", text)
