"""Word block parsing: each failure mode listed before the parser was written."""

import io

from app.services.docx_blocks import parse_docx
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

W = nsdecls("w")


def save(document) -> bytes:
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def rich_document() -> bytes:
    document = Document()
    document.sections[0].header.paragraphs[0].text = "Synthetic header"
    document.add_heading("第一章 招标公告", level=1)
    document.add_paragraph("Opening paragraph.")
    document.add_paragraph("   ")  # empty paragraphs never consume a number
    document.add_heading("一、资格要求", level=2)
    document.add_paragraph("★ 投标人须具备有效的营业执照。")
    document.add_paragraph("Second paragraph in section.")
    table = document.add_table(rows=3, cols=3)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Merged across"
    table.cell(0, 2).text = "Top right"
    table.cell(1, 0).merge(table.cell(2, 0)).text = "Merged down"
    table.cell(1, 1).text = "★ 内存不低于 64GB"
    inner = table.cell(2, 2).add_table(rows=1, cols=2)
    inner.cell(0, 1).text = "Nested value"
    document.add_heading("第二章 技术要求", level=1)
    body = document.element.body
    sdt = parse_xml(
        f"<w:sdt {W}><w:sdtContent><w:p><w:r><w:t>Inside a content control.</w:t></w:r></w:p>"
        "</w:sdtContent></w:sdt>"
    )
    body.insert(len(body) - 1, sdt)
    textbox = parse_xml(
        f"<w:p {W}><w:r><w:pict><w:txbxContent><w:p><w:r><w:t>Text box words</w:t></w:r></w:p>"
        "</w:txbxContent></w:pict></w:r></w:p>"
    )
    body.insert(len(body) - 1, textbox)
    offset = document.add_table(rows=1, cols=2)
    row = offset.rows[0]._tr
    row.insert(0, parse_xml(f'<w:trPr {W}><w:gridBefore w:val="1"/></w:trPr>'))
    row.remove(row.findall(qn("w:tc"))[0])
    row.findall(qn("w:tc"))[0].append(
        parse_xml(f"<w:p {W}><w:r><w:t>Shifted cell</w:t></w:r></w:p>")
    )
    return save(document)


def by_text(blocks, fragment):
    [block] = [b for b in blocks if fragment in b.text]
    return block


def test_blocks_follow_headings_merges_nesting_and_controls():
    sections, warnings = parse_docx(rich_document())
    blocks = [b for s in sections for b in s.blocks]
    ids = [b.block_id for b in blocks]
    assert len(ids) == len(set(ids))
    # Sections start at level 1 and level 2 headings.
    assert [s.blocks[0].text for s in sections] == [
        "第一章 招标公告",
        "一、资格要求",
        "第二章 技术要求",
    ]

    star = by_text(blocks, "营业执照")
    assert star.kind == "paragraph" and star.paragraph == 1
    assert star.section_path == ["第一章 招标公告", "一、资格要求"]
    assert star.label == "第一章 招标公告 > 一、资格要求 > 第 1 段"
    assert by_text(blocks, "Second paragraph").paragraph == 2
    assert by_text(blocks, "Opening paragraph").label == "第一章 招标公告 > 第 1 段"

    assert by_text(blocks, "Merged across").block_id == "t1r1c1"
    assert by_text(blocks, "Top right").block_id == "t1r1c3"
    assert by_text(blocks, "Merged down").block_id == "t1r2c1"
    assert sum("Merged down" in b.text for b in blocks) == 1
    memory = by_text(blocks, "64GB")
    assert (memory.kind, memory.table, memory.row, memory.column) == ("cell", 1, 2, 2)
    assert memory.label == "第一章 招标公告 > 一、资格要求 > 表 1 第 2 行第 2 列"
    nested = by_text(blocks, "Nested value")
    assert nested.block_id == "t1r3c3/t1r1c2"
    assert nested.label.endswith("表 1 第 3 行第 3 列 > 表 1 第 1 行第 2 列")

    assert by_text(blocks, "content control").section_path == ["第二章 技术要求"]
    assert by_text(blocks, "Shifted cell").block_id == "t2r1c2"
    assert not any("Text box words" in b.text or "Synthetic header" in b.text for b in blocks)
    assert warnings == ["Word content not parsed for citations: 1 text boxes, headers/footers."]


def test_numbering_fallback_without_heading_styles():
    document = Document()
    for text in (
        "第一章 总则",
        "一、适用范围",
        "Body line.",
        "（一）细则",
        "Detail line.",
        "第二章 附则",
    ):
        document.add_paragraph(text)
    sections, warnings = parse_docx(save(document))
    blocks = [b for s in sections for b in s.blocks]
    assert (
        by_text(blocks, "Detail line").label == "第一章 总则 > 一、适用范围 > （一）细则 > 第 1 段"
    )
    assert by_text(blocks, "第二章").section_path == ["第二章 附则"]
    assert "No heading styles found" in warnings[0]


def test_long_sections_split_at_block_boundaries():
    document = Document()
    document.add_heading("第一章 长章节", level=1)
    for index in range(40):
        document.add_paragraph(f"段落 {index} " + "内容" * 100)
    sections, _ = parse_docx(save(document), section_chars=2000)
    assert len(sections) > 1
    assert all(len(s.text) <= 2000 + 210 for s in sections)
    assert [b.paragraph for s in sections for b in s.blocks if b.paragraph] == list(range(1, 41))


def test_empty_document_is_reported():
    sections, warnings = parse_docx(save(Document()))
    assert sections == [] and warnings[-1] == "The Word document contains no text."
