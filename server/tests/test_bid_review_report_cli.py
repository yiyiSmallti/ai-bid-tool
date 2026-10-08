"""Real CLI-to-HTTP report admission and verified, private Word transfers."""

import hashlib
import io
import json
import os
import zipfile
from pathlib import Path

import pytest
from app.schemas import bid_review_report as reports
from app.schemas.contracts import Result
from app.schemas.export_contracts import DOCX_MEDIA_TYPE
from bid_cli import main as cli
from bid_cli.schema import command_schema
from docx import Document
from fastapi import FastAPI, Request
from fastapi.responses import JSONResponse, Response
from test_bid_review_cli import ID, NOW, SHA, InterfaceClient, invoke

REVIEW = "00000000-0000-4000-8000-000000000004"
SNAPSHOT = "00000000-0000-4000-8000-000000000005"
ARTIFACT = "00000000-0000-4000-8000-000000000006"
OTHER = "00000000-0000-4000-8000-000000000007"
DECISIONS = "b" * 64
RENDER_INPUT = "c" * 64


def document_bytes():
    document = Document()
    for _, heading in reports.SECTIONS:
        document.add_heading(heading, level=1)
    document.add_paragraph(reports.ADVISORY)
    document.add_paragraph("未评分")
    document.add_paragraph("签章位置未解决；Synthetic tender quote；Synthetic bid quote；confirmed")
    source = io.BytesIO()
    document.save(source)
    output = io.BytesIO()
    with zipfile.ZipFile(source) as original, zipfile.ZipFile(output, "w") as canonical:
        for name in sorted(original.namelist()):
            entry = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            entry.compress_type = zipfile.ZIP_DEFLATED
            canonical.writestr(entry, original.read(name))
    return output.getvalue()


def interface():
    application = FastAPI()
    content = document_bytes()
    artifact = reports.BidReviewReportArtifact(
        id=ARTIFACT,
        org_id=ID,
        task_id=ID,
        report_id=REVIEW,
        snapshot_id=SNAPSHOT,
        format="docx",
        sha256=hashlib.sha256(content).hexdigest(),
        size_bytes=len(content),
        report_input_hash=SHA,
        decisions_snapshot_sha256=DECISIONS,
        renderer_identity="synthetic-report-renderer-v1",
        created_at=NOW,
    ).model_dump(mode="json")
    state = {
        "calls": [],
        "malformed": False,
        "preview_parent": REVIEW,
        "preview_decisions": DECISIONS,
        "link": f"/bid-review-artifacts/{ARTIFACT}/download?signature=synthetic-report-only",
        "artifact": artifact,
        "content": content,
        "served": content,
        "media_type": DOCX_MEDIA_TYPE,
        "redirect": False,
        "expired": False,
        "link_command": "review report download",
        "queue_failure": False,
        "queued_on_preview": False,
        "preview_on_submit": False,
    }

    @application.post("/v4/bid-reviews/{review_id}/artifacts")
    async def render(request: Request, review_id: str):
        assert review_id == REVIEW and request.headers["x-org-id"] == ID
        body = reports.BidReportRenderRequest.model_validate(await request.json())
        assert str(body.report_id) == review_id
        assert body.expected_decisions_snapshot_sha256 == DECISIONS
        state["calls"].append("preview" if body.dry_run else "submit")
        if state["malformed"]:
            return Result(ok=True, command="review report", data={"unexpected": True})
        if state["queued_on_preview"]:
            return Result(
                ok=True,
                command="review report",
                data={"job_id": ID, "status": "queued", "cached": False},
            )
        if not body.dry_run and not state["preview_on_submit"]:
            assert body.expected_input_hash == RENDER_INPUT
            assert body.preflight_token == "synthetic-report-receipt"
            if state["queue_failure"]:
                return JSONResponse(
                    status_code=503,
                    content=Result(
                        ok=False,
                        command="review report",
                        data={
                            "error": {
                                "code": "queue_unavailable",
                                "message": "Synthetic queue failure",
                                "exit_code": 3,
                            }
                        },
                    ).model_dump(mode="json"),
                )
            return Result(
                ok=True,
                command="review report",
                data={"job_id": ID, "status": "queued", "cached": False},
            )
        zero = Result(ok=True, command="zero").cost.model_dump(mode="json")
        preview = reports.BidReportRenderPreview(
            report_id=state["preview_parent"],
            input_hash=RENDER_INPUT,
            report_input_hash=SHA,
            decisions_snapshot_sha256=state["preview_decisions"],
            renderer_identity="synthetic-report-renderer-v1",
            budget={
                "dry_run": True,
                "command": "review report",
                "task_id": ID,
                "input_hash": RENDER_INPUT,
                "as_of": NOW,
                "task_budget": None,
                "planned_calls": 0,
                "maximum_calls": 0,
                "estimate": zero,
                "next_call": None,
                "admission_blocker": None,
                "first_pass_fits": True,
                "full_run_guaranteed": False,
            },
            expires_at="2026-10-08T00:15:00Z",
            preflight_token="synthetic-report-receipt",
        )
        return Result(ok=True, command="review report", data=preview.model_dump(mode="json"))

    @application.get("/v4/bid-review-artifacts/{artifact_id}/download-link")
    async def link(request: Request, artifact_id: str):
        assert artifact_id == ARTIFACT and request.headers["x-org-id"] == ID
        state["calls"].append("link")
        return Result(
            ok=True,
            command=state["link_command"],
            data={"url": state["link"], "expires_in": 300, "artifact": state["artifact"]},
        )

    @application.get("/bid-review-artifacts/{artifact_id}/download")
    async def download(request: Request, artifact_id: str):
        assert artifact_id == ARTIFACT
        assert request.headers["x-org-id"] == ID
        assert request.headers["authorization"] == "Bearer synthetic-report-session"
        assert request.query_params["signature"] == "synthetic-report-only"
        state["calls"].append("bytes")
        if state["expired"]:
            return JSONResponse(
                status_code=403,
                content=Result(
                    ok=False,
                    command="review report download",
                    data={
                        "error": {
                            "code": "download_link_expired",
                            "message": "Synthetic expired link",
                            "exit_code": 4,
                        }
                    },
                ).model_dump(mode="json"),
            )
        if state["redirect"]:
            return Response(status_code=302, headers={"Location": "https://foreign.example.test"})
        return Response(
            content=state["served"],
            media_type=state["media_type"],
            headers={"Cache-Control": "no-store"},
        )

    return application, state


def setup(monkeypatch, tmp_path):
    application, state = interface()
    monkeypatch.setenv("BID_SESSION", "synthetic-report-session")
    monkeypatch.setenv("BID_ORG", ID)
    monkeypatch.setattr(
        cli, "Client", lambda mode, server, saved: InterfaceClient(mode, server, saved, application)
    )
    path = tmp_path / "report.json"
    path.write_text(
        json.dumps(
            {
                "request_id": ID,
                "report_id": REVIEW,
                "expected_decisions_snapshot_sha256": DECISIONS,
            }
        )
    )
    return state, path


def render_flags(path, *extra):
    return ["review", "report", "--id", REVIEW, "--input", str(path), *extra]


def download_flags(path):
    return ["review", "report", "download", "--artifact", ARTIFACT, "--output", str(path)]


def test_report_cli_http_snapshots_and_verified_download_both_modes(monkeypatch, tmp_path, capsys):
    state, path = setup(monkeypatch, tmp_path)
    actual = {}
    for mode in ("local", "remote"):
        commands = {
            "dry-run": render_flags(path, "--dry-run"),
            "submit": render_flags(
                path,
                "--expected-input-hash",
                RENDER_INPUT,
                "--preflight-token",
                "synthetic-report-receipt",
            ),
            "download": download_flags(tmp_path / f"{mode}.docx"),
        }
        for name, flags in commands.items():
            code, body = invoke(["--mode", mode, *flags], capsys)
            assert code == 0
            if name == "download":
                output = tmp_path / f"{mode}.docx"
                assert output.read_bytes() == state["content"]
                assert output.stat().st_mode & 0o777 == 0o600
                reopened = Document(output)
                assert reports.ADVISORY in [row.text for row in reopened.paragraphs]
                body["data"]["output_path"] = f"<OUTPUT>/{mode}.docx"
                assert "url" not in body["data"] and "signature" not in json.dumps(body)
            actual[f"{mode}: review report {name}"] = body
    assert state["calls"] == ["preview", "submit", "link", "bytes"] * 2
    assert not list(tmp_path.glob(".bid-download-*"))
    snapshot = Path(__file__).with_name("snapshots") / "bid-review-report-cli-v4.json"
    if os.environ.get("BID_UPDATE_SNAPSHOTS") == "1":
        snapshot.write_text(json.dumps(actual, indent=2, sort_keys=True) + "\n")
    assert actual == json.loads(snapshot.read_text())


@pytest.mark.parametrize(
    "updates,flags",
    [
        ({"expected_decisions_snapshot_sha256": "bad"}, ["--dry-run"]),
        ({"report_id": OTHER}, ["--dry-run"]),
        ({}, []),
        ({"dry_run": True, "retry": True}, []),
        ({"dry_run": True, "preflight_token": "submit-only"}, []),
        ({"dry_run": True, "expected_input_hash": RENDER_INPUT}, []),
    ],
)
def test_invalid_report_requests_never_dispatch(monkeypatch, tmp_path, capsys, updates, flags):
    state, path = setup(monkeypatch, tmp_path)
    body = json.loads(path.read_text())
    body.update(updates)
    path.write_text(json.dumps(body))
    code, result = invoke(render_flags(path, *flags), capsys)
    assert code == 2 and not result["ok"] and not state["calls"]


def test_report_requires_noninteractive_ids_and_input(monkeypatch, tmp_path, capsys):
    state, _ = setup(monkeypatch, tmp_path)
    code, result = invoke(["review", "report"], capsys)
    assert code == 2 and result["data"]["error"]["code"] == "invalid_input"
    assert not state["calls"]


@pytest.mark.parametrize(
    "state_field,value",
    [("malformed", True), ("preview_parent", OTHER), ("preview_decisions", SHA)],
)
def test_report_invalid_or_changed_snapshot_response_fails(
    monkeypatch, tmp_path, capsys, state_field, value
):
    state, path = setup(monkeypatch, tmp_path)
    state[state_field] = value
    code, result = invoke(render_flags(path, "--dry-run"), capsys)
    assert code == 4 and result["data"]["error"]["code"] == "invalid_server_response"


@pytest.mark.parametrize("dry_run", [True, False])
def test_report_preview_and_submit_responses_cannot_be_interchanged(
    monkeypatch, tmp_path, capsys, dry_run
):
    state, path = setup(monkeypatch, tmp_path)
    state["queued_on_preview" if dry_run else "preview_on_submit"] = True
    flags = (
        ["--dry-run"]
        if dry_run
        else [
            "--expected-input-hash",
            RENDER_INPUT,
            "--preflight-token",
            "synthetic-report-receipt",
        ]
    )
    code, result = invoke(render_flags(path, *flags), capsys)
    assert code == 4 and result["data"]["error"]["code"] == "invalid_server_response"


def test_report_queue_failure_is_retryable(monkeypatch, tmp_path, capsys):
    state, path = setup(monkeypatch, tmp_path)
    state["queue_failure"] = True
    code, result = invoke(
        render_flags(
            path,
            "--expected-input-hash",
            RENDER_INPUT,
            "--preflight-token",
            "synthetic-report-receipt",
        ),
        capsys,
    )
    assert code == 3 and result["data"]["error"]["code"] == "queue_unavailable"


@pytest.mark.parametrize(
    "link",
    [
        "https://foreign.example.test/file",
        "//foreign.example.test/file",
        f"/bid-review-artifacts/{OTHER}/download?signature=synthetic-report-only",
        f"/bid-review-artifacts/{ARTIFACT}/download?signature=one&signature=two",
        f"/bid-review-artifacts/{ARTIFACT}/download?signature=synthetic-report-only&extra=true",
        f"/bid-review-artifacts/{ARTIFACT}/download?signature=synthetic-report-only#fragment",
    ],
)
def test_report_download_rejects_unsafe_routes(monkeypatch, tmp_path, capsys, link):
    state, _ = setup(monkeypatch, tmp_path)
    state["link"] = link
    output = tmp_path / "refused.docx"
    code, result = invoke(download_flags(output), capsys)
    assert code == 4 and result["data"]["error"]["code"] == "invalid_download_link"
    assert state["calls"] == ["link"] and not output.exists()
    assert not list(tmp_path.glob(".bid-download-*"))


@pytest.mark.parametrize("field,value", [("id", OTHER), ("org_id", OTHER), ("format", "console")])
def test_report_download_rejects_wrong_descriptor_scope(
    monkeypatch, tmp_path, capsys, field, value
):
    state, _ = setup(monkeypatch, tmp_path)
    state["artifact"][field] = value
    output = tmp_path / "refused.docx"
    code, result = invoke(download_flags(output), capsys)
    assert code == 4 and result["data"]["error"]["code"] == "invalid_server_response"
    assert state["calls"] == ["link"] and not output.exists()


def test_report_download_rejects_another_command_envelope(monkeypatch, tmp_path, capsys):
    state, _ = setup(monkeypatch, tmp_path)
    state["link_command"] = "export download"
    output = tmp_path / "refused.docx"
    code, result = invoke(download_flags(output), capsys)
    assert code == 4 and result["data"]["error"]["code"] == "invalid_server_response"
    assert state["calls"] == ["link"] and not output.exists()


@pytest.mark.parametrize(
    "failure", ["length", "hash", "package", "media_type", "redirect", "expired"]
)
def test_report_download_failure_never_publishes_local_file(monkeypatch, tmp_path, capsys, failure):
    state, _ = setup(monkeypatch, tmp_path)
    if failure == "length":
        state["served"] = state["content"][:-1]
    elif failure == "hash":
        state["artifact"]["sha256"] = SHA
    elif failure == "package":
        state["served"] = b"unreadable-docx"
        state["artifact"].update(
            sha256=hashlib.sha256(state["served"]).hexdigest(), size_bytes=len(state["served"])
        )
    elif failure == "media_type":
        state["media_type"] = "text/plain"
    else:
        state[failure] = True
    output = tmp_path / "refused.docx"
    code, result = invoke(download_flags(output), capsys)
    expected = (
        "download_link_expired"
        if failure == "expired"
        else "invalid_server_response"
        if failure == "redirect"
        else "bid_report_file_integrity"
    )
    assert code == 4 and result["data"]["error"]["code"] == expected
    assert not output.exists() and not list(tmp_path.glob(".bid-download-*"))


def test_report_download_refuses_existing_path_or_wrong_suffix_before_http(
    monkeypatch, tmp_path, capsys
):
    state, _ = setup(monkeypatch, tmp_path)
    output = tmp_path / "existing.docx"
    output.write_bytes(b"preserve existing file")
    for target in (output, tmp_path / "wrong.pdf"):
        code, result = invoke(download_flags(target), capsys)
        assert code == 2 and not result["ok"] and not state["calls"]
    assert output.read_bytes() == b"preserve existing file"


def test_report_discovery_matches_shared_models():
    schema = command_schema(cli.app)
    for name in ("review report", "review report download"):
        assert name in schema["commands"] and name not in command_schema(cli.app, "3.0")["commands"]
    assert schema["commands"]["review report"]["input"]["title"] == "BidReportRenderRequest"
    assert (
        schema["commands"]["review report download"]["output"]["title"]
        == "BidReportDownloadReceipt"
    )
    parameters = {item["name"] for item in schema["commands"]["review report"]["cli_parameters"]}
    assert {
        "id",
        "input",
        "dry_run",
        "expected_input_hash",
        "preflight_token",
        "retry",
    } <= parameters
