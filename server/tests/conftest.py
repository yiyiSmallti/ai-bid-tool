import hashlib
import os
from uuid import uuid4

import httpx
import pytest
from alembic import command
from alembic.config import Config
from app.api.main import create_app
from app.core.config import Settings
from app.providers.llm import HTTPExtractor
from cryptography.fernet import Fernet
from sqlalchemy import create_engine, text
from sqlalchemy.engine import make_url


def prepare_database(url: str) -> None:
    """Create the application role if missing and migrate the database at `url` to head.

    Roles are cluster-wide; per-worker setup uses the maintenance advisory lock
    before creating the shared role. Existing role passwords are never changed.
    """
    if not (make_url(url).database or "").startswith("bid_test"):
        raise RuntimeError("Tests refuse to mutate a database without the bid_test prefix")
    engine = create_engine(url, hide_parameters=True)
    with engine.begin() as connection:
        if not connection.scalar(text("SELECT 1 FROM pg_roles WHERE rolname='bid_app'")):
            from psycopg import sql

            statement = sql.SQL(
                "CREATE ROLE bid_app LOGIN NOSUPERUSER NOBYPASSRLS NOCREATEDB NOCREATEROLE"
            )
            if os.environ.get("BID_DATABASE_PASSWORD"):
                statement += sql.SQL(" PASSWORD {} ").format(
                    sql.Literal(os.environ["BID_DATABASE_PASSWORD"])
                )
            connection.execute(text(statement.as_string()))
    engine.dispose()
    os.environ["BID_MIGRATION_DATABASE_URL"] = url
    command.upgrade(Config("alembic.ini"), "head")


# Tests that run alone on shard 1 whenever they are selected: the longest scale test,
# and the board latency bounds, which must not inherit the database bloat and vacuum
# load of the scale tests that would otherwise run before them on the same runner.
ISOLATED_SHARD_TESTS = (
    "server/tests/test_draft_batch_reads.py::test_large_draft_reads_match_reference",
    "server/tests/test_team_workflow_stream_acceptance.py::test_board_5000_requirements",
    "server/tests/test_team_workflow_stream_acceptance.py::test_board_complete_bounds",
)
# Scale tests that each take minutes on a CI runner, pinned to shards 2, 3 and 4 so the
# wall-clock time of a sharded run is not the sum of several of them.
PINNED_SHARD_TESTS = (
    "server/tests/test_runtime_integration.py::test_real_background_worker_and_both_cli_modes",
    "server/tests/test_score_request_scope_db.py::test_large_draft_score_preview_run_scopes",
    "server/tests/test_fast_citation.py::test_large_check_publish_includes_deferred_commit",
)


def shard_of(nodeid: str, count: int, reserve_first: bool) -> int:
    """1-based shard for a test ID; isolated and pinned tests keep their shard when it exists."""
    if reserve_first and count > 1 and nodeid.startswith(ISOLATED_SHARD_TESTS):
        return 1
    for position, prefix in enumerate(PINNED_SHARD_TESTS, start=2):
        if nodeid.startswith(prefix) and position <= count:
            return position
    first = 2 if reserve_first and count > 1 else 1
    digest = hashlib.sha256(nodeid.encode()).digest()
    return first + int.from_bytes(digest[:8], "big") % (count - first + 1)


# After -m and -k deselection, so shards split only the tests this run selects.
@pytest.hookimpl(trylast=True)
def pytest_collection_modifyitems(config, items):
    """Keep only this CI shard's tests when BID_TEST_SHARD is set as "index/count".

    Each test ID maps to exactly one shard, identically in every xdist worker and every
    runner, so the shards together run the whole selection.
    """
    shard = os.environ.get("BID_TEST_SHARD")
    if not shard:
        return
    index, count = (int(part) for part in shard.split("/"))
    if not 1 <= index <= count:
        raise pytest.UsageError(f"BID_TEST_SHARD must be index/count, got {shard!r}")
    # The selection is identical on every shard, so this decision is too.
    reserve_first = any(item.nodeid.startswith(ISOLATED_SHARD_TESTS) for item in items)
    selected, deselected = [], []
    for item in items:
        shard_index = shard_of(item.nodeid, count, reserve_first)
        (selected if shard_index == index else deselected).append(item)
    config.hook.pytest_deselected(items=deselected)
    items[:] = selected


def pytest_configure(config):
    """Give each pytest-xdist worker a copy of the migrated test database.

    Tests truncate shared tables, so workers cannot share one database. Roles are
    cluster-wide and migrations create them, so one worker at a time migrates the base
    database under an advisory lock taken in the maintenance database, then copies it.
    """
    worker = os.environ.get("PYTEST_XDIST_WORKER")
    admin = os.environ.get("BID_TEST_ADMIN_URL")
    if not worker or not admin:
        return
    base = make_url(admin)
    name = f"{base.database}_{worker}"
    maintenance = create_engine(
        base.set(database="postgres"), isolation_level="AUTOCOMMIT", hide_parameters=True
    )
    with maintenance.connect() as connection:
        connection.execute(text("SELECT pg_advisory_lock(hashtext('bid_test_workers'))"))
        try:
            prepare_database(admin)
            # A database with open sessions cannot be a template; the migration's pool may
            # still hold one, and no other worker uses the base database under this lock.
            connection.execute(
                text(
                    "SELECT pg_terminate_backend(pid) FROM pg_stat_activity"
                    " WHERE datname = :name AND pid <> pg_backend_pid()"
                ),
                {"name": base.database},
            )
            connection.execute(text(f'DROP DATABASE IF EXISTS "{name}" WITH (FORCE)'))
            connection.execute(text(f'CREATE DATABASE "{name}" TEMPLATE "{base.database}"'))
        finally:
            connection.execute(text("SELECT pg_advisory_unlock(hashtext('bid_test_workers'))"))
    maintenance.dispose()
    os.environ["BID_TEST_ADMIN_URL"] = base.set(database=name).render_as_string(hide_password=False)
    for variable in ("BID_DATABASE_URL",):
        if os.environ.get(variable):
            app = make_url(os.environ[variable]).set(database=name)
            os.environ[variable] = app.render_as_string(hide_password=False)


@pytest.fixture(autouse=True)
def no_retry_waits(monkeypatch):
    # Transient-failure retries are covered by their own test; elsewhere one vendor
    # reply per call keeps scripted responses simple.
    monkeypatch.setattr(HTTPExtractor, "retry_delays", ())


@pytest.fixture(scope="session")
def admin_engine():
    url = os.environ.get("BID_TEST_ADMIN_URL")
    if not url:
        # Failing instead of skipping keeps a run without the database from looking green.
        pytest.fail("BID_TEST_ADMIN_URL must point at an isolated PostgreSQL test runtime")
    # Under xdist, pytest_configure already prepared this worker's copy under a lock.
    prepare_database(url)
    engine = create_engine(url, hide_parameters=True)
    yield engine
    engine.dispose()


@pytest.fixture
def database(admin_engine, monkeypatch):
    if not os.environ.get("BID_SECRETS_KEY"):
        monkeypatch.setenv("BID_SECRETS_KEY", Fernet.generate_key().decode())
    from psycopg import sql

    with admin_engine.begin() as connection:
        tables = connection.scalars(
            text(
                "SELECT tablename FROM pg_tables WHERE schemaname='public' AND tablename <> 'alembic_version'"
            )
        ).all()
        if tables:
            statement = sql.SQL("TRUNCATE {} CASCADE").format(
                sql.SQL(", ").join(sql.Identifier("public", name) for name in tables)
            )
            connection.exec_driver_sql(statement.as_string())
    return admin_engine


class FakeQueue:
    def __init__(self):
        self.calls = []
        self.processor = None
        self.wakes = []
        self.annotation_cleanups = []
        self.attachment_invalidations = []

    async def enqueue(self, job_id: str):
        self.calls.append(job_id)
        return len(self.calls)

    async def enqueue_in_transaction(self, session, job_id: str):
        return await self.enqueue(job_id)

    async def enqueue_attachment_invalidation_in_transaction(self, session, **arguments):
        self.attachment_invalidations.append(arguments)
        return len(self.attachment_invalidations)

    async def enqueue_agent_wake(self, session, session_id: str, *, delay=30):
        self.wakes.append(session_id)
        return len(self.wakes)

    async def enqueue_annotation_cleanup_in_transaction(self, session, job_id: str, *, delay=300):
        self.annotation_cleanups.append((job_id, max(300, delay)))
        return len(self.annotation_cleanups)

    async def ensure_process_delivery(self, session, job):
        if job.queue_id is None:
            job.queue_id = await self.enqueue(str(job.id))


@pytest.fixture
def application(database, tmp_path):
    settings = Settings(data_dir=tmp_path)
    return create_app(settings, queue=FakeQueue())


@pytest.fixture
async def api(application):
    async with (
        application.router.lifespan_context(application),
        httpx.AsyncClient(
            transport=httpx.ASGITransport(app=application), base_url="http://test"
        ) as client,
    ):
        yield client


@pytest.fixture
async def headers(api):
    response = await api.post("/auth/session", json={"access_key": os.environ["BID_ACCESS_KEY"]})
    assert response.status_code == 200
    return {"Authorization": "Bearer " + response.json()["data"]["session"]}


@pytest.fixture
def token_headers():
    from datetime import UTC, datetime, timedelta

    async def create(api, scopes):
        response = await api.post(
            "/tokens",
            headers={"Authorization": "Bearer " + os.environ["BID_ACCESS_KEY"]},
            json={
                "name": "synthetic token",
                "scopes": sorted(scopes),
                "expires_at": (datetime.now(UTC) + timedelta(hours=1)).isoformat(),
            },
        )
        assert response.status_code == 200, response.text
        return {"Authorization": "Bearer " + response.json()["data"]["token"]}

    return create


@pytest.fixture
def pdf_bytes():
    import pymupdf

    with pymupdf.open() as document:
        page = document.new_page()
        page.insert_text(
            (40, 60),
            "Synthetic test fixture, not a real tender.\nRequired delivery within 30 days.\nMinimum memory is 64 GB.",
        )
        second = document.new_page()
        second.insert_text((40, 60), "A valid certificate must be provided.")
        return document.tobytes()


@pytest.fixture
def docx_bytes():
    from io import BytesIO

    from docx import Document

    document = Document()
    document.add_heading("Synthetic template only", level=1)
    document.add_paragraph("Synthetic content, not a real bid or proof.")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


async def seed_platform_credential(
    settings,
    *,
    name="main",
    provider="anthropic",
    endpoint="https://api.anthropic.com",
    key="synthetic-platform-key",
    active=True,
    purpose="standalone_llm",
):
    """Persist a synthetic local credential through the human settings service."""
    from app.core.db import Database
    from app.schemas.platform_credentials import CredentialCreate
    from app.services.auth import HUMAN_SCOPES, Identity
    from app.services.platform_credentials import PlatformCredentialService
    from pydantic import SecretStr

    db = Database(settings)
    try:
        return (
            await PlatformCredentialService(settings, db).create(
                Identity(scopes=set(HUMAN_SCOPES)),
                CredentialCreate(
                    name=name,
                    purpose=purpose,
                    provider=provider,
                    endpoint=endpoint,
                    api_key=SecretStr(key),
                    active=active,
                ),
            )
        ).credential
    finally:
        await db.engine.dispose()


async def credential_app(settings, *, llm: HTTPExtractor, queue):
    """Bind an injected mock adapter to an encrypted local provider revision."""
    from app.core.provider_secrets import ProviderSecrets
    from app.models.provider_configs import ProviderConfig
    from app.providers.configured import configured_llm
    from app.schemas.platform_credentials import CredentialSpec
    from app.services.auth import HUMAN_SCOPES, Identity, set_actor_context
    from sqlalchemy import func, select

    key = llm.settings.llm_api_key
    assert key is not None and isinstance(llm.transport, httpx.MockTransport)
    provider = "anthropic" if llm.name == "anthropic" else "openai"
    endpoint = CredentialSpec.safe_endpoint(
        llm.settings.llm_base_url
        or ("https://api.anthropic.com" if llm.name == "anthropic" else "https://api.openai.com/v1")
    )
    clean_settings = settings.model_copy(update={"llm_api_key": None})
    app = create_app(clean_settings, llm=llm, queue=queue)
    config_id = uuid4()
    async with app.state.db.transaction() as session:
        await set_actor_context(session, Identity(scopes=set(HUMAN_SCOPES)))
        revision = (await session.scalar(select(func.max(ProviderConfig.revision)))) or 0
        config = ProviderConfig(
            id=config_id,
            capability="llm_extract",
            revision=revision + 1,
            encrypted_key=ProviderSecrets(clean_settings).encrypt(
                key.get_secret_value(), config_id
            ),
            key_last4=key.get_secret_value()[-4:],
            data={
                "provider": provider,
                "model": llm.model,
                "base_url": endpoint,
                "json_mode": llm.settings.llm_json_mode,
                "reasoning": list(llm.reasoning_levels.values()),
                "default_reasoning": llm.default_reasoning,
            },
        )
        session.add(config)
        await session.flush()
    configured = configured_llm(clean_settings, config, llm.transport)
    llm.settings = configured.settings
    llm.provider_config_id = config.id
    llm.model_revision = config.revision
    llm.version = configured.version
    llm.credential_resolver = None
    llm.credential_target = None
    return app
