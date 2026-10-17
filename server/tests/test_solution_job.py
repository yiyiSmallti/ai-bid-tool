"""API-to-worker acceptance with synthetic model transport and durable usage.

Failure modes: agents cannot draft, objective scoring rows enter model calls,
needs leak past the chapter/budget, receipts disappear, retries redraft completed
chapters, a queued or late result replaces a human edit, and job reads leak the
encrypted snapshot. Artifacts contain only synthetic provider payloads/results.
"""

import asyncio
import json
import os
from contextlib import asynccontextmanager
from io import BytesIO
from typing import cast
from uuid import UUID

import httpx
import pytest
from app.api.main import create_app
from app.core.config import Settings
from app.jobs import solution_draft
from app.models.entities import Job, UsageRecord
from app.providers.base import ProviderFailure
from app.providers.solution import HTTPSolutionProvider
from app.schemas.solution import SolutionProviderRequest, SolutionProviderResult
from app.services import solution_generation
from docx import Document
from pydantic import SecretStr
from sqlalchemy import select
from tests.conftest import FakeQueue
from tests.test_check_combined import SemanticVendor, semantic_llm
from tests.test_response_cards import login
from tests.test_solution_provider import synthetic_output


class SolutionVendor:
    def __init__(self):
        self.requests = []
        self.calls = 0
        self.fail_at: int | None = None
        self.entered: asyncio.Event | None = None
        self.release: asyncio.Event | None = None

    async def __call__(self, request):
        body = json.loads(request.content)
        name = body["response_format"]["json_schema"]["name"]
        sent = json.loads(body["messages"][-1]["content"])
        self.calls += 1
        if name == "solution_plan":
            # Plan requests carry the item inputs; section calls follow each plan.
            self.requests.append(sent)
            if self.entered is not None:
                self.entered.set()
                assert self.release is not None
                await self.release.wait()
            if len(self.requests) == self.fail_at:
                raise httpx.ConnectError("Synthetic temporary model outage", request=request)
        return httpx.Response(
            200,
            json={
                "model": "synthetic-check-model",
                "choices": [
                    {
                        "finish_reason": "stop",
                        "message": {"content": json.dumps(synthetic_output(name))},
                    }
                ],
                "usage": {"prompt_tokens": 120, "completion_tokens": 80},
            },
        )


def tender_bytes():
    document = Document()
    document.add_paragraph("项目名称：合成服务测试项目")
    document.add_paragraph("采购人：合成采购单位")
    document.add_paragraph("第一章 采购需求")
    document.add_paragraph("电话：13800138000。交接台账应留存。" + "需求细节。" * 300)
    document.add_paragraph("第二章 评分标准")
    table = document.add_table(rows=0, cols=3)
    for title, score, rule in (
        ("整体服务方案", "0~25", "根据组织、执行、检查、闭环内容的完整性、针对性综合评审。"),
        ("快速切换的机制", "0~15", "根据快速切换机制的合理性、可操作性综合评审。"),
        ("企业业绩", "0~5", "每份合同得1分，最高5分。"),
        ("价格", "0~20", "价格按基准价计算。"),
    ):
        row = table.add_row()
        for cell, text in zip(row.cells, (title, score, rule), strict=True):
            cell.text = text
    document.add_paragraph("第三章 投标文件格式")
    document.add_paragraph("附件1 投标函")
    document.add_paragraph("格式中的机密文字不属于采购需求。")
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


@asynccontextmanager
async def case(database, tmp_path, monkeypatch):
    vendor = SolutionVendor()
    settings = Settings(
        data_dir=tmp_path,
        solution_needs_max_chars=1000,
        database_url=SecretStr(os.environ["BID_DATABASE_URL"]),
        access_key=SecretStr(os.environ["BID_ACCESS_KEY"]),
        encryption_key=SecretStr(os.environ["BID_ENCRYPTION_KEY"]),
        token_key=SecretStr(os.environ["BID_TOKEN_KEY"]),
    )
    llm = semantic_llm(settings, cast(SemanticVendor, vendor))

    async def resolved(session, settings, transport=None, job=None, model=None):
        return llm

    monkeypatch.setattr(solution_generation.llm_providers, "resolve_llm", resolved)
    app = create_app(settings, llm=llm, queue=FakeQueue())
    async with (
        app.router.lifespan_context(app),
        httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test") as api,
    ):
        header = await login(api)
        task = (await api.post("/tasks", headers=header, json={"name": "合成方案测试"})).json()[
            "data"
        ]["id"]
        uploaded = await api.post(
            f"/tasks/{task}/documents",
            headers=header,
            files={
                "file": (
                    "synthetic.docx",
                    tender_bytes(),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )
        assert uploaded.status_code == 200, uploaded.text
        document = uploaded.json()["data"]["id"]
        parsed = await api.post(f"/documents/{document}/parse", headers=header, json={})
        assert parsed.status_code == 200, parsed.text
        await app.state.processor(parsed.json()["data"]["job_id"])
        yield api, app, header, task, document, vendor


async def submit(api, header, task, document, titles=None):
    response = await api.post(
        f"/tasks/{task}/solutions/draft",
        headers=header,
        json={
            "document_id": document,
            "titles": titles,
        },
    )
    assert response.status_code == 200, response.text
    return response.json()["data"]["job_id"]


async def chapters(api, header, task):
    response = await api.get(f"/tasks/{task}/solutions", headers=header)
    assert response.status_code == 200, response.text
    return response.json()["items"]


async def test_token_solution_draft_job_persists_usage_and_scoped_inputs(
    database, tmp_path, monkeypatch, token_headers
):
    async with case(database, tmp_path, monkeypatch) as (api, app, header, task, document, vendor):
        token = await token_headers(
            api, {"task:read", "solution:draft", "solution:read", "job:read"}
        )
        job_id = await submit(api, token, task, document)
        await app.state.processor(job_id)
        status = await api.get(f"/jobs/{job_id}", headers=token)
        assert status.status_code == 200, status.text
        assert status.json()["data"]["status"] == "succeeded", status.text
        assert "submission" not in status.json()["data"]["result"]
        assert status.json()["cost"]["calls"] == 10
        assert status.json()["cost"]["llm_tokens"] == 2000
        rows = await chapters(api, token, task)
        assert {row["title"] for row in rows} == {"整体服务方案", "快速切换的机制"}
        assert all(
            row["current"]["origin"] == "model" and row["current"]["state"] == "draft"
            for row in rows
        )
        assert all(
            row["current"]["actor_kind"] == "worker" and row["current"]["job_id"] == job_id
            for row in rows
        )
        assert all(row["placeholder_count"] == 4 for row in rows)
        assert all(len(req["needs_text"]) <= 1000 for req in vendor.requests)
        assert all("13800138000" not in req["needs_text"] for req in vendor.requests)
        assert all("格式中的机密" not in req["needs_text"] for req in vendor.requests)
        assert all(
            req["project_name"] == "合成服务测试项目" and req["tenderee"] == "合成采购单位"
            for req in vendor.requests
        )
        denied = await api.post(
            f"/solutions/{rows[0]['id']}/confirm", headers=token, json={"expected_revision": 1}
        )
        assert denied.status_code == 403
        async with app.state.db.transaction() as session:
            usage = (
                await session.scalars(select(UsageRecord).where(UsageRecord.job_id == UUID(job_id)))
            ).all()
            assert len(usage) == 10
            assert sum(row.tokens for row in usage) == 2000
        (tmp_path / "solution-job-acceptance.json").write_text(
            json.dumps(
                {
                    "job_status": status.json(),
                    "chapters": rows,
                    "provider_requests": vendor.requests,
                    "token_confirmation_status": denied.status_code,
                },
                ensure_ascii=False,
                indent=2,
            )
        )


async def test_solution_queued_and_late_results_preserve_human_edits(
    database, tmp_path, monkeypatch
):
    async with case(database, tmp_path, monkeypatch) as (api, app, header, task, document, vendor):
        first = await submit(api, header, task, document, ["整体服务方案"])
        await app.state.processor(first)
        chapter = (await chapters(api, header, task))[0]
        second = await submit(api, header, task, document, ["整体服务方案"])
        edited = await api.put(
            f"/solutions/{chapter['id']}",
            headers=header,
            json={
                "expected_revision": 1,
                "sections": [{"heading": "人工方案", "paragraphs": ["人工保留的内容"]}],
            },
        )
        assert edited.status_code == 200, edited.text
        await app.state.processor(second)
        status = (await api.get(f"/jobs/{second}", headers=header)).json()
        assert status["data"]["status"] == "failed"
        assert status["data"]["error"]["code"] == "solution_revision_conflict"
        assert len(vendor.requests) == 1
        third = await submit(api, header, task, document, ["整体服务方案"])
        vendor.entered, vendor.release = asyncio.Event(), asyncio.Event()
        running = asyncio.create_task(app.state.processor(third))
        await asyncio.wait_for(vendor.entered.wait(), timeout=10)
        edit = await api.put(
            f"/solutions/{chapter['id']}",
            headers=header,
            json={
                "expected_revision": 2,
                "sections": [{"heading": "最新人工方案", "paragraphs": ["后到模型不能覆盖"]}],
            },
        )
        assert edit.status_code == 200, edit.text
        vendor.release.set()
        await running
        status = (await api.get(f"/jobs/{third}", headers=header)).json()
        assert status["data"]["status"] == "failed"
        assert status["cost"]["calls"] == 1
        current = (await chapters(api, header, task))[0]["current"]
        assert current["revision"] == 3 and current["origin"] == "human"
        assert current["sections"][0]["paragraphs"] == ["后到模型不能覆盖"]
        (tmp_path / "solution-human-edit-fences.json").write_text(
            json.dumps(
                {"queued": second, "late": status, "current": current}, ensure_ascii=False, indent=2
            )
        )


async def test_solution_retry_resumes_after_published_chapter(database, tmp_path, monkeypatch):
    async with case(database, tmp_path, monkeypatch) as (api, app, header, task, document, vendor):
        vendor.fail_at = 2
        job_id = await submit(api, header, task, document)
        with pytest.raises(ProviderFailure):
            await app.state.processor(job_id)
        assert len(await chapters(api, header, task)) == 1
        vendor.fail_at = None
        await app.state.processor(job_id)
        status = (await api.get(f"/jobs/{job_id}", headers=header)).json()
        assert status["data"]["status"] == "succeeded", status
        assert len(vendor.requests) == 3
        assert vendor.requests[0]["title"] != vendor.requests[1]["title"]
        assert vendor.requests[1]["title"] == vendor.requests[2]["title"]
        rows = await chapters(api, header, task)
        assert len(rows) == 2 and all(row["current"]["revision"] == 1 for row in rows)
        async with app.state.db.transaction() as session:
            job = await session.get(Job, UUID(job_id))
            assert job.attempts == 2
        (tmp_path / "solution-retry.json").write_text(
            json.dumps(
                {"status": status, "requests": vendor.requests}, ensure_ascii=False, indent=2
            )
        )


@pytest.mark.parametrize("failure", ["missing_receipt", "reused_receipt"])
async def test_solution_publish_rejects_missing_or_reused_accounting(
    database, tmp_path, monkeypatch, failure
):
    async with case(database, tmp_path, monkeypatch) as (api, app, header, task, document, vendor):

        class InvalidAccountingAdapter:
            def __init__(self, provider):
                self.provider = HTTPSolutionProvider(provider)
                self.previous: SolutionProviderResult | None = None

            async def draft(self, request: SolutionProviderRequest) -> SolutionProviderResult:
                if self.previous is None:
                    self.previous = await self.provider.draft(request)
                if failure == "missing_receipt":
                    return self.previous.model_copy(update={"usages": []})
                return self.previous

        monkeypatch.setattr(solution_draft, "solution_provider", InvalidAccountingAdapter)
        job_id = await submit(api, header, task, document)
        await app.state.processor(job_id)
        status = (await api.get(f"/jobs/{job_id}", headers=header)).json()
        assert status["data"]["status"] == "failed", status
        assert status["data"]["error"]["code"] == "solution_draft_usage_integrity"
        rows = await chapters(api, header, task)
        assert len(rows) == (0 if failure == "missing_receipt" else 1)
        assert len(vendor.requests) == 1
        async with app.state.db.transaction() as session:
            usages = (
                await session.scalars(select(UsageRecord).where(UsageRecord.job_id == UUID(job_id)))
            ).all()
            assert len(usages) == 5
        (tmp_path / "solution-accounting-fence.json").write_text(
            json.dumps(
                {"failure": failure, "status": status, "published_chapters": len(rows)},
                ensure_ascii=False,
                indent=2,
            )
        )
