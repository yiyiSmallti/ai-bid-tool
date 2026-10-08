"""HTTP → encrypted upload → fenced worker → inventory acceptance.

Failure modes were enumerated first in data/work/bid-review-upload/failure-modes.md.
Only synthetic originals are used. Running this file requires the main session's
isolated PostgreSQL runtime; it never starts a service or calls an external vendor.
"""

import asyncio
import hashlib
import io
import json
import zipfile
from datetime import UTC, datetime, timedelta
from pathlib import Path
from uuid import UUID, uuid4

import httpx
import pymupdf
import pytest
from app.core.errors import ServiceError
from app.core.security import TokenSigner
from app.models.bid_review import (
    BidDocumentPage,
    BidPreparation,
    BidPreparationPublication,
    BidPreparedDocument,
    BidSubmission,
    BidSubmissionDocument,
)
from app.models.entities import AuditLog, Document, Job, UsageRecord
from app.services import bid_preparation
from sqlalchemy import func, select, text
from sqlalchemy.exc import DBAPIError

OUTPUT = Path(__file__).resolve().parents[2] / "data/work/bid-review-upload"
PDF = "application/pdf"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
SENTINEL = "SYNTHETIC-CONFIDENTIAL-PRICE-937462"
TABLES = (
    BidSubmission,
    BidSubmissionDocument,
    BidPreparation,
    BidPreparedDocument,
    BidDocumentPage,
    BidPreparationPublication,
    Job,
    UsageRecord,
    AuditLog,
)


def pdf_bytes(*, image=False, signature=False, encrypted=False, empty_password=False):
    with pymupdf.open() as pdf:
        page = pdf.new_page(width=300, height=400)
        page.insert_text((20, 30), SENTINEL)
        if signature:
            field = pymupdf.Widget()
            field.field_type = pymupdf.PDF_WIDGET_TYPE_SIGNATURE
            field.field_name = "synthetic-private-signature-field"
            field.rect = pymupdf.Rect(20, 80, 120, 130)
            page.add_widget(field)
        if image:
            # An image-only page and a scan with a native page number distinguish
            # page kind from mere existence of any native text.
            pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 40, 40), False)
            pix.clear_with(127)
            image_bytes = pix.tobytes("png")
            for with_header in (False, True):
                page = pdf.new_page(width=300, height=400)
                page.insert_image(pymupdf.Rect(0, 50, 300, 400), stream=image_bytes)
                if with_header:
                    page.insert_text((20, 25), "3")
        if encrypted:
            return pdf.tobytes(
                encryption=pymupdf.PDF_ENCRYPT_AES_256,
                owner_pw="synthetic-owner",
                user_pw="" if empty_password else "synthetic-user",
            )
        return pdf.tobytes()


def docx_bytes():
    from docx import Document as WordDocument

    document = WordDocument()
    document.add_heading("Synthetic heading", level=1)
    document.add_paragraph(SENTINEL)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def metadata(files, **changes):
    return {
        "request_id": str(uuid4()),
        "files": [
            {
                "role": "tender" if index == 0 else "bid",
                "kind": "tender" if index == 0 else "other",
                "media_type": media,
                "sha256": hashlib.sha256(content).hexdigest(),
                "size_bytes": len(content),
            }
            for index, (_, media, content) in enumerate(files)
        ],
        **changes,
    }


def result(response):
    assert response.status_code == 200, response.text
    value = response.json()
    assert set(value) == {"ok", "command", "data", "items", "warnings", "cost", "duration_ms"}
    assert value["ok"] is True
    assert value["cost"]["charge"] == value["cost"]["task_amount"] == "0"
    assert SENTINEL not in response.text
    assert "synthetic-private-signature-field" not in response.text
    return value["data"]


def error(response, code=None, status=None):
    assert response.status_code == status if status else response.status_code >= 400
    value = response.json()["data"]["error"]
    if code:
        assert value["code"] == code, value
    assert SENTINEL not in response.text
    return value


async def task(api, header):
    return result(
        await api.post("/v4/tasks", headers=header, json={"name": "Synthetic review task"})
    )["id"]


async def upload(api, header, task_id, files, body=None):
    return await api.post(
        f"/v4/tasks/{task_id}/bid-submissions",
        headers=header,
        data={"metadata": json.dumps(body or metadata(files))},
        files=[("files", (name, content, media)) for name, media, content in files],
    )


async def preview(api, header, task_id, submission_id, **changes):
    return await api.post(
        f"/v4/tasks/{task_id}/bid-submissions/{submission_id}/prepare",
        headers=header,
        json={
            "request_id": str(uuid4()),
            "submission_id": submission_id,
            "dry_run": True,
            **changes,
        },
    )


async def submit(api, header, task_id, submission_id, receipt, **changes):
    body = {
        "request_id": str(uuid4()),
        "submission_id": submission_id,
        "expected_input_hash": receipt["input_hash"],
        "preflight_token": receipt["preflight_token"],
        **changes,
    }
    return await api.post(
        f"/v4/tasks/{task_id}/bid-submissions/{submission_id}/prepare", headers=header, json=body
    )


async def counts(application, org_id):
    async with application.state.db.transaction(org_id) as session:
        return {
            model.__tablename__: await session.scalar(select(func.count()).select_from(model))
            for model in TABLES
        }


def originals():
    return [
        ("private-tender.pdf", PDF, pdf_bytes()),
        ("private-bid.pdf", PDF, pdf_bytes(image=True, signature=True)),
    ]


async def run(application, org_id, job_id):
    await application.state.processor(str(org_id), str(job_id))


async def test_upload_prepare_inventory_and_encryption(api, headers, tenants, application, caplog):
    task_id, files = await task(api, headers[0]), originals()
    before = await counts(application, tenants["orgs"][0])
    objects_before = list((application.state.storage.root / "org").rglob("*"))
    dry = result(await upload(api, headers[0], task_id, files, metadata(files, dry_run=True)))
    assert dry["limits"]["files"] == 20 and len(dry["files"]) == 2
    assert await counts(application, tenants["orgs"][0]) == before
    assert list((application.state.storage.root / "org").rglob("*")) == objects_before
    body = metadata(files)
    uploaded = result(await upload(api, headers[0], task_id, files, body))
    assert result(await upload(api, headers[0], task_id, files, body)) == uploaded
    before = await counts(application, tenants["orgs"][0])
    receipt = result(await preview(api, headers[0], task_id, uploaded["id"]))
    assert receipt["local_only"] is True and receipt["budget"]["planned_calls"] == 0
    assert (
        receipt["budget"]["maximum_calls"] == 0
        and receipt["budget"]["full_run_guaranteed"] is False
    )
    assert (
        datetime.fromisoformat(receipt["expires_at"])
        - datetime.fromisoformat(receipt["budget"]["as_of"])
    ).total_seconds() == 900
    assert await counts(application, tenants["orgs"][0]) == before
    request_id = str(uuid4())
    accepted = result(
        await submit(api, headers[0], task_id, uploaded["id"], receipt, request_id=request_id)
    )
    replay = result(
        await submit(api, headers[0], task_id, uploaded["id"], receipt, request_id=request_id)
    )
    assert replay["job_id"] == accepted["job_id"] and replay["cached"]
    error(
        await submit(
            api, headers[0], task_id, uploaded["id"], receipt, request_id=request_id, retry=True
        ),
        "idempotency_conflict",
        409,
    )
    assert len(application.state.queue.calls) == 1
    await run(application, tenants["orgs"][0], accepted["job_id"])
    detail = result(await api.get(f"/v4/bid-submissions/{uploaded['id']}", headers=headers[0]))
    assert detail["preparation"]["status"] == "succeeded"
    assert detail["submission"]["state"] == "prepared"
    assert sum(row["page_count"] for row in detail["inventory"]) == 4
    assert sum(row["text_pages"] for row in detail["inventory"]) == 2
    assert sum(row["image_pages"] for row in detail["inventory"]) == 2
    assert sum(row["signature_fields"] for row in detail["inventory"]) == 1
    job = result(await api.get(f"/v4/jobs/{accepted['job_id']}", headers=headers[0]))
    assert "submission" not in job.get("result", {})
    async with application.state.db.transaction(tenants["orgs"][0]) as session:
        assert await session.scalar(select(func.count()).select_from(Document)) == 0
        assert await session.scalar(select(func.count()).select_from(UsageRecord)) == 0
        docs = list((await session.scalars(select(BidSubmissionDocument))).all())
        pages = list((await session.scalars(select(BidDocumentPage))).all())
        for document in docs:
            assert "private" not in document.upload_name_encrypted
            stored = application.state.storage.path(
                document.org_id, document.storage_key
            ).read_bytes()
            assert not stored.startswith(b"%PDF-") and SENTINEL.encode() not in stored
        assert any(page.page_kind == "image" and page.text_status == "native" for page in pages)
        assert all(SENTINEL not in (page.text_encrypted or "") for page in pages)
        logs = (await session.scalars(select(AuditLog.details))).all()
        assert SENTINEL not in json.dumps(logs)
    assert SENTINEL not in caplog.text
    OUTPUT.mkdir(parents=True, exist_ok=True)
    (OUTPUT / "http-inventory.json").write_text(
        json.dumps({"submission": detail, "job": job}, indent=2) + "\n"
    )


@pytest.mark.parametrize(
    "variant",
    [
        "hash",
        "length",
        "mime",
        "extension",
        "zip",
        "doc",
        "docm",
        "password",
        "encrypted-empty",
        "repaired",
        "many",
        "missing-role",
        "external-docx",
    ],
)
async def test_rejections_store_nothing(variant, api, headers, tenants, application):
    task_id, files = await task(api, headers[0]), originals()
    if variant in {"doc", "docm", "extension"}:
        files[1] = ("private." + (variant if variant != "extension" else "txt"), PDF, files[1][2])
    elif variant == "mime":
        files[1] = (files[1][0], "application/octet-stream", files[1][2])
    elif variant == "zip":
        files[1] = ("fake.docx", DOCX, b"PK\x03\x04not-a-docx")
    elif variant in {"password", "encrypted-empty"}:
        files[1] = (
            "encrypted.pdf",
            PDF,
            pdf_bytes(encrypted=True, empty_password=variant == "encrypted-empty"),
        )
    elif variant == "repaired":
        content = files[1][2]
        content = content[: content.rfind(b"startxref")] + b"startxref\n0\n%%EOF\n"
        with pymupdf.open(stream=content, filetype="pdf") as repaired:
            assert repaired.is_repaired
        files[1] = ("repaired.pdf", PDF, content)
    elif variant == "many":
        files = [files[0], *[files[1]] * 20]
    elif variant == "external-docx":
        output = io.BytesIO()
        with (
            zipfile.ZipFile(io.BytesIO(docx_bytes())) as source,
            zipfile.ZipFile(output, "w") as target,
        ):
            for item in source.infolist():
                content = source.read(item)
                if item.filename == "word/_rels/document.xml.rels":
                    content = content.replace(
                        b"</Relationships>",
                        b'<Relationship Id="bad" Type="external" TargetMode="External" Target="https://example.invalid/x"/></Relationships>',
                    )
                target.writestr(item, content)
        files[1] = ("external.docx", DOCX, output.getvalue())
    body = metadata(files)
    if variant == "hash":
        body["files"][1]["sha256"] = "0" * 64
    elif variant == "length":
        body["files"][1]["size_bytes"] += 1
    elif variant == "mime":
        body["files"][1]["media_type"] = PDF
    elif variant == "missing-role":
        body["files"][1].update(role="tender", kind="tender")
    before = await counts(application, tenants["orgs"][0])
    stored = list((application.state.storage.root / "org").rglob("*"))
    error(await upload(api, headers[0], task_id, files, body))
    assert await counts(application, tenants["orgs"][0]) == before
    assert list((application.state.storage.root / "org").rglob("*")) == stored


async def test_deployment_aggregate_limit_and_replay_conflict(api, headers, tenants, application):
    task_id, files = await task(api, headers[0]), originals()
    body = metadata(files)
    result(await upload(api, headers[0], task_id, files, body))
    changed = [
        ("different-private.pdf", media, content) if index == 0 else (name, media, content)
        for index, (name, media, content) in enumerate(files)
    ]
    error(await upload(api, headers[0], task_id, changed, body), "idempotency_conflict", 409)
    application.state.processor.settings.max_upload_bytes = (
        max(len(content) for _, _, content in files) + 1
    )
    before = await counts(application, tenants["orgs"][0])
    error(await upload(api, headers[0], task_id, files), "bid_upload_limit", 413)
    assert await counts(application, tenants["orgs"][0]) == before


async def test_two_org_task_token_and_generic_original_gates(api, headers, tenants, application):
    task_a, task_b = await task(api, headers[0]), await task(api, headers[1])
    a = result(await upload(api, headers[0], task_a, originals()))
    b = result(await upload(api, headers[1], task_b, originals()))
    error(await api.get(f"/v4/bid-submissions/{b['id']}", headers=headers[0]), status=404)
    error(await api.get(f"/v4/tasks/{task_b}/bid-submissions", headers=headers[0]), status=404)
    error(await upload(api, headers[0], task_b, originals()), status=404)
    error(await preview(api, headers[0], task_a, b["id"]), status=404)
    other_task = await task(api, headers[0])
    error(await preview(api, headers[0], other_task, a["id"]), status=404)
    issue = {
        "name": "Synthetic metadata token",
        "expires_at": "2030-01-01T00:00:00Z",
        "scopes": ["task:read", "bid-review:read", "job:read"],
    }
    issued = result(await api.post("/v4/tokens", headers=headers[0], json=issue))
    token_header = {**headers[0], "Authorization": "Bearer " + issued["token"]}
    result(await api.get(f"/v4/bid-submissions/{a['id']}", headers=token_header))
    for header, org_id, task_id, submission in (
        (headers[0], tenants["orgs"][0], task_a, a),
        (headers[1], tenants["orgs"][1], task_b, b),
    ):
        receipt = result(await preview(api, header, task_id, submission["id"]))
        accepted = result(await submit(api, header, task_id, submission["id"], receipt))
        await run(application, org_id, accepted["job_id"])
        async with application.state.db.transaction(org_id) as session:
            for model in TABLES[:6]:
                assert await session.scalar(select(func.count()).select_from(model)) > 0
    result(await api.get(f"/v4/bid-submissions/{a['id']}", headers=token_header))
    error(await upload(api, token_header, task_a, originals()), status=403)
    error(await preview(api, token_header, task_a, a["id"]), status=403)
    for scope in ("bid-review:upload", "bid-review:prepare", "bid-review:original:read"):
        error(
            await api.post(
                "/v4/tokens",
                headers=headers[0],
                json={**issue, "scopes": ["bid-review:read", scope]},
            ),
            status=403,
        )
    for file in a["files"]:
        for identifier in (file["id"], file["file_id"]):
            for suffix in ("", "/download-link", "/chunks", "/pages/1/preview"):
                response = await api.get(
                    f"/v4/documents/{identifier}{suffix}", headers=token_header
                )
                assert response.status_code in {403, 404}
    async with application.state.db.transaction(tenants["orgs"][0]) as session:
        for model in TABLES[:6]:
            assert (
                await session.scalar(
                    select(func.count())
                    .select_from(model)
                    .where(model.org_id == tenants["orgs"][1])
                )
                == 0
            )
    async with application.state.db.transaction() as session:
        for model in TABLES[:6]:
            assert await session.scalar(select(func.count()).select_from(model)) == 0


async def test_receipt_expiry_mismatch_and_pagination(api, headers, application):
    task_id = await task(api, headers[0])
    a = result(await upload(api, headers[0], task_id, originals()))
    b = result(await upload(api, headers[0], task_id, originals()))
    receipt = result(await preview(api, headers[0], task_id, a["id"]))
    signer = TokenSigner.for_tokens(application.state.processor.settings)
    payload = signer.open(receipt["preflight_token"])
    expired = signer.issue(
        payload, 0, expires_at=int((datetime.now(UTC) - timedelta(seconds=1)).timestamp())
    )
    error(
        await submit(api, headers[0], task_id, a["id"], receipt, preflight_token=expired),
        "bid_preflight_expired",
    )
    error(
        await submit(api, headers[0], task_id, a["id"], receipt, expected_input_hash="0" * 64),
        "bid_input_changed",
    )
    payload["actor_user_id"] = str(uuid4())
    mismatched = signer.issue(payload, 900)
    error(
        await submit(api, headers[0], task_id, a["id"], receipt, preflight_token=mismatched),
        "bid_preflight_mismatch",
    )
    error(await submit(api, headers[0], task_id, b["id"], receipt), "bid_input_changed")
    page = await api.get(f"/v4/tasks/{task_id}/bid-submissions?limit=1", headers=headers[0])
    assert result(page)["total"] == 2
    cursor = page.json()["data"]["next_cursor"]
    next_page = await api.get(
        f"/v4/tasks/{task_id}/bid-submissions",
        headers=headers[0],
        params={"limit": 1, "cursor": cursor},
    )
    assert result(next_page)["next_cursor"] is None
    assert next_page.json()["items"][0]["id"] != page.json()["items"][0]["id"]
    error(
        await api.get(f"/v4/tasks/{task_id}/bid-submissions?limit=101", headers=headers[0]),
        status=422,
    )
    error(await api.get(f"/tasks/{task_id}/bid-submissions", headers=headers[0]), status=404)


@pytest.mark.parametrize("conversion_fails", [False, True])
async def test_docx_conversion_atomic_inventory(
    conversion_fails, api, headers, tenants, application
):
    settings = application.state.processor.settings
    settings.converter_url = "http://converter.invalid"
    settings.review_office_profile = "synthetic-converter-fonts-v1"
    calls = []

    def convert(request):
        calls.append(request)
        return httpx.Response(
            422 if conversion_fails else 200, content=b"" if conversion_fails else pdf_bytes()
        )

    application.state.processor.converter_transport = httpx.MockTransport(convert)
    task_id = await task(api, headers[0])
    files = [("tender.pdf", PDF, pdf_bytes()), ("bid.docx", DOCX, docx_bytes())]
    root = result(await upload(api, headers[0], task_id, files))
    receipt = result(await preview(api, headers[0], task_id, root["id"]))
    assert calls == []
    accepted = result(await submit(api, headers[0], task_id, root["id"], receipt))
    await run(application, tenants["orgs"][0], accepted["job_id"])
    detail = result(await api.get(f"/v4/bid-submissions/{root['id']}", headers=headers[0]))
    assert len(calls) == 1
    if conversion_fails:
        assert detail["preparation"]["status"] == "failed" and detail["inventory"] == []
        assert detail["submission"]["state"] == "uploaded"
    else:
        assert detail["preparation"]["status"] == "succeeded"
        word = next(
            document
            for document in detail["submission"]["documents"]
            if document["media_type"] == DOCX
        )
        assert word["citation_mode"] == "block" and word["sha256"] != word["rendered_pdf_sha256"]
        assert "docx_page_map_unknown" in word["parsing_warnings"]
    async with application.state.db.transaction(tenants["orgs"][0]) as session:
        for model in (BidPreparedDocument, BidDocumentPage, BidPreparationPublication):
            count = await session.scalar(select(func.count()).select_from(model))
            assert (count == 0) == conversion_fails


async def test_cancel_mid_preparation_no_publication(
    api, headers, tenants, application, monkeypatch
):
    task_id = await task(api, headers[0])
    root = result(await upload(api, headers[0], task_id, originals()))
    receipt = result(await preview(api, headers[0], task_id, root["id"]))
    accepted = result(await submit(api, headers[0], task_id, root["id"], receipt))
    original, reached, release = bid_preparation.pdf_pages, asyncio.Event(), asyncio.Event()

    async def paused(*args, **kwargs):
        async for page in original(*args, **kwargs):
            reached.set()
            await release.wait()
            yield page

    monkeypatch.setattr(bid_preparation, "pdf_pages", paused)
    worker = asyncio.create_task(run(application, tenants["orgs"][0], accepted["job_id"]))
    try:
        await asyncio.wait_for(reached.wait(), 30)
        result(await api.post(f"/v4/jobs/{accepted['job_id']}/cancel", headers=headers[0]))
    finally:
        release.set()
        await worker
    detail = result(await api.get(f"/v4/bid-submissions/{root['id']}", headers=headers[0]))
    assert detail["preparation"]["status"] == "cancelled" and detail["inventory"] == []
    async with application.state.db.transaction(tenants["orgs"][0]) as session:
        assert await session.scalar(select(func.count()).select_from(BidDocumentPage)) == 0


async def test_submitted_versions_immutable(api, headers, tenants, application):
    task_id = await task(api, headers[0])
    root = result(await upload(api, headers[0], task_id, originals()))
    receipt = result(await preview(api, headers[0], task_id, root["id"]))
    accepted = result(await submit(api, headers[0], task_id, root["id"], receipt))
    await run(application, tenants["orgs"][0], accepted["job_id"])
    for table, field in (
        ("bid_submissions", "revision=revision+1"),
        ("bid_submission_documents", "size_bytes=size_bytes+1"),
        ("bid_document_pages", "page=page+1"),
        ("bid_prepared_documents", "page_count=page_count+1"),
        ("bid_preparation_publications", "page_count=page_count+1"),
    ):
        with pytest.raises(DBAPIError):
            async with application.state.db.transaction(tenants["orgs"][0]) as session:
                await session.execute(text(f"UPDATE {table} SET {field}"))
        with pytest.raises(DBAPIError):
            async with application.state.db.transaction(tenants["orgs"][0]) as session:
                await session.execute(text(f"DELETE FROM {table}"))
    result(await api.get(f"/v4/bid-submissions/{root['id']}", headers=headers[0]))


async def test_composite_parent_and_database_token_constraints(api, headers, tenants, application):
    from app.models.entities import ApiToken
    from app.services.auth import ROLE_SCOPES, Identity, set_actor_context

    tasks = [await task(api, header) for header in headers]
    prepared = []
    for header, task_id, org_id in zip(headers, tasks, tenants["orgs"], strict=True):
        root = result(await upload(api, header, task_id, originals()))
        receipt = result(await preview(api, header, task_id, root["id"]))
        accepted = result(await submit(api, header, task_id, root["id"], receipt))
        await run(application, org_id, accepted["job_id"])
        prepared.append(root)
    actor = Identity(tenants["users"][0], tenants["orgs"][0], set(ROLE_SCOPES["admin"]), "admin")
    # Copy B values through B's own tenant transaction, then attempt to bind those
    # real parents under A. Every A insert has a legitimate authenticated actor.
    foreign = {}
    async with application.state.db.transaction(tenants["orgs"][1]) as session:
        for model in TABLES[:6]:
            row = await session.scalar(select(model).limit(1))
            assert row is not None
            foreign[model] = {
                column.name: getattr(row, column.name) for column in model.__table__.columns
            }
    for model, source in foreign.items():
        values = {**source, "id": uuid4(), "org_id": tenants["orgs"][0]}
        for name in ("created_by", "published_by"):
            if name in values:
                values[name] = actor.user_id
        if model is BidSubmission:
            values.update(revision=1000, request_id=uuid4())
        for name in ("storage_key", "rendered_pdf_storage_key"):
            if name in values:
                values[name] = values[name].replace(
                    f"org/{tenants['orgs'][1]}/", f"org/{actor.org_id}/", 1
                )
        with pytest.raises(DBAPIError) as rejected:
            async with application.state.db.transaction(actor.org_id) as session:
                await set_actor_context(session, actor)
                await session.execute(model.__table__.insert().values(**values))
        assert rejected.value.orig.sqlstate == "23503", (
            model.__tablename__,
            rejected.value.orig.sqlstate,
        )
    # A legitimate same-org task is still not the source document's task.
    other_task = UUID(await task(api, headers[0]))
    async with application.state.db.transaction(actor.org_id) as session:
        row = await session.scalar(select(BidPreparedDocument).limit(1))
        values = {
            column.name: getattr(row, column.name)
            for column in BidPreparedDocument.__table__.columns
        }
    values.update(id=uuid4(), task_id=other_task)
    with pytest.raises(DBAPIError) as rejected:
        async with application.state.db.transaction(actor.org_id) as session:
            await set_actor_context(session, actor)
            await session.execute(BidPreparedDocument.__table__.insert().values(**values))
    assert rejected.value.orig.sqlstate == "23503"
    with pytest.raises(DBAPIError) as rejected:
        async with application.state.db.transaction(actor.org_id) as session:
            await set_actor_context(session, actor)
            session.add(
                ApiToken(
                    id=uuid4(),
                    org_id=actor.org_id,
                    user_id=actor.user_id,
                    name="Synthetic forbidden grant",
                    digest="f" * 64,
                    scopes=["bid-review:read", "bid-review:prepare"],
                    expires_at=datetime.now(UTC) + timedelta(days=1),
                )
            )
            await session.flush()
    assert rejected.value.orig.sqlstate == "23514"
    assert rejected.value.orig.diag.constraint_name == "token_forbidden_bid_review_scopes"


async def test_storage_failure_keeps_inventory_unpublished(
    api, headers, tenants, application, monkeypatch
):
    task_id = await task(api, headers[0])
    root = result(await upload(api, headers[0], task_id, originals()))
    receipt = result(await preview(api, headers[0], task_id, root["id"]))
    accepted = result(await submit(api, headers[0], task_id, root["id"], receipt))
    original_put, written = application.state.storage.put, 0

    async def fail_second_page(org_id, key, content):
        nonlocal written
        if "/pages/" in key:
            written += 1
            if written == 2:
                raise ServiceError("synthetic_storage_refusal", "Synthetic storage refusal", 500, 4)
        await original_put(org_id, key, content)

    monkeypatch.setattr(application.state.storage, "put", fail_second_page)
    await run(application, tenants["orgs"][0], accepted["job_id"])
    detail = result(await api.get(f"/v4/bid-submissions/{root['id']}", headers=headers[0]))
    assert written == 2
    assert detail["preparation"]["status"] == "failed" and detail["inventory"] == []
    async with application.state.db.transaction(tenants["orgs"][0]) as session:
        for model in (BidPreparedDocument, BidDocumentPage, BidPreparationPublication):
            assert await session.scalar(select(func.count()).select_from(model)) == 0


async def test_combined_thousand_page_limit_has_no_partial_publication(
    api, headers, tenants, application
):
    with pymupdf.open() as pdf:
        for _ in range(1000):
            pdf.new_page(width=20, height=20)
        thousand = pdf.tobytes()
    task_id = await task(api, headers[0])
    files = [("tender.pdf", PDF, pdf_bytes()), ("thousand-pages.pdf", PDF, thousand)]
    root = result(await upload(api, headers[0], task_id, files))
    receipt = result(await preview(api, headers[0], task_id, root["id"]))
    accepted = result(await submit(api, headers[0], task_id, root["id"], receipt))
    # The first document renders once; the second exceeds the remaining 999-page
    # bound before any of its pages are rendered. Each original alone is valid.
    await run(application, tenants["orgs"][0], accepted["job_id"])
    detail = result(await api.get(f"/v4/bid-submissions/{root['id']}", headers=headers[0]))
    assert detail["preparation"]["status"] == "failed" and detail["inventory"] == []
    async with application.state.db.transaction(tenants["orgs"][0]) as session:
        for model in (BidPreparedDocument, BidDocumentPage, BidPreparationPublication):
            assert await session.scalar(select(func.count()).select_from(model)) == 0
