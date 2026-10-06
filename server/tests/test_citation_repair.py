"""Citation repair through extraction, review cards, drafts, and audit.

Failure modes are fixed before the tests, per repository policy:

* extraction must save the unique, contiguous source span instead of the
  normalized model spelling, while keeping the model spelling separately;
* a normalized quote with more than one possible source span must be rejected;
* preview must be read-only, deterministic, and bind requirements, chunks, and
  response-card revisions so a stale preview cannot execute;
* only a current human organization administrator may execute repair: API
  tokens, agents, other roles, and foreign organizations must fail closed;
* unlocatable and ambiguous legacy rows must remain byte-for-byte unchanged;
* repair may update only quote provenance, never requirement text, category,
  location, extraction job, fingerprint, or any other protected field;
* every changed quote or newly preserved model quote needs one audit event, and
  audit details must contain identifiers and hashes rather than tender text or
  the operator's reason;
* a card based on a changed quote must retain its state and immutable history but
  become needs_reconfirmation, including a comply-only decision and draft gap;
* a confirmed card must be reopened, a pending card withdrawn, and both edited
  and reviewed again; comply-only must be decided again by a human reviewer;
* an unlocatable citation must still fail the strict confirmation gate after a
  repair run.
"""

from __future__ import annotations

import hashlib
import io
import json
import re
from contextlib import asynccontextmanager
from datetime import UTC, datetime, timedelta
from pathlib import Path
from uuid import UUID

import httpx
import pymupdf
import pytest
from app.core.config import Settings
from app.providers.llm import AnthropicExtractor
from conftest import PASSWORD, FakeQueue, credential_app
from docx import Document
from sqlalchemy import text
from task_fixtures import confirm_requirements_async
from test_llm_providers import Vendor, anthropic_reply

SYNTHETIC_KEY = "synthetic-citation-repair-key-not-real"
WORD_SOURCE_QUOTE = "单卡显存不低于 80 GB"
WORD_MODEL_QUOTE = "单卡显存不低于80GB"
SOURCE_QUOTES = {
    # A line break survives PDF text extraction; runs of spaces are rebuilt from layout.
    1: "Unique\nnormalized citation.",
    2: "Confirmed card citation.",
    3: "Pending card citation.",
    4: "Comply only citation.",
    5: "Unlocatable legacy citation.",
    6: "Ambiguous target citation.",
    7: "Already exact unchanged citation.",
    8: "Duplicate extraction citation. Duplicate extraction citation.",
}
MODEL_QUOTES = {
    1: "Unique normalized citation.",
    2: SOURCE_QUOTES[2],
    3: SOURCE_QUOTES[3],
    4: SOURCE_QUOTES[4],
    5: SOURCE_QUOTES[5],
    6: SOURCE_QUOTES[6],
    7: SOURCE_QUOTES[7],
    8: "Duplicate extraction citation.",
}
REPAIRED_QUOTES = {
    2: "Confirmed   card citation.",
    3: "Pending   card citation.",
    4: "Comply   only citation.",
}
PROTECTED_REQUIREMENT_FIELDS = {
    "task_id",
    "document_id",
    "chunk_id",
    "page",
    "location",
    "text",
    "category",
    "starred",
    "condition",
    "fingerprint",
    "job_id",
}


def labelled_pdf() -> bytes:
    """Create public-test-style material that cannot be mistaken for a real tender."""
    with pymupdf.open() as document:
        for page_number, quote in SOURCE_QUOTES.items():
            page = document.new_page()
            page.insert_text((40, 60), "SYNTHETIC TEST MATERIAL - NOT A REAL TENDER")
            page.insert_text((40, 90), quote)
            assert page_number == len(document)
        return document.tobytes()


def extraction_items() -> list[dict]:
    return [
        {
            "category": "technical",
            "starred": False,
            "text": f"Synthetic requirement {page}",
            "ref": str(page),
            "quote": MODEL_QUOTES[page],
            "condition": None,
        }
        for page in SOURCE_QUOTES
    ]


def acceptance_word() -> bytes:
    document = Document()
    document.add_heading("合成技术要求", level=1)  # p1
    document.add_paragraph(WORD_SOURCE_QUOTE)  # p2
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def acceptance_item() -> dict:
    return {
        "category": "technical",
        "starred": False,
        "text": WORD_MODEL_QUOTE,
        "ref": "p2",
        "quote": WORD_MODEL_QUOTE,
        "condition": {"param": "单卡显存", "op": ">=", "value": "80", "unit": "GB"},
    }


def settings_for(tmp_path: Path) -> Settings:
    return Settings(
        data_dir=tmp_path,
        llm_provider="anthropic",
        llm_api_key=SYNTHETIC_KEY,
        llm_model=None,
        llm_concurrency=1,
    )


async def login(api: httpx.AsyncClient, org_id: UUID, label: str) -> dict[str, str]:
    response = await api.post(
        "/auth/login",
        json={"email": f"{label}@example.test", "password": PASSWORD, "org_id": str(org_id)},
    )
    assert response.status_code == 200, response.text
    return {
        "Authorization": "Bearer " + response.json()["data"]["session"],
        "X-Org-Id": str(org_id),
    }


@asynccontextmanager
async def repair_client(tenants, tmp_path: Path, *, vendor_replies: int = 1):
    vendor = Vendor(*(anthropic_reply(extraction_items()) for _ in range(vendor_replies)))
    settings = settings_for(tmp_path)
    app = await credential_app(
        settings,
        llm=AnthropicExtractor(settings, transport=vendor.transport()),
        queue=FakeQueue(),
    )
    async with (
        app.router.lifespan_context(app),
        httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test") as api,
    ):
        headers = [
            await login(api, org, label)
            for org, label in zip(tenants["orgs"], ("a", "b"), strict=True)
        ]
        yield api, app, headers, vendor


async def run_document_job(api, app, header, document_id: str, action: str) -> tuple[str, dict]:
    response = await api.post(
        f"/documents/{document_id}/{action}", headers=header, json={"dry_run": False}
    )
    assert response.status_code == 200, response.text
    job_id = response.json()["data"]["job_id"]
    await app.state.processor(header["X-Org-Id"], job_id)
    status = (await api.get(f"/jobs/{job_id}", headers=header)).json()["data"]
    return job_id, status


async def extract_fixture(api, app, header, tmp_path: Path, suffix: str):
    task_response = await api.post(
        "/tasks", headers=header, json={"name": f"Synthetic citation repair {suffix}"}
    )
    assert task_response.status_code == 200, task_response.text
    task_id = task_response.json()["data"]["id"]
    path = tmp_path / f"synthetic-citation-repair-{suffix}.pdf"
    path.write_bytes(labelled_pdf())
    upload = await api.post(
        f"/tasks/{task_id}/documents",
        headers=header,
        files={"file": (path.name, path.read_bytes(), "application/pdf")},
    )
    assert upload.status_code == 200, upload.text
    document_id = upload.json()["data"]["id"]
    _, parse_status = await run_document_job(api, app, header, document_id, "parse")
    assert parse_status["status"] == "succeeded", parse_status
    extraction_job_id, extraction_status = await run_document_job(
        api, app, header, document_id, "extract"
    )
    assert extraction_status["status"] == "succeeded", extraction_status
    response = await api.get(
        f"/tasks/{task_id}/requirements",
        headers=header,
        params={"job": extraction_job_id},
    )
    assert response.status_code == 200, response.text
    rows = response.json()["items"]
    return task_id, document_id, extraction_job_id, extraction_status, rows


def set_role(admin_engine, org_id: UUID, user_id: UUID, role: str) -> None:
    from task_fixtures import set_role as set_task_actor_role

    set_task_actor_role(admin_engine, org_id, user_id, role)


async def create_commitment_card(api, header, task_id, job_id, requirement, label):
    response = await api.post(
        f"/tasks/{task_id}/cards",
        headers=header,
        json={
            "extraction_job_id": job_id,
            "requirement_id": requirement["id"],
            "content": {
                "response_kind": "commitment",
                "response_text": f"We commit to the synthetic {label} requirement.",
                "deviation": "none",
                "deviation_note": f"The synthetic {label} response matches the requirement.",
                "evidence": [],
            },
        },
    )
    assert response.status_code == 200, response.text
    return response.json()["data"]


async def card_action(api, header, card, action: str, **extra):
    return await api.post(
        f"/cards/{card['id']}/actions",
        headers=header,
        json={"expected_revision": card["revision"], "action": action, **extra},
    )


async def require_card_action(api, header, card, action: str, **extra):
    response = await card_action(api, header, card, action, **extra)
    assert response.status_code == 200, response.text
    return response.json()["data"]


async def update_card_without_changing_business_text(api, header, card):
    response = await api.put(
        f"/cards/{card['id']}",
        headers=header,
        json={
            "expected_revision": card["revision"],
            "content": card["content"],
        },
    )
    assert response.status_code == 200, response.text
    return response.json()["data"]


def requirement_rows(admin_engine, task_id: str) -> dict[int, dict]:
    with admin_engine.connect() as connection:
        rows = connection.execute(
            text(
                "SELECT id::text, task_id::text, document_id::text, chunk_id::text, page, "
                "location, quote, model_quote, text, category, starred, condition, "
                "fingerprint, job_id::text FROM requirements WHERE task_id = :task "
                "ORDER BY page"
            ),
            {"task": task_id},
        ).mappings()
        return {row["page"]: dict(row) for row in rows}


def seed_legacy_corruption(admin_engine, requirements: dict[int, dict]) -> None:
    """Simulate rows accepted by the pre-exact-span implementation."""
    with admin_engine.begin() as connection:
        connection.execute(
            text("UPDATE requirements SET model_quote = NULL WHERE id = :id"),
            {"id": requirements[1]["id"]},
        )
        for page, quote in REPAIRED_QUOTES.items():
            connection.execute(
                text("UPDATE chunks SET text = :source WHERE id = :chunk"),
                {"source": quote, "chunk": requirements[page]["source"]["chunk_id"]},
            )
        connection.execute(
            text("UPDATE requirements SET quote = :quote, model_quote = NULL WHERE id = :id"),
            {"quote": "Missing legacy quote.", "id": requirements[5]["id"]},
        )
        connection.execute(
            text("UPDATE chunks SET text = :source WHERE id = :chunk"),
            {
                "source": f"{SOURCE_QUOTES[6]} / {SOURCE_QUOTES[6]}",
                "chunk": requirements[6]["source"]["chunk_id"],
            },
        )
        connection.execute(
            text("UPDATE requirements SET quote = :quote, model_quote = NULL WHERE id = :id"),
            {"quote": "Ambiguous  target citation.", "id": requirements[6]["id"]},
        )


def stable_hashes(value):
    """Hash UUID-shaped access handles before writing the reproducible artifact."""
    if isinstance(value, dict):
        return {stable_hashes(key): stable_hashes(item) for key, item in value.items()}
    if isinstance(value, list):
        return [stable_hashes(item) for item in value]
    if isinstance(value, str):
        try:
            UUID(value)
        except ValueError:
            return value
        return "uuid-sha256:" + hashlib.sha256(value.encode()).hexdigest()[:16]
    return value


def assert_hashes_and_ids_only(value) -> None:
    """Audit scalar strings may be UUIDs or SHA-256 values, never business prose."""
    if isinstance(value, dict):
        for item in value.values():
            assert_hashes_and_ids_only(item)
        return
    if isinstance(value, list):
        for item in value:
            assert_hashes_and_ids_only(item)
        return
    if isinstance(value, str):
        if value == "session":
            return
        try:
            UUID(value)
            return
        except ValueError:
            assert re.fullmatch(r"[0-9a-f]{64}", value), value


async def test_extraction_persists_exact_span_and_rejects_ambiguous_match(
    tenants, tmp_path, admin_engine
):
    async with repair_client(tenants, tmp_path) as (api, app, headers, vendor):
        task_id, _, job_id, status, rows = await extract_fixture(
            api, app, headers[0], tmp_path, "exact-span"
        )

    assert status["result"]["created"] == 7
    assert status["result"]["rejected"] == [
        {
            "position": "第 8 页",
            "quote": MODEL_QUOTES[8],
            "reason": "ambiguous_quote",
        }
    ]
    assert [row["source"]["page"] for row in rows] == [1, 2, 3, 4, 5, 6, 7]
    first = rows[0]
    assert first["source"]["quote"] == SOURCE_QUOTES[1]
    assert first["model_quote"] == MODEL_QUOTES[1]
    assert first["source"]["quote"] != first["model_quote"]
    assert task_id and job_id
    assert len(vendor.requests) == 1

    # Acceptance sample: the real Anthropic adapter returns normalized Chinese
    # Word text, while persistence and confirmation use the exact paragraph span.
    word_vendor = Vendor(anthropic_reply([acceptance_item()]))
    word_settings = settings_for(tmp_path)
    word_app = await credential_app(
        word_settings,
        llm=AnthropicExtractor(word_settings, transport=word_vendor.transport()),
        queue=FakeQueue(),
    )
    async with (
        word_app.router.lifespan_context(word_app),
        httpx.AsyncClient(
            transport=httpx.ASGITransport(app=word_app), base_url="http://test"
        ) as word_api,
    ):
        header = await login(word_api, tenants["orgs"][0], "a")
        word_task = (
            await word_api.post(
                "/tasks", headers=header, json={"name": "Synthetic Chinese exact citation"}
            )
        ).json()["data"]["id"]
        upload = await word_api.post(
            f"/tasks/{word_task}/documents",
            headers=header,
            files={"file": ("synthetic-exact-citation.docx", acceptance_word())},
        )
        assert upload.status_code == 200, upload.text
        document_id = upload.json()["data"]["id"]
        _, parsed = await run_document_job(word_api, word_app, header, document_id, "parse")
        assert parsed["status"] == "succeeded", parsed
        word_job, extracted = await run_document_job(
            word_api, word_app, header, document_id, "extract"
        )
        assert extracted["status"] == "succeeded", extracted
        word_rows = (
            await word_api.get(
                f"/tasks/{word_task}/requirements", headers=header, params={"job": word_job}
            )
        ).json()["items"]
        assert len(word_rows) == 1
        [word_requirement] = word_rows
        assert word_requirement["source"]["location"]["block_id"] == "p2"
        assert word_requirement["source"]["quote"] == WORD_SOURCE_QUOTE
        assert word_requirement["model_quote"] == WORD_MODEL_QUOTE

        async with word_app.state.db.transaction(tenants["orgs"][0]) as session:
            await confirm_requirements_async(
                session,
                tenants["orgs"][0],
                UUID(word_task),
                settings=word_app.state.processor.settings,
            )
        word_card = await create_commitment_card(
            word_api,
            header,
            word_task,
            word_job,
            word_requirement,
            "Chinese exact citation",
        )
        set_role(admin_engine, tenants["orgs"][0], tenants["users"][0], "technical")
        word_card = await require_card_action(word_api, header, word_card, "submit")
        word_card = await require_card_action(
            word_api, header, word_card, "confirm", reviewed_evidence_ids=[]
        )
        assert word_card["state"] == "confirmed"
        assert word_card["eligibility"] == "eligible"
        assert word_card["source"]["quote"] == WORD_SOURCE_QUOTE
        assert len(word_vendor.requests) == 1


async def test_preview_execute_permissions_cards_drafts_and_audit(tenants, tmp_path, admin_engine):
    from task_fixtures import reviewer_header

    reason = "Synthetic operator reason that must not enter the audit log."
    async with repair_client(tenants, tmp_path, vendor_replies=2) as (api, app, headers, _):
        header, foreign_header = headers
        task_id, _, job_id, status, listed = await extract_fixture(
            api, app, header, tmp_path, "workflow"
        )
        assert status["result"]["created"] == 7
        requirements = {row["source"]["page"]: row for row in listed}
        async with app.state.db.transaction(tenants["orgs"][0]) as session:
            await confirm_requirements_async(
                session, tenants["orgs"][0], UUID(task_id), settings=app.state.processor.settings
            )

        foreign_task_id, _, foreign_job_id, _, _ = await extract_fixture(
            api, app, foreign_header, tmp_path, "foreign"
        )

        # Cards are created while all citations are still exact. That lets the
        # test prove repair stales already-reviewed decisions without rewriting
        # their current state or immutable revision history.
        normalized = await create_commitment_card(
            api, header, task_id, job_id, requirements[1], "normalized source"
        )
        confirmed = await create_commitment_card(
            api, header, task_id, job_id, requirements[2], "confirmed"
        )
        pending = await create_commitment_card(
            api, header, task_id, job_id, requirements[3], "pending"
        )
        unlocatable = await create_commitment_card(
            api, header, task_id, job_id, requirements[5], "unlocatable"
        )
        hash_bound = await create_commitment_card(
            api, header, task_id, job_id, requirements[7], "preview hash"
        )
        _, reviewer = await reviewer_header(
            api, admin_engine, tenants["orgs"][0], UUID(task_id), "technical"
        )
        normalized = await require_card_action(api, reviewer, normalized, "submit")
        normalized = await require_card_action(
            api, reviewer, normalized, "confirm", reviewed_evidence_ids=[]
        )
        assert normalized["source"]["quote"] == SOURCE_QUOTES[1]
        assert normalized["state"] == "confirmed" and normalized["eligibility"] == "eligible"
        confirmed = await require_card_action(api, reviewer, confirmed, "submit")
        confirmed = await require_card_action(
            api, reviewer, confirmed, "confirm", reviewed_evidence_ids=[]
        )
        pending = await require_card_action(api, reviewer, pending, "submit")
        unlocatable = await require_card_action(api, reviewer, unlocatable, "submit")
        disposition = await api.post(
            f"/tasks/{task_id}/cards/dispositions",
            headers=reviewer,
            json={
                "extraction_job_id": job_id,
                "items": [
                    {
                        "requirement_id": requirements[4]["id"],
                        "expected_revision": None,
                        "disposition": "comply_only",
                        "reason": "This synthetic clause only needs compliance.",
                    }
                ],
            },
        )
        assert disposition.status_code == 200, disposition.text
        [comply] = disposition.json()["data"]["cards"]

        pre_repair_draft = await api.post(
            f"/tasks/{task_id}/drafts",
            headers=header,
            json={"extraction_job_id": job_id},
        )
        assert pre_repair_draft.status_code == 200, pre_repair_draft.text
        pre_repair_job_id = pre_repair_draft.json()["data"]["job_id"]
        await app.state.processor(header["X-Org-Id"], pre_repair_job_id)
        pre_repair_job = (await api.get(f"/jobs/{pre_repair_job_id}", headers=header)).json()[
            "data"
        ]
        assert pre_repair_job["status"] == "succeeded"
        pre_repair_draft_id = pre_repair_job["result"]["draft_id"]
        assert (await api.get(f"/drafts/{pre_repair_draft_id}", headers=header)).json()["data"][
            "validity"
        ] == "current"

        card_before = {
            card["requirement_id"]: {
                "id": card["id"],
                "revision": card["revision"],
                "revision_id": card["revision_id"],
                "state": card["state"],
                "content": card["content"],
            }
            for card in (confirmed, pending, comply, unlocatable)
        }

        seed_legacy_corruption(admin_engine, requirements)
        before = requirement_rows(admin_engine, task_id)
        with admin_engine.connect() as connection:
            audit_before = connection.scalar(
                text("SELECT count(*) FROM audit_logs WHERE action = 'requirement.repair_citation'")
            )
            revision_before = connection.scalar(
                text("SELECT count(*) FROM response_card_revisions")
            )

        set_role(admin_engine, tenants["orgs"][0], tenants["users"][0], "admin")
        preview_response = await api.get(
            f"/tasks/{task_id}/requirements/repair", headers=header, params={"job": job_id}
        )
        assert preview_response.status_code == 200, preview_response.text
        preview_result = preview_response.json()
        preview = preview_result["data"]
        assert preview_result["command"] == "req repair-citations"
        assert preview == {
            "task_id": task_id,
            "extraction_job_id": job_id,
            "preview_hash": preview["preview_hash"],
            "execute": False,
            "changed": 0,
            "repairable": 4,
            "unlocatable": 2,
            "unchanged": 1,
        }
        assert re.fullmatch(r"[0-9a-f]{64}", preview["preview_hash"])
        items = {item["source"]["page"]: item for item in preview_result["items"]}
        assert items[1]["status"] == "repairable"
        assert items[1]["source"]["quote"] == SOURCE_QUOTES[1]
        assert items[1]["model_quote"] is None
        assert items[1]["proposed_quote"] == SOURCE_QUOTES[1]
        assert items[1]["quote_changed"] is False
        for page in REPAIRED_QUOTES:
            assert items[page]["status"] == "repairable"
            assert items[page]["model_quote"] == SOURCE_QUOTES[page]
            assert items[page]["proposed_quote"] == REPAIRED_QUOTES[page]
            assert items[page]["quote_changed"] is True and items[page]["reason"] is None
        assert items[5]["status"] == "unlocatable"
        assert items[5]["reason"] == "quote_not_at_position"
        assert items[5]["proposed_quote"] is None and items[5]["quote_changed"] is False
        assert items[6]["status"] == "unlocatable"
        assert items[6]["reason"] == "ambiguous_quote"
        assert items[7]["status"] == "unchanged"

        # GET is deterministic and has no database side effects.
        repeated_preview = await api.get(
            f"/tasks/{task_id}/requirements/repair", headers=header, params={"job": job_id}
        )
        # Everything except the request timing repeats exactly.
        assert {**repeated_preview.json(), "duration_ms": 0} == {
            **preview_response.json(),
            "duration_ms": 0,
        }
        assert requirement_rows(admin_engine, task_id) == before
        with admin_engine.connect() as connection:
            assert (
                connection.scalar(
                    text(
                        "SELECT count(*) FROM audit_logs "
                        "WHERE action = 'requirement.repair_citation'"
                    )
                )
                == audit_before
            )
            assert (
                connection.scalar(text("SELECT count(*) FROM response_card_revisions"))
                == revision_before
            )

        expiry = (datetime.now(UTC) + timedelta(hours=1)).isoformat()
        token_record = (
            await api.post(
                "/tokens",
                headers=header,
                json={
                    "name": "Synthetic repair reader",
                    "scopes": ["task:read"],
                    "expires_at": expiry,
                },
            )
        ).json()["data"]
        token_header = {**header, "Authorization": "Bearer " + token_record["token"]}
        execute_body = {
            "extraction_job_id": job_id,
            "expected_preview": preview["preview_hash"],
            "reason": reason,
        }
        blank_reason = await api.post(
            f"/tasks/{task_id}/requirements/repair",
            headers=header,
            json={**execute_body, "reason": "   "},
        )
        assert blank_reason.status_code == 400
        assert blank_reason.json()["data"]["error"]["code"] == "invalid_repair_reason"
        token_attempt = await api.post(
            f"/tasks/{task_id}/requirements/repair", headers=token_header, json=execute_body
        )
        assert token_attempt.status_code == 403
        assert token_attempt.json()["data"]["error"]["code"] == "forbidden"
        assert (
            await api.get(
                f"/tasks/{task_id}/requirements/repair",
                headers=token_header,
                params={"job": job_id},
            )
        ).status_code == 403

        set_role(admin_engine, tenants["orgs"][0], tenants["users"][0], "technical")
        role_attempt = await api.post(
            f"/tasks/{task_id}/requirements/repair", headers=header, json=execute_body
        )
        assert role_attempt.status_code == 403
        assert role_attempt.json()["data"]["error"]["code"] == "forbidden"
        assert (
            await api.get(
                f"/tasks/{task_id}/requirements/repair",
                headers=header,
                params={"job": job_id},
            )
        ).status_code == 403
        set_role(admin_engine, tenants["orgs"][0], tenants["users"][0], "admin")

        for method, body in (("GET", None), ("POST", execute_body)):
            foreign = await api.request(
                method,
                f"/tasks/{task_id}/requirements/repair",
                headers=foreign_header,
                params={"job": job_id} if method == "GET" else None,
                json=body,
            )
            assert foreign.status_code == 404
            assert foreign.json()["data"]["error"]["code"] == "not_found"

        for task, job in ((task_id, foreign_job_id), (foreign_task_id, foreign_job_id)):
            foreign_scope = await api.get(
                f"/tasks/{task}/requirements/repair", headers=header, params={"job": job}
            )
            assert foreign_scope.status_code == 404
            assert foreign_scope.json()["data"]["error"]["code"] == "not_found"

        from app.core.errors import ServiceError
        from app.schemas.citation_repair_contracts import CitationRepairRequest
        from app.services import citation_repair
        from app.services.auth import ROLE_SCOPES, Identity

        agent = Identity(
            user_id=tenants["users"][0],
            org_id=tenants["orgs"][0],
            scopes=set(ROLE_SCOPES["admin"]),
            role="admin",
            actor_kind="agent",
        )
        with pytest.raises(ServiceError) as agent_error:
            async with app.state.db.transaction(tenants["orgs"][0]) as session:
                await citation_repair.repair_citations(
                    session,
                    agent,
                    UUID(task_id),
                    UUID(job_id),
                    CitationRepairRequest.model_validate(execute_body),
                )
        assert agent_error.value.code == "forbidden"

        # The preview hash covers protected requirement fields. A direct legacy
        # data correction between review and execution must force a fresh review.
        with admin_engine.begin() as connection:
            connection.execute(
                text("UPDATE requirements SET text = text || ' (legacy note)' WHERE id = :id"),
                {"id": requirements[1]["id"]},
            )
        set_role(admin_engine, tenants["orgs"][0], tenants["users"][0], "admin")
        stale = await api.post(
            f"/tasks/{task_id}/requirements/repair", headers=header, json=execute_body
        )
        assert stale.status_code == 409
        assert stale.json()["data"]["error"]["code"] == "repair_preview_changed"
        assert all(
            requirement_rows(admin_engine, task_id)[page]["quote"] == before[page]["quote"]
            for page in before
        )

        chunk_bound_preview = (
            await api.get(
                f"/tasks/{task_id}/requirements/repair", headers=header, params={"job": job_id}
            )
        ).json()["data"]
        with admin_engine.begin() as connection:
            connection.execute(
                text("UPDATE chunks SET text = text || E'\\n' WHERE id = :id"),
                {"id": requirements[7]["source"]["chunk_id"]},
            )
        chunk_stale = await api.post(
            f"/tasks/{task_id}/requirements/repair",
            headers=header,
            json={**execute_body, "expected_preview": chunk_bound_preview["preview_hash"]},
        )
        assert chunk_stale.status_code == 409
        assert chunk_stale.json()["data"]["error"]["code"] == "repair_preview_changed"

        card_bound_preview = (
            await api.get(
                f"/tasks/{task_id}/requirements/repair", headers=header, params={"job": job_id}
            )
        ).json()["data"]
        hash_bound = await update_card_without_changing_business_text(api, header, hash_bound)
        card_stale = await api.post(
            f"/tasks/{task_id}/requirements/repair",
            headers=header,
            json={**execute_body, "expected_preview": card_bound_preview["preview_hash"]},
        )
        assert card_stale.status_code == 409
        assert card_stale.json()["data"]["error"]["code"] == "repair_preview_changed"

        # The final protected-field baseline includes the intentional legacy note
        # and chunk/card changes above. Citation repair itself must preserve it.
        before = requirement_rows(admin_engine, task_id)
        with admin_engine.connect() as connection:
            revision_before = connection.scalar(
                text("SELECT count(*) FROM response_card_revisions")
            )

        fresh_response = await api.get(
            f"/tasks/{task_id}/requirements/repair", headers=header, params={"job": job_id}
        )
        fresh = fresh_response.json()["data"]
        executed_response = await api.post(
            f"/tasks/{task_id}/requirements/repair",
            headers=header,
            json={**execute_body, "expected_preview": fresh["preview_hash"]},
        )
        assert executed_response.status_code == 200, executed_response.text
        executed_result = executed_response.json()
        assert executed_result["command"] == "req repair-citations"
        assert executed_result["data"] == {
            **fresh,
            "execute": True,
            "changed": 4,
        }
        assert executed_result["items"] == fresh_response.json()["items"]

        after = requirement_rows(admin_engine, task_id)
        for page in before:
            assert {key: after[page][key] for key in PROTECTED_REQUIREMENT_FIELDS} == {
                key: before[page][key] for key in PROTECTED_REQUIREMENT_FIELDS
            }
        for page, quote in REPAIRED_QUOTES.items():
            assert after[page]["quote"] == quote
            assert after[page]["model_quote"] == SOURCE_QUOTES[page]
        assert after[1]["quote"] == SOURCE_QUOTES[1]
        assert after[1]["model_quote"] == SOURCE_QUOTES[1]
        for page in (5, 6, 7):
            assert (after[page]["quote"], after[page]["model_quote"]) == (
                before[page]["quote"],
                before[page]["model_quote"],
            )

        with admin_engine.connect() as connection:
            audits = [
                dict(row)
                for row in connection.execute(
                    text(
                        "SELECT object_id::text, details FROM audit_logs "
                        "WHERE action = 'requirement.repair_citation' ORDER BY object_id"
                    )
                ).mappings()
            ]
            assert (
                connection.scalar(text("SELECT count(*) FROM response_card_revisions"))
                == revision_before
            )
        assert len(audits) - audit_before == 4
        repair_audits = audits[-4:]
        assert {row["object_id"] for row in repair_audits} == {
            requirements[page]["id"] for page in (1, *REPAIRED_QUOTES)
        }
        for row in repair_audits:
            assert_hashes_and_ids_only(row)
        audit_json = json.dumps(repair_audits, ensure_ascii=False)
        assert reason not in audit_json
        assert all(
            quote not in audit_json
            for quote in [
                *SOURCE_QUOTES.values(),
                *MODEL_QUOTES.values(),
                *REPAIRED_QUOTES.values(),
                *(row["text"] for row in before.values()),
            ]
        )

        # This draft predates the legacy corruption and repair. It is already
        # stale here, before any card is reopened, withdrawn, edited, or re-decided.
        repaired_old_draft = (
            await api.get(f"/drafts/{pre_repair_draft_id}", headers=header)
        ).json()["data"]
        assert repaired_old_draft["validity"] == "stale"
        assert {requirements[page]["id"] for page in REPAIRED_QUOTES} <= set(
            repaired_old_draft["invalidated_requirements"]
        )

        # Repair changes only requirement provenance. The card pointer, state,
        # and history remain, while eligibility blocks every old decision.
        changed_cards = {}
        for page, card in ((2, confirmed), (3, pending), (4, comply)):
            response = await api.get(
                f"/cards/{card['id']}", headers=header, params={"history": True}
            )
            assert response.status_code == 200, response.text
            data = response.json()["data"]
            current = data["card"]
            original = card_before[requirements[page]["id"]]
            assert (current["revision"], current["revision_id"], current["state"]) == (
                original["revision"],
                original["revision_id"],
                original["state"],
            )
            assert data["history"][-1]["revision_id"] == original["revision_id"]
            assert current["eligibility"] == "needs_reconfirmation"
            changed_cards[page] = current

        invalid = (await api.get(f"/cards/{unlocatable['id']}", headers=header)).json()["data"]
        assert invalid["state"] == "pending_review"
        assert invalid["eligibility"] == "needs_reconfirmation"
        invalid_source = await api.get(
            f"/v4/requirements/{requirements[5]['id']}/review", headers=header
        )
        assert invalid_source.status_code == 200, invalid_source.text
        assert invalid_source.json()["data"]["requirement"]["state"] == "invalidated"
        assert invalid_source.json()["data"]["requirement"]["citation_valid"] is False
        normalized_after = (await api.get(f"/cards/{normalized['id']}", headers=header)).json()[
            "data"
        ]
        assert normalized_after["state"] == "confirmed"
        # The quote stayed exact, but this reviewed round also pinned the full
        # requirement before its provenance and protected text were changed.
        assert normalized_after["eligibility"] == "needs_reconfirmation"
        assert normalized_after["revision_id"] == normalized["revision_id"]
        strict_gate = await card_action(api, reviewer, invalid, "confirm", reviewed_evidence_ids=[])
        assert strict_gate.status_code == 409
        assert strict_gate.json()["data"]["error"]["code"] == "requirement_invalidated"
        # Reaccept only currently locatable sources; this never approves response
        # cards or overrides the unresolved legacy citations on pages 5 and 6.
        async with app.state.db.transaction(tenants["orgs"][0]) as session:
            await confirm_requirements_async(
                session,
                tenants["orgs"][0],
                UUID(task_id),
                [UUID(requirements[page]["id"]) for page in (1, 2, 3, 4, 7)],
                settings=app.state.processor.settings,
            )
        pending_stale_confirm = await card_action(
            api, reviewer, changed_cards[3], "confirm", reviewed_evidence_ids=[]
        )
        assert pending_stale_confirm.status_code == 409
        assert pending_stale_confirm.json()["data"]["error"]["code"] == "needs_reconfirmation"

        first_draft = await api.post(
            f"/tasks/{task_id}/drafts",
            headers=header,
            json={"extraction_job_id": job_id},
        )
        assert first_draft.status_code == 200, first_draft.text
        await app.state.processor(header["X-Org-Id"], first_draft.json()["data"]["job_id"])
        draft_job = (
            await api.get(f"/jobs/{first_draft.json()['data']['job_id']}", headers=header)
        ).json()["data"]
        stale_card_draft = (
            await api.get(f"/drafts/{draft_job['result']['draft_id']}", headers=header)
        ).json()["data"]
        gaps = {row["requirement_id"]: row["reasons"] for row in stale_card_draft["gaps"]}
        for page in REPAIRED_QUOTES:
            assert "needs_reconfirmation" in gaps[requirements[page]["id"]]

        # The recovery flow remains explicitly human: leave the protected state,
        # append an edited draft revision, then submit and confirm it again.
        normalized_reopened = await require_card_action(
            api,
            reviewer,
            normalized_after,
            "reopen",
            reason="Review the changed requirement text and provenance with the same exact quote.",
        )
        normalized_reopened = await require_card_action(
            api, reviewer, normalized_reopened, "submit"
        )
        normalized_reconfirmed = await require_card_action(
            api, reviewer, normalized_reopened, "confirm", reviewed_evidence_ids=[]
        )
        assert normalized_reconfirmed["eligibility"] == "eligible"
        reopened = await require_card_action(
            api,
            reviewer,
            changed_cards[2],
            "reopen",
            reason="The tender citation was repaired and must be reviewed again.",
        )
        assert reopened["eligibility"] == "unconfirmed"
        reopened = await update_card_without_changing_business_text(api, reviewer, reopened)
        reopened = await require_card_action(api, reviewer, reopened, "submit")
        reconfirmed = await require_card_action(
            api, reviewer, reopened, "confirm", reviewed_evidence_ids=[]
        )
        assert reconfirmed["state"] == "confirmed" and reconfirmed["eligibility"] == "eligible"

        protected_pending_edit = await api.put(
            f"/cards/{changed_cards[3]['id']}",
            headers=reviewer,
            json={
                "expected_revision": changed_cards[3]["revision"],
                "content": changed_cards[3]["content"],
            },
        )
        assert protected_pending_edit.status_code == 409
        withdrawn = await require_card_action(
            api,
            reviewer,
            changed_cards[3],
            "withdraw",
            reason="The repaired citation requires a fresh edit and review.",
        )
        assert withdrawn["eligibility"] == "unconfirmed"
        withdrawn = await update_card_without_changing_business_text(api, reviewer, withdrawn)
        withdrawn = await require_card_action(api, reviewer, withdrawn, "submit")
        reconfirmed_pending = await require_card_action(
            api, reviewer, withdrawn, "confirm", reviewed_evidence_ids=[]
        )
        assert reconfirmed_pending["eligibility"] == "eligible"

        redisposition = await api.post(
            f"/tasks/{task_id}/cards/dispositions",
            headers=reviewer,
            json={
                "extraction_job_id": job_id,
                "items": [
                    {
                        "requirement_id": requirements[4]["id"],
                        "expected_revision": changed_cards[4]["revision"],
                        "disposition": "comply_only",
                        "reason": "The repaired citation was reviewed for compliance again.",
                    }
                ],
            },
        )
        assert redisposition.status_code == 200, redisposition.text
        [redisposed] = redisposition.json()["data"]["cards"]
        assert redisposed["eligibility"] == "comply_only"

        second_draft = await api.post(
            f"/tasks/{task_id}/drafts",
            headers=header,
            json={"extraction_job_id": job_id},
        )
        assert second_draft.status_code == 200, second_draft.text
        await app.state.processor(header["X-Org-Id"], second_draft.json()["data"]["job_id"])
        second_job = (
            await api.get(f"/jobs/{second_draft.json()['data']['job_id']}", headers=header)
        ).json()["data"]
        current_draft = (
            await api.get(f"/drafts/{second_job['result']['draft_id']}", headers=header)
        ).json()["data"]
        assert {
            row["requirement_id"] for rows in current_draft["tables"].values() for row in rows
        } == {requirements[1]["id"], requirements[2]["id"], requirements[3]["id"]}
        assert [row["requirement_id"] for row in current_draft["comply_only"]] == [
            requirements[4]["id"]
        ]
        assert current_draft["validity"] == "current"
        assert (await api.get(f"/drafts/{draft_job['result']['draft_id']}", headers=header)).json()[
            "data"
        ]["validity"] == "stale"

        artifact = tmp_path / "citation-repair-e2e-artifact.json"
        artifact.write_text(
            json.dumps(
                stable_hashes(
                    {
                        "fixture": "synthetic-labelled-materials-only",
                        "command": executed_result["command"],
                        "preview": fresh,
                        "execute": executed_result["data"],
                        "items": [
                            {
                                "requirement_id": item["requirement_id"],
                                "status": item["status"],
                                "reason": item["reason"],
                                "quote_changed": item["quote_changed"],
                                "source_quote_sha256": hashlib.sha256(
                                    item["source"]["quote"].encode()
                                ).hexdigest(),
                                "model_quote_sha256": hashlib.sha256(
                                    item["model_quote"].encode()
                                ).hexdigest()
                                if item["model_quote"] is not None
                                else None,
                                "proposed_quote_sha256": hashlib.sha256(
                                    item["proposed_quote"].encode()
                                ).hexdigest()
                                if item["proposed_quote"] is not None
                                else None,
                            }
                            for item in executed_result["items"]
                        ],
                        "card_eligibility_after_repair": {
                            str(page): changed_cards[page]["eligibility"]
                            for page in REPAIRED_QUOTES
                        },
                        "draft_gap_reasons": gaps,
                        "pre_repair_draft_after_repair": {
                            "validity": repaired_old_draft["validity"],
                            "invalidated_requirements": repaired_old_draft[
                                "invalidated_requirements"
                            ],
                        },
                        "audit": repair_audits,
                        "rerun": (
                            ".venv/bin/python -m pytest -q -p no:cacheprovider "
                            "server/tests/test_citation_repair.py::"
                            "test_preview_execute_permissions_cards_drafts_and_audit"
                        ),
                    }
                ),
                ensure_ascii=False,
                indent=2,
                sort_keys=True,
            )
        )
        artifact_data = json.loads(artifact.read_text())
        assert artifact_data["execute"]["changed"] == 4
        artifact_text = artifact.read_text()
        assert "Authorization" not in artifact_text
        assert token_record["token"] not in artifact_text
        assert reason not in artifact_text
        assert all(quote not in artifact_text for quote in SOURCE_QUOTES.values())
        assert all(row["id"] not in artifact_text for row in listed)
