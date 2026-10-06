"""Model drafting through API + processor with synthetic MockTransport vendors.

Failure modes identified before implementation:
* Boundary: foreign org/task/job/selection, revoked grants, unsent fields/pages,
  credentials/paths and prior card text must never enter the vendor request.
* Privacy: every outbound string (including location labels) needs redaction;
  only an org human admin may disable it; manifests/logs/errors expose no text.
* Citations: placeholder, joined, ambiguous, unknown, unsent and stale references
  are rejected independently, never repaired; commitments discard extra refs.
* Review: no valid evidence means needs_material, never commitment conversion;
  recorded negative deviations cannot be weakened; workers cannot confirm.
* Concurrency: confirmed/pending/comply_only and newly confirmed cards stay
  untouched; input changes, lease loss and cancellation cannot publish old work.
* Cost: each retry/split call settles before parsing; balance, charge and call
  ceilings fence subsequent calls; first-pass planning scales the call ceiling.
* Idempotency: cache includes model/catalog/reasoning/prompt/redaction/inputs;
  retries cannot duplicate revisions or billing; dry-run makes no calls or writes.
* Workflow: extraction -> generation -> human review -> immutable draft, with
  two-org isolation and a repeatable sanitized artifact outside the repository.
* Output limits: legacy/environment options must not inject a conflicting output
  alias into either adapter's drafting request, reserve too little or spend funds.
"""

import asyncio
import hashlib
import json
import re
from contextlib import asynccontextmanager
from datetime import UTC, datetime, timedelta
from decimal import Decimal
from uuid import UUID, uuid4

import httpx
import pytest
from app.api.main import create_app
from app.models.entities import AuditLog, BalanceEntry, Job, OrgBalance, UsageRecord, VendorCall
from app.models.response_cards import (
    CardGenerationRun,
    Evidence,
    ResponseCard,
    ResponseCardRevision,
)
from app.providers.llm import AnthropicExtractor, OpenAICompatibleExtractor
from conftest import FakeQueue, credential_app
from sqlalchemy import func, select, update
from task_fixtures import reviewer_header
from test_api import create_document, run_job
from test_llm_providers import provider_reply, settings_for
from test_response_cards import (
    PRODUCT_DATA,
    TENDER_LINES,
    card_action,
    create_card,
    create_tender,
    labelled_pdf,
    login,
    require_action,
    sanitized_artifact,
    select_real_materials,
    set_role,
)


class DraftVendor:
    """Synthetic responses only; records request bodies, never credential headers."""

    def __init__(self, provider="openai", lines=None):
        self.provider = provider
        self.lines = lines or TENDER_LINES
        self.drafts = []
        self.bodies = []
        self.respond = None

    async def __call__(self, request):
        body = json.loads(request.content)
        user = body["messages"][-1]["content"]
        if user.startswith("{"):
            sent = json.loads(user)
            self.drafts.append(sent)
            self.bodies.append(body)
            if self.respond:
                return await self.respond(sent)
            return provider_reply(self.provider, self.proposals(sent))
        pages = re.findall(r'<page number="(\d+)">', user)
        items = [
            {
                "category": "qualification" if int(page) == 2 else "technical",
                "starred": False,
                "text": self.lines[int(page) - 1],
                "ref": page,
                "quote": self.lines[int(page) - 1],
                "condition": None,
            }
            for page in pages
        ]
        return provider_reply(self.provider, items)

    @staticmethod
    def proposals(sent):
        items = []
        for requirement in sent["requirements"]:
            page = requirement["location"]["page"]
            refs = []
            if page == 1:
                refs = [
                    {"ref": m["ref"], "quote": m["text"]}
                    for m in sent["materials"]
                    if m.get("field_path") == "model"
                ][:1]
            if page == 2:
                refs = [
                    {"ref": m["ref"], "quote": m["text"].splitlines()[-1]}
                    for m in sent["materials"]
                    if m["kind"] == "certificate_pdf_page"
                ][:1]
            items.append(
                {
                    "requirement_id": requirement["requirement_id"],
                    "response_kind": "evidence" if page in {1, 2} else "commitment",
                    "suggested_disposition": "respond",
                    "response_text": "Synthetic model proposal for human review.",
                    "deviation": "none",
                    "deviation_note": "The synthetic proposal describes the specified obligation.",
                    "evidence": refs,
                }
            )
        return items


@asynccontextmanager
async def drafting_client(
    tenants, tmp_path, provider="openai", *, paid=True, lines=None, **options
):
    vendor = DraftVendor(provider, lines)
    # Tiny batch budgets (one requirement per batch) are below the configurable minimum.
    batch = options.pop("llm_batch_chars", None)
    settings = settings_for(tmp_path, provider, **options)
    if batch is not None:
        settings = settings.model_copy(update={"llm_batch_chars": batch})
    adapter = AnthropicExtractor if provider == "anthropic" else OpenAICompatibleExtractor
    llm = adapter(
        settings,
        httpx.MockTransport(vendor),
        platform_model_id="draft-test" if paid else None,
        sale_usd_per_mtok=(1, 1) if paid else None,
    )
    llm.model_revision = 1 if paid else None
    app = await credential_app(settings, llm=llm, queue=FakeQueue())
    async with (
        app.router.lifespan_context(app),
        httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test") as api,
    ):
        headers = [
            await login(api, org, label)
            for org, label in zip(tenants["orgs"], ("a", "b"), strict=True)
        ]
        if paid:
            for org in tenants["orgs"]:
                async with app.state.db.transaction(org) as session:
                    session.add(OrgBalance(org_id=org, currency="USD", balance=Decimal("10")))
                    session.add(
                        BalanceEntry(
                            org_id=org,
                            kind="adjust",
                            currency="USD",
                            amount=Decimal("10"),
                            balance_after=Decimal("10"),
                            actor="test",
                            reason="Synthetic initial funds",
                        )
                    )
        yield api, app, headers, vendor, llm


async def submit(api, header, task, extraction, **body):
    response = await api.post(
        f"/tasks/{task}/cards/generations",
        headers=header,
        json={"extraction_job_id": extraction, **body},
    )
    assert response.status_code == 200, response.text
    return response.json()


async def execute(api, app, header, receipt):
    job = receipt["data"]["job_id"]
    await app.state.processor(header["X-Org-Id"], job)
    response = await api.get(f"/jobs/{job}", headers=header)
    assert response.status_code == 200, response.text
    return response.json()


@pytest.mark.parametrize("provider", ["anthropic", "openai"])
@pytest.mark.parametrize(
    "options",
    [
        {"max_tokens": 64000},
        {"max_completion_tokens": 1},
        {"max_output_tokens": 1},
        {"max_new_tokens": 1},
        {"generation_config": {"maxOutputTokens": 1}},
    ],
)
async def test_drafting_rejects_conflicting_output_limits(tenants, tmp_path, provider, options):
    async with drafting_client(tenants, tmp_path, provider) as (api, app, headers, vendor, llm):
        header = headers[0]
        task, _, extraction, _ = await create_tender(api, app, header, tmp_path)
        # Reproduce a legacy level/environment option after the unrelated extraction.
        llm.settings = llm.settings.model_copy(update={"llm_request_options": json.dumps(options)})
        receipt = await submit(api, header, task, extraction)
        result = (await execute(api, app, header, receipt))["data"]
        assert result["status"] == "failed", result
        assert result["error"]["code"] == "invalid_provider_options"
        assert vendor.drafts == []
        async with app.state.db.transaction(tenants["orgs"][0]) as session:
            assert (
                await session.scalar(
                    select(func.count())
                    .select_from(VendorCall)
                    .where(VendorCall.job_id == UUID(receipt["data"]["job_id"]))
                )
                == 0
            )
            assert await session.scalar(select(func.count()).select_from(CardGenerationRun)) == 0
        (tmp_path / "drafting-limit-rejection.json").write_text(
            json.dumps(
                sanitized_artifact({"provider": provider, "options": options, "job": result}),
                indent=2,
            )
        )


async def slots(api, header, task, extraction):
    response = await api.get(f"/tasks/{task}/cards", headers=header, params={"job": extraction})
    assert response.status_code == 200, response.text
    return response.json()["items"]


async def counts(app, org):
    async with app.state.db.transaction(org) as session:
        return {
            model.__tablename__: await session.scalar(select(func.count()).select_from(model))
            for model in (
                Job,
                CardGenerationRun,
                ResponseCardRevision,
                ResponseCard,
                Evidence,
                UsageRecord,
                VendorCall,
                AuditLog,
            )
        }


async def token_header(api, header, *, scopes=None):
    response = await api.post(
        "/tokens",
        headers=header,
        json={
            "name": "Synthetic drafting agent",
            "scopes": scopes
            or [
                "task:read",
                "job:read",
                "card:read",
                "card:write",
                "card:generate",
                "memory:read",
                "memory:retrieve",
                "resource:read",
                "certificate:read",
                "certificate:file:read",
                "evidence:source:read",
                "profile:read",
            ],
            "expires_at": (datetime.now(UTC) + timedelta(hours=1)).isoformat(),
        },
    )
    assert response.status_code == 200, response.text
    return {
        "Authorization": "Bearer " + response.json()["data"]["token"],
        "X-Org-Id": header["X-Org-Id"],
    }


def same_instants(value):
    """Replace ISO timestamps with aware datetimes so offsets do not affect equality."""
    if isinstance(value, dict):
        return {key: same_instants(item) for key, item in value.items()}
    if isinstance(value, list):
        return [same_instants(item) for item in value]
    if isinstance(value, str) and len(value) >= 20 and value[10:11] == "T":
        try:
            return datetime.fromisoformat(value.replace("Z", "+00:00"))
        except ValueError:
            return value
    return value


@pytest.mark.parametrize("provider", ["openai", "anthropic"])
async def test_generation_review_and_draft_full_chain(tenants, tmp_path, admin_engine, provider):
    async with drafting_client(tenants, tmp_path, provider) as (api, app, headers, vendor, _):
        header = headers[0]
        task, _, extraction, requirements = await create_tender(
            api, app, header, tmp_path, confirmed=True
        )
        await select_real_materials(api, header, task, tmp_path)
        agent = await token_header(api, header)
        receipt = await submit(api, agent, task, extraction)
        terminal = await execute(api, app, header, receipt)
        assert terminal["data"]["status"] == "succeeded", terminal
        generated = terminal["data"]["result"]
        assert generated["completion"] == "complete" and len(generated["created_revision_ids"]) == 5
        assert len(generated["usage_record_ids"]) == len(vendor.drafts) == 1
        assert "submission" not in generated and "encrypted_input" not in json.dumps(terminal)
        current = await slots(api, header, task, extraction)
        for slot in current:
            card = slot["card"]
            assert (
                card["state"] == "draft"
                and card["origin"] == "model"
                and card["actor_kind"] == "worker"
            )
            assert card["confirmed_by"] is None and card["disposition"] is None
            assert all(evidence["confirmed_by"] is None for evidence in card["evidence"])
        cached = await submit(api, header, task, extraction)
        assert cached["data"]["cached"] and cached["data"]["job_id"] == receipt["data"]["job_id"]
        await execute(api, app, header, cached)
        assert len(vendor.drafts) == 1
        async with app.state.db.transaction(tenants["orgs"][0]) as session:
            usage = await session.get(UsageRecord, UUID(generated["usage_record_ids"][0]))
            assert usage is not None and usage.call_id == usage.id and usage.charge > 0
            entries = list(
                (
                    await session.scalars(
                        select(BalanceEntry).where(BalanceEntry.usage_record_id == usage.id)
                    )
                ).all()
            )
            assert len(entries) == 1 and entries[0].amount == -usage.charge
            run = await session.scalar(select(CardGenerationRun))
            assert run and run.actor_kind == "worker" and run.encrypted_input
            assert "Synthetic model proposal" not in json.dumps(run.input_manifest)
        first_pending = await require_action(api, agent, current[0]["card"], "submit")
        denied = await card_action(api, agent, first_pending, "confirm")
        assert denied.status_code == 403
        from app.core.errors import ServiceError
        from app.schemas.response_card_contracts import CardAction
        from app.services import response_cards as card_service
        from app.services.auth import ROLE_SCOPES, Identity

        for actor_kind in ("agent", "worker"):
            with pytest.raises(ServiceError) as rejected:
                async with app.state.db.transaction(tenants["orgs"][0]) as session:
                    await card_service.card_action(
                        session,
                        Identity(
                            tenants["users"][0],
                            tenants["orgs"][0],
                            set(ROLE_SCOPES["admin"]),
                            "admin",
                            actor_kind=actor_kind,
                        ),
                        UUID(first_pending["id"]),
                        CardAction(expected_revision=first_pending["revision"], action="confirm"),
                        app.state.storage,
                    )
            assert rejected.value.code == "forbidden"
        review_headers = {}
        for domain in ("technical", "commercial"):
            _, review_headers[domain] = await reviewer_header(
                api, admin_engine, tenants["orgs"][0], UUID(task), domain
            )
        for index, slot in enumerate(current):
            review_header = review_headers[slot["card"]["review_domain"]]
            pending = (
                first_pending
                if index == 0
                else await require_action(api, review_header, slot["card"], "submit")
            )
            extra = {"reviewed_evidence_ids": [item["id"] for item in pending["evidence"]]}
            if pending["warning_codes"]:
                extra |= {
                    "reviewed_warning_codes": pending["warning_codes"],
                    "reason": "Reviewed synthetic source and proof obligation.",
                }
            confirmed = await require_action(api, review_header, pending, "confirm", **extra)
            assert confirmed["state"] == "confirmed" and confirmed["content"] == pending["content"]
            for evidence in confirmed["evidence"]:
                if evidence["source_archive"]:
                    assert evidence["quote_check"] == "human_page_review"
                    assert evidence["source_archive"]["confirmed_by"] is None
        draft = await api.post(
            f"/tasks/{task}/drafts", headers=header, json={"extraction_job_id": extraction}
        )
        assert draft.status_code == 200, draft.text
        assembled = await execute(api, app, header, draft.json())
        assert assembled["data"]["status"] == "succeeded", assembled
        shown = await api.get(f"/drafts/{assembled['data']['result']['draft_id']}", headers=header)
        view = shown.json()["data"]
        assert view["completion"] == "complete" and view["validity"] == "current"
        assert sum(len(rows) for rows in view["tables"].values()) == len(requirements)
        artifact = {
            "case": "generation-review-draft",
            "provider": provider,
            "generation": terminal,
            "draft": shown.json(),
            "rerun": "pytest -q server/tests/test_card_generation.py -k full_chain",
        }
        (tmp_path / f"generation-{provider}.json").write_text(
            json.dumps(sanitized_artifact(artifact), ensure_ascii=False, indent=2)
        )


@pytest.mark.parametrize(
    "case,reason",
    [
        ("placeholder", "redacted_placeholder"),
        ("unsent_field", "quote_not_sent"),
        ("foreign_selection", "unknown_ref"),
        ("joined", "quote_not_sent"),
        ("unknown", "unknown_ref"),
        ("ambiguous", "ambiguous_quote"),
        ("stale_selection", "stale_material"),
        ("commitment_refs", "commitment_evidence_dropped"),
    ],
)
async def test_invalid_model_references_are_dropped(tenants, tmp_path, case, reason):
    async with drafting_client(tenants, tmp_path) as (api, app, headers, vendor, _):
        header = headers[0]
        task, _, extraction, requirements = await create_tender(api, app, header, tmp_path)
        product, selected, _, _, _ = await select_real_materials(api, header, task, tmp_path)
        other, _, _, _ = await create_tender(api, app, header, tmp_path, suffix="other")
        _, foreign, _, _, _ = await select_real_materials(api, header, other, tmp_path)
        profile = await api.post(
            "/resources/profiles",
            headers=header,
            json={
                "data": {
                    "name": "Synthetic profile",
                    "standard_wording": "Alpha proof. Middle excluded. Omega proof.",
                    "performance_summary": "Repeated proof. Repeated proof.",
                }
            },
        )
        assert profile.status_code == 200, profile.text
        picked = await api.post(
            f"/tasks/{task}/profiles",
            headers=header,
            json={"profile_id": profile.json()["data"]["profile_id"]},
        )
        assert picked.status_code == 200, picked.text

        async def response(sent):
            proposal = vendor.proposals(sent)[0]
            ref = next(
                item["ref"] for item in sent["materials"] if item.get("field_path") == "model"
            )
            quote = PRODUCT_DATA["model"]
            if case == "placeholder":
                quote = "[REDACTED_AMOUNT]"
            elif case == "unsent_field":
                quote = "UNSELECTED PRIVATE FIELD"
            elif case == "foreign_selection":
                ref = foreign["id"]
            elif case == "unknown":
                ref = "model-supplied-new-source"
            elif case == "joined":
                ref = next(
                    item["ref"]
                    for item in sent["materials"]
                    if item.get("field_path") == "standard_wording"
                )
                quote = "Alpha proof. Omega proof."
            elif case == "ambiguous":
                ref = next(
                    item["ref"]
                    for item in sent["materials"]
                    if item.get("field_path") == "performance_summary"
                )
                quote = "Repeated proof."
            elif case == "stale_selection":
                changed = await api.post(
                    f"/resources/products/{product['product_id']}/revisions",
                    headers=header,
                    json={
                        "expected_revision": 1,
                        "data": PRODUCT_DATA | {"model": "Replacement declaration"},
                    },
                )
                assert changed.status_code == 200, changed.text
                replaced = await api.post(
                    f"/tasks/{task}/products",
                    headers=header,
                    json={"product_id": product["product_id"]},
                )
                assert replaced.status_code == 200, replaced.text
            elif case == "commitment_refs":
                proposal["response_kind"] = "commitment"
            proposal["evidence"] = [{"ref": ref, "quote": quote}]
            return provider_reply("openai", [proposal])

        vendor.respond = response
        receipt = await submit(
            api, header, task, extraction, requirement_ids=[requirements[0]["id"]]
        )
        terminal = await execute(api, app, header, receipt)
        assert terminal["data"]["status"] == "succeeded", terminal
        result = terminal["data"]["result"]
        assert result["completion"] == "partial" and not terminal["ok"]
        assert any(
            item.endswith(":" + reason)
            for item in result["rejected_references"][requirements[0]["id"]]
        )
        card = (await slots(api, header, task, extraction))[0]["card"]
        assert card["state"] == "draft" and card["evidence"] == [] and card["confirmed_by"] is None
        assert card["review_hint"] == (None if case == "commitment_refs" else "needs_material")
        assert card["content"]["response_kind"] == (
            "commitment" if case == "commitment_refs" else "evidence"
        )
        assert selected["id"] not in json.dumps(vendor.drafts) and foreign["id"] not in json.dumps(
            vendor.drafts
        )
        async with app.state.db.transaction(tenants["orgs"][0]) as session:
            assert await session.scalar(select(func.count()).select_from(Evidence)) == 0


SENSITIVE = {
    "amount": "98765.43",
    "contact": "Synthetic Person",
    "phone": "13800138000",
    "identity": "110101199003071234",
    "bank_account": "6222021234567890123",
}
SENSITIVE_LINES = [
    "Price: USD 98765.43;",
    "Contact: Synthetic Person; Phone: 13800138000;",
    "Identity number: 110101199003071234;",
    "Bank account: 6222021234567890123;",
]


@pytest.mark.parametrize("enabled", [True, False])
async def test_redaction_all_inputs_dry_run_and_admin_switch(
    tenants, tmp_path, admin_engine, enabled, caplog
):
    async with drafting_client(tenants, tmp_path, lines=SENSITIVE_LINES) as (
        api,
        app,
        headers,
        vendor,
        _,
    ):
        header = headers[0]
        task, document = await create_document(api, header, labelled_pdf(SENSITIVE_LINES))
        await run_job(api, app, header, document, "parse")
        extraction, status = await run_job(api, app, header, document, "extract")
        assert status["status"] == "succeeded", status
        profile = await api.post(
            "/resources/profiles",
            headers=header,
            json={
                "data": {
                    "name": "Synthetic sensitive fixture",
                    "standard_wording": "\n".join(SENSITIVE_LINES),
                }
            },
        )
        assert profile.status_code == 200, profile.text
        assert (
            await api.post(
                f"/tasks/{task}/profiles",
                headers=header,
                json={"profile_id": profile.json()["data"]["profile_id"]},
            )
        ).status_code == 200
        agent = await token_header(api, header)
        denied = await api.put(
            f"/tasks/{task}/model-redaction",
            headers=agent,
            json={"expected_revision": 1, "model_redaction_enabled": False},
        )
        assert denied.status_code == 403
        set_role(admin_engine, tenants["orgs"][0], tenants["users"][0], "technical")
        denied = await api.put(
            f"/tasks/{task}/model-redaction",
            headers=header,
            json={"expected_revision": 1, "model_redaction_enabled": False},
        )
        assert denied.status_code == 403
        set_role(admin_engine, tenants["orgs"][0], tenants["users"][0], "admin")
        original = await submit(api, header, task, extraction, dry_run=True)
        if not enabled:
            changed = await api.put(
                f"/tasks/{task}/model-redaction",
                headers=header,
                json={"expected_revision": 1, "model_redaction_enabled": False},
            )
            assert changed.status_code == 200, changed.text
        before, queued = await counts(app, tenants["orgs"][0]), list(app.state.queue.calls)
        preview = await submit(api, header, task, extraction, dry_run=True)
        assert await counts(app, tenants["orgs"][0]) == before and app.state.queue.calls == queued
        assert not vendor.drafts
        data = preview["data"]
        assert data["model_redaction_enabled"] is enabled
        assert all(
            data["redacted_counts"][key] > 0
            for key in ("amount", "contact", "identity", "bank_account")
        )
        assert data["cost_basis"] == "known" and Decimal(data["estimated_charge"]) > 0
        if not enabled:
            assert data["input_hash"] != original["data"]["input_hash"]
            assert any("unredacted" in warning for warning in preview["warnings"])
        receipt = await submit(api, header, task, extraction)
        terminal = await execute(api, app, header, receipt)
        assert terminal["data"]["status"] == "succeeded", terminal
        wire = json.dumps(vendor.drafts, ensure_ascii=False)
        for value in SENSITIVE.values():
            assert (value in wire) is (not enabled)
            assert value not in json.dumps(preview, ensure_ascii=False)
            assert value not in json.dumps(terminal, ensure_ascii=False)
            assert value not in caplog.text
        async with app.state.db.transaction(tenants["orgs"][0]) as session:
            run = await session.scalar(select(CardGenerationRun))
            assert run is not None and run.encrypted_input
            assert all(value not in run.encrypted_input for value in SENSITIVE.values())
            audits = list(
                (
                    await session.scalars(
                        select(AuditLog).where(AuditLog.action.like("card.generate%"))
                    )
                ).all()
            )
            assert audits
            assert all(
                value not in json.dumps([row.details for row in audits])
                for value in SENSITIVE.values()
            )


async def test_state_skips_negative_deviation_and_no_prior_response_leak(
    tenants, tmp_path, admin_engine
):
    async with drafting_client(tenants, tmp_path) as (api, app, headers, vendor, _):
        header = headers[0]
        task, _, extraction, requirements = await create_tender(
            api, app, header, tmp_path, confirmed=True
        )
        content = {
            "response_kind": "commitment",
            "response_text": "PRIVATE PRIOR RESPONSE MUST NEVER BE SENT",
            "deviation": "negative",
            "deviation_note": "Cannot meet the synthetic obligation.",
            "evidence": [],
        }
        negative = await create_card(api, header, task, extraction, requirements[0], content)
        set_role(admin_engine, tenants["orgs"][0], tenants["users"][0], "technical")
        confirmed = await create_card(api, header, task, extraction, requirements[2], content)
        confirmed = await require_action(
            api, header, await require_action(api, header, confirmed, "submit"), "confirm"
        )
        pending = await create_card(api, header, task, extraction, requirements[3], content)
        pending = await require_action(api, header, pending, "submit")
        comply = await api.post(
            f"/tasks/{task}/cards/dispositions",
            headers=header,
            json={
                "extraction_job_id": extraction,
                "items": [
                    {
                        "requirement_id": requirements[4]["id"],
                        "disposition": "comply_only",
                        "reason": "Synthetic procedural obligation.",
                    }
                ],
            },
        )
        assert comply.status_code == 200, comply.text
        before = await slots(api, header, task, extraction)
        receipt = await submit(api, header, task, extraction)
        result = (await execute(api, app, header, receipt))["data"]["result"]
        assert result["skipped"] == {
            requirements[0]["id"]: "negative_deviation_weakened",
            requirements[2]["id"]: "confirmed",
            requirements[3]["id"]: "pending_review",
            requirements[4]["id"]: "comply_only",
        }
        after = await slots(api, header, task, extraction)
        for index in (0, 2, 3, 4):
            assert after[index] == before[index]
        assert "PRIVATE PRIOR RESPONSE" not in json.dumps(vendor.drafts)
        assert negative["revision_id"] == after[0]["card"]["revision_id"]


@pytest.mark.parametrize("change", ["confirm", "edit", "cancel", "takeover", "expire"])
async def test_changes_while_vendor_runs_are_fenced(tenants, tmp_path, admin_engine, change):
    async with drafting_client(tenants, tmp_path) as (api, app, headers, vendor, _):
        header = headers[0]
        task, _, extraction, requirements = await create_tender(
            api, app, header, tmp_path, confirmed=True
        )
        set_role(admin_engine, tenants["orgs"][0], tenants["users"][0], "technical")
        card = await create_card(
            api,
            header,
            task,
            extraction,
            requirements[2],
            {
                "response_kind": "commitment",
                "response_text": "Prior human proposed obligation.",
                "deviation": "none",
                "deviation_note": "Synthetic unchanged obligation.",
                "evidence": [],
            },
        )
        entered, release = asyncio.Event(), asyncio.Event()

        async def response(sent):
            entered.set()
            await release.wait()
            return provider_reply("openai", vendor.proposals(sent))

        vendor.respond = response
        receipt = await submit(
            api, header, task, extraction, requirement_ids=[requirements[2]["id"]]
        )
        job_id = receipt["data"]["job_id"]
        running = asyncio.create_task(app.state.processor(header["X-Org-Id"], job_id))
        await asyncio.wait_for(entered.wait(), 5)
        try:
            if change == "confirm":
                card = await require_action(
                    api, header, await require_action(api, header, card, "submit"), "confirm"
                )
            elif change == "edit":
                changed = await api.put(
                    f"/cards/{card['id']}",
                    headers=header,
                    json={
                        "expected_revision": card["revision"],
                        "content": card["content"] | {"response_text": "Concurrent human edit."},
                    },
                )
                assert changed.status_code == 200, changed.text
                card = changed.json()["data"]
            elif change == "cancel":
                assert (await api.post(f"/jobs/{job_id}/cancel", headers=header)).status_code == 200
            else:
                async with app.state.db.transaction(tenants["orgs"][0]) as session:
                    values = (
                        {"run_id": uuid4()}
                        if change == "takeover"
                        else {"lease_until": datetime.now(UTC) - timedelta(seconds=1)}
                    )
                    await session.execute(
                        update(Job).where(Job.id == UUID(job_id)).values(**values)
                    )
        finally:
            release.set()
        await running
        shown = (await api.get(f"/cards/{card['id']}", headers=header)).json()["data"]
        # Same instants; a write response and a later read may differ only in UTC offset.
        assert same_instants(shown) == same_instants(card)
        terminal = (await api.get(f"/jobs/{job_id}", headers=header)).json()["data"]
        if change in {"confirm", "edit"}:
            assert terminal["status"] == "succeeded", terminal
            assert terminal["result"]["skipped"][requirements[2]["id"]] == (
                "confirmed" if change == "confirm" else "revision_conflict"
            )
        async with app.state.db.transaction(tenants["orgs"][0]) as session:
            paid = list(
                (
                    await session.scalars(
                        select(UsageRecord).where(UsageRecord.job_id == UUID(job_id))
                    )
                ).all()
            )
            assert len(paid) == 1 and paid[0].charge > 0


@pytest.mark.parametrize("failure", ["truncated", "malformed", "transient"])
@pytest.mark.parametrize("provider", ["openai", "anthropic"])
async def test_drafting_halving_and_retries_account_every_call(
    tenants, tmp_path, monkeypatch, provider, failure
):
    adapter = AnthropicExtractor if provider == "anthropic" else OpenAICompatibleExtractor
    monkeypatch.setattr(adapter, "retry_delays", (0, 0))
    async with drafting_client(tenants, tmp_path, provider) as (api, app, headers, vendor, _):
        header = headers[0]
        task, _, extraction, requirements = await create_tender(api, app, header, tmp_path)

        async def response(sent):
            reply = provider_reply(provider, vendor.proposals(sent)).json()
            if len(vendor.drafts) == 1:
                if failure == "transient":
                    return httpx.Response(
                        503,
                        json={"error": {"type": "SENSITIVE-ERROR-ECHO"}, "usage": reply["usage"]},
                    )
                if provider == "anthropic":
                    if failure == "truncated":
                        reply["stop_reason"] = "max_tokens"
                    else:
                        reply["content"] = [{"type": "text", "text": "SENSITIVE-MALFORMED-ECHO"}]
                elif failure == "truncated":
                    reply["choices"][0]["finish_reason"] = "length"
                else:
                    reply["choices"][0]["message"]["content"] = "SENSITIVE-MALFORMED-ECHO"
            return httpx.Response(200, json=reply)

        vendor.respond = response
        receipt = await submit(
            api, header, task, extraction, requirement_ids=[r["id"] for r in requirements[2:]]
        )
        terminal = await execute(api, app, header, receipt)
        assert terminal["data"]["status"] == "succeeded", terminal
        assert len(vendor.drafts) == (2 if failure == "transient" else 3)
        result = terminal["data"]["result"]
        assert result["completion"] == "complete" and len(result["created_revision_ids"]) == 3
        assert len(result["usage_record_ids"]) == len(vendor.drafts)
        assert "SENSITIVE-" not in json.dumps(terminal)
        async with app.state.db.transaction(tenants["orgs"][0]) as session:
            usage = list(
                (
                    await session.scalars(
                        select(UsageRecord).where(
                            UsageRecord.job_id == UUID(receipt["data"]["job_id"])
                        )
                    )
                ).all()
            )
            assert sum(row.tokens for row in usage) == result["cost"]["llm_tokens"]
            charged = await session.scalar(
                select(func.sum(-BalanceEntry.amount)).where(
                    BalanceEntry.usage_record_id.in_([row.id for row in usage])
                )
            )
            assert charged == sum(row.charge for row in usage)


async def test_first_pass_planning_scales_call_ceiling(tenants, tmp_path):
    async with drafting_client(
        tenants, tmp_path, llm_batch_chars=1, job_max_vendor_calls=1, job_vendor_calls_per_batch=1
    ) as (api, app, headers, vendor, _):
        task, _, extraction, requirements = await create_tender(api, app, headers[0], tmp_path)
        receipt = await submit(
            api,
            headers[0],
            task,
            extraction,
            requirement_ids=[row["id"] for row in requirements[2:]],
        )
        terminal = await execute(api, app, headers[0], receipt)
        assert terminal["data"]["status"] == "succeeded", terminal
        assert len(vendor.drafts) == 3
        assert terminal["data"]["result"]["completion"] == "complete"


async def test_ceiling_keeps_valid_partial_results_and_cache_does_not_rebill(
    tenants, tmp_path, monkeypatch
):
    monkeypatch.setattr(OpenAICompatibleExtractor, "retry_delays", (0, 0))
    async with drafting_client(
        tenants, tmp_path, llm_batch_chars=1, job_max_vendor_calls=1, job_vendor_calls_per_batch=1
    ) as (api, app, headers, vendor, _):
        header = headers[0]
        task, _, extraction, requirements = await create_tender(api, app, header, tmp_path)

        async def response(sent):
            if len(vendor.drafts) == 1:
                return httpx.Response(
                    503,
                    json={
                        "error": {"type": "busy"},
                        "usage": {"prompt_tokens": 10, "completion_tokens": 1},
                    },
                )
            return provider_reply("openai", vendor.proposals(sent))

        vendor.respond = response
        selected = [row["id"] for row in requirements[2:4]]
        receipt = await submit(api, header, task, extraction, requirement_ids=selected)
        terminal = await execute(api, app, header, receipt)
        assert terminal["data"]["status"] == "succeeded", terminal
        result = terminal["data"]["result"]
        assert (
            result["completion"] == "partial" and result["stop_reason"] == "job_call_limit_exceeded"
        )
        assert len(result["created_revision_ids"]) == 1 and len(result["usage_record_ids"]) == 2
        # Both batches start together; whichever met the 503 is the one the ceiling stopped.
        stopped = [key for key in selected if result["skipped"].get(key) == "generation_stopped"]
        assert len(vendor.drafts) == 2 and len(stopped) == 1
        repeated = await submit(api, header, task, extraction, requirement_ids=selected, retry=True)
        assert repeated["data"]["job_id"] == receipt["data"]["job_id"]
        await execute(api, app, header, repeated)
        assert len(vendor.drafts) == 2


@pytest.mark.parametrize("guard", ["balance_at_submission", "balance_at_call", "charge_ceiling"])
async def test_funds_and_charge_guards_prevent_any_drafting_call(tenants, tmp_path, guard):
    async with drafting_client(tenants, tmp_path) as (api, app, headers, vendor, _):
        header = headers[0]
        task, _, extraction, _ = await create_tender(api, app, header, tmp_path)
        if guard == "balance_at_submission":
            async with app.state.db.transaction(tenants["orgs"][0]) as session:
                await session.execute(update(OrgBalance).values(balance=0))
            before = await counts(app, tenants["orgs"][0])
            preview = await submit(api, header, task, extraction, dry_run=True)
            assert "insufficient_balance" in preview["warnings"]
            denied = await api.post(
                f"/tasks/{task}/cards/generations",
                headers=header,
                json={"extraction_job_id": extraction},
            )
            assert (
                denied.status_code == 402
                and denied.json()["data"]["error"]["code"] == "insufficient_balance"
            )
            assert await counts(app, tenants["orgs"][0]) == before
        else:
            receipt = await submit(api, header, task, extraction)
            if guard == "balance_at_call":
                async with app.state.db.transaction(tenants["orgs"][0]) as session:
                    await session.execute(update(OrgBalance).values(balance=0))
            else:
                app.state.processor.settings = app.state.processor.settings.model_copy(
                    update={"job_max_charge": Decimal("0.000001")}
                )
            terminal = await execute(api, app, header, receipt)
            assert terminal["data"]["status"] == "failed", terminal
            assert terminal["data"]["error"]["code"] == (
                "insufficient_balance"
                if guard == "balance_at_call"
                else "job_charge_limit_exceeded"
            )
        assert not vendor.drafts
        async with app.state.db.transaction(tenants["orgs"][0]) as session:
            assert await session.scalar(select(func.count()).select_from(ResponseCard)) == 0


async def test_unknown_prices_and_no_catalog_reasoning_do_not_invent_estimates(tenants, tmp_path):
    async with drafting_client(tenants, tmp_path, paid=False) as (api, app, headers, vendor, _):
        task, _, extraction, _ = await create_tender(api, app, headers[0], tmp_path)
        preview = await submit(api, headers[0], task, extraction, reasoning="high", dry_run=True)
        assert preview["data"]["reasoning"] is None
        assert (
            preview["data"]["cost_basis"] == "unknown"
            and preview["data"]["estimated_cost"]["usd"] is None
        )
        assert preview["data"]["estimated_charge"] is None
        assert any("ignored" in warning for warning in preview["warnings"]) and not vendor.drafts


async def test_reasoning_cache_and_changed_configuration_fail_closed(tenants, tmp_path):
    async with drafting_client(tenants, tmp_path) as (api, app, headers, vendor, llm):
        header = headers[0]
        task, _, extraction, requirements = await create_tender(api, app, header, tmp_path)
        llm.reasoning_levels = {
            "low": {
                "name": "low",
                "request_options": {"reasoning_effort": "low"},
                "batch_chars": 8000,
            },
            "high": {
                "name": "high",
                "request_options": {"reasoning_effort": "high"},
                "batch_chars": 1,
            },
        }
        llm.default_reasoning = "low"
        denied = await api.post(
            f"/tasks/{task}/cards/generations",
            headers=header,
            json={"extraction_job_id": extraction, "reasoning": "invented"},
        )
        assert (
            denied.status_code == 400
            and denied.json()["data"]["error"]["code"] == "unsupported_reasoning"
        )
        low = await submit(api, header, task, extraction, dry_run=True)
        high = await submit(api, header, task, extraction, reasoning="high", dry_run=True)
        assert low["data"]["input_hash"] != high["data"]["input_hash"] and not vendor.drafts
        selected = [row["id"] for row in requirements[2:]]
        receipt = await submit(
            api, header, task, extraction, reasoning="high", requirement_ids=selected
        )
        terminal = await execute(api, app, header, receipt)
        assert terminal["data"]["status"] == "succeeded", terminal
        assert [body["reasoning_effort"] for body in vendor.bodies] == ["high"] * 3
        pending = await submit(
            api, header, task, extraction, reasoning="low", requirement_ids=selected
        )
        llm.model_revision = 2
        failed = await execute(api, app, header, pending)
        assert (
            failed["data"]["status"] == "failed"
            and failed["data"]["error"]["code"] == "generation_model_changed"
        )
        assert len(vendor.drafts) == 3


async def test_batches_run_concurrently_within_the_limit_and_publish_in_order(tenants, tmp_path):
    async with drafting_client(tenants, tmp_path, llm_batch_chars=1, llm_concurrency=2) as (
        api,
        app,
        headers,
        vendor,
        _,
    ):
        header = headers[0]
        task, _, extraction, requirements = await create_tender(api, app, header, tmp_path)
        flight = {"now": 0, "most": 0}
        both = asyncio.Event()

        async def response(sent):
            flight["now"] += 1
            flight["most"] = max(flight["most"], flight["now"])
            if flight["now"] == 2:
                both.set()
            # The first two calls wait for each other, which only concurrent batches can do.
            await asyncio.wait_for(both.wait(), 5)
            await asyncio.sleep(0)
            flight["now"] -= 1
            return provider_reply("openai", vendor.proposals(sent))

        vendor.respond = response
        selected = [row["id"] for row in requirements[2:5]]
        receipt = await submit(api, header, task, extraction, requirement_ids=selected)
        terminal = await execute(api, app, header, receipt)
        assert terminal["data"]["status"] == "succeeded", terminal
        assert terminal["data"]["result"]["completion"] == "complete"
        assert flight["most"] == 2 and len(vendor.drafts) == 3
        cards = await api.get(f"/tasks/{task}/cards", headers=header, params={"job": extraction})
        revisions = {
            slot["requirement_id"]: slot["card"]["revision_id"]
            for slot in cards.json()["items"]
            if slot["card"]
        }
        assert terminal["data"]["result"]["created_revision_ids"] == [
            revisions[key] for key in selected
        ]


@pytest.mark.parametrize("when", ["before_first_call", "between_batches"])
async def test_changed_redaction_switch_stops_new_calls(tenants, tmp_path, when):
    # One call at a time, so the switch changes strictly between two batches.
    async with drafting_client(tenants, tmp_path, llm_batch_chars=1, llm_concurrency=1) as (
        api,
        app,
        headers,
        vendor,
        _,
    ):
        header = headers[0]
        task, _, extraction, requirements = await create_tender(api, app, header, tmp_path)
        receipt = await submit(
            api, header, task, extraction, requirement_ids=[r["id"] for r in requirements[2:]]
        )

        async def change_switch():
            response = await api.put(
                f"/tasks/{task}/model-redaction",
                headers=header,
                json={"expected_revision": 1, "model_redaction_enabled": False},
            )
            assert response.status_code == 200, response.text

        if when == "before_first_call":
            await change_switch()
        else:

            async def response(sent):
                await change_switch()
                return provider_reply("openai", vendor.proposals(sent))

            vendor.respond = response
        terminal = await execute(api, app, header, receipt)
        if when == "before_first_call":
            assert (
                terminal["data"]["status"] == "failed"
                and terminal["data"]["error"]["code"] == "generation_input_changed"
            )
            assert not vendor.drafts
        else:
            assert terminal["data"]["status"] == "succeeded", terminal
            assert (
                terminal["data"]["result"]["stop_reason"] == "generation_input_changed"
                and len(vendor.drafts) == 1
            )


async def test_two_orgs_and_same_org_tasks_cannot_cross_input_scope(tenants, tmp_path):
    async with drafting_client(tenants, tmp_path) as (api, app, headers, vendor, _):
        first = await create_tender(api, app, headers[0], tmp_path, suffix="A")
        other = await create_tender(api, app, headers[0], tmp_path, suffix="A-other")
        foreign = await create_tender(api, app, headers[1], tmp_path, suffix="B")
        await select_real_materials(api, headers[1], foreign[0], tmp_path)
        secret = await api.post(
            "/resources/profiles", headers=headers[1], json={"data": {"name": "FOREIGN ORG SECRET"}}
        )
        assert secret.status_code == 200, secret.text
        await api.post(
            f"/tasks/{foreign[0]}/profiles",
            headers=headers[1],
            json={"profile_id": secret.json()["data"]["profile_id"]},
        )
        for task, job, req in [
            (foreign[0], foreign[2], None),
            (first[0], foreign[2], None),
            (first[0], other[2], None),
            (first[0], first[2], other[3][0]["id"]),
            (first[0], first[2], foreign[3][0]["id"]),
        ]:
            body = {"extraction_job_id": job, "dry_run": True}
            if req:
                body["requirement_ids"] = [req]
            response = await api.post(
                f"/tasks/{task}/cards/generations", headers=headers[0], json=body
            )
            assert response.status_code == 404, response.text
        receipt = await submit(api, headers[0], first[0], first[2])
        terminal = await execute(api, app, headers[0], receipt)
        assert terminal["data"]["status"] == "succeeded", terminal
        assert "FOREIGN ORG SECRET" not in json.dumps(vendor.drafts)
        assert all(not sent["materials"] for sent in vendor.drafts)
        assert (
            await api.get(f"/jobs/{receipt['data']['job_id']}", headers=headers[1])
        ).status_code == 404
        card = (await slots(api, headers[0], first[0], first[2]))[0]["card"]
        assert (await api.get(f"/cards/{card['id']}", headers=headers[1])).status_code == 404
        async with app.state.db.transaction(tenants["orgs"][1]) as session:
            assert await session.scalar(select(func.count()).select_from(CardGenerationRun)) == 0
            assert await session.scalar(select(func.count()).select_from(ResponseCard)) == 0
            assert (
                await session.scalar(
                    select(func.count())
                    .select_from(UsageRecord)
                    .where(UsageRecord.job_id == UUID(receipt["data"]["job_id"]))
                )
                == 0
            )


async def test_uncited_model_inputs_become_stale_and_block_confirmation_and_drafts(
    tenants, tmp_path, admin_engine
):
    async with drafting_client(tenants, tmp_path) as (api, app, headers, _, _):
        header = headers[0]
        task, _, extraction, requirements = await create_tender(
            api, app, header, tmp_path, confirmed=True
        )
        product, _, _, _, _ = await select_real_materials(api, header, task, tmp_path)
        receipt = await submit(
            api, header, task, extraction, requirement_ids=[requirements[2]["id"]]
        )
        terminal = await execute(api, app, header, receipt)
        assert terminal["data"]["status"] == "succeeded", terminal
        set_role(admin_engine, tenants["orgs"][0], tenants["users"][0], "technical")
        card = (await slots(api, header, task, extraction))[2]["card"]
        assert card["evidence"] == [] and card["content"]["response_kind"] == "commitment"
        card = await require_action(
            api, header, await require_action(api, header, card, "submit"), "confirm"
        )
        draft = await api.post(
            f"/tasks/{task}/drafts", headers=header, json={"extraction_job_id": extraction}
        )
        assert draft.status_code == 200, draft.text
        assembled = await execute(api, app, header, draft.json())
        assert assembled["data"]["status"] == "succeeded", assembled
        changed = await api.post(
            f"/resources/products/{product['product_id']}/revisions",
            headers=header,
            json={"expected_revision": 1, "data": PRODUCT_DATA | {"model": "New declared model"}},
        )
        assert changed.status_code == 200, changed.text
        assert (
            await api.post(
                f"/tasks/{task}/products",
                headers=header,
                json={"product_id": product["product_id"]},
            )
        ).status_code == 200
        shown = (await api.get(f"/cards/{card['id']}", headers=header)).json()["data"]
        assert shown["eligibility"] == "stale_material"
        old = (
            await api.get(f"/drafts/{assembled['data']['result']['draft_id']}", headers=header)
        ).json()["data"]
        assert (
            old["validity"] == "stale" and requirements[2]["id"] in old["invalidated_requirements"]
        )
        reopened = await require_action(
            api, header, shown, "reopen", reason="Selected material replaced."
        )
        pending = await require_action(api, header, reopened, "submit")
        denied = await card_action(api, header, pending, "confirm")
        assert (
            denied.status_code == 409 and denied.json()["data"]["error"]["code"] == "stale_material"
        )
        fresh = await api.post(
            f"/tasks/{task}/drafts", headers=header, json={"extraction_job_id": extraction}
        )
        assert fresh.status_code == 200, fresh.text
        refreshed = await execute(api, app, header, fresh.json())
        assert refreshed["data"]["status"] == "succeeded", refreshed
        current = (
            await api.get(f"/drafts/{refreshed['data']['result']['draft_id']}", headers=header)
        ).json()["data"]
        assert any(
            gap["requirement_id"] == requirements[2]["id"] and "stale_material" in gap["reasons"]
            for gap in current["gaps"]
        )


async def test_certificate_snapshot_redacts_text_skips_scan_pages_and_omits_other_pages(
    tenants, tmp_path
):
    async with drafting_client(tenants, tmp_path) as (api, app, headers, vendor, _):
        header = headers[0]
        task, _, extraction, _ = await create_tender(api, app, header, tmp_path)
        metadata = {
            "kind": "qualification",
            "name": "Synthetic certificate only",
            "number": "SYNTHETIC-CERT",
        }
        certificate = await api.post(
            "/resources/certificates", headers=header, json={"data": metadata}
        )
        assert certificate.status_code == 200, certificate.text
        certificate_id = certificate.json()["data"]["certificate_id"]
        import pymupdf

        with pymupdf.open() as pdf:
            first = pdf.new_page()
            first.insert_text(
                (40, 60), "SYNTHETIC CERTIFICATE INPUT ONLY\n" + "\n".join(SENSITIVE_LINES)
            )
            pdf.new_page()  # No text layer: must remain unavailable, with no OCR or image request.
            last = pdf.new_page()
            last.insert_text((40, 60), "UNSELECTED CERTIFICATE PAGE MUST NEVER BE SENT")
            content = pdf.tobytes()
        uploaded = await api.post(
            f"/resources/certificates/{certificate_id}/file-revisions",
            headers=header,
            data={"metadata": json.dumps({"expected_revision": 1, "data": metadata})},
            files={"file": ("synthetic.pdf", content, "application/pdf")},
        )
        assert uploaded.status_code == 200, uploaded.text
        selected = await api.post(
            f"/tasks/{task}/certificates", headers=header, json={"certificate_id": certificate_id}
        )
        assert selected.status_code == 200, selected.text
        source_ids = []
        for page in (1, 2):
            source = await api.post(
                f"/tasks/{task}/evidence-sources",
                headers=header,
                json={"task_certificate_id": selected.json()["data"]["id"], "page": page},
            )
            assert source.status_code == 200, source.text
            source_ids.append(source.json()["data"]["source"]["id"])
        preview = await submit(api, header, task, extraction, dry_run=True)
        assert any(
            warning == f"page_text_unavailable:{source_ids[1]}" for warning in preview["warnings"]
        )
        pages = [
            entry
            for entry in preview["data"]["input_manifest"]["materials"]
            if entry["kind"] == "certificate_pdf_page"
        ]
        assert (
            len(pages) == 1
            and pages[0]["page"] == 1
            and pages[0]["evidence_source_id"] == source_ids[0]
        )
        assert pages[0]["original_sha256"] == hashlib.sha256(content).hexdigest()
        receipt = await submit(api, header, task, extraction)
        terminal = await execute(api, app, header, receipt)
        assert terminal["data"]["status"] == "succeeded", terminal
        wire = json.dumps(vendor.drafts)
        assert all(
            value not in wire
            for value in [
                *SENSITIVE.values(),
                "UNSELECTED CERTIFICATE PAGE",
                "storage_key",
                "image_url",
                "base64",
                "org/",
            ]
        )
        assert "[REDACTED_" in wire
        assert all(
            material["page"] == 1
            for sent in vendor.drafts
            for material in sent["materials"]
            if material["kind"] == "certificate_pdf_page"
        )


async def test_docx_location_text_is_redacted_and_citations_stay_fixed(tenants, tmp_path):
    from io import BytesIO

    from docx import Document as WordDocument

    async with drafting_client(tenants, tmp_path) as (api, app, headers, vendor, _):
        header = headers[0]
        document = WordDocument()
        document.add_heading("Contact: Synthetic Person; Phone: 13800138000;", 1)
        document.add_paragraph("Synthetic obligation shall be performed within thirty days.")
        buffer = BytesIO()
        document.save(buffer)
        task = (await api.post("/tasks", headers=header, json={"name": "Synthetic Word"})).json()[
            "data"
        ]["id"]
        uploaded = await api.post(
            f"/tasks/{task}/documents",
            headers=header,
            files={"file": ("synthetic.docx", buffer.getvalue())},
        )
        assert uploaded.status_code == 200, uploaded.text
        document_id = uploaded.json()["data"]["id"]
        await run_job(api, app, header, document_id, "parse")
        # The extraction vendor still uses its normal API; only fixture synthesis changes.
        original_vendor = type(vendor).__call__

        async def word_vendor(self, request):
            body = json.loads(request.content)
            user = body["messages"][-1]["content"]
            if user.startswith("{"):
                return await original_vendor(self, request)
            blocks = re.findall(r'<block id="([^\"]+)">\n(.*?)\n</block>', user, re.S)
            items = [
                {
                    "category": "technical",
                    "starred": False,
                    "text": text,
                    "ref": ref,
                    "quote": text,
                    "condition": None,
                }
                for ref, text in blocks
                if "Synthetic obligation" in text
            ]
            return provider_reply("openai", items)

        from unittest.mock import patch

        with patch.object(DraftVendor, "__call__", word_vendor):
            extraction, status = await run_job(api, app, header, document_id, "extract")
        assert status["status"] == "succeeded", status
        receipt = await submit(api, header, task, extraction)
        terminal = await execute(api, app, header, receipt)
        assert terminal["data"]["status"] == "succeeded", terminal
        sent = vendor.drafts[0]["requirements"][0]
        assert sent["location"]["page"] is None and sent["location"]["block"]["block_id"]
        assert SENSITIVE["contact"] not in json.dumps(sent) and SENSITIVE[
            "phone"
        ] not in json.dumps(sent)
        assert "[REDACTED_" in json.dumps(sent)
        source = (await slots(api, header, task, extraction))[0]["source"]
        assert source["page"] is None and SENSITIVE["contact"] in source["location"]["label"]


async def test_platform_default_catalog_is_used_without_changing_extraction(
    tenants, tmp_path, admin_engine
):
    from app.core.config import Settings
    from app.models.entities import PlatformModel
    from sqlalchemy.orm import Session

    # Long catalog IDs exercise the 0018 adapter identity width as well as real resolution.
    catalog_id = "synthetic-drafting-catalog-long-id-12345"  # the 40-character maximum
    from conftest import seed_platform_credential

    await seed_platform_credential(
        Settings(data_dir=tmp_path),
        name="synthetic_unset",
        provider="openai",
        endpoint="https://vendor.example.test/v1",
    )
    with Session(admin_engine) as session, session.begin():
        session.add(
            PlatformModel(
                id=catalog_id,
                capability="llm_extract",
                provider="openai",
                model="synthetic-model",
                base_url="https://vendor.example.test/v1",
                credential="synthetic_unset",
                vendor_input_usd_per_mtok=1,
                vendor_output_usd_per_mtok=2,
                sale_input_per_mtok=1,
                sale_output_per_mtok=2,
                is_default=True,
                enabled=True,
                reasoning=[
                    {
                        "name": "low",
                        "request_options": {"reasoning_effort": "low"},
                        "batch_chars": 8000,
                    }
                ],
                default_reasoning="low",
                revision=1,
                updated_by="test@example.test",
            )
        )
    vendor = DraftVendor()
    app = create_app(
        Settings(data_dir=tmp_path, llm_provider="disabled"),
        queue=FakeQueue(),
        llm_transport=httpx.MockTransport(vendor),
    )
    async with (
        app.router.lifespan_context(app),
        httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test") as api,
    ):
        header = await login(api, tenants["orgs"][0], "a")
        async with app.state.db.transaction(tenants["orgs"][0]) as session:
            session.add(OrgBalance(org_id=tenants["orgs"][0], currency="USD", balance=10))
        task, _, extraction, requirements = await create_tender(api, app, header, tmp_path)
        preview = await submit(api, header, task, extraction, dry_run=True)
        assert (
            preview["data"]["platform_model_id"] == catalog_id
            and preview["data"]["model_revision"] == 1
        )
        assert preview["data"]["reasoning"] == "low"
        receipt = await submit(
            api, header, task, extraction, requirement_ids=[r["id"] for r in requirements[2:]]
        )
        terminal = await execute(api, app, header, receipt)
        assert terminal["data"]["status"] == "succeeded", terminal
        assert terminal["data"]["reasoning"] == "low"
        async with app.state.db.transaction(tenants["orgs"][0]) as session:
            run = await session.scalar(select(CardGenerationRun))
            assert run and len(run.adapter_version) > 40 and run.platform_model_id == catalog_id
            usages = list(
                (
                    await session.scalars(
                        select(UsageRecord).where(
                            UsageRecord.job_id == UUID(receipt["data"]["job_id"])
                        )
                    )
                ).all()
            )
            assert len(usages) == 1 and usages[0].platform_model_id == catalog_id
            assert usages[0].usd is not None and usages[0].charge is not None


async def test_live_token_grants_are_rechecked_before_vendor_and_result_reads(tenants, tmp_path):
    from app.models.entities import ApiToken

    async with drafting_client(tenants, tmp_path) as (api, app, headers, vendor, _):
        header = headers[0]
        task, _, extraction, _ = await create_tender(api, app, header, tmp_path)
        await select_real_materials(api, header, task, tmp_path)
        restricted = await token_header(
            api, header, scopes=["task:read", "job:read", "card:read", "card:generate"]
        )
        denied = await api.post(
            f"/tasks/{task}/cards/generations",
            headers=restricted,
            json={"extraction_job_id": extraction, "dry_run": True},
        )
        assert denied.status_code == 403 and not vendor.drafts
        agent = await token_header(api, header)
        receipt = await submit(api, agent, task, extraction)
        async with app.state.db.transaction(tenants["orgs"][0]) as session:
            await session.execute(update(ApiToken).values(revoked=True))
        terminal = await execute(api, app, header, receipt)
        assert (
            terminal["data"]["status"] == "failed"
            and terminal["data"]["error"]["code"] == "forbidden"
        )
        assert not vendor.drafts
        denied = await api.get(f"/jobs/{receipt['data']['job_id']}", headers=restricted)
        assert denied.status_code in {401, 403}


@pytest.mark.parametrize(
    "reported,expected,job_state",
    [
        ("claude-sonnet-4-5", "claude-sonnet-4-5", "succeeded"),
        ("Contact: PRIVATE MODEL METADATA ECHO", "unverified-model", "failed"),
    ],
)
async def test_drafting_actual_fallback_identity_and_unsafe_metadata(
    tenants, tmp_path, reported, expected, job_state, caplog
):
    async with drafting_client(tenants, tmp_path, "anthropic") as (api, app, headers, vendor, _):
        header = headers[0]
        task, _, extraction, requirements = await create_tender(api, app, header, tmp_path)

        async def response(sent):
            payload = provider_reply("anthropic", vendor.proposals(sent)).json()
            payload["model"] = reported
            return httpx.Response(200, json=payload)

        vendor.respond = response
        receipt = await submit(
            api, header, task, extraction, requirement_ids=[requirements[2]["id"]]
        )
        terminal = await execute(api, app, header, receipt)
        assert terminal["data"]["status"] == job_state, terminal
        if job_state == "failed":
            assert terminal["data"]["error"]["code"] == "invalid_provider_model"
            assert reported not in json.dumps(terminal) and reported not in caplog.text
        async with app.state.db.transaction(tenants["orgs"][0]) as session:
            rows = list(
                (
                    await session.scalars(
                        select(UsageRecord).where(
                            UsageRecord.job_id == UUID(receipt["data"]["job_id"])
                        )
                    )
                ).all()
            )
            assert len(rows) == 1 and rows[0].model == expected and rows[0].charge > 0
            if job_state == "failed":
                assert await session.scalar(select(func.count()).select_from(ResponseCard)) == 0
