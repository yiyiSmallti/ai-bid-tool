"""Parameter gap filling through document upload, real adapters, jobs and stored citations."""

import io
import json
import re

import httpx
import pymupdf
import pytest
from app.api.main import create_app
from app.providers.llm import AnthropicExtractor, OpenAICompatibleExtractor
from conftest import FakeQueue
from docx import Document
from test_api import run_job
from test_docx_extraction import item
from test_job_boundaries import session_for
from test_llm_providers import Vendor, anthropic_reply, not_json, settings_for, usage_rows

PARAMETERS = [
    "（2）触摸显示屏屏幕尺寸（对角）：≥23.8 英寸",
    "屏幕宽高比：16:9",
    "分辨率：≥1920x1080",
    "面板类型：IPS 技术",
    "亮度（典型值）：≥250 cd/m²",
    "对比度：≥1000:1",
]
CELL_TEXT = "；".join(PARAMETERS[:3]) + ";" + "\n".join(PARAMETERS[3:])


def parameter_word(text=CELL_TEXT):
    document = Document()
    document.add_heading("技术参数", level=1)
    document.add_paragraph("设备须提供操作手册。")
    document.add_table(rows=1, cols=1).cell(0, 0).text = text
    document.add_paragraph("★3.合同的终止：")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


async def extract_document(api, app, header, content, name="parameters.docx"):
    task = (await api.post("/tasks", headers=header, json={"name": "Parameter coverage"})).json()[
        "data"
    ]["id"]
    uploaded = await api.post(
        f"/tasks/{task}/documents", headers=header, files={"file": (name, content)}
    )
    assert uploaded.status_code == 200, uploaded.text
    document = uploaded.json()["data"]["id"]
    _, parsed = await run_job(api, app, header, document, "parse")
    assert parsed["status"] == "succeeded", parsed
    _, status = await run_job(api, app, header, document, "extract")
    rows = (await api.get(f"/tasks/{task}/requirements", headers=header)).json()["items"]
    return status, rows


def vendor_reply(provider, items):
    if provider == "anthropic":
        return anthropic_reply(items)
    return httpx.Response(
        200,
        json={
            "model": "synthetic-model",
            "choices": [
                {
                    "finish_reason": "stop",
                    "message": {"content": json.dumps({"items": items}, ensure_ascii=False)},
                }
            ],
            "usage": {"prompt_tokens": 1200, "completion_tokens": 300},
        },
    )


def user_content(request):
    return json.loads(request.content)["messages"][-1]["content"]


@pytest.mark.parametrize("provider", ["anthropic", "openai"])
async def test_word_parameter_gaps_are_filled_and_each_call_is_recorded(
    provider, tenants, tmp_path
):
    first = [PARAMETERS[0], PARAMETERS[3]]
    missing = [p for p in PARAMETERS if p not in first]
    vendor = Vendor(
        vendor_reply(provider, [item("t1r1c1", p) for p in first]),
        vendor_reply(provider, [item("t1r1c1", p) for p in missing]),
    )
    settings = settings_for(
        tmp_path, provider, llm_input_usd_per_mtok=4, llm_output_usd_per_mtok=20
    )
    adapter = AnthropicExtractor if provider == "anthropic" else OpenAICompatibleExtractor
    app = create_app(settings, llm=adapter(settings, vendor.transport()), queue=FakeQueue())
    async with app.router.lifespan_context(app), session_for(app, tenants) as (api, header):
        status, rows = await extract_document(api, app, header, parameter_word())
        usages = await usage_rows(app, tenants)

    assert status["status"] == "succeeded", status
    assert status["result"]["gap_fill"] == {
        "segments": 4,
        "calls": 1,
        "added": 4,
        "remaining": 0,
    }
    assert status["result"]["created"] == 6
    assert {r["source"]["quote"] for r in rows} == set(PARAMETERS)
    assert all(r["source"]["location"]["block_id"] == "t1r1c1" for r in rows)
    assert len(usages) == 2 and all(u.tokens == 1500 for u in usages)
    assert status["result"]["cost"]["llm_tokens"] == 3000
    assert status["result"]["cost"]["usd"] == pytest.approx(0.0216)
    assert float(sum(u.usd for u in usages)) == pytest.approx(0.0216)
    assert len(vendor.requests) == 2
    gap = user_content(vendor.requests[1])
    assert '<block id="t1r1c1">' in gap
    assert all(p in gap for p in missing) and all(p not in gap for p in first)
    assert "<page" not in gap and '<block id="p2">' not in gap


async def test_gap_fill_rejects_invented_and_joined_quotes_and_counts_only_new_saved_items(
    tenants, tmp_path
):
    source = "内存：≥16 GB；容量：≥512 GB；功率：≤100 W"
    first = item("t1r1c1", "容量：≥512 GB")
    vendor = Vendor(
        anthropic_reply([first]),
        anthropic_reply(
            [
                item("t1r1c1", "内存：≥16 GB"),
                item("t1r1c1", "内存：≥16 GB"),  # duplicate must not inflate added
                first,  # absent from this request but present in the full stored block
                item("t1r1c1", "功率：≤50 W"),
                item("t1r1c1", "内存：≥16 GB\n功率：≤100 W"),
            ]
        ),
    )
    settings = settings_for(tmp_path, "anthropic")
    app = create_app(
        settings, llm=AnthropicExtractor(settings, vendor.transport()), queue=FakeQueue()
    )
    async with app.router.lifespan_context(app), session_for(app, tenants) as (api, header):
        status, rows = await extract_document(api, app, header, parameter_word(source))
        assert len(await usage_rows(app, tenants)) == 2
    assert status["status"] == "succeeded", status
    assert status["result"]["gap_fill"] == {
        "segments": 2,
        "calls": 1,
        "added": 1,
        "remaining": 1,
    }
    assert {r["source"]["quote"] for r in rows} == {"内存：≥16 GB", "容量：≥512 GB"}
    assert [r["reason"] for r in status["result"]["rejected"]] == [
        "quote_not_at_position",
        "quote_not_at_position",
    ]
    assert "容量：≥512 GB" not in user_content(vendor.requests[1])
    assert len(vendor.requests) == 2  # no recursive gap-filling round


@pytest.mark.parametrize("complete", [True, False])
async def test_no_gap_call_when_specs_are_covered_or_no_specs_exist(complete, tenants, tmp_path):
    source = "★（1）内存：≥16 GB；屏幕比例：16:9" if complete else "设备资料；3.合同的终止："
    quotes = ["内存：≥16 GB", "屏幕比例：16:9"] if complete else []
    vendor = Vendor(anthropic_reply([item("t1r1c1", q) for q in quotes]))
    settings = settings_for(tmp_path, "anthropic")
    app = create_app(
        settings, llm=AnthropicExtractor(settings, vendor.transport()), queue=FakeQueue()
    )
    async with app.router.lifespan_context(app), session_for(app, tenants) as (api, header):
        status, rows = await extract_document(api, app, header, parameter_word(source))
        assert len(await usage_rows(app, tenants)) == 1
    assert status["status"] == "succeeded", status
    assert len(vendor.requests) == 1
    assert status["result"]["gap_fill"] == {
        "segments": 0,
        "calls": 0,
        "added": 0,
        "remaining": 0,
    }
    assert len(rows) == len(quotes)  # the starred heading must not be added by the rule


async def test_uncited_first_pass_does_not_hide_gaps_in_another_block(tenants, tmp_path):
    source = "内存：≥16 GB；容量：≥512 GB"
    vendor = Vendor(
        anthropic_reply([item("p2", source)]),  # source is in the cell, not this paragraph
        anthropic_reply([item("t1r1c1", q) for q in source.split("；")]),
    )
    settings = settings_for(tmp_path, "anthropic")
    app = create_app(
        settings, llm=AnthropicExtractor(settings, vendor.transport()), queue=FakeQueue()
    )
    async with app.router.lifespan_context(app), session_for(app, tenants) as (api, header):
        status, rows = await extract_document(api, app, header, parameter_word(source))
    assert status["status"] == "succeeded", status
    assert status["result"]["gap_fill"] == {
        "segments": 2,
        "calls": 1,
        "added": 2,
        "remaining": 0,
    }
    assert len(rows) == 2 and len(status["result"]["rejected"]) == 1


@pytest.mark.parametrize(
    "first,second",
    [
        ("3.5mm插孔：≥2个", "5mm插孔：≥2个"),
        ("内存：≥16 GB", "扩展内存：≥16 GB"),
    ],
)
async def test_parameter_coverage_uses_the_actual_segment_not_a_name_substring(
    first, second, tenants, tmp_path
):
    # Failure mode: one parameter name is contained in another parameter's name.
    source = first + "；" + second
    vendor = Vendor(
        anthropic_reply([item("t1r1c1", second)]),
        anthropic_reply([item("t1r1c1", first)]),
    )
    settings = settings_for(tmp_path, "anthropic")
    app = create_app(
        settings, llm=AnthropicExtractor(settings, vendor.transport()), queue=FakeQueue()
    )
    async with app.router.lifespan_context(app), session_for(app, tenants) as (api, header):
        status, rows = await extract_document(api, app, header, parameter_word(source))
    assert status["status"] == "succeeded", status
    assert status["result"]["gap_fill"] == {
        "segments": 1,
        "calls": 1,
        "added": 1,
        "remaining": 0,
    }
    assert {r["source"]["quote"] for r in rows} == set(source.split("；"))


@pytest.mark.parametrize("provider", ["anthropic", "openai"])
async def test_a_whole_list_quote_covers_no_segments_but_a_single_segment_does(
    provider, tenants, tmp_path
):
    source = "内存：≥16 GB；容量：≥512 GB；功率：≤100 W"
    whole_list = item("t1r1c1", source)
    whole_list["text"] = "内存至少 16 GB"
    covered = item("t1r1c1", "容量：≥512 GB")
    vendor = Vendor(
        vendor_reply(provider, [whole_list, covered]),
        vendor_reply(
            provider,
            [item("t1r1c1", "内存：≥16 GB"), item("t1r1c1", "功率：≤100 W")],
        ),
    )
    settings = settings_for(tmp_path, provider)
    adapter = AnthropicExtractor if provider == "anthropic" else OpenAICompatibleExtractor
    app = create_app(settings, llm=adapter(settings, vendor.transport()), queue=FakeQueue())
    async with app.router.lifespan_context(app), session_for(app, tenants) as (api, header):
        status, rows = await extract_document(api, app, header, parameter_word(source))

    assert status["status"] == "succeeded", status
    assert status["result"]["gap_fill"] == {
        "segments": 2,
        "calls": 1,
        "added": 2,
        "remaining": 0,
    }
    assert {row["source"]["quote"] for row in rows} == {
        source,
        "内存：≥16 GB",
        "容量：≥512 GB",
        "功率：≤100 W",
    }
    assert len(vendor.requests) == 2
    assert next(row for row in rows if row["source"]["quote"] == source)["text"] == "内存至少 16 GB"
    gap = user_content(vendor.requests[1])
    assert "内存：≥16 GB" in gap and "功率：≤100 W" in gap
    assert "容量：≥512 GB" not in gap


@pytest.mark.parametrize("failure", ["truncated", "malformed", "refused", "bad_ref"])
async def test_gap_call_outcomes_keep_all_billed_usage(failure, tenants, tmp_path):
    source = "内存：≥16 GB；容量：≥512 GB；功率：≤100 W"
    replies = {
        "truncated": anthropic_reply([], stop="max_tokens"),
        "malformed": not_json(),
        "refused": anthropic_reply([], stop="refusal"),
        "bad_ref": anthropic_reply([item("p999", "容量：≥512 GB")]),
    }
    recoverable = failure in {"truncated", "malformed"}
    vendor = Vendor(
        anthropic_reply([item("t1r1c1", "内存：≥16 GB")]),
        replies[failure],
        *(
            [
                anthropic_reply([item("t1r1c1", "容量：≥512 GB")]),
                anthropic_reply([item("t1r1c1", "功率：≤100 W")]),
            ]
            if recoverable
            else []
        ),
    )
    settings = settings_for(tmp_path, "anthropic", llm_concurrency=1)
    app = create_app(
        settings, llm=AnthropicExtractor(settings, vendor.transport()), queue=FakeQueue()
    )
    async with app.router.lifespan_context(app), session_for(app, tenants) as (api, header):
        status, rows = await extract_document(api, app, header, parameter_word(source))
        usages = await usage_rows(app, tenants)
    assert len(usages) == len(vendor.requests) == (4 if recoverable else 2)
    assert sum(u.tokens for u in usages) == (4506 if failure == "malformed" else len(usages) * 1500)
    if recoverable:
        assert status["status"] == "succeeded", status
        assert len(rows) == 3
        assert status["result"]["gap_fill"] == {
            "segments": 2,
            "calls": 3,
            "added": 2,
            "remaining": 0,
        }
        assert all('<block id="t1r1c1">' in user_content(r) for r in vendor.requests)
    elif failure == "bad_ref":
        assert status["status"] == "succeeded", status
        assert [row["source"]["quote"] for row in rows] == ["内存：≥16 GB"]
        assert status["result"]["gap_fill"] == {
            "segments": 2,
            "calls": 1,
            "added": 0,
            "remaining": 2,
        }
        assert status["result"]["rejected"] == [
            {
                "position": "p999",
                "quote": "容量：≥512 GB",
                "reason": "unknown_position",
            }
        ]
    else:
        assert status["status"] == "failed", status
        assert status["error"]["code"] == "provider_refused"
        assert rows == []


async def test_pdf_gaps_keep_page_numbers_and_do_not_borrow_coverage_from_other_pages(
    tenants, tmp_path
):
    with pymupdf.open() as document:
        for _ in range(2):
            document.new_page().insert_text((40, 60), "Memory: 16 GB;Storage: 512 GB")
        content = document.tobytes()
    vendor = Vendor(
        anthropic_reply([item("1", "Memory: 16 GB"), item("1", "Storage: 512 GB")]),
        anthropic_reply([item("2", "Memory: 16 GB"), item("2", "Storage: 512 GB")]),
    )
    settings = settings_for(tmp_path, "anthropic")
    app = create_app(
        settings, llm=AnthropicExtractor(settings, vendor.transport()), queue=FakeQueue()
    )
    async with app.router.lifespan_context(app), session_for(app, tenants) as (api, header):
        status, rows = await extract_document(api, app, header, content, "parameters.pdf")
    assert status["status"] == "succeeded", status
    assert status["result"]["gap_fill"] == {
        "segments": 2,
        "calls": 1,
        "added": 2,
        "remaining": 0,
    }
    assert sorted(r["source"]["page"] for r in rows) == [1, 1, 2, 2]
    assert re.findall(r'<page number="(\d+)">', user_content(vendor.requests[1])) == ["2"]
