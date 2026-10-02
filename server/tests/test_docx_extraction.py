"""Word tenders end to end: upload, located parsing, block-cited extraction and its failures."""

import io
import json
import re

import httpx
import pytest
from app.api.main import create_app
from app.core.config import Settings
from app.providers.llm import AnthropicExtractor
from conftest import FakeQueue
from docx import Document
from fakes import FakeLLM
from sqlalchemy import text
from sqlalchemy.orm import Session
from test_api import run_job
from test_job_boundaries import session_for
from test_llm_providers import Vendor, anthropic_reply, settings_for


def word_tender() -> bytes:
    document = Document()
    document.add_heading("第一章 总则", level=1)  # p1
    document.add_paragraph("投标人须具备有效的营业执照。")  # p2
    document.add_heading("一、技术参数", level=2)  # p3
    document.add_paragraph("服务器内存不低于 64GB。")  # p4
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "参数", "要求"
    table.cell(1, 0).text, table.cell(1, 1).text = "CPU", "★ 核心数不少于 32 核"
    document.add_paragraph("投标文件须注明“响应内容”及具体数值。")  # p5, curly quotes
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


async def upload_word(api, header) -> tuple[str, str]:
    task = (await api.post("/tasks", headers=header, json={"name": "Word tender"})).json()["data"][
        "id"
    ]
    response = await api.post(
        f"/tasks/{task}/documents", headers=header, files={"file": ("tender.docx", word_tender())}
    )
    assert response.status_code == 200, response.text
    return task, response.json()["data"]["id"]


def item(ref, quote, category="technical", starred=False):
    return {
        "category": category,
        "starred": starred,
        "text": quote,
        "ref": ref,
        "quote": quote,
        "condition": None,
    }


async def test_word_tender_parses_into_located_chunks(tenants, tmp_path):
    app = create_app(Settings(data_dir=tmp_path), llm=FakeLLM(), queue=FakeQueue())
    async with app.router.lifespan_context(app), session_for(app, tenants) as (api, header):
        _, document = await upload_word(api, header)
        _, status = await run_job(api, app, header, document, "parse")
        assert status["status"] == "succeeded", status
        assert status["result"]["sections"] == 2 and status["result"]["pages"] == 0
        meta = (await api.get(f"/documents/{document}", headers=header)).json()["data"]
        assert (meta["citation_mode"], meta["page_count"]) == ("block", None)
        chunks = (await api.get(f"/documents/{document}/chunks", headers=header)).json()["items"]
        assert [(c["seq"], c["page"], c["citation_verified"]) for c in chunks] == [
            (1, None, True),
            (2, None, True),
        ]
        labels = [b["label"] for c in chunks for b in c["blocks"]]
        assert "第一章 总则 > 第 1 段" in labels
        assert "第一章 总则 > 一、技术参数 > 表 1 第 2 行第 2 列" in labels

        # The fake provider cites each chunk's last block through the same pipeline.
        _, extracted = await run_job(api, app, header, document, "extract")
        assert extracted["status"] == "succeeded", extracted


async def test_word_extraction_cites_blocks_and_adds_starred_cells(tenants, tmp_path):
    vendor = Vendor(
        anthropic_reply([item("p2", "投标人须具备有效的营业执照。", "qualification")]),
        anthropic_reply(
            [item("p4", "服务器内存不低于 64GB。"), item("t1r2c2", "核心数不少于 32 核")]
        ),
    )
    settings = settings_for(tmp_path, "anthropic")
    # A tiny batch budget (below the configurable minimum) sends each section separately.
    tiny = settings.model_copy(update={"llm_batch_chars": 40, "llm_concurrency": 1})
    app = create_app(
        settings,
        llm=AnthropicExtractor(tiny, transport=vendor.transport()),
        queue=FakeQueue(),
    )
    async with app.router.lifespan_context(app), session_for(app, tenants) as (api, header):
        task, document = await upload_word(api, header)
        await run_job(api, app, header, document, "parse")
        _, status = await run_job(api, app, header, document, "extract")
        assert status["status"] == "succeeded", status
        rows = (await api.get(f"/tasks/{task}/requirements", headers=header)).json()["items"]

    sources = [
        (
            r["source"]["page"],
            r["source"]["location"]["block_id"],
            r["source"]["quote"],
            r["starred"],
        )
        for r in rows
    ]
    assert sources == [
        (None, "p2", "投标人须具备有效的营业执照。", False),
        (None, "p4", "服务器内存不低于 64GB。", False),
        (None, "t1r2c2", "核心数不少于 32 核", True),
    ]
    cell = rows[2]["source"]["location"]
    assert (cell["kind"], cell["table"], cell["row"], cell["column"]) == ("cell", 1, 2, 2)
    assert cell["label"] == "第一章 总则 > 一、技术参数 > 表 1 第 2 行第 2 列"
    first = json.loads(vendor.requests[0].content)["messages"][0]["content"]
    assert '<block id="p2">' in first and '<section path="第一章 总则">' in first
    assert "<page" not in first


@pytest.mark.parametrize(
    "items,code",
    [
        # Every item fails, so nothing is saved. The first quote comes from another block in
        # the same chunk, so chunk-level checks cannot pass it.
        ([item("p4", "核心数不少于 32 核")], "invalid_citation"),
        ([item("p4", "服务器内存不低于 64GB。参数")], "invalid_citation"),  # spans two blocks
        (
            [item("p999", "投标人须具备有效的营业执照。")],
            "invalid_citation",
        ),  # unknown block
    ],
)
async def test_word_citations_must_sit_inside_the_cited_block(items, code, tenants, tmp_path):
    vendor = Vendor(
        anthropic_reply(items),
        *([anthropic_reply([])] if code == "invalid_citation" else []),
    )
    settings = settings_for(tmp_path, "anthropic")
    app = create_app(
        settings, llm=AnthropicExtractor(settings, transport=vendor.transport()), queue=FakeQueue()
    )
    async with app.router.lifespan_context(app), session_for(app, tenants) as (api, header):
        task, document = await upload_word(api, header)
        await run_job(api, app, header, document, "parse")
        _, status = await run_job(api, app, header, document, "extract")
        assert (status["status"], status["error"]["code"]) == ("failed", code)
        assert (await api.get(f"/tasks/{task}/requirements", headers=header)).json()["items"] == []
        if items[0]["ref"] == "p999":
            assert status["result"]["rejected"] == [
                {
                    "position": "p999",
                    "quote": "投标人须具备有效的营业执照。",
                    "reason": "unknown_position",
                }
            ]


async def test_reparse_replaces_the_old_unverified_word_chunk(tenants, tmp_path, admin_engine):
    app = create_app(Settings(data_dir=tmp_path), llm=FakeLLM(), queue=FakeQueue())
    async with app.router.lifespan_context(app), session_for(app, tenants) as (api, header):
        task, document = await upload_word(api, header)
        org = tenants["orgs"][0]
        with Session(admin_engine) as session, session.begin():
            session.execute(
                text(
                    "INSERT INTO chunks (id, org_id, task_id, document_id, page, seq, text, ocr, citation_verified) "
                    "VALUES (gen_random_uuid(), :o, :t, :d, 1, 1, 'legacy flat text', false, false)"
                ),
                {"o": org, "t": task, "d": document},
            )
        await run_job(api, app, header, document, "parse")
        chunks = (await api.get(f"/documents/{document}/chunks", headers=header)).json()["items"]
        assert [c["text"] for c in chunks if c["text"] == "legacy flat text"] == []
        assert all(c["blocks"] and c["page"] is None for c in chunks)


async def test_straight_quotes_match_curly_source_quotes(tenants, tmp_path):
    model_quote = '投标文件须注明"响应内容"及具体数值。'
    source_quote = "投标文件须注明“响应内容”及具体数值。"
    vendor = Vendor(
        anthropic_reply([item("p5", model_quote)]),
        anthropic_reply([]),
    )
    settings = settings_for(tmp_path, "anthropic")
    app = create_app(
        settings, llm=AnthropicExtractor(settings, transport=vendor.transport()), queue=FakeQueue()
    )
    async with app.router.lifespan_context(app), session_for(app, tenants) as (api, header):
        task, document = await upload_word(api, header)
        await run_job(api, app, header, document, "parse")
        _, status = await run_job(api, app, header, document, "extract")
        assert status["status"] == "succeeded", status
        rows = (await api.get(f"/tasks/{task}/requirements", headers=header)).json()["items"]
        assert [row["source"]["location"]["block_id"] for row in rows] == [
            "t1r2c2",
            "p5",
        ]  # source order
        normalized = rows[1]
        assert normalized["source"]["quote"] == source_quote
        assert normalized["model_quote"] == model_quote
        assert rows[0]["model_quote"] is None  # source rule, not model output


async def test_unknown_word_ref_is_rejected_without_discarding_a_verified_item(tenants, tmp_path):
    vendor = Vendor(
        anthropic_reply(
            [
                item("p2", "投标人须具备有效的营业执照。", "qualification"),
                item("p999", "投标人须具备虚构位置的证明。", "qualification"),
            ]
        ),
        anthropic_reply([]),
    )
    settings = settings_for(tmp_path, "anthropic")
    app = create_app(
        settings, llm=AnthropicExtractor(settings, transport=vendor.transport()), queue=FakeQueue()
    )
    async with app.router.lifespan_context(app), session_for(app, tenants) as (api, header):
        task, document = await upload_word(api, header)
        await run_job(api, app, header, document, "parse")
        _, status = await run_job(api, app, header, document, "extract")
        rows = (await api.get(f"/tasks/{task}/requirements", headers=header)).json()["items"]

    assert status["status"] == "succeeded", status
    assert [row["source"]["location"]["block_id"] for row in rows] == ["p2", "t1r2c2"]
    assert status["result"]["rejected"] == [
        {
            "position": "p999",
            "quote": "投标人须具备虚构位置的证明。",
            "reason": "unknown_position",
        }
    ]


async def test_star_markers_apply_per_segment_and_rule_add_only_missing_starred_segments(
    tenants, tmp_path
):
    cell_text = (
        "★内存：≥16 GB，支持 ECC；容量：≥512 GB；★功率：≤100 W\n"
        "★接口数量：≥2 个；保修期：≥3 年\n"
        "★3.合同的终止："
    )
    document = Document()
    document.add_heading("技术参数", level=1)
    document.add_table(rows=1, cols=1).cell(0, 0).text = cell_text
    output = io.BytesIO()
    document.save(output)
    vendor = Vendor(
        anthropic_reply(
            [
                item("t1r1c1", "内存：≥16 GB"),
                item("t1r1c1", "支持 ECC"),
                item("t1r1c1", "容量：≥512 GB"),
                item("t1r1c1", "接口数量：≥2 个"),
                item("t1r1c1", "保修期：≥3 年"),
            ]
        ),
        # The first starred segment is returned as a whole; the missing power segment
        # is deliberately left for the deterministic starred-source rule.
        anthropic_reply([item("t1r1c1", "★内存：≥16 GB，支持 ECC")]),
    )
    settings = settings_for(tmp_path, "anthropic")
    app = create_app(
        settings, llm=AnthropicExtractor(settings, transport=vendor.transport()), queue=FakeQueue()
    )
    async with app.router.lifespan_context(app), session_for(app, tenants) as (api, header):
        task = (await api.post("/tasks", headers=header, json={"name": "Star segments"})).json()[
            "data"
        ]["id"]
        uploaded = await api.post(
            f"/tasks/{task}/documents",
            headers=header,
            files={"file": ("stars.docx", output.getvalue())},
        )
        document_id = uploaded.json()["data"]["id"]
        await run_job(api, app, header, document_id, "parse")
        _, status = await run_job(api, app, header, document_id, "extract")
        rows = (await api.get(f"/tasks/{task}/requirements", headers=header)).json()["items"]

    assert status["status"] == "succeeded", status
    assert status["result"]["gap_fill"] == {
        "segments": 2,
        "calls": 1,
        "added": 1,
        "remaining": 0,
    }
    by_text = {row["text"]: row for row in rows}
    assert by_text["内存：≥16 GB"]["starred"] is True
    assert by_text["支持 ECC"]["starred"] is True
    assert by_text["★内存：≥16 GB，支持 ECC"]["starred"] is True
    assert by_text["容量：≥512 GB"]["starred"] is False
    assert by_text["接口数量：≥2 个"]["starred"] is True
    assert by_text["保修期：≥3 年"]["starred"] is False
    power = next(row for row in rows if "功率：≤100 W" in row["source"]["quote"])
    assert power["starred"] is True and power["model_quote"] is None
    assert all("合同的终止" not in row["text"] for row in rows)


async def test_uncited_items_are_dropped_and_reported_while_the_rest_are_saved(tenants, tmp_path):
    vendor = Vendor(
        anthropic_reply(
            [
                item("p2", "投标人须具备有效的营业执照。", "qualification"),
                item("p4", "核心数不少于 32 核"),  # quote belongs to the table cell
            ]
        ),
        anthropic_reply([]),
    )
    settings = settings_for(tmp_path, "anthropic")
    app = create_app(
        settings, llm=AnthropicExtractor(settings, transport=vendor.transport()), queue=FakeQueue()
    )
    async with app.router.lifespan_context(app), session_for(app, tenants) as (api, header):
        task, document = await upload_word(api, header)
        await run_job(api, app, header, document, "parse")
        _, status = await run_job(api, app, header, document, "extract")
        assert status["status"] == "succeeded", status
        rows = (await api.get(f"/tasks/{task}/requirements", headers=header)).json()["items"]

    # The ★ cell the model mis-cited is still added from the source by the rule.
    assert [r["source"]["location"]["block_id"] for r in rows] == ["p2", "t1r2c2"]
    assert status["result"]["rejected"] == [
        {
            "position": "第一章 总则 > 一、技术参数 > 第 1 段",
            "quote": "核心数不少于 32 核",
            "reason": "quote_not_at_position",
        }
    ]
    assert any("1 extracted requirements were not saved" in w for w in status["result"]["warnings"])


async def test_truncated_word_batches_are_halved_until_the_output_fits(tenants, tmp_path):
    answers = {
        "p2": item("p2", "投标人须具备有效的营业执照。", "qualification"),
        "p4": item("p4", "服务器内存不低于 64GB。"),
        "p5": item("p5", "投标文件须注明“响应内容”及具体数值。"),
    }
    seen: list[list[str]] = []

    def vendor(request: httpx.Request) -> httpx.Response:
        content = json.loads(request.content)["messages"][0]["content"]
        blocks = re.findall(r'<block id="([^"]+)">', content)
        seen.append(blocks)
        # Any request with more than two blocks overflows the output limit.
        if len(blocks) > 2:
            return anthropic_reply([], stop="max_tokens")
        return anthropic_reply([answers[b] for b in blocks if b in answers])

    settings = settings_for(tmp_path, "anthropic")
    app = create_app(
        settings,
        llm=AnthropicExtractor(settings, transport=httpx.MockTransport(vendor)),
        queue=FakeQueue(),
    )
    async with app.router.lifespan_context(app), session_for(app, tenants) as (api, header):
        task, document = await upload_word(api, header)
        await run_job(api, app, header, document, "parse")
        _, status = await run_job(api, app, header, document, "extract")
        assert status["status"] == "succeeded", status
        rows = (await api.get(f"/tasks/{task}/requirements", headers=header)).json()["items"]

    assert [r["source"]["location"]["block_id"] for r in rows] == ["p2", "p4", "t1r2c2", "p5"]
    # 9 blocks overflow; then the sections (2 and 7), then halves down to two blocks.
    # One final gap batch asks for the uncovered starred cell; the rule still adds it
    # when the vendor leaves that batch empty.
    assert sorted(map(len, seen), reverse=True) == [9, 7, 4, 3, 2, 2, 2, 2, 1, 1]
    assert status["result"]["rejected"] == []


async def test_a_long_cell_is_split_by_lines_and_cited_as_the_whole_cell(tenants, tmp_path):
    lines = [f"{n}.功能项 {n} 须支持在线办理。" for n in range(1, 7)]
    document = Document()
    document.add_heading("第三章 采购需求", level=1)
    document.add_table(rows=1, cols=1).cell(0, 0).text = "\n".join(lines)
    output = io.BytesIO()
    document.save(output)
    seen: list[int] = []

    def vendor(request: httpx.Request) -> httpx.Response:
        content = json.loads(request.content)["messages"][0]["content"]
        cell = re.search(r'<block id="t1r1c1">\n(.*?)\n</block>', content, re.S)
        shown = cell.group(1).splitlines() if cell else []
        seen.append(len(shown))
        # More than two requirement lines overflow the output limit.
        if len(shown) > 2:
            return anthropic_reply([], stop="max_tokens")
        return anthropic_reply([item("t1r1c1", line) for line in shown])

    settings = settings_for(tmp_path, "anthropic")
    app = create_app(
        settings,
        llm=AnthropicExtractor(settings, transport=httpx.MockTransport(vendor)),
        queue=FakeQueue(),
    )
    async with app.router.lifespan_context(app), session_for(app, tenants) as (api, header):
        task = (await api.post("/tasks", headers=header, json={"name": "Long cell"})).json()[
            "data"
        ]["id"]
        uploaded = await api.post(
            f"/tasks/{task}/documents",
            headers=header,
            files={"file": ("tender.docx", output.getvalue())},
        )
        document_id = uploaded.json()["data"]["id"]
        await run_job(api, app, header, document_id, "parse")
        _, status = await run_job(api, app, header, document_id, "extract")
        assert status["status"] == "succeeded", status
        rows = (await api.get(f"/tasks/{task}/requirements", headers=header)).json()["items"]

    assert sorted(r["source"]["quote"] for r in rows) == sorted(lines)
    assert {r["source"]["location"]["label"] for r in rows} == {
        "第三章 采购需求 > 表 1 第 1 行第 1 列"
    }
    # Heading and cell; the heading alone (no cell) and the cell alone; then the cell's
    # lines three and three, then two and one of each half.
    assert sorted(seen, reverse=True) == [6, 6, 3, 3, 2, 2, 1, 1, 0]
