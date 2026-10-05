"""Simulated product proposals through the real API, worker, search, fetch broker and exports.

Failure modes enumerated before implementation:
- Only a human with resource:write and task:resource may simulate; tokens, bidders and
  viewers are refused; another org's task is a 404.
- The dry run sends nothing to the model, search or fetch, and a changed preview is refused.
- Pages on a proposed vendor's domain are fetched before others; a failed fetch moves on
  to the next vendor; a catalog page is followed one level only through links it really
  contains; a product name the page does not contain is never stored, so the item gets no
  product; software and service items never get one.
- A malformed proposal batch is split and retried instead of failing the run.
- With Perplexity configured, search is narrowed to the proposed domain and carries the
  key only to that service; a page that cannot be fetched is read from the search extract,
  quotes are checked against that extract, and the outcome names it as the source.
- A model quote that is not verbatim in the fetched page is dropped before storage.
- Kept products and statements are selected on the task and registered as simulated; a
  later human revision keeps the mark, including after rebinding to an ordinary product.
- A confirmed card citing a simulated statement blocks a final section and needs an
  acknowledgment in a review copy. Either the feature or its product suffices; unrelated
  simulated resources must not block ordinary declarations.
- The generic job reader does not expose the submission.
"""

import io
import json
from decimal import Decimal
from urllib.parse import urlsplit
from uuid import UUID, uuid4

import httpx
import pytest
from app.models.entities import BalanceEntry, OrgBalance, SimulatedResource
from app.providers.llm import OpenAICompatibleExtractor
from conftest import FakeQueue, credential_app
from docx import Document
from sqlalchemy import select
from test_exports import draft, setup_template
from test_llm_providers import settings_for
from test_response_cards import login, run_document_job, set_role

PAGE = (
    "<html><head><script>ignore()</script></head><body><h1>SP-100 智能打印终端规格</h1>"
    "<p>内存：8GB DDR4</p><p>打印：支持A4双面自动打印</p>"
    + "<p>"
    + "产品适用于校园自助打印与文印服务场景，" * 12
    + "</p></body></html>"
)
CATALOG = (
    "<html><body><p>"
    + "终端产品目录，" * 40
    + '</p><a href="/sp100">SP-100 打印终端</a><a href="/about">关于我们</a></body></html>'
)
OTHER_PAGE = "<html><body><p>" + "交换机系列产品介绍，" * 40 + "</p></body></html>"


def tender_docx() -> bytes:
    document = Document()
    document.add_paragraph("第三章 采购需求")
    table = document.add_table(rows=4, cols=3)
    for row, cells in enumerate(
        [
            ("序号", "名称", "技术参数"),
            ("1", "智能打印终端", "内存不少于8GB，支持A4双面打印"),
            ("2", "学生管理系统定制开发", "支持学生信息管理与统计"),
            ("3", "接入交换机", "端口不少于24个千兆电口"),
        ]
    ):
        for column, text in enumerate(cells):
            table.cell(row, column).text = text
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


class Vendor:
    """Synthetic model replies: extraction by block, proposals by item name, quotes."""

    def __init__(self):
        self.kinds: list[str] = []

    @staticmethod
    def reply(payload: dict) -> httpx.Response:
        return httpx.Response(
            200,
            json={
                "model": "synthetic-model",
                "choices": [
                    {
                        "finish_reason": "stop",
                        "message": {"content": json.dumps(payload, ensure_ascii=False)},
                    }
                ],
                "usage": {"prompt_tokens": 500, "completion_tokens": 50},
            },
        )

    def __call__(self, request: httpx.Request) -> httpx.Response:
        body = json.loads(request.content)
        system, user = body["messages"][0]["content"], body["messages"][-1]["content"]
        if "拟投产品的模拟助手" in system:
            self.kinds.append("propose")
            vendors = {
                "打印": [
                    ("BadVendor", "bad.example", "打印终端"),
                    ("SynVendor", "www.vendor.example", "打印终端"),
                ],
                "交换机": [("SynNet", "net.example", "交换机")],
            }
            asked = json.loads(user)["items"]
            if len(asked) > 2:
                return self.reply({"items": "not a list"})
            items = []
            for item in asked:
                found = next((rows for key, rows in vendors.items() if key in item["name"]), [])
                items.append(
                    {
                        "item_key": item["item_key"],
                        "kind": "hardware" if found else "software",
                        "candidates": [
                            {"vendor": name, "domain": domain, "query": f"{name} {query}"}
                            for name, domain, query in found
                        ],
                    }
                )
            return self.reply({"items": items})
        if "参数摘录助手" in system:
            self.kinds.append("quote")
            sent = json.loads(user)
            page = sent["page_text"]
            if "终端产品目录" in page:
                # Only a link the page really contains may be followed.
                links = [row["url"] for row in sent["links"] if "sp100" in row["url"]]
                return self.reply(
                    {
                        "model": "",
                        "parameters": [],
                        "next_urls": ["https://vendor.example/invented", *links],
                    }
                )
            # The switch page never names SW-24, so that remembered model must be refused.
            return self.reply(
                {
                    "next_urls": [],
                    "model": "SP-100" if "SP-100" in page else "SW-24",
                    "parameters": [
                        {"label": "内存", "quote": "内存：8GB DDR4"},
                        {"label": "臆造", "quote": "内存：16GB"},
                    ],
                }
            )
        self.kinds.append("extract")
        items = []
        for ref, text in __import__("re").findall(
            r'<block id="([^"]+)">\n(.*?)\n</block>', user, 2
        ):
            if "不少于" in text or "支持" in text:
                items.append(
                    {
                        "category": "technical",
                        "starred": False,
                        "text": text,
                        "ref": ref,
                        "quote": text,
                        "condition": None,
                    }
                )
        return self.reply({"items": items})


FETCHED: list[str] = []
SEARCH_RESULTS = {
    "BadVendor": ["https://bad.example/terminal"],
    "SynVendor": ["https://market.example/sp100", "https://vendor.example/catalog"],
    "SynNet": ["https://net.example/sw24"],
}


def search_handler(request: httpx.Request) -> httpx.Response:
    query = request.url.params["q"]
    urls = next((rows for name, rows in SEARCH_RESULTS.items() if name in query), [])
    return httpx.Response(
        200,
        json={
            "results": [{"url": url, "title": query, "engines": ["synthetic"]} for url in urls],
            "unresponsive_engines": [],
        },
    )


# Perplexity extracts the page our own fetch cannot reach.
EXTRACTS = {
    "https://bad.example/terminal": "SP-100 智能打印终端规格\n内存：8GB DDR4\n打印：支持A4双面自动打印\n"
    + "产品适用于校园自助打印与文印服务场景，" * 12
}
PERPLEXITY_REQUESTS: list[dict] = []


def perplexity_handler(request: httpx.Request) -> httpx.Response:
    assert str(request.url) == "https://api.perplexity.ai/search"
    assert request.headers["authorization"] == "Bearer synthetic-pplx-key"
    body = json.loads(request.content)
    PERPLEXITY_REQUESTS.append(body)
    urls = next((rows for name, rows in SEARCH_RESULTS.items() if name in body["query"]), [])
    domains = body.get("search_domain_filter")
    if domains:
        urls = [url for url in urls if urlsplit(url).hostname in domains]
    return httpx.Response(
        200,
        json={
            "results": [
                {"url": url, "title": body["query"], "snippet": EXTRACTS.get(url, "")}
                for url in urls
            ]
        },
    )


def fetch_handler(request: httpx.Request) -> httpx.Response:
    FETCHED.append(request.headers["host"] + request.url.path)
    if request.headers["host"] == "bad.example":
        return httpx.Response(503)
    page = {"/sp100": PAGE, "/catalog": CATALOG}.get(request.url.path, OTHER_PAGE)
    return httpx.Response(
        200, content=page.encode(), headers={"content-type": "text/html; charset=utf-8"}
    )


async def public_dns(host):
    return ("93.184.216.34",)


@pytest.mark.parametrize("provider", ["searxng", "perplexity"])
@pytest.mark.parametrize(
    "material",
    ["feature", "rebound_feature", "linked_feature", "product", "normal_feature", "normal_product"],
)
async def test_simulation_records_only_verbatim_marked_parameters(
    tenants, tmp_path, admin_engine, monkeypatch, provider, material
):
    FETCHED.clear()
    PERPLEXITY_REQUESTS.clear()
    policy = tmp_path / "open-policy.json"
    policy.write_text(
        json.dumps(
            {
                "policies": [{"revision": "dev-open-v1", "open_public_https": True}],
                "revoked_revisions": [],
            }
        )
    )
    monkeypatch.setenv("BID_SANDBOX_POLICY_FILE", str(policy))
    monkeypatch.setenv("BID_SANDBOX_DEV_OPEN_EGRESS", "1")
    monkeypatch.setenv("BID_SANDBOX_FETCH_QUOTA", str(tmp_path / "quota.sqlite3"))
    vendor = Vendor()
    search = (
        {"search_url": "https://search.test", "search_provider": "searxng"}
        if provider == "searxng"
        else {"search_provider": "perplexity"}
    )
    if provider == "perplexity":
        from app.schemas.platform_credentials import ResolvedCredential, ServiceResolveTarget
        from pydantic import SecretStr

        selected = ServiceResolveTarget(service="vendor_search", credential_id=uuid4())

        class SearchResolver:
            def __init__(self, settings):
                pass

            async def select_service(self, service):
                return selected

            async def resolve_for_call(self, target):
                return ResolvedCredential(
                    selected.credential_id,
                    1,
                    1,
                    "perplexity",
                    "https://api.perplexity.ai",
                    SecretStr("synthetic-pplx-key"),
                )

        monkeypatch.setattr(
            "app.services.platform_credentials.PlatformCredentialResolver", SearchResolver
        )
    settings = settings_for(tmp_path, "openai", **search)
    llm = OpenAICompatibleExtractor(
        settings,
        httpx.MockTransport(vendor),
        platform_model_id="sim-test",
        sale_usd_per_mtok=(1, 1),
    )
    llm.model_revision = 1
    app = await credential_app(settings, llm=llm, queue=FakeQueue())
    app.state.processor.search_transport = httpx.MockTransport(
        search_handler if provider == "searxng" else perplexity_handler
    )
    app.state.processor.sandbox_fetch_transport = httpx.MockTransport(fetch_handler)
    app.state.processor.sandbox_resolver = public_dns
    async with (
        app.router.lifespan_context(app),
        httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test") as api,
    ):
        for org in tenants["orgs"]:
            async with app.state.db.transaction(org) as session:
                session.add(OrgBalance(org_id=org, currency="USD", balance=Decimal("10")))
                session.add(
                    BalanceEntry(
                        org_id=org,
                        kind="adjust",
                        currency="USD",
                        amount=Decimal("10"),
                        balance_after=Decimal("10"),
                        actor="test",
                        reason="Synthetic initial funds",
                    )
                )
        header, foreign = [
            await login(api, org, label)
            for org, label in zip(tenants["orgs"], ("a", "b"), strict=True)
        ]
        task = (
            await api.post("/tasks", headers=header, json={"name": "Synthetic simulation"})
        ).json()["data"]["id"]
        uploaded = await api.post(
            f"/tasks/{task}/documents",
            headers=header,
            files={
                "file": (
                    "tender.docx",
                    tender_docx(),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )
        document = uploaded.json()["data"]["id"]
        await run_document_job(api, app, header, document, "parse")
        extraction = await run_document_job(api, app, header, document, "extract")
        route = f"/tasks/{task}/product-simulations"

        before = len(vendor.kinds)
        preview = await api.post(
            route, headers=header, json={"extraction_job_id": extraction, "dry_run": True}
        )
        assert preview.status_code == 200, preview.text
        data = preview.json()["data"]
        assert [item["name"] for item in data["items"]] == [
            "智能打印终端",
            "学生管理系统定制开发",
            "接入交换机",
        ]
        assert data["admission_blocker"] is None and len(vendor.kinds) == before
        stale = await api.post(
            route,
            headers=header,
            json={"extraction_job_id": extraction, "expected_input_hash": "0" * 64},
        )
        assert stale.status_code == 409
        assert (
            await api.post(
                route, headers=foreign, json={"extraction_job_id": extraction, "dry_run": True}
            )
        ).status_code == 404
        set_role(admin_engine, tenants["orgs"][0], tenants["users"][0], "bidder")
        denied = await api.post(
            route, headers=header, json={"extraction_job_id": extraction, "dry_run": True}
        )
        assert denied.status_code == 403
        set_role(admin_engine, tenants["orgs"][0], tenants["users"][0], "admin")

        submitted = await api.post(
            route,
            headers=header,
            json={"extraction_job_id": extraction, "expected_input_hash": data["input_hash"]},
        )
        assert submitted.status_code == 200, submitted.text
        job = submitted.json()["data"]["job_id"]
        await app.state.processor(header["X-Org-Id"], job)
        status = (await api.get(f"/jobs/{job}", headers=header)).json()["data"]
        assert status["status"] == "succeeded", status
        result = status["result"]
        assert "submission" not in result and result["simulated"] is True
        assert result["counts"] == {"items": 3, "hardware": 2, "products": 1, "parameters": 1}, (
            json.dumps(result["items"], ensure_ascii=False)
        )
        outcome = {item["name"]: item for item in result["items"]}
        assert outcome["学生管理系统定制开发"]["status"] == "not_hardware"
        switch = outcome["接入交换机"]
        assert switch["status"] == "no_page" and "model" not in switch
        assert [row["result"] for row in switch["tried"]] == ["no_product"]
        printer = outcome["智能打印终端"]
        assert printer["status"] == "quoted" and printer["parameters"] == [
            {"label": "内存", "quote": "内存：8GB DDR4"}
        ]
        tried = [(row["url"], row["source"], row["result"]) for row in printer["tried"]]
        if provider == "searxng":
            assert (printer["vendor"], printer["model"], printer["source"]) == (
                "SynVendor",
                "SP-100",
                "page",
            )
            assert tried == [
                ("https://bad.example/terminal", "page", "http_status_denied"),
                ("https://vendor.example/catalog", "page", "no_product"),
                ("https://vendor.example/sp100", "page", "quoted"),
            ]
            assert printer["url"] == "https://vendor.example/sp100"
        else:
            assert (printer["vendor"], printer["source"]) == ("BadVendor", "search_extract")
            assert tried == [("https://bad.example/terminal", "search_extract", "quoted")]
            assert printer["tried"][0]["fetch"] == "http_status_denied"
            assert [
                (row["query"], row.get("search_domain_filter")) for row in PERPLEXITY_REQUESTS
            ] == [
                ("BadVendor 打印终端", ["bad.example"]),
                ("SynNet 交换机", ["net.example"]),
            ]
        assert not any(url.startswith("market.example") or "invented" in url for url in FETCHED)
        assert vendor.kinds.count("propose") == 3

        products = (await api.get(f"/tasks/{task}/products", headers=header)).json()["items"]
        features = (await api.get(f"/tasks/{task}/features", headers=header)).json()["items"]
        assert [row["data"]["name"] for row in products] == ["【模拟】智能打印终端"]
        assert [row["data"]["description"] for row in features] == ["内存：8GB DDR4"]
        marks = (await api.get(f"/tasks/{task}/simulated-resources", headers=header)).json()["data"]
        assert sorted(marks["selection_ids"]) == sorted([products[0]["id"], features[0]["id"]])
        assert (
            await api.get(f"/tasks/{task}/simulated-resources", headers=foreign)
        ).status_code == 404

        # A human revision of the simulated product cannot clear the mark.
        product_id = printer["product_id"]
        revised = await api.post(
            f"/resources/products/{product_id}/revisions",
            headers=header,
            json={
                "expected_revision": 1,
                "data": {"name": "改名后的打印终端", "vendor": "SynVendor", "model": "SP-100"},
            },
        )
        assert revised.status_code == 200, revised.text
        async with app.state.db.transaction(tenants["orgs"][0]) as session:
            marked = (await session.scalars(select(SimulatedResource))).all()
            assert {str(row.product_id) for row in marked if row.product_id} == {product_id}
            assert {str(row.feature_id) for row in marked if row.feature_id} == set(
                printer["feature_ids"]
            )

        target_product = product_id
        if material in {"rebound_feature", "normal_feature", "normal_product"}:
            ordinary = await api.post(
                "/resources/products",
                headers=header,
                json={"data": {"name": "普通打印终端", "vendor": "SynVendor", "model": "SP-100"}},
            )
            assert ordinary.status_code == 200, ordinary.text
            target_product = ordinary.json()["data"]["product_id"]
            assert target_product != product_id

        selected_material = features[0]
        evidence = {
            "kind": "feature",
            "selection_id": selected_material["id"],
            "field_path": "description",
            "quote": "内存：8GB DDR4",
        }
        if material in {"rebound_feature", "linked_feature", "normal_feature"}:
            feature_data = {**features[0]["data"], "product_id": target_product}
            if material == "rebound_feature":
                feature_id = printer["feature_ids"][0]
                changed = await api.post(
                    f"/resources/features/{feature_id}/revisions",
                    headers=header,
                    json={"expected_revision": 1, "data": feature_data},
                )
            else:
                changed = await api.post(
                    "/resources/features", headers=header, json={"data": feature_data}
                )
            assert changed.status_code == 200, changed.text
            revision = changed.json()["data"]
            if material == "rebound_feature":
                assert revision["feature_id"] == feature_id and revision["revision"] == 2
                assert revision["data"]["product_id"] == target_product
            chosen = await api.post(
                f"/tasks/{task}/features",
                headers=header,
                json={"feature_id": revision["feature_id"], "revision": revision["revision"]},
            )
            assert chosen.status_code == 200, chosen.text
            selected_material = chosen.json()["data"]
            assert selected_material["feature_revision_id"] == revision["id"]
            evidence["selection_id"] = selected_material["id"]
            if material == "rebound_feature":
                marks = (
                    await api.get(f"/tasks/{task}/simulated-resources", headers=header)
                ).json()["data"]
                assert selected_material["id"] in marks["selection_ids"]
        elif material in {"product", "normal_product"}:
            chosen = await api.post(
                f"/tasks/{task}/products",
                headers=header,
                json={"product_id": target_product},
            )
            assert chosen.status_code == 200, chosen.text
            selected_material = chosen.json()["data"]
            evidence = {
                "kind": "product",
                "selection_id": selected_material["id"],
                "field_path": "model",
                "quote": "SP-100",
            }

        # Confirmation cannot clear either root's simulation mark.
        requirements = (
            await api.get(f"/tasks/{task}/requirements", headers=header, params={"job": extraction})
        ).json()["items"]
        memory = next(row for row in requirements if "内存" in row["text"])
        card = await api.post(
            f"/tasks/{task}/cards",
            headers=header,
            json={
                "extraction_job_id": extraction,
                "requirement_id": memory["id"],
                "content": {
                    "response_kind": "evidence",
                    "response_text": "引用所选产品声明。",
                    "deviation": "none",
                    "deviation_note": "模拟参数与要求对应。",
                    "evidence": [evidence],
                },
            },
        )
        assert card.status_code == 200, card.text
        card = card.json()["data"]
        set_role(admin_engine, tenants["orgs"][0], tenants["users"][0], "technical")
        for action, extra in (("submit", {}), ("confirm", None)):
            body = {"expected_revision": card["revision"], "action": action}
            if extra is None:
                body |= {
                    "reviewed_evidence_ids": [item["id"] for item in card["evidence"]],
                    "reviewed_warning_codes": card["warning_codes"],
                    "reason": "合成审阅。" if card["warning_codes"] else None,
                }
            response = await api.post(f"/cards/{card['id']}/actions", headers=header, json=body)
            assert response.status_code == 200, response.text
            card = response.json()["data"]
        assert card["state"] == "confirmed"
        # Complete the unrelated requirements so only the simulation gate can block.
        disposed = await api.post(
            f"/tasks/{task}/cards/dispositions",
            headers=header,
            json={
                "extraction_job_id": extraction,
                "items": [
                    {
                        "requirement_id": row["id"],
                        "expected_revision": None,
                        "disposition": "comply_only",
                        "reason": "合成场景中其余要求按须遵守条款处置。",
                    }
                    for row in requirements
                    if row["id"] != memory["id"]
                ],
            },
        )
        assert disposed.status_code == 200, disposed.text
        set_role(admin_engine, tenants["orgs"][0], tenants["users"][0], "admin")
        selected, binding = await setup_template(api, header, task)
        set_role(admin_engine, tenants["orgs"][0], tenants["users"][0], "bidder")
        draft_id = await draft(api, app, header, task, extraction)
        simulated = material not in {"normal_feature", "normal_product"}
        for mode, severity in (("final_section", "block"), ("review_copy", "acknowledge")):
            body = {
                "draft_id": draft_id,
                "task_template_id": selected["id"],
                "binding_id": binding["id"],
                "mode": mode,
            }
            preview = await api.post(
                f"/tasks/{task}/export-runs",
                headers=header,
                json={**body, "dry_run": True},
            )
            blocked = simulated and mode == "final_section"
            assert preview.status_code == (400 if blocked else 200), preview.text
            data = preview.json()["data"]
            assert data["gap_count"] == 0 and data["ready"] == (not blocked)
            issues = [
                item for item in data["issues"] if item["code"] == "export_simulated_material"
            ]
            assert [item["severity"] for item in issues] == ([severity] if simulated else []), (
                preview.text
            )
            if simulated:
                assert issues[0]["requirement_ids"] == [memory["id"]]
                assert issues[0]["evidence_ids"] == [card["evidence"][0]["id"]]
            submitted = await api.post(
                f"/tasks/{task}/export-runs",
                headers=header,
                json={
                    **body,
                    "expected_input_hash": data["input_hash"],
                    "acknowledged_issue_ids": [
                        item["issue_id"]
                        for item in data["issues"]
                        if item["severity"] == "acknowledge"
                    ],
                },
            )
            assert submitted.status_code == (400 if blocked else 200), submitted.text
            if blocked:
                assert submitted.json()["data"]["error"]["code"] == "export_simulated_material"
        assert UUID(job)
