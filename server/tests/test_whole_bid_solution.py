"""Synthetic tender-to-DOCX and download scenarios for solution assembly."""

import asyncio
import io
from datetime import UTC, datetime
from pathlib import Path
from types import SimpleNamespace
from typing import Literal, cast
from unittest.mock import AsyncMock, Mock
from uuid import uuid4

import httpx
import pytest
from app.api import tenders
from app.core.config import Settings
from app.core.db import Database
from app.core.errors import ServiceError
from app.core.security import TokenSigner
from app.models.entities import Task
from app.providers.storage import Storage
from app.schemas.solution import (
    SolutionChapter,
    SolutionRevision,
    SolutionSection,
    SolutionSubsection,
    SolutionTable,
)
from app.schemas.whole_bid import WholeBidReceipt
from app.services import bid_outline, solution, task_access, whole_bid
from app.services.auth import Identity
from bid_cli import whole_bid as whole_bid_cli
from bid_cli.client import Client
from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from fastapi import FastAPI, Request
from fastapi.responses import JSONResponse
from sqlalchemy.ext.asyncio import AsyncSession

TITLES = [
    "整体服务方案",
    "服务质量保证措施",
    "对人防车库平战结合的快速切换的机制",
    "项目组织机构、负责人及成员配置情况",
    "项目物力配置情况",
    "合理化建议及特色服务",
]


def tender(extra_forms="附件2 服务方案\n按要求提供方案"):
    scoring = "\n".join(f"{title}\t0~10\t提供具体方法与执行安排" for title in TITLES)
    return [
        {
            "page": 1,
            "text": (
                "第一章 评标办法\n" + scoring + "\n第二章 采购需求\n提供管理服务。\n"
                "第三章 投标文件格式\n附件1 投标函\n投标人：\n" + extra_forms
            ),
        }
    ]


def chapter(
    document_id,
    title,
    *,
    state: Literal["draft", "confirmed"] = "confirmed",
    body="按计划开展现场检查。",
):
    return SolutionChapter(
        id=uuid4(),
        task_id=uuid4(),
        document_id=document_id,
        title=title,
        score_range="0~10",
        current=SolutionRevision(
            id=uuid4(),
            revision=1,
            state=state,
            origin="model",
            sections=[
                SolutionSection(
                    heading="一、执行安排",
                    paragraphs=[body],
                    subsections=[
                        SolutionSubsection(
                            heading="检查频次",
                            paragraphs=["每周巡查一次。"],
                            table=SolutionTable(columns=["事项", "频次"], rows=[["巡查", "每周"]]),
                        )
                    ],
                )
            ],
            actor_kind="session",
            actor_token_id=None,
            agent_session_id=None,
            job_id=None,
            created_at=datetime.now(UTC),
            confirmed_at=datetime.now(UTC) if state == "confirmed" else None,
        ),
        placeholder_count=body.count("【待填："),
    )


def test_outline_retains_rule_continuations_and_subjective_items():
    chunks = tender()
    chunks[0]["text"] = (
        chunks[0]["text"]
        .replace(
            "整体服务方案\t0~10\t提供具体方法与执行安排",
            "整体服务方案\t0~10\t应结合\t采购需求\n  提供   详细内容。\n进一步明确执行步骤。",
        )
        .replace(
            "第二章 采购需求",
            (
                "投标报价\t0~10\t价格评分\n类似项目业绩\t0~10\t按合同数量得分\n"
                "文件相符性\t0~10\t文件相符检查\n投标文件编制\t0~10\t编制规范\n"
                "证书数量\t0~10\t依证书计算\n第二章 采购需求"
            ),
        )
    )
    outline = bid_outline.build_outline(uuid4(), chunks)
    assert outline.scoring_items[0].rule == "应结合 采购需求 提供 详细内容。 进一步明确执行步骤。"
    assert [item.title for item in outline.scoring_items if item.subjective] == TITLES
    assert len(outline.scoring_items) == 11


def test_outline_subjective_rule_and_objective_exclusions():
    outline = bid_outline.build_outline(
        uuid4(),
        [
            {
                "page": 1,
                "text": (
                    "第一章 评标办法\n执行安排\t0~10\t结合采购需求提供服务方案\n"
                    "服务保障\t0~5\t保障落实\n管理制度\t0~5\t执行规范\n项目特色\t0~5\t特色内容\n"
                    "证明评价\t0~5\t按业绩数量评价\n第二章 投标文件格式\n附件1 投标函\n正文"
                ),
            }
        ],
    )
    assert [(item.title, item.subjective) for item in outline.scoring_items] == [
        ("执行安排", True),
        ("服务保障", True),
        ("管理制度", True),
        ("项目特色", True),
        ("证明评价", False),
    ]


@pytest.mark.parametrize(
    "solution_read,has_chapters", [(True, True), (False, True), (False, False)]
)
async def test_assembly_loads_current_chapters_and_rejects_stale_score_range(
    monkeypatch, solution_read, has_chapters
):
    document_id, task_id = uuid4(), uuid4()
    chunks = tender()
    outline = bid_outline.build_outline(document_id, chunks)
    chapters = [chapter(document_id, title) for title in TITLES]
    chapters[0] = chapter(document_id, TITLES[0], body="旧分值范围的已确认正文。")
    chapters[0] = chapters[0].model_copy(update={"score_range": "0~20"})
    current = AsyncMock(return_value=chapters if has_chapters else [])
    monkeypatch.setattr(solution, "current_chapters", current)
    monkeypatch.setattr(
        whole_bid.documents,
        "require_document",
        AsyncMock(
            return_value=SimpleNamespace(
                id=document_id,
                task_id=task_id,
                media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        ),
    )
    monkeypatch.setattr(bid_outline, "outline", AsyncMock(return_value=outline))
    monkeypatch.setattr(bid_outline, "load_chunks", AsyncMock(return_value=chunks))
    monkeypatch.setattr(whole_bid, "selected_bidder", AsyncMock(return_value="合成投标人"))
    monkeypatch.setattr(whole_bid, "card_responses", AsyncMock(return_value=[]))
    actor = Identity({"task:read", "solution:read"} if solution_read else {"task:read"})

    async def check_access(session, actor, task_id, scope):
        actor.require(scope)

    access = AsyncMock(side_effect=check_access)
    monkeypatch.setattr(task_access, "access", access)
    session = SimpleNamespace(
        info={"actor": actor},
        get_one=AsyncMock(return_value=Task(id=task_id, name="合成项目", tender_number=None)),
    )
    if has_chapters and not solution_read:
        with pytest.raises(ServiceError, match="Permission denied"):
            await whole_bid.assemble(
                cast(AsyncSession, session), document_id, cast(Storage, Mock(spec=Storage))
            )
        access.assert_awaited_once_with(session, actor, task_id, scope="solution:read")
        return
    content, tally = await whole_bid.assemble(
        cast(AsyncSession, session), document_id, cast(Storage, Mock(spec=Storage))
    )
    current.assert_awaited_once_with(session, document_id)
    if has_chapters:
        access.assert_awaited_once_with(session, actor, task_id, scope="solution:read")
    else:
        access.assert_not_awaited()
    artifact = Path("data/work/solution-assembly") / (
        "stale-range.docx" if has_chapters else "form-only.docx"
    )
    await asyncio.to_thread(artifact.parent.mkdir, parents=True, exist_ok=True)
    await asyncio.to_thread(artifact.write_bytes, content)
    text = "\n".join(p.text for p in Document(io.BytesIO(content)).paragraphs)
    assert "【待起草：整体服务方案】" in text
    assert "旧分值范围的已确认正文。" not in text
    assert ("按计划开展现场检查。" in text) is has_chapters
    assert tally.unconfirmed_chapters == (1 if has_chapters else len(TITLES))
    assert tally.draft


@pytest.mark.parametrize("placement", ["technical_cover", "solution", "last"])
def test_solution_docx_order_styles_placeholders_and_missing_chapters(placement):
    forms = {
        "technical_cover": "附件2 技术文件\n技\n术\n文\n件\n附件3 服务方案\n提供方案\n附件4 其他材料\n正文",
        "solution": "附件2 服务方案\n提供方案\n附件3 其他材料\n正文",
        "last": "附件2 其他材料\n正文",
    }[placement]
    document_id = uuid4()
    chunks = tender(forms)
    outline = bid_outline.build_outline(document_id, chunks)
    chapters = [chapter(document_id, title) for title in reversed(TITLES[:-1])]
    chapters[-1] = chapter(
        document_id, TITLES[0], state="draft", body="安排【待填：项目负责人】开展检查。"
    )
    content, tally = whole_bid.build(
        outline,
        chunks,
        None,
        Task(name="合成服务项目", tender_number=None),
        "合成投标人",
        [],
        chapters,
    )
    artifact = Path("data/work/solution-assembly")
    artifact.mkdir(parents=True, exist_ok=True)
    output = artifact / f"{placement}.docx"
    output.write_bytes(content)
    doc = Document(io.BytesIO(content))
    paragraphs = doc.paragraphs
    headings = [
        p
        for p in paragraphs
        if p.style is not None
        and p.style.name == "Form Start Heading"
        and p.text in [f"{n}、{title}" for n, title in enumerate(TITLES, 1)]
    ]
    assert [p.text for p in headings] == [f"{n}、{title}" for n, title in enumerate(TITLES, 1)]
    assert all(
        p.style is not None
        and p.style.base_style is not None
        and p.style.base_style.name == "Heading 1"
        and p.style.paragraph_format.page_break_before
        for p in headings
    )
    texts = [p.text for p in paragraphs]
    first = paragraphs.index(headings[0])
    if placement == "technical_cover":
        assert first < texts.index("附件3")
    elif placement == "solution":
        assert texts.index("附件2") < first < texts.index("附件3")
    else:
        assert texts.index("附件2") < first
    number = texts[first].split("、")[0]
    section = next(p for p in paragraphs if p.text == f"{number}.1 执行安排")
    assert section.style is not None
    assert section.style.name == "Heading 2"
    sub = next(p for p in paragraphs if p.text == f"{number}.1.1 检查频次")
    assert sub.style is not None and sub.style.name == "Heading 3"
    grid = next(t for t in doc.tables if t.cell(0, 0).text == "事项")
    assert [c.text for c in grid.rows[1].cells] == ["巡查", "每周"]
    body = next(p for p in paragraphs if "项目负责人】开展检查" in p.text)
    indent = body.paragraph_format.first_line_indent
    assert indent is not None and indent.pt == 24
    marked = next(r for r in body.runs if r.text == "【待填：项目负责人】")
    assert marked.font.highlight_color == WD_COLOR_INDEX.YELLOW
    missing = next(p for p in paragraphs if p.text == f"【待起草：{TITLES[-1]}】")
    assert missing.runs[0].font.highlight_color == WD_COLOR_INDEX.YELLOW
    assert tally.unconfirmed_chapters == 2
    assert tally.placeholders.count("项目负责人") == 1
    assert tally.draft and "审阅稿" in texts
    assert tally.headers()["X-Whole-Bid-Unconfirmed-Chapters"] == "2"


def test_confirmed_chapter_with_missing_fact_still_marks_review_draft():
    document_id = uuid4()
    chunks = tender()
    outline = bid_outline.build_outline(document_id, chunks)
    chapters = [chapter(document_id, title, body="由【待填：成员姓名】执行。") for title in TITLES]
    _, tally = whole_bid.build(
        outline,
        chunks,
        None,
        Task(name="合成项目", tender_number=None),
        "合成投标人",
        [],
        chapters,
    )
    assert tally.unconfirmed_chapters == 0
    assert tally.placeholders.count("成员姓名") == 6
    assert tally.draft


def test_legacy_build_defaults_missing_chapters_to_review_draft():
    chunks = tender()
    content, tally = whole_bid.build(
        bid_outline.build_outline(uuid4(), chunks),
        chunks,
        None,
        Task(name="合成项目", tender_number=None),
        "合成投标人",
        [],
    )
    assert tally.unconfirmed_chapters == 6
    assert "【待起草：整体服务方案】" in "\n".join(
        p.text for p in Document(io.BytesIO(content)).paragraphs
    )
    receipt = WholeBidReceipt(
        document_id=uuid4(),
        output_path="synthetic.docx",
        size_bytes=1,
        draft=True,
        placeholders=0,
        unconfirmed_rows=0,
        copies=0,
    )
    assert receipt.unconfirmed_chapters == 0


async def test_cli_download_preserves_solution_review_count(monkeypatch, tmp_path):
    document_id = uuid4()
    path = f"/documents/{document_id}/whole-bid"

    def reply(request):
        return httpx.Response(
            200,
            content=b"synthetic-docx",
            headers={
                "X-Whole-Bid-Draft": "1",
                "X-Whole-Bid-Unconfirmed-Chapters": "3",
            },
        )

    client = Client("remote", "http://test", contract_version="6.0")
    monkeypatch.setattr(
        client, "request", AsyncMock(return_value={"data": {"url": path + "?signature=synthetic"}})
    )
    monkeypatch.setattr(
        client,
        "transport",
        lambda: httpx.AsyncClient(transport=httpx.MockTransport(reply), base_url="http://test"),
    )
    monkeypatch.setattr(whole_bid_cli, "credential", lambda: "synthetic")
    result = await whole_bid_cli.download(client, document_id, tmp_path / "receipt.docx")
    assert result["data"]["unconfirmed_chapters"] == 3
    assert result["data"]["draft"] is True
    assert (tmp_path / "receipt.docx").read_bytes() == b"synthetic-docx"


@pytest.mark.parametrize(
    "human,draft,status", [(False, False, 403), (False, True, 200), (True, False, 200)]
)
async def test_download_enforces_human_only_submission_copy(monkeypatch, human, draft, status):
    actor = Identity({"task:read", "export"}, actor_kind="session" if human else "token")

    async def context():
        return None, actor

    async def assemble(*args):
        return b"synthetic-docx", whole_bid.Tally(unconfirmed_chapters=int(draft))

    monkeypatch.setattr(whole_bid, "assemble", assemble)
    monkeypatch.setattr(tenders, "check_signature", lambda *args, **kwargs: None)
    app = FastAPI()

    @app.exception_handler(ServiceError)
    async def error(request: Request, exc: ServiceError):
        return JSONResponse({"code": exc.code}, status_code=exc.status)

    app.include_router(
        tenders.create_router(
            context,
            Settings.model_construct(),
            cast(Database, Mock(spec=Database)),
            cast(Storage, Mock(spec=Storage)),
            None,
            cast(TokenSigner, Mock(spec=TokenSigner)),
            None,
            None,
            None,
        )
    )
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as http:
        response = await http.get(f"/documents/{uuid4()}/whole-bid?signature=synthetic")
    assert response.status_code == status
    if status == 200:
        assert response.headers["X-Whole-Bid-Unconfirmed-Chapters"] == str(int(draft))
