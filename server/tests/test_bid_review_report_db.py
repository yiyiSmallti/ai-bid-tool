"""Report HTTP → real worker acceptance; run only in the main isolated DB session.

Failure inventory: data/work/bid-review-report/failure-modes.md. All documents and
provider replies are synthetic. --basetemp retains the opened DOCX and receipts.
"""

import hashlib
import io
import json
from datetime import UTC, datetime, timedelta
from urllib.parse import parse_qs, urlsplit
from uuid import UUID, uuid4

import pytest
from app.core.security import TokenSigner
from app.models.entities import Membership
from docx import Document
from sqlalchemy import select
from sqlalchemy.orm import Session
from test_bid_review_findings_db import (
    BID_RESPONSE,
    CLAUSES,
    SIGNATURE_CLAUSE,
    FindingsVendor,
    classify,
    decision_body,
    execute,
    make_case,
    reviewer,
)
from test_bid_review_run_db import assert_private_absent, counts, data, preview, submit
from test_bid_review_upload_db import run
from test_team_workflow_membership import add_member, person, workflow

SECTIONS = [
    ("overall", "一、总体结论"),
    ("basic_information", "二、基本信息"),
    ("compliance", "三、废标判定"),
    ("signatures", "签章校验"),
    ("risks", "四、高风险缺陷"),
    ("scores", "五、得分预估"),
    ("evidence", "六、证据核对"),
    ("remediation", "七、补救清单"),
    ("methodology", "八、检验说明"),
]
ADVISORY = "评标委员会决定最终评审结果；本报告仅供辅助审查。"
REASON = "Synthetic human dismissal preserved in report snapshot."


async def report(api, header, case, **params):
    response = await api.get(
        f"/v4/bid-reviews/{case['review']}/report", headers=header, params=params
    )
    return data(response), response.json()["items"]


async def render_preview(api, header, case, *, request_id=None):
    current, _ = await report(api, header, case)
    body = {
        "request_id": request_id or str(uuid4()),
        "report_id": case["review"],
        "dry_run": True,
        "expected_decisions_snapshot_sha256": current["decisions_snapshot_sha256"],
    }
    receipt = data(
        await api.post(f"/v4/bid-reviews/{case['review']}/artifacts", headers=header, json=body)
    )
    return body, receipt


async def render_submit(api, header, case, body, receipt):
    return await api.post(
        f"/v4/bid-reviews/{case['review']}/artifacts",
        headers=header,
        json={
            **body,
            "dry_run": False,
            "expected_input_hash": receipt["input_hash"],
            "preflight_token": receipt["preflight_token"],
        },
    )


async def render(api, header, application, org, case):
    body, receipt = await render_preview(api, header, case)
    accepted = data(await render_submit(api, header, case, body, receipt))
    await run(application, org, accepted["job_id"])
    job = data(await api.get(f"/v4/jobs/{accepted['job_id']}", headers=header))
    assert job["status"] == "succeeded", job
    return job, receipt


def docx_text(content):
    opened = Document(io.BytesIO(content))
    return "\n".join(
        [paragraph.text for paragraph in opened.paragraphs]
        + [cell.text for table in opened.tables for row in table.rows for cell in row.cells]
    )


async def published_case(api, headers, application, tenants, tmp_path, monkeypatch):
    case = await make_case(
        api,
        headers,
        application,
        tenants,
        tmp_path,
        monkeypatch,
        vendor=FindingsVendor(clauses=(*CLAUSES, SIGNATURE_CLAUSE), outcome="deviation"),
        price=True,
    )
    return await execute(api, headers, application, tenants, case)


async def test_snapshot_word_pipeline_decisions_pagination_and_later_drift(
    api, headers, application, tenants, tmp_path, monkeypatch, admin_engine
):
    case = await published_case(api, headers, application, tenants, tmp_path, monkeypatch)
    finding = case["findings"][0]
    bidder = await reviewer(api, headers, tenants, admin_engine, case, "bidder")
    classified = data(await classify(api, headers[0], case, finding))
    decision = data(
        await api.post(
            f"/v4/bid-reviews/{case['review']}/findings/{finding['id']}/decisions",
            headers=bidder,
            json=decision_body(case, classified, reason=REASON),
        )
    )
    before_counts = await counts(application, tenants["orgs"][0])
    storage_before = sorted(str(path) for path in application.state.storage.root.rglob("*"))
    calls_before = len(case["vendor"].requests)
    body, receipt = await render_preview(api, bidder, case)
    assert receipt["budget"]["planned_calls"] == 0
    assert await counts(application, tenants["orgs"][0]) == before_counts
    assert sorted(str(path) for path in application.state.storage.root.rglob("*")) == storage_before
    accepted = data(await render_submit(api, bidder, case, body, receipt))
    replay = data(await render_submit(api, bidder, case, body, receipt))
    assert replay["job_id"] == accepted["job_id"]
    # The queued renderer must use its admitted snapshot, even before it starts.
    data(await classify(api, headers[0], case, case["findings"][1], "technical"))
    await run(application, tenants["orgs"][0], accepted["job_id"])
    finished = data(await api.get(f"/v4/jobs/{accepted['job_id']}", headers=headers[0]))
    assert finished["status"] == "succeeded", finished
    assert len(case["vendor"].requests) == calls_before
    snapshot_id = finished["result"]["snapshot_id"]
    artifacts = finished["result"]["artifacts"]
    assert {entry["format"] for entry in artifacts} == {"docx", "console"}
    assert {entry["decisions_snapshot_sha256"] for entry in artifacts} == {
        receipt["decisions_snapshot_sha256"]
    }
    all_sections = {}
    for section, title in SECTIONS:
        cursor, rows = None, []
        while True:
            args = {"section": section, "snapshot_id": snapshot_id, "limit": 1}
            if cursor:
                args["cursor"] = cursor
            meta, page = await report(api, headers[0], case, **args)
            assert meta["snapshot_id"] == snapshot_id
            assert meta["completion"] == "partial"  # excluded price page remains uncovered
            assert len(page) <= 1
            rows.extend(page)
            cursor = meta["next_cursor"]
            if not cursor:
                break
        assert title in [entry["title"] for entry in meta["sections"]]
        all_sections[section] = rows
    compliance = json.dumps(all_sections["compliance"], ensure_ascii=False)
    assert CLAUSES[0] in compliance and BID_RESPONSE in compliance and REASON in compliance
    signatures = json.dumps(all_sections["signatures"], ensure_ascii=False)
    assert "unresolved" in signatures and "presence_not_checked" in signatures
    descriptor = next(entry for entry in artifacts if entry["format"] == "docx")
    link = data(
        await api.get(
            f"/v4/bid-review-artifacts/{descriptor['id']}/download-link", headers=headers[0]
        )
    )
    download = await api.get(link["url"], headers=headers[0])
    assert download.status_code == 200, download.text
    assert download.headers["cache-control"] == "no-store"
    assert "attachment" in download.headers["content-disposition"]
    assert_private_absent(download.headers["content-disposition"])
    assert len(download.content) == descriptor["size_bytes"]
    assert hashlib.sha256(download.content).hexdigest() == descriptor["sha256"]
    rendered = docx_text(download.content)
    for _, title in SECTIONS:
        assert title in rendered
    for required in (ADVISORY, CLAUSES[0], BID_RESPONSE, REASON, "未评分", "unresolved"):
        assert required in rendered
    (tmp_path / "report-snapshot.docx").write_bytes(download.content)
    new_decision = data(
        await api.post(
            f"/v4/bid-reviews/{case['review']}/findings/{finding['id']}/decisions",
            headers=bidder,
            json=decision_body(case, decision, "reopen", reason="Later decision must not drift."),
        )
    )
    assert new_decision["revision"] > decision["revision"]
    fixed, fixed_rows = await report(
        api, headers[0], case, snapshot_id=snapshot_id, section="compliance"
    )
    assert fixed_rows == all_sections["compliance"]
    assert fixed["decisions_snapshot_sha256"] == receipt["decisions_snapshot_sha256"]
    current, _ = await report(api, headers[0], case)
    assert current["decisions_snapshot_sha256"] != receipt["decisions_snapshot_sha256"]
    later, _ = await render(api, headers[0], application, tenants["orgs"][0], case)
    assert later["result"]["snapshot_id"] != snapshot_id
    unchanged = await api.get(link["url"], headers=headers[0])
    assert unchanged.content == download.content
    (tmp_path / "report-snapshot-receipt.json").write_text(
        json.dumps({"snapshot_id": snapshot_id, "artifacts": artifacts}, indent=2) + "\n"
    )


async def test_report_projections_scopes_cross_org_and_live_download_authority(
    api, headers, application, tenants, tmp_path, monkeypatch, admin_engine
):
    case = await published_case(api, headers, application, tenants, tmp_path, monkeypatch)
    job, receipt = await render(api, headers[0], application, tenants["orgs"][0], case)
    snapshot = job["result"]["snapshot_id"]
    word = next(a for a in job["result"]["artifacts"] if a["format"] == "docx")
    link_path = f"/v4/bid-review-artifacts/{word['id']}/download-link"
    link = data(await api.get(link_path, headers=headers[0]))
    assert (await api.get(link_path.removeprefix("/v4"), headers=headers[0])).status_code == 404
    for scope in ("bid-review:report:render", "bid-review:report:download"):
        refused = await api.post(
            "/v4/tokens",
            headers=headers[0],
            json={
                "name": "Forbidden report token",
                "expires_at": (datetime.now(UTC) + timedelta(hours=1)).isoformat(),
                "scopes": ["bid-review:read", scope],
            },
        )
        assert refused.status_code in {400, 403, 422}
    issued = data(
        await api.post(
            "/v4/tokens",
            headers=headers[0],
            json={
                "name": "Safe report token",
                "expires_at": (datetime.now(UTC) + timedelta(hours=1)).isoformat(),
                "scopes": ["task:read", "job:read", "bid-review:read"],
            },
        )
    )
    token = {**headers[0], "Authorization": "Bearer " + issued["token"]}
    for header, projection in [(token, "safe")]:
        meta, rows = await report(api, header, case, snapshot_id=snapshot, section="compliance")
        serialized = json.dumps([meta, rows])
        assert meta["projection"] == projection
        assert "quote" not in serialized and CLAUSES[0] not in serialized
        assert not meta["artifacts"]
        assert_private_absent(serialized)
        assert (await api.get(link_path, headers=header)).status_code == 403
        assert (await api.get(link["url"], headers=header)).status_code == 403
        body = {
            "request_id": str(uuid4()),
            "report_id": case["review"],
            "dry_run": True,
            "expected_decisions_snapshot_sha256": receipt["decisions_snapshot_sha256"],
        }
        assert (
            await api.post(f"/v4/bid-reviews/{case['review']}/artifacts", headers=header, json=body)
        ).status_code == 403
        token_job = await api.get(f"/v4/jobs/{job['id']}", headers=header)
        assert token_job.status_code == 403
        assert "quote" not in token_job.text and "artifacts" not in token_job.text
    for role in ("technical", "viewer"):
        if role == "viewer":
            user, human = await person(api, admin_engine, tenants["orgs"][0], role="viewer")
            membership = await workflow(api, headers[0], case["task"])
            data(
                await add_member(
                    api, headers[0], case["task"], user, membership["revision"], role="observer"
                )
            )
        else:
            human = await reviewer(api, headers, tenants, admin_engine, case, role)
        meta, rows = await report(api, human, case, snapshot_id=snapshot, section="compliance")
        assert meta["projection"] == "cleared"
        serialized = json.dumps([meta, rows])
        assert "quote" not in serialized and CLAUSES[0] not in serialized
        assert_private_absent(serialized)
        assert not meta["artifacts"]
        assert (await api.get(link_path, headers=human)).status_code == 403
    for path in (
        f"/v4/bid-reviews/{case['review']}/report",
        f"/v4/bid-reviews/{case['review']}/reports",
        link_path,
        link["url"],
    ):
        assert (await api.get(path, headers=headers[1])).status_code == 404
    stranger_task_case = dict(case, review=str(uuid4()))
    assert (
        await api.get(
            f"/v4/bid-reviews/{stranger_task_case['review']}/report",
            headers=headers[0],
            params={"snapshot_id": snapshot},
        )
    ).status_code == 404
    with Session(admin_engine) as session, session.begin():
        member = session.scalar(
            select(Membership).where(
                Membership.org_id == tenants["orgs"][0],
                Membership.user_id == tenants["users"][0],
            )
        )
        member.role = "viewer"
    assert (await api.get(link["url"], headers=headers[0])).status_code == 403


async def test_expired_tampered_links_and_stale_preview_are_rejected(
    api, headers, application, tenants, tmp_path, monkeypatch
):
    case = await published_case(api, headers, application, tenants, tmp_path, monkeypatch)
    body, receipt = await render_preview(api, headers[0], case)
    data(await classify(api, headers[0], case, case["findings"][0]))
    assert (await render_submit(api, headers[0], case, body, receipt)).status_code == 409
    job, _ = await render(api, headers[0], application, tenants["orgs"][0], case)
    word = next(a for a in job["result"]["artifacts"] if a["format"] == "docx")
    link = data(
        await api.get(f"/v4/bid-review-artifacts/{word['id']}/download-link", headers=headers[0])
    )
    signer = TokenSigner.for_tokens(application.state.processor.settings)
    parsed = urlsplit(link["url"])
    payload = signer.open(parse_qs(parsed.query)["signature"][0])
    for changed, ttl in ((payload, -1), ({**payload, "kind": "attachment"}, 300)):
        response = await api.get(
            parsed.path,
            headers=headers[0],
            params={
                "signature": signer.issue(changed, ttl),
            },
        )
        assert response.status_code == 404
    assert (await api.get(parsed.path, headers=headers[0])).status_code in {404, 422}
    assert (await api.get(link["url"])).status_code in {401, 403}


async def test_unpublished_run_cannot_create_report_snapshot(
    api, headers, application, tenants, tmp_path, monkeypatch
):
    case = await make_case(api, headers, application, tenants, tmp_path, monkeypatch)
    receipt = data(await preview(api, headers[0], case))
    accepted = data(await submit(api, headers[0], case, receipt))
    from app.models.bid_review_run import BidReviewRun

    async with application.state.db.transaction(tenants["orgs"][0]) as session:
        row = await session.scalar(
            select(BidReviewRun).where(BidReviewRun.job_id == UUID(accepted["job_id"]))
        )
        review_id = str(row.id)
    assert (
        await api.get(f"/v4/bid-reviews/{review_id}/report", headers=headers[0])
    ).status_code == 409
    rejected = await api.post(
        f"/v4/bid-reviews/{review_id}/artifacts",
        headers=headers[0],
        json={
            "request_id": str(uuid4()),
            "report_id": review_id,
            "dry_run": True,
            "expected_decisions_snapshot_sha256": "0" * 64,
        },
    )
    assert rejected.status_code == 409


@pytest.mark.parametrize("failure", ["renderer", "storage", "database"])
async def test_failed_render_never_publishes_partial_artifacts(
    failure, api, headers, application, tenants, tmp_path, monkeypatch
):
    from app.core.errors import ServiceError
    from app.models.bid_review_report import BidReviewReportArtifact, BidReviewReportSnapshot
    from sqlalchemy import event, func

    case = await published_case(api, headers, application, tenants, tmp_path, monkeypatch)
    body, receipt = await render_preview(api, headers[0], case)
    accepted = data(await render_submit(api, headers[0], case, body, receipt))
    before = len(case["vendor"].requests)
    written = []
    from app.jobs import bid_review_report as report_worker

    real_render = report_worker.render_process
    real_put = application.state.storage.put

    async def storage_write(org_id, key, content):
        if "/reports/" in key:
            written.append(key)
            if failure == "storage" and len(written) == 2:
                raise ServiceError("synthetic_storage_refusal", "Synthetic storage refusal", 500, 4)
        return await real_put(org_id, key, content)

    async def broken_renderer(execution, snapshot):
        raise ServiceError("synthetic_render_refusal", "Synthetic render refusal", 500, 4)

    def broken_insert(mapper, connection, target):
        raise ServiceError("synthetic_commit_refusal", "Synthetic commit refusal", 500, 4)

    monkeypatch.setattr(application.state.storage, "put", storage_write)
    if failure == "renderer":
        monkeypatch.setattr("app.jobs.bid_review_report.render_process", broken_renderer)
    if failure == "database":
        event.listen(BidReviewReportArtifact, "before_insert", broken_insert)
    try:
        await run(application, tenants["orgs"][0], accepted["job_id"])
    finally:
        if failure == "database":
            event.remove(BidReviewReportArtifact, "before_insert", broken_insert)
    job = data(await api.get(f"/v4/jobs/{accepted['job_id']}", headers=headers[0]))
    assert job["status"] == "failed", job
    assert len(case["vendor"].requests) == before
    async with application.state.db.transaction(tenants["orgs"][0]) as session:
        assert await session.scalar(select(func.count()).select_from(BidReviewReportArtifact)) == 0
        snapshot = await session.scalar(
            select(BidReviewReportSnapshot).where(
                BidReviewReportSnapshot.job_id == UUID(accepted["job_id"])
            )
        )
        snapshot_id = str(snapshot.id)
    unpublished = await api.get(
        f"/v4/bid-reviews/{case['review']}/report",
        headers=headers[0],
        params={"snapshot_id": snapshot_id},
    )
    assert unpublished.status_code == 409
    assert unpublished.json()["data"]["error"]["code"] == "bid_review_report_unpublished"
    if failure in {"storage", "database"}:
        assert written  # encrypted unreachable objects are not published or blindly purged
    (tmp_path / f"report-failed-{failure}.json").write_text(
        json.dumps(
            {
                "snapshot_id": snapshot_id,
                "status": job["status"],
                "published_artifacts": 0,
                "staged_count": len(written),
            }
        )
        + "\n"
    )

    if failure == "renderer":
        history_response = await api.get(
            f"/v4/bid-reviews/{case['review']}/reports", headers=headers[0]
        )
        data(history_response)
        entry = next(
            row for row in history_response.json()["items"] if row["snapshot_id"] == snapshot_id
        )
        assert entry["can_retry"] and entry["input_hash"] == receipt["input_hash"]
        monkeypatch.setattr(report_worker, "render_process", real_render)
        retry_body, retry_receipt = await render_preview(api, headers[0], case)
        assert retry_body["request_id"] != body["request_id"]
        retry_body["retry"] = True
        retried = data(await render_submit(api, headers[0], case, retry_body, retry_receipt))
        assert retried["job_id"] == accepted["job_id"]
        replay = data(await render_submit(api, headers[0], case, retry_body, retry_receipt))
        assert replay["job_id"] == accepted["job_id"]
        await run(application, tenants["orgs"][0], retried["job_id"])
        finished = data(await api.get(f"/v4/jobs/{retried['job_id']}", headers=headers[0]))
        assert finished["status"] == "succeeded", finished
        assert finished["result"]["snapshot_id"] == snapshot_id


async def test_snapshot_artifact_rls_immutability_and_cross_task_binding(
    api, headers, application, tenants, tmp_path, monkeypatch
):
    from app.models.bid_review_report import BidReviewReportArtifact, BidReviewReportSnapshot
    from sqlalchemy import func, text
    from sqlalchemy.exc import DBAPIError

    case = await published_case(api, headers, application, tenants, tmp_path, monkeypatch)
    job, _ = await render(api, headers[0], application, tenants["orgs"][0], case)
    models = (BidReviewReportSnapshot, BidReviewReportArtifact)
    async with application.state.db.transaction(tenants["orgs"][0]) as session:
        for model in models:
            assert await session.scalar(select(func.count()).select_from(model)) > 0
    for org in (None, tenants["orgs"][1]):
        async with application.state.db.transaction(org) as session:
            if org is None:
                await session.execute(text("SELECT set_config('app.current_org','',true)"))
            for model in models:
                assert await session.scalar(select(func.count()).select_from(model)) == 0
    for model in models:
        for statement in (
            f"UPDATE {model.__tablename__} SET id=id",
            f"DELETE FROM {model.__tablename__}",
        ):
            with pytest.raises(DBAPIError):
                async with application.state.db.transaction(tenants["orgs"][0]) as session:
                    await session.execute(text(statement))
    another = await published_case(api, headers, application, tenants, tmp_path, monkeypatch)
    assert another["task"] != case["task"]
    mismatched = await api.get(
        f"/v4/bid-reviews/{another['review']}/report",
        headers=headers[0],
        params={"snapshot_id": job["result"]["snapshot_id"]},
    )
    assert mismatched.status_code == 404
    foreign_post = await api.post(
        f"/v4/bid-reviews/{case['review']}/artifacts",
        headers=headers[1],
        json={
            "request_id": str(uuid4()),
            "report_id": case["review"],
            "dry_run": True,
            "expected_decisions_snapshot_sha256": "0" * 64,
        },
    )
    assert foreign_post.status_code == 404
