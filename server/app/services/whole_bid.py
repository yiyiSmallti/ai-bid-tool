"""Assemble a whole bid (整本标书) from a tender's format chapter.

Every form of the outline is reproduced in order from the tender's own layout,
blanks the installation can answer are filled from the tender and the task's
selected company profile, and the procurement deviation table is filled with the
needs chapter's clauses. Whatever is unknown stays a highlighted placeholder, and
the file says it is a review draft (审阅稿) while any remain.
"""

import io
import math
import re
from dataclasses import dataclass, field, replace
from typing import cast
from uuid import UUID

from docx import Document as WordDocument
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.styles.style import CharacterStyle, ParagraphStyle
from docx.text.paragraph import Paragraph as DocxParagraph
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.errors import ServiceError, not_found
from app.models.entities import OrgProfileRevision, Requirement, Task, TaskOrgProfile
from app.models.response_cards import ResponseCard, ResponseCardRevision
from app.providers.storage import Storage
from app.schemas.profile_contracts import OrgProfileData
from app.schemas.solution import SolutionChapter, SolutionTable
from app.schemas.whole_bid import BidOutline, OutlineForm, OutlineScoringItem
from app.services import bid_outline, documents, task_access
from app.services.form_layout import Blank, Cell, Grid, Paragraph, Piece, Text, form_blocks, plain

BODY_FONT, HEADING_FONT = "宋体", "黑体"
HEADING = re.compile(r"^([一二三四五六七八九十]+、|（[一二三四五六七八九十]+）)\s*(.{1,30})$")
SENTENCE_END = re.compile(r"[。；]")
DRAFT_RESPONSE, DRAFT_DEVIATION = "完全响应", "无偏离"
PRICE_SCORE = re.compile(r"(投标)?(报价|价格)(分|得分)?")
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


@dataclass(frozen=True)
class Facts:
    """Values a blank may be filled with; None means unknown and stays a placeholder."""

    project_name: str | None
    project_number: str | None
    tenderee: str | None
    bidder: str | None
    address: str | None = None
    phone: str | None = None
    legal_rep: str | None = None
    credit_code: str | None = None


@dataclass
class Tally:
    """What still needs a human before the file can be submitted."""

    placeholders: list[str] = field(default_factory=list)
    unconfirmed_rows: int = 0
    unconfirmed_chapters: int = 0
    copies: list[str] = field(default_factory=list)

    @property
    def draft(self) -> bool:
        return bool(
            self.placeholders or self.unconfirmed_rows or self.unconfirmed_chapters or self.copies
        )

    def headers(self) -> dict[str, str]:
        return {
            "X-Whole-Bid-Draft": "1" if self.draft else "0",
            "X-Whole-Bid-Placeholders": str(len(self.placeholders)),
            "X-Whole-Bid-Unconfirmed-Rows": str(self.unconfirmed_rows),
            "X-Whole-Bid-Unconfirmed-Chapters": str(self.unconfirmed_chapters),
            "X-Whole-Bid-Copies": str(len(self.copies)),
        }


@dataclass(frozen=True)
class Clause:
    heading: str
    text: str


@dataclass(frozen=True)
class Response:
    """The current response card of one extracted requirement."""

    quote: str
    text: str
    deviation: str
    confirmed: bool


DEVIATION_LABELS = {"none": "无偏离", "positive": "正偏离", "negative": "负偏离"}
SECRET_FIELD = re.compile(r"\{\{secret\.[^}]+\}\}")
SOLUTION_PLACEHOLDER = re.compile(r"【待填：([^】]+)】")
# Numbering a model may still put on a heading; the renderer numbers headings itself.
HEADING_NUMBER = re.compile(
    r"^\s*(第[一二三四五六七八九十\d]+[章节]|[一二三四五六七八九十]+[、.．]|\d+(\.\d+)*(、|[.．](?!\d)|\s)|[（(][一二三四五六七八九十\d]+[）)])\s*"
)


def unnumbered(heading: str) -> str:
    return HEADING_NUMBER.sub("", heading, count=1) or heading


def compact(text: str) -> str:
    return re.sub(r"\s", "", text)


def clause_response(clause: Clause, responses: list[Response]) -> tuple[str, str] | None:
    """The confirmed response and deviation of a clause, or None while any requirement
    quoted from it is unconfirmed or none is. Several requirements in one clause keep
    their responses in order and the clause takes the worst deviation."""
    text = compact(clause.text)
    matched = [
        response
        for response in responses
        if (quote := compact(response.quote)) and (quote in text or text in quote)
    ]
    if not matched or not all(response.confirmed for response in matched):
        return None
    # A requirement extracted twice, or once inside a longer one, answers once.
    quotes = [compact(response.quote) for response in matched]
    matched = [
        response
        for position, response in enumerate(matched)
        if not any(
            quotes[position] in other and (quotes[position] != other or earlier < position)
            for earlier, other in enumerate(quotes)
            if earlier != position
        )
    ]
    worst = max(matched, key=lambda response: list(DEVIATION_LABELS).index(response.deviation))
    joined = "\n".join(dict.fromkeys(response.text for response in matched))
    # Confidential values are filled at export from the encrypted library, never here.
    return SECRET_FIELD.sub("【待填：保密字段】", joined), DEVIATION_LABELS[worst.deviation]


# Field rules match the label just before a blank, or the hint printed on it.
# The bidder's name is typed even where a seal goes over it, so it comes before signatures.
FIELD_RULES: tuple[tuple[str, re.Pattern[str]], ...] = (
    ("bidder", re.compile(r"^(投标人|供应商|企业|单位)(全称|名称)?(（[^）]*）)?[：:]?$")),
    ("signature", re.compile(r"签字|签名|盖章|公章|签章|印刷体")),
    ("price", re.compile(r"总价|报价|保证金|金额|人民币|万元|元整")),
    ("project_both", re.compile(r"项目名称[、及和]((项目|招标|采购)?编号)")),
    ("project_number", re.compile(r"(项目|招标|采购)编号[：:]?$")),
    ("project_name", re.compile(r"(采购项目|项目名称[：:]?)$")),
    ("tenderee", re.compile(r"招标人|采购人")),
    ("bidder", re.compile(r"(投标人|供应商|单位)(全称|名称)[：:]?$|^投标人名称|投标人名称、地址")),
    ("address", re.compile(r"地\s*址[：:]?$")),
    ("phone", re.compile(r"电话[：:]?$")),
    ("legal_rep", re.compile(r"法定代表人[：:]?$|^姓名$")),
    ("credit_code", re.compile(r"统一社会信用代码[：:]?$")),
    ("date", re.compile(r"(日期|时\s*间)[：:]?$")),
)
# A signature line printed without a blank, like 投标人名称（盖章） or 投标人名称：（盖章）.
SIGN_LINE = re.compile(r"^(?P<label>[^（）：:\s]{2,10}(（[^）]*）)?)[：:]?(?P<tail>（[^）]*）)?$")
FIELD_NAMES = {
    "project_both": "项目名称、项目编号",
    "project_number": "项目编号",
    "project_name": "项目名称",
    "tenderee": "招标人",
    "bidder": "投标人名称",
    "address": "地址",
    "phone": "电话",
    "legal_rep": "法定代表人",
    "credit_code": "统一社会信用代码",
    "date": "日期",
}


BRACKET_HINT = re.compile(r"^\s*[（(]([^（）()]{2,20})[）)]")
BRACKET_FIELD = re.compile(r"[（(]([^（）()]{2,20})[）)]")
# Fields a bracketed hint inside running text may be replaced by; a person's name is not
# among them, since the text around it decides whose name is meant.
NAMED_FIELDS = {"project_both", "project_number", "project_name", "tenderee", "bidder"}


def short_title(title: str) -> str:
    """A long form title such as a list of acceptable copies, cut at its first clause."""
    return re.split(r"[；;。]", title)[0]


def field_of(label: str, hint: str) -> str | None:
    for source in (hint, label):
        text = source.strip()
        if not text:
            continue
        for name, pattern in FIELD_RULES:
            if pattern.search(text):
                return name
    return None


def value_of(name: str | None, facts: Facts) -> str | None:
    if name == "project_both":
        if facts.project_name and facts.project_number:
            return f"{facts.project_name}（{facts.project_number}）"
        return None
    if name in {
        "project_name",
        "project_number",
        "tenderee",
        "bidder",
        "address",
        "phone",
        "legal_rep",
        "credit_code",
    }:
        return getattr(facts, name)
    return None


def tender_facts(text: str, task: Task, bidder: OrgProfileData | None) -> Facts:
    """Project and tenderee as the tender states them; the task's own fields are fallbacks."""

    def first(pattern: str) -> str | None:
        match = re.search(pattern, text)
        return match.group(1).strip(" ：:，。") if match else None

    return Facts(
        project_name=first(r"项目名称[：:]\s*([^\n：:]{4,80}?)\s*(?:\n|项目编号|$)") or task.name,
        project_number=first(r"项目编号[：:]\s*([^\s，。；、：:]{4,40})")
        or first(r"招标编号[：:]\s*([^\s，。；、：:]{4,40})")
        or task.tender_number,
        tenderee=first(r"(?:采购人|采购单位|招标人)(?:名称)?[：:]\s*([^\n，,。；]{2,60})"),
        bidder=bidder.name if bidder else None,
        address=bidder.address if bidder else None,
        phone=bidder.phone if bidder else None,
        legal_rep=bidder.legal_representative if bidder else None,
        credit_code=bidder.credit_code if bidder else None,
    )


def needs_clauses(blocks: list[Paragraph | Grid], chapter_title: str) -> list[Clause]:
    """One clause per paragraph under its nearest heading; a table is one clause."""
    clauses: list[Clause] = []
    heading = ""
    caption = ""
    for block in blocks:
        if isinstance(block, Grid):
            rows = {cell.row for cell in block.cells if re.fullmatch(r"\d+", cell.text)}
            if clauses and clauses[-1].heading == heading and clauses[-1].text.startswith(caption):
                counted = re.search(r"共(\d+)项", clauses[-1].text)
                count = (int(counted.group(1)) if counted else 0) + len(rows)
                clauses[-1] = Clause(heading, f"{caption}（共{count}项，详见招标文件）")
            else:
                clauses.append(Clause(heading, f"{caption}（共{len(rows)}项，详见招标文件）"))
            continue
        text = plain(block).strip()
        if not text or re.sub(r"\s", "", text) == chapter_title:
            continue
        match = HEADING.match(text)
        if match and not SENTENCE_END.search(text):
            heading = match.group(2).strip()
            caption = ""
            continue
        if len(text) <= 40 and not SENTENCE_END.search(text):
            caption = text
            continue
        clauses.append(Clause(heading, text))
    return clauses


# The share of a subsection heading's weight a clause must cover to be answered there.
ANSWER_SHARE = 0.3


def bigrams(text: str) -> set[str]:
    han = re.sub(r"[^\u4e00-\u9fff]", "", text)
    return {han[i : i + 2] for i in range(len(han) - 1)}


def answering_subsections(
    clauses: list[Clause], items: list[OutlineScoringItem], chapters: list[SolutionChapter]
) -> list[str | None]:
    """The solution subsection (`N.M.K`) whose heading a needs clause best covers, or None.

    Section bodies restate the whole needs chapter, so only headings are compared.
    Bigrams common to many headings and clauses (公用, 人防, 工程) weigh little; a
    clause covering too little of any heading, such as a payment term, has no answer.
    """
    by_item = {(chapter.title, chapter.score_range): chapter for chapter in chapters}
    heads: list[tuple[str, set[str]]] = []
    for chapter_number, item in enumerate((item for item in items if item.subjective), 1):
        chapter = by_item.get((item.title, item.score_range))
        for position, section in enumerate(chapter.current.sections if chapter else [], 1):
            for sub_position, sub in enumerate(section.subsections, 1):
                heads.append(
                    (
                        f"{chapter_number}.{position}.{sub_position}",
                        bigrams(section.heading + sub.heading),
                    )
                )
    texts = [bigrams(clause.text) for clause in clauses]
    documents = [grams for _, grams in heads] + texts
    frequency: dict[str, int] = {}
    for grams in documents:
        for gram in grams:
            frequency[gram] = frequency.get(gram, 0) + 1

    def weight(gram: str) -> float:
        return math.log((len(documents) + 1) / (frequency.get(gram, 0) + 1))

    answers: list[str | None] = []
    for grams in texts:
        best, number = 0.0, None
        for candidate, head in heads:
            total = sum(weight(gram) for gram in head) or 1
            share = sum(weight(gram) for gram in grams & head) / total
            if share > best:
                best, number = share, candidate
        answers.append(number if best >= ANSWER_SHARE else None)
    return answers


COLUMN_WEIGHTS = {
    "number": 1.0,
    "topic": 1.6,
    "requirement": 5.0,
    "response": 5.0,
    "deviation": 1.4,
    "other": 1.6,
}


def deviation_column(header: str) -> str:
    if re.search(r"序号", header):
        return "number"
    if re.search(r"页|所在|对应投标文件", header):
        return "other"
    if re.search(r"偏离", header):
        return "deviation"
    if re.search(r"响应|应答", header):
        return "response"
    if re.search(r"需求|要求|招标文件", header):
        return "requirement"
    if re.search(r"内容|条款|项目", header):
        return "topic"
    return "other"


def form_anchor(form: OutlineForm) -> str:
    return f"form_{form.index}"


def solution_anchor(number: str) -> str:
    return "solution_" + number.replace(".", "_")


TEXT_WIDTH = Cm(16)


class Renderer:
    """Writes outline forms into one DOCX in the house style: A4, 宋体 body, 黑体 headings."""

    def __init__(self, facts: Facts, tally: Tally):
        self.facts, self.tally = facts, tally
        self.pending_break = False
        # A bookmark for the next page's first paragraph, which contents entries point at.
        self.pending_anchor: str | None = None
        self.bookmarks = 0
        self.word = WordDocument()
        section = self.word.sections[0]
        section.page_width, section.page_height = Cm(21), Cm(29.7)
        section.left_margin = section.right_margin = Cm(2.5)
        section.top_margin = section.bottom_margin = Cm(2.5)
        self._style("Normal", BODY_FONT, 12)
        for level, size in ((1, 16), (2, 14), (3, 12)):
            style = self._style(f"Heading {level}", HEADING_FONT, size)
            style.font.bold = True
            # The linked character style carries the template's theme blue too, and
            # renderers apply it to heading runs.
            char = cast(CharacterStyle, self.word.styles[f"Heading {level} Char"])
            for target in (style, char):
                target.font.color.rgb = RGBColor(0, 0, 0)
                color = target.element.get_or_add_rPr().find(qn("w:color"))
                for name in ("w:themeColor", "w:themeShade", "w:themeTint"):
                    color.attrib.pop(qn(name), None)
            style.paragraph_format.space_before = Pt(6)
            style.paragraph_format.space_after = Pt(6)
        for name, base in (("Form Start", "Normal"), ("Form Start Heading", "Heading 1")):
            style = cast(ParagraphStyle, self.word.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH))
            style.base_style = self.word.styles[base]
            style.paragraph_format.page_break_before = True
        self._page_numbers(section)

    def _style(self, name: str, font: str, size: float) -> ParagraphStyle:
        style = cast(ParagraphStyle, self.word.styles[name])
        style.font.name = font
        style.font.size = Pt(size)
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), font)
        return style

    def _page_numbers(self, section) -> None:
        paragraph = section.footer.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._field(paragraph, "PAGE")

    @staticmethod
    def _field(paragraph, instruction: str) -> None:
        run = paragraph.add_run()
        for kind, text in (("begin", None), (None, instruction), ("separate", None), ("end", None)):
            if kind:
                element = OxmlElement("w:fldChar")
                element.set(qn("w:fldCharType"), kind)
            else:
                element = OxmlElement("w:instrText")
                element.set(qn("xml:space"), "preserve")
                element.text = f" {text} "
            run._r.append(element)

    def page_break(self, anchor: str | None = None) -> None:
        self.pending_break = True
        self.pending_anchor = anchor

    def bookmark(self, paragraph, name: str) -> None:
        start, end = OxmlElement("w:bookmarkStart"), OxmlElement("w:bookmarkEnd")
        start.set(qn("w:id"), str(self.bookmarks))
        start.set(qn("w:name"), name)
        end.set(qn("w:id"), str(self.bookmarks))
        self.bookmarks += 1
        paragraph._p.append(start)
        paragraph._p.append(end)

    def contents_entry(self, text: str, anchor: str, level: int = 1) -> None:
        """A contents line whose page number is a PAGEREF field; Word and WPS fill it in
        when fields update, which the document asks for on opening."""
        entry = self._new()
        entry.paragraph_format.left_indent = Cm(0.5 * level)
        entry.paragraph_format.tab_stops.add_tab_stop(
            Cm(16), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS
        )
        self.run(entry, f"{text}\t")
        self._field(entry, f"PAGEREF {anchor} \\h")

    def _new(self, text: str = "", style: str | None = None):
        """A paragraph that starts a new page after `page_break`. The break is a style's
        page-break-before, which Word, WPS and the docx-preview renderer all honour."""
        if self.pending_break:
            style = "Form Start Heading" if style == "Heading 1" else "Form Start"
            self.pending_break = False
        paragraph = self.word.add_paragraph(text, style=style)
        if self.pending_anchor:
            self.bookmark(paragraph, self.pending_anchor)
            self.pending_anchor = None
        return paragraph

    def run(self, paragraph, text: str, *, size=None, bold=False, underline=False, mark=False):
        run = paragraph.add_run(text)
        if size:
            run.font.size = Pt(size)
        run.font.bold = bold or None
        run.font.underline = underline or None
        if mark:
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        return run

    def placeholder(self, paragraph, label: str, *, prefix: str = "待填") -> None:
        self.tally.placeholders.append(label)
        self.run(paragraph, f"【{prefix}：{label}】", mark=True, underline=prefix == "待填")

    def cover(
        self,
        outline: BidOutline,
        forms: list[OutlineForm],
        solution_after: int | None = None,
        chapters: list[SolutionChapter] | None = None,
    ) -> None:
        for _ in range(6):
            self._new()
        for text, size in (
            (self.facts.project_name, 22),
            (f"项目编号：{self.facts.project_number}" if self.facts.project_number else None, 14),
        ):
            if text:
                paragraph = self._new()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                self.run(paragraph, text, size=size, bold=size > 14)
        paragraph = self._new()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self.run(paragraph, "投 标 文 件", size=36, bold=True)
        for _ in range(4):
            self._new()
        for label, value in (("投标人", self.facts.bidder), ("日期", None)):
            paragraph = self._new()
            paragraph.paragraph_format.left_indent = Cm(4)
            self.run(paragraph, f"{label}：", size=14)
            if value:
                self.run(paragraph, value, size=14, underline=True)
            else:
                self.placeholder(paragraph, label)
        self.page_break()
        heading = self._new()
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self.run(heading, "目  录", size=16, bold=True)
        for position, form in enumerate(forms):
            self.contents_entry(f"{form.label}  {short_title(form.title)}", form_anchor(form))
            if position == solution_after:
                by_item = {(c.title, c.score_range): c for c in chapters or []}
                for number, item in enumerate(
                    (item for item in outline.scoring_items if item.subjective), 1
                ):
                    self.contents_entry(f"{number}、{item.title}", f"solution_{number}", 2)
                    chapter = by_item.get((item.title, item.score_range))
                    for position, section in enumerate(
                        chapter.current.sections if chapter else [], 1
                    ):
                        self.contents_entry(
                            f"{number}.{position} {unnumbered(section.heading)}",
                            f"solution_{number}_{position}",
                            3,
                        )

    def solution(self, items: list[OutlineScoringItem], chapters: list[SolutionChapter]) -> None:
        """Render the current revision in scoring order, keeping absent drafts visible."""
        by_item = {(chapter.title, chapter.score_range): chapter for chapter in chapters}
        for chapter_number, item in enumerate((item for item in items if item.subjective), 1):
            self.page_break(f"solution_{chapter_number}")
            self._new(f"{chapter_number}、{item.title}", style="Heading 1")
            chapter = by_item.get((item.title, item.score_range))
            if chapter is None:
                self.tally.unconfirmed_chapters += 1
                self.run(self._new(), f"【待起草：{item.title}】", mark=True)
                continue
            if chapter.current.state != "confirmed":
                self.tally.unconfirmed_chapters += 1
            for position, section in enumerate(chapter.current.sections, 1):
                number = f"{chapter_number}.{position}"
                heading = self._new(style="Heading 2")
                self.bookmark(heading, f"solution_{chapter_number}_{position}")
                self.solution_text(heading, f"{number} {unnumbered(section.heading)}")
                self.solution_paragraphs(section.paragraphs)
                for sub_position, sub in enumerate(section.subsections, 1):
                    sub_heading = self._new(style="Heading 3")
                    self.bookmark(sub_heading, solution_anchor(f"{number}.{sub_position}"))
                    self.solution_text(
                        sub_heading, f"{number}.{sub_position} {unnumbered(sub.heading)}"
                    )
                    self.solution_paragraphs(sub.paragraphs)
                    if sub.table:
                        self.solution_table(sub.table)

    def solution_paragraphs(self, texts: list[str]) -> None:
        for text in texts:
            paragraph = self._new()
            paragraph.paragraph_format.first_line_indent = Pt(24)
            self.solution_text(paragraph, text)

    def solution_table(self, grid: SolutionTable) -> None:
        table = self.word.add_table(rows=1 + len(grid.rows), cols=len(grid.columns))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        # The header row repeats on each page a long table spans.
        repeat = OxmlElement("w:tblHeader")
        repeat.set(qn("w:val"), "true")
        table.rows[0]._tr.get_or_add_trPr().append(repeat)
        for row, values in enumerate([grid.columns, *grid.rows]):
            for column, value in enumerate(values):
                cell = table.cell(row, column)
                if row == 0:
                    shading = OxmlElement("w:shd")
                    shading.set(qn("w:val"), "clear")
                    shading.set(qn("w:fill"), "D9D9D9")
                    cell._tc.get_or_add_tcPr().append(shading)
                self.solution_text(cell.paragraphs[0], value, size=10.5, bold=row == 0)
        self._new()

    def solution_text(self, paragraph, text: str, *, size=None, bold=False) -> None:
        start = 0
        for match in SOLUTION_PLACEHOLDER.finditer(text):
            self.run(paragraph, text[start : match.start()], size=size, bold=bold)
            self.placeholder(paragraph, match.group(1))
            start = match.end()
        self.run(paragraph, text[start:], size=size, bold=bold)

    def blank(self, paragraph, piece: Blank, label: str) -> None:
        name = field_of(label, piece.hint)
        value = value_of(name, self.facts)
        if value:
            self.run(paragraph, f" {value} ", underline=True)
            return
        if name in FIELD_NAMES and name != "date":
            self.placeholder(paragraph, FIELD_NAMES[name])
            return
        # Signatures, prices, dates and unlabelled blanks stay underlined space for the human.
        width = max(2, round(piece.width / 6))
        if piece.hint:
            self.run(paragraph, piece.hint, underline=True)
        else:
            self.run(paragraph, " " * width, underline=True)

    def paragraph(self, block: Paragraph, *, heading: bool) -> None:
        paragraph = self._new(style="Heading 1" if heading else None)
        self._fill(paragraph, block, heading=heading)
        fmt = paragraph.paragraph_format
        start = (fmt.left_indent or 0) + (fmt.first_line_indent or 0)
        width = Pt(sum(12 if ord(char) > 0xFF else 6 for char in paragraph.text))
        if start > Pt(100) and start + width > TEXT_WIDTH >= width:
            # A signature line indented to the right outgrows the line once a name
            # fills its blank; move it left so it stays one line.
            fmt.left_indent = Pt(0)
            fmt.first_line_indent = TEXT_WIDTH - width

    def _fill(self, paragraph, block: Paragraph, *, heading: bool) -> None:
        if block.align == "center" or heading:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            fmt = paragraph.paragraph_format
            fmt.left_indent = Pt(min(block.indent, 300))
            fmt.first_line_indent = Pt(max(-block.indent, min(block.first_line, 300)))
        line = SIGN_LINE.match(re.sub(r"\s", "", plain(block)))
        if line and not heading and not any(isinstance(p, Blank) for p in block.pieces):
            value = value_of(field_of(line["label"], ""), self.facts)
            if value:
                self.run(paragraph, f"{line['label']}：")
                self.run(paragraph, f" {value} ", underline=True)
                self.run(paragraph, line["tail"] or "")
                return
        label = ""
        pieces = list(block.pieces)
        for position, piece in enumerate(pieces):
            if isinstance(piece, Text):
                size = piece.size if piece.size and abs(piece.size - 12) > 1.5 else None
                self.hinted_text(
                    paragraph, piece.text, size=None if heading else size, bold=piece.bold
                )
                label = (label + piece.text)[-16:]
                continue
            # A hint printed in brackets after the blank, like ____（投标人名称）, names it.
            after = pieces[position + 1] if position + 1 < len(pieces) else None
            hint = BRACKET_HINT.match(after.text) if isinstance(after, Text) else None
            if (
                isinstance(after, Text)
                and hint
                and not piece.hint
                and field_of("", hint.group(1)) in FIELD_NAMES
            ):
                pieces[position + 1] = Text(after.text[hint.end() :], after.size, after.bold)
                piece = Blank(width=piece.width, hint=hint.group(1))
            self.blank(paragraph, piece, label)
            label = ""

    def hinted_text(self, paragraph, text: str, *, size=None, bold=False) -> None:
        """Text whose bracketed hint names a project or party, like 致：（招标人名称）, gets the value."""
        start = 0
        for match in BRACKET_FIELD.finditer(text):
            name = field_of("", match.group(1))
            value = value_of(name, self.facts)
            if match.group(1).endswith(("名称", "编号")) and value and name in NAMED_FIELDS:
                self.run(paragraph, text[start : match.start()], size=size, bold=bold)
                self.run(paragraph, f" {value} ", size=size, bold=bold, underline=True)
                start = match.end()
        self.run(paragraph, text[start:], size=size, bold=bold)

    def grid(
        self,
        grid: Grid,
        rows: list[list[str]] | None = None,
        weights: list[float] | None = None,
        drafted: set[int] | None = None,
    ) -> None:
        """A tender table with its spans; `rows` replaces the template rows under the header,
        `weights` the tender's column widths, and `drafted` rows are highlighted."""
        header_rows = 1 if rows is not None else len(grid.heights)
        columns = len(grid.widths)
        if self.pending_break:
            self._new()
        table = self.word.add_table(rows=header_rows + len(rows or []), cols=columns)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        widths = grid.widths if weights is None else weights
        total = sum(widths) or 1
        for position, width in enumerate(widths):
            column = table.columns[position]
            column.width = Cm(16 * width / total)
            for cell in column.cells:
                cell.width = Cm(16 * width / total)
        cells: list[Cell] = [cell for cell in grid.cells if cell.row < header_rows]
        for cell in cells:
            target = table.cell(cell.row, cell.column)
            last_row = min(cell.row + cell.rows, header_rows) - 1
            last_column = min(cell.column + cell.columns, columns) - 1
            if (last_row, last_column) != (cell.row, cell.column):
                target = target.merge(table.cell(last_row, last_column))
            self._cell(target, cell.text, self._cell_label(grid, cell))
        for offset, values in enumerate(rows or []):
            for column, value in enumerate(values[:columns]):
                target = table.cell(header_rows + offset, column)
                draft = offset in (drafted or set()) and value in {DRAFT_RESPONSE, DRAFT_DEVIATION}
                self._cell(target, value, None, mark=draft)
        self._new()

    @staticmethod
    def _cell_label(grid: Grid, cell: Cell) -> str | None:
        """The label of an empty cell: the cell to its left, or above it in a two-row table."""
        if cell.text:
            return None
        for other in grid.cells:
            if other.row == cell.row and other.column + other.columns == cell.column and other.text:
                return other.text
        if len(grid.heights) == 2 and cell.row == 1:
            for other in grid.cells:
                if other.row == 0 and other.column == cell.column and other.text:
                    return other.text
        return None

    def _cell(self, target, text: str, label: str | None, *, mark: bool = False) -> None:
        paragraph = target.paragraphs[0]
        if text:
            self.run(paragraph, text, size=10.5, mark=mark)
            return
        if label:
            name = field_of(re.sub(r"\s", "", label), "")
            value = value_of(name, self.facts)
            if value:
                self.run(paragraph, value, size=10.5)

    def form(
        self,
        form: OutlineForm,
        blocks: list[Paragraph | Grid],
        deviation: list[Clause] | None,
        scoring: list[OutlineScoringItem],
        responses: list[Response],
        answers: list[str | None],
    ) -> None:
        title = re.sub(r"[\s：:]", "", form.title)
        label = re.compile(f"^{re.escape(form.label.rstrip('、.．'))}[、.．]?")

        def stripped(block: Paragraph) -> str:
            return label.sub("", re.sub(r"[\s：:]", "", plain(block)))

        heading = next(
            (
                index
                for index, block in enumerate(blocks)
                if isinstance(block, Paragraph)
                and len(text := stripped(block)) >= 2
                and (text.startswith(title[:30]) or title.startswith(text))
            ),
            None,
        )
        if heading is not None:
            blocks = _fit_title(blocks, heading, title, stripped)
        if heading is None and form.kind != "cover":
            self._new(f"{form.label} {form.title}", style="Heading 1")
        start = len(self.word.element.body)
        bare_label = False
        for index, block in enumerate(blocks):
            if isinstance(block, Grid):
                if form.kind == "deviation" and deviation is not None:
                    rows, weights, drafted = self.deviation_rows(block, deviation, responses)
                    self.grid(block, rows, weights, drafted)
                    self.deviation_pages(answers)
                    self.tally.unconfirmed_rows += len(drafted)
                    deviation = None
                elif form.kind == "scoring_index" and scoring:
                    self.grid(block, self.scoring_rows(block, scoring))
                    self.scoring_pages(scoring)
                    scoring = []
                else:
                    self.grid(block)
                continue
            if (
                (index < heading if heading is not None else form.kind != "cover")
                and form.label
                and not stripped(block)
                and plain(block).strip()
            ):
                # 附件11 printed on its own line joins the title heading.
                bare_label = True
                continue
            if index == heading and bare_label:
                block = replace(block, pieces=[Text(f"{form.label}："), *block.pieces])
            self.paragraph(block, heading=index == heading)
        # The closing address and signature lines move to the next page together
        # rather than leaving 日期 alone on it.
        closing = []
        for element in reversed(self.word.element.body[start:]):
            if element.tag == qn("w:sectPr"):
                continue
            if element.tag != qn("w:p") or len(closing) == 10:
                break
            paragraph = DocxParagraph(element, self.word._body)
            if len(paragraph.text.strip()) > 40:
                break
            closing.append(paragraph)
        for paragraph in closing[1:]:
            paragraph.paragraph_format.keep_with_next = True
        if form.kind == "copy":
            paragraph = self._new()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            wanted = short_title(form.title)
            self.tally.copies.append(wanted)
            self.run(paragraph, f"【待附：{wanted}】", mark=True)

    def scoring_pages(self, items: list[OutlineScoringItem]) -> None:
        """Point each solution item's page cell at its chapter's first page; other
        items keep an empty cell for the bidder."""
        table = self.word.tables[-1]
        header = [cell.text for cell in table.rows[0].cells]
        page = next((index for index, text in enumerate(header) if "页" in text), None)
        if page is None:
            return
        chapters = {
            item.title: number
            for number, item in enumerate((item for item in items if item.subjective), 1)
        }
        scored = [item for item in items if not PRICE_SCORE.fullmatch(item.title)]
        for row, item in zip(table.rows[1:], scored, strict=False):
            if item.title in chapters:
                paragraph = row.cells[page].paragraphs[0]
                self._field(paragraph, f"PAGEREF solution_{chapters[item.title]} \\h")

    def deviation_pages(self, answers: list[str | None]) -> None:
        """Point each clause's page cell at the solution subsection answering it; a clause
        with no answering subsection keeps an empty cell for the bidder."""
        table = self.word.tables[-1]
        header = [cell.text for cell in table.rows[0].cells]
        page = next((index for index, text in enumerate(header) if "页" in text), None)
        if page is None:
            return
        for row, number in zip(table.rows[1:], answers, strict=False):
            if number:
                paragraph = row.cells[page].paragraphs[0]
                self.run(paragraph, f"{number} 节第 ")
                self._field(paragraph, f"PAGEREF {solution_anchor(number)} \\h")
                self.run(paragraph, " 页")

    @staticmethod
    def scoring_rows(grid: Grid, items: list[OutlineScoringItem]) -> list[list[str]]:
        """Scored items other than the price score, in the item column; the page
        column stays empty until the final layout is known."""
        columns = len(grid.widths)
        return [
            [f"{item.title}（{item.score_range}分）"] + [""] * (columns - 1)
            for item in items
            if not PRICE_SCORE.fullmatch(item.title)
        ]

    @staticmethod
    def deviation_rows(
        grid: Grid, clauses: list[Clause], responses: list[Response]
    ) -> tuple[list[list[str]], list[float], set[int]]:
        """Rows under the tender's own header, column weights that give the requirement
        text room, and the rows still drafted; `deviation_pages` fills the page column."""
        headers = [
            cell.text for cell in sorted(grid.cells, key=lambda c: c.column) if cell.row == 0
        ]
        roles = [deviation_column(header) for header in headers]
        weights = [COLUMN_WEIGHTS[role] for role in roles]
        rows = []
        drafted: set[int] = set()
        for number, clause in enumerate(clauses, start=1):
            confirmed = clause_response(clause, responses)
            if confirmed is None:
                drafted.add(number - 1)
            response, deviation = confirmed or (DRAFT_RESPONSE, DRAFT_DEVIATION)
            values = {
                "number": str(number),
                "topic": clause.heading,
                "requirement": clause.text,
                "response": response,
                "deviation": deviation,
                "other": "",
            }
            rows.append([values[role] for role in roles])
        return rows, weights, drafted

    def finish(self) -> bytes:
        # Contents page numbers are fields; ask Word and WPS to compute them on opening.
        update = OxmlElement("w:updateFields")
        update.set(qn("w:val"), "true")
        self.word.settings.element.append(update)
        buffer = io.BytesIO()
        self.word.save(buffer)
        return buffer.getvalue()


def split_by_form(blocks: list[Paragraph | Grid], forms: list[OutlineForm]):
    """Blocks of each form, cut at the paragraph that starts with the form's label.

    A paragraph that also starts with the title is preferred, since list-numbered
    forms repeat labels (1、 in each group) and their clauses reuse the numbers.
    """
    texts = [
        re.sub(r"\s", "", plain(block)) if isinstance(block, Paragraph) else None
        for block in blocks
    ]
    starts: list[int] = []
    cursor = 0
    for form in forms:
        # The outline writes 1、 where the tender may print "1 投标函", so the
        # separator after the number is optional.
        titled = re.compile(
            re.escape(form.label.rstrip("、.．"))
            + r"[、.．]?"
            + re.escape(re.sub(r"\s", "", form.title)[:4])
        )
        found = next(
            (p for p in range(cursor, len(blocks)) if titled.match(texts[p] or "")),
            None,
        )
        if found is None:
            found = next(
                (p for p in range(cursor, len(blocks)) if (texts[p] or "").startswith(form.label)),
                None,
            )
        if found is None:
            starts.append(-1)
        else:
            starts.append(found)
            cursor = found + 1
    out = []
    for position, start in enumerate(starts):
        if start < 0:
            out.append([])
            continue
        later = [value for value in starts[position + 1 :] if value >= 0]
        out.append(blocks[start : later[0] if later else len(blocks)])
    return out


def text_blocks(form: OutlineForm) -> list[Paragraph | Grid]:
    """Word tenders: paragraphs and tab-separated rows from the outline text."""
    blocks: list[Paragraph | Grid] = []
    rows: list[list[str]] = []

    def flush():
        if rows:
            width = max(len(row) for row in rows)
            cells = [
                Cell(text=text, row=r, column=c)
                for r, row in enumerate(rows)
                for c, text in enumerate(row + [""] * (width - len(row)))
            ]
            blocks.append(
                Grid(cells=cells, widths=[1.0] * width, heights=[1.0] * len(rows), page=0)
            )
            rows.clear()

    for line in [f"{form.label}"] + form.text.splitlines():
        if "\t" in line:
            rows.append(line.split("\t"))
            continue
        flush()
        pieces: list[Text | Blank] = []
        for part in re.split(r"([_＿]{2,})", line):
            if re.fullmatch(r"[_＿]{2,}", part):
                pieces.append(Blank(width=len(part) * 6))
            elif part:
                pieces.append(Text(part))
        if pieces and line.rstrip().endswith(("：", ":")):
            pieces.append(Blank(width=120))
        if pieces:
            blocks.append(Paragraph(pieces=pieces, page=0))
    flush()
    return blocks


async def selected_bidder(session: AsyncSession, task_id: UUID) -> OrgProfileData | None:
    data = await session.scalar(
        select(OrgProfileRevision.data)
        .join(TaskOrgProfile, TaskOrgProfile.profile_revision_id == OrgProfileRevision.id)
        .where(TaskOrgProfile.task_id == task_id, TaskOrgProfile.active.is_(True))
        .order_by(TaskOrgProfile.created_at.desc())
        .limit(1)
    )
    return OrgProfileData.model_validate(data) if data else None


async def card_responses(session: AsyncSession, document_id: UUID) -> list[Response]:
    rows = (
        await session.execute(
            select(
                Requirement.quote,
                ResponseCardRevision.response_text,
                ResponseCardRevision.deviation,
                ResponseCardRevision.state,
            )
            .join(ResponseCard, ResponseCard.requirement_id == Requirement.id)
            .join(ResponseCardRevision, ResponseCardRevision.id == ResponseCard.current_revision_id)
            .where(Requirement.document_id == document_id)
        )
    ).all()
    return [
        Response(quote, text or "", deviation or "none", state == "confirmed")
        for quote, text, deviation, state in rows
    ]


def build(
    outline: BidOutline,
    chunks: list[dict],
    content: bytes | None,
    task: Task,
    bidder: OrgProfileData | None,
    responses: list[Response],
    chapters: list[SolutionChapter] | None = None,
) -> tuple[bytes, Tally]:
    """The DOCX for an outline; `content` is the PDF, or None for a Word tender."""
    if not outline.forms:
        raise ServiceError(
            "format_chapter_missing", "The tender has no format chapter with numbered forms", 409, 2
        )
    full_text = "\n".join(chunk["text"] or "" for chunk in chunks)
    facts = tender_facts(full_text, task, bidder)
    tally = Tally()
    renderer = Renderer(facts, tally)
    needs = next((c for c in outline.chapters if c.role == "needs"), None)
    clauses: list[Clause] = []
    if content and needs and needs.page_start and needs.page_end:
        pages = list(range(needs.page_start, needs.page_end + 1))
        clauses = needs_clauses(form_blocks(content, pages), needs.title)
    starts = [form.page_start for form in outline.forms if form.page_start]
    ends = [form.page_end for form in outline.forms if form.page_end]
    if content and starts and ends:
        blocks = form_blocks(content, list(range(min(starts), max(ends) + 1)))
        per_form = split_by_form(blocks, outline.forms)
    else:
        per_form = [text_blocks(form) for form in outline.forms]
    technical_cover = next(
        (
            index
            for index, form in enumerate(outline.forms)
            if form.kind == "cover" and "技术" in form.title
        ),
        None,
    )
    solution_form = next(
        (index for index, form in enumerate(outline.forms) if form.kind == "solution"),
        len(outline.forms) - 1,
    )
    solution_after = technical_cover if technical_cover is not None else solution_form
    answers = answering_subsections(clauses, outline.scoring_items, chapters or [])
    renderer.cover(outline, outline.forms, solution_after, chapters)
    for index, (form, blocks) in enumerate(zip(outline.forms, per_form, strict=True)):
        renderer.page_break(form_anchor(form))
        renderer.form(
            form, blocks or text_blocks(form), clauses, outline.scoring_items, responses, answers
        )
        if index == solution_after:
            renderer.solution(outline.scoring_items, chapters or [])
    first_paragraph = renderer.word.paragraphs[0]
    first_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    renderer.run(
        first_paragraph, "审阅稿" if tally.draft else "提交稿", size=14, bold=True, mark=tally.draft
    )
    renderer.word.core_properties.title = facts.project_name or "投标文件"
    return renderer.finish(), tally


async def assemble(
    session: AsyncSession, document_id: UUID, storage: Storage
) -> tuple[bytes, Tally]:
    from app.services import solution

    document = await documents.require_document(session, document_id)
    outline = await bid_outline.outline(session, document_id)
    task = await session.get_one(Task, document.task_id)
    chapters = await solution.current_chapters(session, document.id)
    if chapters:
        actor = session.info.get("actor")
        if actor is None:
            raise not_found()
        await task_access.access(session, actor, task.id, scope="solution:read")
    chunks = await bid_outline.load_chunks(session, document.id)
    is_pdf = document.media_type == "application/pdf"
    content = await storage.read(document.storage_key) if is_pdf else None
    bidder = await selected_bidder(session, task.id)
    return build(
        outline,
        chunks,
        content,
        task,
        bidder,
        await card_responses(session, document.id),
        chapters,
    )


def _fit_title(
    blocks: list[Paragraph | Grid], heading: int, title: str, stripped
) -> list[Paragraph | Grid]:
    """Make the heading block hold exactly the form title: a title wrapped onto
    the next lines takes them back, and body text joined after it is split off."""
    blocks = list(blocks)
    block = blocks[heading]
    assert isinstance(block, Paragraph)
    while (
        heading + 1 < len(blocks)
        and isinstance(after := blocks[heading + 1], Paragraph)
        and len(text := stripped(block)) < len(title)
        and title.startswith(text + stripped(after))
    ):
        block = replace(block, pieces=[*block.pieces, *after.pieces])
        del blocks[heading + 1]
    rest: list[Piece] = []
    if len(stripped(block)) > len(title):
        for index, piece in enumerate(block.pieces):
            if not isinstance(piece, Text):
                continue
            for cut in range(len(piece.text) + 1):
                head = [*block.pieces[:index], replace(piece, text=piece.text[:cut])]
                if stripped(replace(block, pieces=head)) == title:
                    tail = replace(piece, text=piece.text[cut:].lstrip())
                    rest = (
                        [tail, *block.pieces[index + 1 :]]
                        if tail.text
                        else block.pieces[index + 1 :]
                    )
                    block = replace(block, pieces=head)
                    break
            if rest:
                break
    blocks[heading] = block
    if rest:
        blocks.insert(heading + 1, replace(block, pieces=rest, align="left"))
    return blocks
