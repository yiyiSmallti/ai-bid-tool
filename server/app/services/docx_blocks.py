"""Split a Word document into addressable blocks: paragraphs and table cells.

Word files have no reliable pages, so every block gets a stable identifier, the
heading path above it, and a human-readable label. Quotes are later checked
against the single block they cite.
"""

import io
import re
from dataclasses import dataclass, field
from typing import Any

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

from app.schemas.contracts import Block, SectionText

SECTION_CHARS = 8000
HEADING_STYLE = re.compile(r"^(?:heading|标题)\s*(\d)", re.IGNORECASE)
FALLBACK_HEADINGS = (
    (1, re.compile(r"^第[一二三四五六七八九十百零〇\d]+[章篇部]")),
    (2, re.compile(r"^[一二三四五六七八九十]+、")),
    (3, re.compile(r"^[（(][一二三四五六七八九十]+[）)]")),
)


def style_outline(style) -> int | None:
    # Walk the style inheritance chain for an explicit outline level.
    while style is not None:
        ppr = style.element.pPr
        node = ppr.find(qn("w:outlineLvl")) if ppr is not None else None
        if node is not None:
            return int(node.get(qn("w:val")))
        style = style.base_style
    return None


def heading_level(paragraph: Paragraph) -> int | None:
    style = paragraph.style
    if style is not None:
        match = HEADING_STYLE.match(style.name or "")
        if match:
            return int(match.group(1))
    ppr = paragraph._p.pPr
    node = ppr.find(qn("w:outlineLvl")) if ppr is not None else None
    level = int(node.get(qn("w:val"))) if node is not None else style_outline(style)
    # Outline level 9 means body text.
    return level + 1 if level is not None and level < 9 else None


def fallback_level(text: str) -> int | None:
    for level, pattern in FALLBACK_HEADINGS:
        if pattern.match(text):
            return level
    return None


@dataclass
class Walker:
    body: Any
    use_fallback: bool
    blocks: list[Block] = field(default_factory=list)
    path: list[tuple[int, str]] = field(default_factory=list)
    paragraph_total: int = 0
    paragraph_in_section: int = 0
    tables: int = 0

    def section(self) -> list[str]:
        return [text for _, text in self.path]

    def label(self, tail: str) -> str:
        return " > ".join([*self.section(), tail]) if tail else " > ".join(self.section())

    def walk(self, container) -> None:
        for child in container.iterchildren():
            if child.tag == qn("w:p"):
                self.paragraph(Paragraph(child, self.body))
            elif child.tag == qn("w:tbl"):
                self.tables += 1
                self.table(child, f"t{self.tables}", f"表 {self.tables}", self.tables)
            elif child.tag == qn("w:sdt"):
                # Content controls wrap ordinary paragraphs and tables.
                content = child.find(qn("w:sdtContent"))
                if content is not None:
                    self.walk(content)

    def paragraph(self, paragraph: Paragraph) -> None:
        text = paragraph.text.strip()
        if not text:
            return
        self.paragraph_total += 1
        level = fallback_level(text) if self.use_fallback else heading_level(paragraph)
        if level is not None:
            while self.path and self.path[-1][0] >= level:
                self.path.pop()
            self.path.append((level, text))
            self.paragraph_in_section = 0
            self.blocks.append(
                Block(
                    block_id=f"p{self.paragraph_total}",
                    kind="paragraph",
                    section_path=self.section(),
                    label=self.label(""),
                    text=text,
                )
            )
            return
        self.paragraph_in_section += 1
        self.blocks.append(
            Block(
                block_id=f"p{self.paragraph_total}",
                kind="paragraph",
                section_path=self.section(),
                paragraph=self.paragraph_in_section,
                label=self.label(f"第 {self.paragraph_in_section} 段"),
                text=text,
            )
        )

    def table(self, tbl, prefix: str, caption: str, number: int) -> None:
        for row_index, row in enumerate(tbl.findall(qn("w:tr")), start=1):
            trpr = row.find(qn("w:trPr"))
            before = trpr.find(qn("w:gridBefore")) if trpr is not None else None
            column = 1 + (int(before.get(qn("w:val"))) if before is not None else 0)
            for cell in row.findall(qn("w:tc")):
                tcpr = cell.find(qn("w:tcPr"))
                span_node = tcpr.find(qn("w:gridSpan")) if tcpr is not None else None
                span = int(span_node.get(qn("w:val"))) if span_node is not None else 1
                merge = tcpr.find(qn("w:vMerge")) if tcpr is not None else None
                # A vMerge without "restart" continues the cell above; it is not a new cell.
                continued = merge is not None and merge.get(qn("w:val")) != "restart"
                if not continued:
                    self.cell(cell, prefix, caption, number, row_index, column)
                column += span

    def cell(self, cell, prefix: str, caption: str, number: int, row: int, column: int) -> None:
        block_id = f"{prefix}r{row}c{column}"
        label = f"{caption} 第 {row} 行第 {column} 列"
        text = "\n".join(Paragraph(p, self.body).text for p in cell.findall(qn("w:p"))).strip()
        if text:
            self.blocks.append(
                Block(
                    block_id=block_id,
                    kind="cell",
                    section_path=self.section(),
                    table=number,
                    row=row,
                    column=column,
                    label=self.label(label),
                    text=text,
                )
            )
        # Nested tables become their own blocks instead of joining the cell text.
        for index, inner in enumerate(cell.findall(qn("w:tbl")), start=1):
            self.table(inner, f"{block_id}/t{index}", f"{label} > 表 {index}", number)


def skipped_content(document) -> list[str]:
    body = document.element.body
    counts = {
        "text boxes": len(body.findall(".//" + qn("w:txbxContent"))),
        "images": len(body.findall(".//" + qn("w:drawing"))),
    }
    notes = [f"{count} {name}" for name, count in counts.items() if count]
    part = document.part
    for kind, label in (
        ("footnotes", "footnotes"),
        ("endnotes", "endnotes"),
        ("comments", "comments"),
    ):
        for rel in part.rels.values():
            if rel.reltype.endswith("/" + kind) and not rel.is_external:
                # Separators carry a w:type (their ids differ between Word and WPS);
                # only untyped notes and comments are real content.
                element = kind[:-1].encode()
                pattern = rb"<w:" + element + rb" (?![^>]*w:type=)[^>]*>"
                real = len(re.findall(pattern, rel.target_part.blob))
                if real:
                    notes.append(f"{real} {label}")
    if any(
        section.header.paragraphs and any(p.text.strip() for p in section.header.paragraphs)
        for section in document.sections
    ) or any(
        any(p.text.strip() for p in section.footer.paragraphs) for section in document.sections
    ):
        notes.append("headers/footers")
    return notes


def parse_docx(
    content: bytes, section_chars: int = SECTION_CHARS
) -> tuple[list[SectionText], list[str]]:
    document = Document(io.BytesIO(content))
    body = document._body
    has_headings = any(
        heading_level(Paragraph(p, body)) is not None
        for p in document.element.body.iter(qn("w:p"))
        if Paragraph(p, body).text.strip()
    )
    walker = Walker(body=body, use_fallback=not has_headings)
    walker.walk(document.element.body)
    warnings = []
    if not has_headings:
        warnings.append(
            "No heading styles found; sections were inferred from numbering such as 第X章 and 一、."
        )
    skipped = skipped_content(document)
    if skipped:
        warnings.append("Word content not parsed for citations: " + ", ".join(skipped) + ".")
    if not walker.blocks:
        return [], [*warnings, "The Word document contains no text."]

    ids = [block.block_id for block in walker.blocks]
    assert len(ids) == len(set(ids)), "block identifiers must be unique"
    sections: list[list[Block]] = [[]]
    size = 0
    for block in walker.blocks:
        top_heading = (
            block.kind == "paragraph" and block.paragraph is None and len(block.section_path) <= 2
        )
        if sections[-1] and (top_heading or size + len(block.text) > section_chars):
            sections.append([])
            size = 0
        sections[-1].append(block)
        size += len(block.text)
    return [
        SectionText(seq=index, text="\n".join(b.text for b in group), blocks=group)
        for index, group in enumerate(sections, start=1)
    ], warnings
