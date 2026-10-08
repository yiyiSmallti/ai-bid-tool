"""Local editable report rendering using the existing deterministic Word stack."""

import io
import json
from datetime import UTC, datetime
from typing import cast

from docx import Document
from docx.shared import Pt
from docx.styles.style import ParagraphStyle

from app.schemas.bid_review_report import ADVISORY, SECTIONS
from app.services.export_renderer import RENDERER_PROFILE, RenderLimits, deterministic_package

RENDERER_IDENTITY = "bid-review-report-v1;" + RENDERER_PROFILE
LIMITS = RenderLimits(max_output_bytes=100 * 1024 * 1024, max_expanded_bytes=200 * 1024 * 1024)
LABELS = {
    "responded": "已响应",
    "deviation": "偏离",
    "missing": "缺失",
    "unknown": "未知",
    "fatal": "致命",
    "high": "高风险",
    "medium": "中风险",
    "open": "待审查",
    "dismissed": "已驳回",
    "confirmed": "已确认",
    "classify": "分类",
    "dismiss": "驳回",
    "confirm": "确认",
    "reopen": "重新打开",
    "commercial": "商务",
    "technical": "技术",
    "rejection": "废标风险",
    "lost_points": "扣分风险",
    "both": "废标及扣分风险",
    "uncertain": "影响未确定",
    "rule": "本地规则",
    "model": "模型判断",
    "complete": "已完成指定范围",
    "partial": "部分完成",
    "applies": "适用",
    "not_applicable": "不适用",
    "alternative": "允许替代签章",
    "valid": "有效",
    "invalid": "无效",
    "unresolved": "未核实",
    "trusted": "链到本地信任根",
    "untrusted": "未链到本地信任根",
    "unsigned": "未发现数字签名",
    "company_seal": "单位公章",
    "legal_representative_signature": "法定代表人签字",
    "authorized_agent_signature": "授权代理人签字",
    "personal_seal": "个人印章",
    "date": "日期",
    "seam_seal": "骑缝章",
    "every_page_electronic_seal": "逐页电子章",
    "pdf_digital_signature": "PDF 数字签名",
    "every_page": "每页",
    "seam_group": "骑缝组",
    "specified": "指定位置",
    "company": "单位",
    "legal_representative": "法定代表人",
    "authorized_agent": "授权代理人",
}


def display(value):
    if value is None:
        return "未知 / 未提供"
    if isinstance(value, bool):
        return "是" if value else "否"
    if isinstance(value, list):
        return "、".join(display(item) for item in value) or "无"
    text = str(value)
    return f"{LABELS[text]}（{text}）" if text in LABELS else text


def fields_table(document, fields):
    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for label, value in fields:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = label, display(value)
    return table


def citation_text(citations):
    lines = []
    for item in citations:
        page_kind = "固定转换页" if item.get("page_label") == "rendered_docx" else "原件页"
        lines.append(f"文件 {item['document_id']} · {page_kind} {item['page']}：{item['quote']}")
        location = item.get("location")
        if location:
            path = " / ".join(location.get("section_path", []))
            lines.append(f"Word 结构位置：{path} {location.get('block_id', '')}".strip())
    return "\n".join(lines)


def finding_table(document, row):
    finding, decision = row["finding"], row["decision"]
    basis = finding["basis"]
    fields = [
        ("审查项目", finding["title"]),
        ("发现编号", finding["id"]),
        ("机器原始响应状态", finding["outcome"]),
        ("原始风险等级", finding["severity"]),
        ("影响", finding["impact"]),
        ("招标页码与原文", citation_text(finding["tender_support"])),
        (
            "投标页码与原文",
            citation_text(finding["bid_support"])
            or "没有可引用响应原文；依据以下缺失检索或本地文件校验记录。",
        ),
        ("依据类型", basis["kind"]),
        ("规则 / 提示版本", basis["rule_or_prompt_version"]),
        ("判断依据", finding["explanation"]),
        ("人工处置状态", finding["state"]),
        ("职责", finding["review_domain"]),
        ("补救动作", finding["remediation"]),
    ]
    obligation = row.get("obligation")
    if obligation:
        fields[1:1] = [
            ("条款类别", obligation["category"]),
            (
                "强制标记",
                "、".join(
                    mark
                    for mark, key in (("★", "starred"), ("▲", "triangle"))
                    if obligation.get(key)
                )
                or "无 ★/▲ 标记",
            ),
        ]
    if basis.get("model"):
        fields.append(("模型与置信度", f"{basis['model']}；{basis.get('confidence', '未知')}"))
    if decision:
        fields += [
            ("人工决定", decision["action"]),
            ("决定原因", decision["reason"]),
            ("决定时间", decision["decided_at"]),
            ("决定人", decision["decided_by"]),
            ("决定修订", decision["revision"]),
        ]
    else:
        fields.append(("人工决定", "尚无人工决定"))
    search = finding.get("absence_search")
    if search:
        pages = [
            f"文件 {page['document_id']} 第 {page['page']} 页" for page in search["searched_pages"]
        ]
        fields += [
            ("缺失检索方法与范围", f"{search['method']}；{search['coverage']}"),
            ("实际检索页", "\n".join(pages) or "无页级检索；见文件清单或范围限制"),
            ("实际检索文件", search["inspected_bid_document_ids"]),
            ("检索限制", search["limitation_codes"]),
        ]
    local = finding.get("rule_evidence")
    if local:
        fields += [
            ("本地校验文件", local["document_ids"]),
            ("本地校验记录", local["validation_ids"]),
        ]
    fields.append(("未确定范围", finding["limitation_codes"]))
    fields_table(document, fields)


def basic_information(document, row, snapshot):
    fields_table(
        document,
        [
            ("任务", row["task_name"]),
            ("招标编号", row["tender_number"]),
            ("评估日期", row["assessment_date"]),
            ("任务编号", row["task_id"]),
            ("提交版本", row["submission_id"]),
            ("审查运行", row["review_id"]),
            ("运行作业", row["review_job_id"]),
            ("运行时间", row.get("review_created_at")),
            ("本地准备版本", row.get("preparation_id")),
            ("完成范围", row["completion"]),
            ("审查输入哈希", snapshot["report_input_hash"]),
            ("人工决定快照哈希", snapshot["decisions_snapshot_sha256"]),
        ],
    )
    for key, title in (("tender_documents", "招标文件"), ("bid_documents", "投标文件")):
        document.add_heading(title, 2)
        for item in row[key]:
            fields_table(
                document,
                [
                    ("文件", item.get("name", item["id"])),
                    ("文件编号", item["id"]),
                    ("材料类型", item["kind"]),
                    ("字节数", item["size_bytes"]),
                    ("文件哈希", item["sha256"]),
                ],
            )


def signing_requirement(document, requirement):
    fields_table(
        document,
        [
            ("签章要求", requirement["mark_types"]),
            ("适用状态", requirement["applicability"]),
            ("要求主体", requirement["owner_roles"]),
            ("必须填写日期", requirement["date_required"]),
            ("位置规则", requirement["location_rule"]),
            (
                "招标原文",
                citation_text([requirement["citation"]])
                if requirement.get("citation")
                else "无已确认引用",
            ),
        ],
    )
    for location in requirement["required_locations"]:
        fields_table(
            document,
            [
                ("所需文件", location.get("document_id")),
                ("所需页码", location.get("page")),
                ("骑缝组", location.get("group_id")),
                ("位置状态", "未核实（unresolved）；未判定存在"),
                ("未知原因", location["reason_code"]),
            ],
        )


def pdf_validation(document, value):
    fields_table(
        document,
        [
            ("PDF 校验文件", value["document_id"]),
            ("校验状态", value["status"]),
            ("校验器", value["validator_identity"]),
            ("校验时间", value["validation_time"]),
            ("信任根快照", value["trust_store_sha256"]),
            ("最终修订覆盖", value["final_revision"]["status"]),
            ("最后签名后修改", value["final_revision"]["modified_after_last_signature"]),
        ],
    )
    for signature in value["signatures"]:
        fields_table(
            document,
            [
                ("签名序号", signature["signature_index"]),
                ("字节范围覆盖", signature["coverage_status"]),
                ("密码学校验", signature["crypto_status"]),
                ("内容摘要", signature["content_digest_status"]),
                ("签名值", signature["signature_value_status"]),
                ("证书有效期", signature["certificate_validity_status"]),
                ("信任链", signature["trust_status"]),
                ("时间戳", signature["timestamp_status"]),
                ("声明签署时间（非可信时间）", signature["claimed_signing_time"]),
                ("签署后修改", signature["modified_after_signing"]),
                ("吊销状态", "未知；未联网查询"),
                ("限制", signature["reason_codes"]),
            ],
        )
    document.add_paragraph("数字签名校验不证明可见签章位置、主体或日期填写完整。")


def render_report(snapshot):
    """Return the same frozen report as deterministic DOCX and canonical console JSON."""
    document = Document()
    normal = cast(ParagraphStyle, document.styles["Normal"])
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    properties = document.core_properties
    properties.author, properties.last_modified_by = "", ""
    properties.created = properties.modified = datetime(2000, 1, 1, tzinfo=UTC)
    properties.title = "投标文件审查报告"
    document.add_heading("投标文件审查报告", 0)
    document.add_paragraph(ADVISORY)
    for key, title in SECTIONS:
        document.add_heading(title, 1)
        for row in snapshot["rows"][key]:
            kind = row["kind"]
            if kind == "finding":
                finding_table(document, row)
            elif kind == "basic_information":
                basic_information(document, row, snapshot)
            elif kind == "notice":
                document.add_paragraph(f"{row['label']}：{display(row['text'])}")
                for value in row.get("values", []):
                    document.add_paragraph(display(value))
            elif kind == "signing_requirement":
                signing_requirement(document, row["requirement"])
            elif kind == "pdf_validation":
                pdf_validation(document, row["validation"])
            elif kind == "obligation":
                document.add_paragraph("未覆盖条款：" + row["obligation"]["text"])
                document.add_paragraph(citation_text([row["obligation"]["citation"]]))
                document.add_paragraph("响应状态：未知（unknown）；未完成符合性审查")
            elif kind == "usage":
                usage = row["usage"]
                fields_table(
                    document,
                    [
                        ("服务提供者", usage["provider"]),
                        ("模型", usage["model"]),
                        ("模型版本", usage["version"]),
                        ("耗时（毫秒）", usage["duration_ms"]),
                        (
                            "输入 / 输出 token",
                            f"{usage['input_tokens']} / {usage['output_tokens']}",
                        ),
                        ("供应商成本 USD", usage["usd"]),
                    ],
                )
            elif kind == "cost":
                cost = row["cost"]
                fields_table(
                    document,
                    [
                        ("费用统计", "审查运行实际用量；本地报告生成不调用模型"),
                        ("模型 token", cost["llm_tokens"]),
                        ("供应商成本 USD", cost["usd"]),
                        ("计费币种", cost["billing_currency"]),
                        ("平台收费", cost["charge"]),
                        ("任务计入费用", cost["task_amount"]),
                        ("未定价调用", cost["unpriced_calls"]),
                        ("未解决调用 / 预留", cost["unresolved_calls"]),
                    ],
                )
            else:
                raise ValueError("Unsupported report row kind")
    output = io.BytesIO()
    document.save(output)
    docx_bytes = deterministic_package(output.getvalue(), LIMITS)
    checked = Document(io.BytesIO(docx_bytes))
    paragraphs = {paragraph.text for paragraph in checked.paragraphs}
    if ADVISORY not in paragraphs or any(title not in paragraphs for _, title in SECTIONS):
        raise ValueError("Rendered report lost a required section or advisory statement")
    console = json.dumps(
        snapshot, sort_keys=True, separators=(",", ":"), ensure_ascii=True
    ).encode()
    if len(console) > LIMITS.max_output_bytes:
        raise ValueError("Console report exceeds output limit")
    return docx_bytes, console
