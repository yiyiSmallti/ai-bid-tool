"""Strict template adaptation and deterministic editable DOCX rendering.

The caller resolves every database and storage reference before entering this
module.  The renderer accepts only a fixed manifest plus local, private files
created for the current worker attempt; it never follows package links or URLs.
"""

from __future__ import annotations

import hashlib
import io
import json
import platform
import re
import zipfile
import zlib
from collections.abc import Mapping, Sequence
from dataclasses import dataclass
from datetime import UTC, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path, PurePosixPath
from typing import Any, NoReturn

import docx
import lxml
from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Pt, RGBColor
from lxml import etree  # pyright: ignore[reportAttributeAccessIssue]

from app.services import redaction

SECTIONS = (
    "substantive",
    "commercial",
    "technical",
    "comply_only",
    "gaps",
    "evidence_appendix",
)
TABLE_SECTIONS = frozenset({"substantive", "commercial", "technical"})
COLUMNS = ("ordinal", "requirement", "response", "compliance")
ANCHORS = {section: f"{{{{bid.{section}}}}}" for section in SECTIONS}
METADATA_MARKERS = frozenset({"{{bid.task_name}}", "{{bid.tender_number}}"})
SECTION_TITLES = {
    "substantive": "实质性响应一览表",
    "commercial": "商务响应偏离表",
    "technical": "技术响应偏离表",
    "comply_only": "须遵守条款清单",
    "gaps": "缺口清单",
    "evidence_appendix": "证据附件索引",
}
COLUMN_TITLES = {
    "ordinal": "序号",
    "requirement": "招标文件要求",
    "response": "投标文件响应内容",
    "compliance": "响应情况",
}
DEVIATION_LABELS = {"none": "无偏离", "positive": "正偏离", "negative": "负偏离"}
# Simple tables: column headings and width shares in percent.
COMPLY_COLUMNS = (("序号", 8), ("招标文件要求", 72), ("响应情况", 20))
GAP_COLUMNS = (("序号", 8), ("招标文件要求", 62), ("缺口原因", 30))
INDEX_COLUMNS = (("编号", 10), ("材料", 30), ("内容", 35), ("对应条款", 25))
GAP_LABELS = {
    "missing_card": "缺少响应卡",
    "unconfirmed": "响应尚未确认",
    "rejected": "响应已驳回",
    "needs_material": "待补材料",
    "unclassified": "响应尚未分类",
    "stale_material": "材料选择已失效",
    "invalid_citation": "招标原文引用失效",
    "needs_reconfirmation": "引用修复后待重新确认",
}
ADAPTER_VERSION = "docx-template-adapter-v1"
# Page height kept for an attachment caption in the template heading style. Its
# revision, hashes and requirement ids take three or four heading lines in Word;
# one inch pushed every page image onto a page of its own.
CAPTION_RESERVE = 2 * 914_400
RENDERER_PROFILE = (
    "docx-export-v3"
    f";implementation={platform.python_implementation()}"
    f";python={platform.python_version()}"
    f";platform={platform.system()}-{platform.machine()}"
    f";python-docx={docx.__version__}"
    f";lxml={lxml.__version__}"
    f";libxml={'.'.join(str(part) for part in etree.LIBXML_VERSION)}"
    f";zlib={zlib.ZLIB_VERSION}/{zlib.ZLIB_RUNTIME_VERSION}"
    ";zip=deflate9;ooxml=sorted-attrs-v1"
)
_DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_ZIP_TIME = (1980, 1, 1, 0, 0, 0)
_SHA256 = re.compile(r"^[0-9a-f]{64}$")
_PLACEHOLDER = re.compile(r"\{\{bid\.[^{}]+\}\}")
_W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
_XML_NS = "http://www.w3.org/XML/1998/namespace"


@dataclass(frozen=True)
class RenderLimits:
    max_requirements: int = 2_000
    max_attachments: int = 300
    max_output_bytes: int = 512 * 1024 * 1024
    max_expanded_bytes: int = 1024 * 1024 * 1024
    max_template_bytes: int = 40 * 1024 * 1024
    max_page_bytes: int = 40 * 1024 * 1024
    max_page_pixels: int = 20_000_000
    max_image_side: int = 8_192


DEFAULT_LIMITS = RenderLimits()


@dataclass(frozen=True)
class AuthorizedExportInputs:
    """Private worker files resolved from the run's authorized descriptors."""

    template_path: Path
    evidence_page_paths: Mapping[int, Path]


@dataclass(frozen=True)
class TemplateInspection:
    template_sha256: str
    static_content_hash: str
    anchors: tuple[dict[str, object], ...]
    metadata_fields: tuple[str, ...]
    adapter_version: str = ADAPTER_VERSION


@dataclass(frozen=True)
class RenderCandidate:
    content: bytes
    sha256: str
    size_bytes: int
    media_type: str
    renderer_profile: str
    manifest_hash: str


class ExportRenderError(ValueError):
    """A stable, non-retryable renderer refusal."""

    def __init__(self, code: str, message: str):
        self.code = code
        super().__init__(message)


def _fail(code: str, message: str) -> NoReturn:
    raise ExportRenderError(code, message)


def canonical_json(value: Mapping[str, object]) -> bytes:
    try:
        return json.dumps(
            value,
            sort_keys=True,
            ensure_ascii=True,
            separators=(",", ":"),
            allow_nan=False,
        ).encode("utf-8")
    except (TypeError, ValueError) as exc:
        raise ExportRenderError(
            "invalid_export_manifest", "Manifest is not canonical JSON"
        ) from exc


def manifest_sha256(manifest: Mapping[str, object]) -> str:
    return hashlib.sha256(canonical_json(manifest)).hexdigest()


def _mapping(value: object, label: str) -> Mapping[str, Any]:
    if not isinstance(value, Mapping):
        _fail("invalid_export_manifest", f"{label} must be an object")
    return value


def _sequence(value: object, label: str) -> Sequence[Any]:
    if not isinstance(value, (list, tuple)):
        _fail("invalid_export_manifest", f"{label} must be an array")
    return value


def _text(value: object, label: str) -> str:
    if not isinstance(value, str) or not value.strip():
        _fail("invalid_export_manifest", f"{label} must be nonblank text")
    return value


def _sha(value: object, label: str) -> str:
    if not isinstance(value, str) or _SHA256.fullmatch(value) is None:
        _fail("invalid_export_manifest", f"{label} must be a lowercase SHA-256")
    return value


def _integer(value: object, label: str, *, minimum: int = 0) -> int:
    if not isinstance(value, int) or isinstance(value, bool) or value < minimum:
        _fail("invalid_export_manifest", f"{label} must be an integer >= {minimum}")
    return value


def _xml(data: bytes, part: str) -> etree._Element:
    parser = etree.XMLParser(resolve_entities=False, no_network=True, huge_tree=False)
    try:
        root = etree.fromstring(data, parser=parser)
    except (etree.XMLSyntaxError, ValueError) as exc:
        raise ExportRenderError("unsupported_template", f"Malformed XML part: {part}") from exc
    if root.getroottree().docinfo.doctype:
        _fail("unsupported_template", f"DOCTYPE is unsupported in {part}")
    return root


_UNSUPPORTED_LOCAL_NAMES = frozenset(
    {
        "altChunk",
        "commentRangeStart",
        "commentRangeEnd",
        "commentReference",
        "customXml",
        "del",
        "delText",
        "embeddedFont",
        "fldData",
        "hyperlink",
        "ins",
        "moveFrom",
        "moveFromRangeStart",
        "moveFromRangeEnd",
        "moveTo",
        "moveToRangeStart",
        "moveToRangeEnd",
        "object",
        "oleObject",
        "pict",
        "smartTag",
        "specVanish",
        "subDoc",
        "txbxContent",
        "vanish",
        "webHidden",
    }
)
_STORY_ALLOWED = frozenset(
    {
        "body",
        "br",
        "b",
        "bCs",
        "bookmarkEnd",
        "bookmarkStart",
        "bottom",
        "cantSplit",
        "cnfStyle",
        "color",
        "cols",
        "contextualSpacing",
        "docGrid",
        "document",
        "drawing",
        "eastAsianLayout",
        "evenAndOddHeaders",
        "fldChar",
        "fldSimple",
        "footerReference",
        "ftr",
        "gridCol",
        "hdr",
        "headerReference",
        "highlight",
        "i",
        "iCs",
        "ind",
        "instrText",
        "jc",
        "keepLines",
        "keepNext",
        "lang",
        "lastRenderedPageBreak",
        "left",
        "noProof",
        "numId",
        "numPr",
        "outlineLvl",
        "p",
        "pBdr",
        "pageBreakBefore",
        "pPr",
        "pStyle",
        "pgMar",
        "pgNumType",
        "pgSz",
        "proofErr",
        "r",
        "right",
        "rPr",
        "rStyle",
        "sectPr",
        "shd",
        "spacing",
        "sz",
        "szCs",
        "t",
        "tab",
        "tabs",
        "tbl",
        "tblBorders",
        "tblCellMar",
        "tblGrid",
        "tblHeader",
        "tblLayout",
        "tblLook",
        "tblPr",
        "tblStyle",
        "tblW",
        "tc",
        "tcBorders",
        "tcMar",
        "tcPr",
        "tcW",
        "textAlignment",
        "titlePg",
        "top",
        "tr",
        "trHeight",
        "trPr",
        "type",
        "u",
        "vertAlign",
        "vMerge",
        "w",
        "widowControl",
    }
)
_ALLOWED_RELATIONSHIP_SUFFIXES = frozenset(
    {
        "/customXml",
        "/customXmlProps",
        "/extended-properties",
        "/fontTable",
        "/footer",
        "/header",
        "/numbering",
        "/officeDocument",
        "/package",
        "/settings",
        "/styles",
        "/stylesWithEffects",
        "/theme",
        "/webSettings",
        "/core-properties",
        "/thumbnail",
    }
)
_ALLOWED_MEMBER_PATTERNS = (
    re.compile(r"^\[Content_Types\]\.xml$"),
    re.compile(r"^_rels/\.rels$"),
    re.compile(r"^docProps/(?:app|core)\.xml$"),
    re.compile(r"^docProps/thumbnail\.jpeg$"),
    re.compile(r"^word/document\.xml$"),
    re.compile(r"^word/_rels/document\.xml\.rels$"),
    re.compile(
        r"^word/(?:styles|stylesWithEffects|settings|webSettings|fontTable|numbering)\.xml$"
    ),
    re.compile(r"^word/theme/theme\d+\.xml$"),
    re.compile(r"^word/(?:header|footer)\d+\.xml$"),
    re.compile(r"^customXml/item1\.xml$"),
    re.compile(r"^customXml/itemProps1\.xml$"),
    re.compile(r"^customXml/_rels/item1\.xml\.rels$"),
)


def _safe_zip(content: bytes, limits: RenderLimits) -> dict[str, bytes]:
    if not content or len(content) > limits.max_template_bytes:
        _fail("unsupported_template", "Template bytes exceed the approved limit")
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            infos = archive.infolist()
            if not infos or len({entry.filename for entry in infos}) != len(infos):
                _fail("unsupported_template", "Template package has duplicate or no members")
            expanded = 0
            files: dict[str, bytes] = {}
            for entry in infos:
                path = PurePosixPath(entry.filename)
                if (
                    entry.flag_bits & 1
                    or path.is_absolute()
                    or ".." in path.parts
                    or "\\" in entry.filename
                    or entry.is_dir()
                ):
                    _fail("unsupported_template", "Template package contains an unsafe member")
                expanded += entry.file_size
                if expanded > limits.max_expanded_bytes:
                    _fail("export_resource_limit", "Template expanded size exceeds the limit")
                files[entry.filename] = archive.read(entry)
    except (zipfile.BadZipFile, RuntimeError, OSError) as exc:
        raise ExportRenderError("unsupported_template", "Template is not a readable DOCX") from exc
    required = {"[Content_Types].xml", "_rels/.rels", "word/document.xml"}
    if not required.issubset(files):
        _fail("unsupported_template", "Template is missing required DOCX parts")
    lowered = {name.casefold() for name in files}
    if any(
        token in name
        for name in lowered
        for token in ("vbaproject", "vbadata", "embeddings/", "activex/", "comments")
    ):
        _fail("unsupported_template", "Template contains macros, OLE, ActiveX, or comments")
    if any(name.startswith("word/media/") for name in lowered):
        _fail("unsupported_template", "Template embedded media is unsupported")
    unknown = [
        name
        for name in files
        if not any(pattern.fullmatch(name) for pattern in _ALLOWED_MEMBER_PATTERNS)
    ]
    if unknown:
        _fail("unsupported_template", "Template package contains an unknown part")
    content_types = files["[Content_Types].xml"].lower()
    if b"macroenabled" in content_types or b"vba" in content_types or b"oleobject" in content_types:
        _fail("unsupported_template", "Template contains an unsupported package content type")
    return files


def _scan_relationships(files: Mapping[str, bytes]) -> None:
    for name, data in files.items():
        if not name.endswith(".rels"):
            continue
        root = _xml(data, name)
        if etree.QName(root).localname != "Relationships":
            _fail("unsupported_template", f"Unknown relationships XML in {name}")
        for relation in root:
            if etree.QName(relation).localname != "Relationship":
                _fail("unsupported_template", f"Unknown relationship element in {name}")
            if relation.get("TargetMode", "Internal") != "Internal":
                _fail("unsupported_template", "External package relationships are unsupported")
            rel_type = relation.get("Type", "")
            if not any(rel_type.endswith(suffix) for suffix in _ALLOWED_RELATIONSHIP_SUFFIXES):
                _fail("unsupported_template", f"Unsupported relationship type in {name}")


def _scan_xml(files: Mapping[str, bytes]) -> None:
    story_names = {
        name
        for name in files
        if name == "word/document.xml"
        or re.fullmatch(r"word/(?:header|footer)\d+\.xml", name) is not None
    }
    for name, data in files.items():
        if not (name.endswith(".xml") or name.endswith(".rels")):
            continue
        root = _xml(data, name)
        for element in root.iter():
            local = etree.QName(element).localname
            if local in _UNSUPPORTED_LOCAL_NAMES:
                _fail("unsupported_template", f"Unsupported {local} element in {name}")
        if name in story_names:
            for element in root.iter():
                if (
                    etree.QName(element).namespace == _W_NS
                    and etree.QName(element).localname not in _STORY_ALLOWED
                ):
                    _fail("unsupported_template", f"Unknown story XML in {name}")
            if root.xpath(".//w:drawing", namespaces={"w": _W_NS}):
                _fail("unsupported_template", f"Template drawing is unsupported in {name}")
            _validate_fields(root, name)
    custom = files.get("customXml/item1.xml")
    if custom is not None:
        root = _xml(custom, "customXml/item1.xml")
        if (
            etree.QName(root).localname != "Sources"
            or any(root)
            or "".join(root.itertext()).strip()
        ):
            _fail("unsupported_template", "Custom XML data is unsupported")
    if any(name.startswith("customXml/") for name in files) and custom is None:
        _fail("unsupported_template", "Unknown custom XML is unsupported")


def _validate_fields(root: etree._Element, part: str) -> None:
    complex_instructions = [
        " ".join(text.split()).upper()
        for text in root.xpath(".//w:instrText/text()", namespaces={"w": _W_NS})
    ]
    simple_fields = root.xpath(".//w:fldSimple", namespaces={"w": _W_NS})
    simple_instructions = [
        " ".join(str(field.get(f"{{{_W_NS}}}instr", "")).split()).upper() for field in simple_fields
    ]
    instructions = [*complex_instructions, *simple_instructions]
    for instruction in instructions:
        command = instruction.split()[0] if instruction else ""
        if command not in {"PAGE", "NUMPAGES"}:
            _fail("unsupported_template", f"Unsupported field instruction in {part}")
    fld_chars = root.xpath(".//w:fldChar", namespaces={"w": _W_NS})
    if fld_chars and not instructions:
        _fail("unsupported_template", f"Field without PAGE/NUMPAGES instruction in {part}")
    field_kinds = [element.get(f"{{{_W_NS}}}fldCharType") for element in fld_chars]
    if any(kind not in {"begin", "separate", "end"} for kind in field_kinds):
        _fail("unsupported_template", f"Malformed field marker in {part}")
    if field_kinds.count("begin") != len(complex_instructions) or field_kinds.count("end") != len(
        complex_instructions
    ):
        _fail("unsupported_template", f"Unbalanced field markers in {part}")
    if (instructions or fld_chars) and part == "word/document.xml":
        _fail("unsupported_template", "Fields are allowed only in headers and footers")


def _style_types(files: Mapping[str, bytes]) -> dict[str, str]:
    data = files.get("word/styles.xml")
    if data is None:
        _fail("unsupported_template", "Template has no styles part")
    root = _xml(data, "word/styles.xml")
    result: dict[str, str] = {}
    for style in root.xpath("./w:style", namespaces={"w": _W_NS}):
        style_id = style.get(f"{{{_W_NS}}}styleId")
        style_type = style.get(f"{{{_W_NS}}}type")
        if style_id and style_type:
            result[style_id] = style_type
    return result


def _validated_sections(
    sections: Sequence[Mapping[str, object]], style_types: Mapping[str, str]
) -> dict[str, Mapping[str, object]]:
    if len(sections) != len(SECTIONS):
        _fail("unsupported_template_binding", "Binding must contain exactly six sections")
    by_name: dict[str, Mapping[str, object]] = {}
    for index, raw in enumerate(sections):
        section = _mapping(raw, f"sections[{index}]")
        name = _text(section.get("section"), f"sections[{index}].section")
        if name != SECTIONS[index] or name in by_name:
            _fail("unsupported_template_binding", "Binding sections must be unique and ordered")
        heading = _text(section.get("heading_style_id"), f"sections[{index}].heading_style_id")
        table = _text(section.get("table_style_id"), f"sections[{index}].table_style_id")
        if style_types.get(heading) != "paragraph" or style_types.get(table) != "table":
            _fail(
                "unsupported_template_binding",
                "Binding style IDs must exist with the required kind",
            )
        columns = _sequence(section.get("columns", []), f"sections[{index}].columns")
        if name in TABLE_SECTIONS:
            if len(columns) != len(COLUMNS):
                _fail("unsupported_template_binding", "Response tables require all four columns")
            seen: list[str] = []
            total = Decimal(0)
            for column_index, raw_column in enumerate(columns):
                column = _mapping(raw_column, f"sections[{index}].columns[{column_index}]")
                key = _text(column.get("key"), "column.key")
                if key not in COLUMNS or key in seen:
                    _fail(
                        "unsupported_template_binding",
                        "Response column keys must be unique and complete",
                    )
                try:
                    width = Decimal(str(column.get("width_percent")))
                except (InvalidOperation, ValueError):
                    _fail("unsupported_template_binding", "Column widths must be decimal numbers")
                if not width.is_finite() or width <= 0 or width > 100:
                    _fail(
                        "unsupported_template_binding", "Column widths must be positive and <= 100"
                    )
                seen.append(key)
                total += width
            if set(seen) != set(COLUMNS) or total != Decimal(100):
                _fail(
                    "unsupported_template_binding", "Response column widths must total exactly 100"
                )
        elif columns:
            _fail("unsupported_template_binding", "Only response tables accept column mappings")
        by_name[name] = section
    return by_name


def _paragraph_text(element: etree._Element) -> str:
    return "".join(element.xpath(".//w:t/text()", namespaces={"w": _W_NS}))


def _paragraph_style(element: etree._Element) -> str | None:
    values = element.xpath("./w:pPr/w:pStyle/@w:val", namespaces={"w": _W_NS})
    return values[0] if values else None


def _story_static_summary(
    files: Mapping[str, bytes], registered_headings: Sequence[str]
) -> tuple[list[dict[str, object]], list[dict[str, object]], list[str]]:
    if len(set(registered_headings)) != len(registered_headings) or any(
        not isinstance(value, str) or not value.strip() for value in registered_headings
    ):
        _fail("unsupported_template_binding", "Registered headings must be unique nonblank strings")
    allowed_headings = set(registered_headings)
    document = _xml(files["word/document.xml"], "word/document.xml")
    body = document.find(f"{{{_W_NS}}}body")
    if body is None:
        _fail("unsupported_template", "Template document body is missing")
    if body.xpath("./w:tbl", namespaces={"w": _W_NS}):
        _fail("unsupported_template", "Static template body tables are unsupported")
    summary: list[dict[str, object]] = []
    anchors: list[dict[str, object]] = []
    metadata_fields: list[str] = []
    section_index = 0
    paragraphs = body.xpath("./w:p", namespaces={"w": _W_NS})
    for index, paragraph in enumerate(paragraphs, 1):
        text = _paragraph_text(paragraph)
        placeholders = _PLACEHOLDER.findall(text)
        anchor_sections = [name for name, marker in ANCHORS.items() if text == marker]
        if anchor_sections:
            if paragraph.xpath("./w:pPr/w:sectPr", namespaces={"w": _W_NS}):
                _fail(
                    "unsupported_template",
                    "Template anchor paragraph cannot carry a section break",
                )
            anchors.append(
                {
                    "section": anchor_sections[0],
                    "paragraph": index,
                    "section_index": section_index,
                }
            )
        elif text in METADATA_MARKERS:
            field = text.removeprefix("{{bid.").removesuffix("}}")
            if field in metadata_fields:
                _fail("unsupported_template", "Template metadata marker is duplicated")
            metadata_fields.append(field)
        elif placeholders:
            _fail("unsupported_template", "Template contains an unknown or mixed-text placeholder")
        elif text.strip() and text not in allowed_headings:
            _fail(
                "unsupported_template", "Static template text is not a registered chapter heading"
            )
        else:
            summary.append(
                {
                    "part": "word/document.xml",
                    "paragraph": index,
                    "style_id": _paragraph_style(paragraph),
                    "text": text,
                }
            )
        section_index += len(paragraph.xpath("./w:pPr/w:sectPr", namespaces={"w": _W_NS}))
    for name in sorted(
        part for part in files if re.fullmatch(r"word/(?:header|footer)\d+\.xml", part) is not None
    ):
        root = _xml(files[name], name)
        for index, paragraph in enumerate(
            root.xpath("./w:p|./w:tbl//w:p", namespaces={"w": _W_NS}), 1
        ):
            text = _paragraph_text(paragraph)
            visible_without_field = "".join(
                paragraph.xpath(
                    ".//w:t[not(ancestor::w:instrText)]/text()", namespaces={"w": _W_NS}
                )
            )
            has_page_field = bool(
                paragraph.xpath(".//w:instrText|.//w:fldSimple", namespaces={"w": _W_NS})
            )
            safe_page_result = has_page_field and re.fullmatch(
                r"[\d\s第页共/\\\-—·|]*", visible_without_field
            )
            if visible_without_field.strip() and not safe_page_result:
                _fail("unsupported_template", "Header/footer static business text is unsupported")
            summary.append(
                {
                    "part": name,
                    "paragraph": index,
                    "style_id": _paragraph_style(paragraph),
                    "text": text,
                }
            )
    if (
        len(anchors) != len(SECTIONS)
        or tuple(str(anchor["section"]) for anchor in anchors) != SECTIONS
    ):
        _fail("unsupported_template", "Six unique body anchors must appear in the fixed order")
    return summary, anchors, metadata_fields


def validate_template(
    content: bytes,
    sections: Sequence[Mapping[str, object]],
    registered_headings: Sequence[str],
    *,
    limits: RenderLimits = DEFAULT_LIMITS,
) -> TemplateInspection:
    """Inspect the strict OOXML subset used by the export adapter."""

    files = _safe_zip(content, limits)
    _scan_relationships(files)
    _scan_xml(files)
    _validated_sections(sections, _style_types(files))
    summary, anchors, metadata_fields = _story_static_summary(files, registered_headings)
    static_hash = hashlib.sha256(canonical_json({"paragraphs": summary})).hexdigest()
    try:
        Document(io.BytesIO(content))
    except Exception as exc:
        raise ExportRenderError(
            "unsupported_template", "Template cannot be opened as DOCX"
        ) from exc
    return TemplateInspection(
        template_sha256=hashlib.sha256(content).hexdigest(),
        static_content_hash=static_hash,
        anchors=tuple(anchors),
        metadata_fields=tuple(metadata_fields),
    )


def _read_private(path: Path, limit: int, code: str) -> bytes:
    try:
        if not path.is_file() or path.is_symlink():
            _fail(code, "Authorized input is missing or is a symlink")
        with path.open("rb") as handle:
            value = handle.read(limit + 1)
    except OSError as exc:
        raise ExportRenderError(code, "Authorized input cannot be read") from exc
    if not value or len(value) > limit:
        _fail(code, "Authorized input exceeds its byte limit")
    return value


# Image evidence is rendered uniformly: the document never reveals which images are
# prototypes, user screenshots or vendor captures.
IMAGE_MATERIALS = frozenset(
    {
        "user_screenshot",
        "browser_screenshot",
        "user_diagram",
        "certificate_image",
        "vendor_web",
        "vendor_pdf",
        "prototype",
    }
)


def _validate_png(content: bytes, width: int, height: int) -> None:
    if len(content) < 24 or content[:8] != b"\x89PNG\r\n\x1a\n" or content[12:16] != b"IHDR":
        _fail("attachment_integrity", "Evidence page is not a PNG")
    encoded_width = int.from_bytes(content[16:20], "big")
    encoded_height = int.from_bytes(content[20:24], "big")
    if encoded_width != width or encoded_height != height:
        _fail("attachment_integrity", "Evidence page dimensions do not match the manifest")


def _validate_manifest(
    manifest: Mapping[str, object], inspection: TemplateInspection, limits: RenderLimits
) -> tuple[Mapping[str, object], list[Mapping[str, Any]], list[Mapping[str, Any]]]:
    if manifest.get("renderer_profile") != RENDERER_PROFILE:
        _fail("renderer_profile_mismatch", "Manifest renderer profile is not available")
    mode = manifest.get("mode")
    if mode not in {"final_section", "review_copy"}:
        _fail("invalid_export_manifest", "Unknown export mode")
    template = _mapping(manifest.get("template"), "template")
    if _sha(template.get("sha256"), "template.sha256") != inspection.template_sha256:
        _fail("template_integrity", "Template bytes do not match the fixed revision")
    if (
        _sha(template.get("static_content_hash"), "template.static_content_hash")
        != inspection.static_content_hash
    ):
        _fail(
            "template_static_content_changed", "Template static content does not match the binding"
        )
    if template.get("adapter_version") not in {None, ADAPTER_VERSION}:
        _fail("template_adapter_mismatch", "Template adapter version is not available")
    items = [
        _mapping(item, f"items[{index}]")
        for index, item in enumerate(_sequence(manifest.get("items"), "items"))
    ]
    if len(items) > limits.max_requirements:
        _fail("export_resource_limit", "Requirement count exceeds the render profile")
    ordinals = [_integer(item.get("ordinal"), "item.ordinal", minimum=1) for item in items]
    if ordinals != list(range(1, len(items) + 1)):
        _fail("invalid_export_manifest", "Item ordinals must be unique, contiguous, and ordered")
    requirement_ids = [_text(item.get("requirement_id"), "item.requirement_id") for item in items]
    if len(set(requirement_ids)) != len(requirement_ids):
        _fail("invalid_export_manifest", "Every requirement must occur exactly once")
    gaps = 0
    evidence_bindings: dict[str, tuple[str, int | None]] = {}
    image_evidence: set[str] = set()
    for item in items:
        kind = item.get("kind")
        if kind not in {"row", "comply_only", "gap"}:
            _fail("invalid_export_manifest", "Item kind is unsupported")
        _mapping(item.get("source"), "item.source")
        _text(item.get("location_label"), "item.location_label")
        if kind == "row":
            if item.get("table") not in TABLE_SECTIONS:
                _fail("invalid_export_manifest", "Response row has an invalid table")
            response_kind = item.get("response_kind")
            if (
                response_kind not in {"evidence", "commitment"}
                or item.get("deviation") not in DEVIATION_LABELS
            ):
                _fail("invalid_export_manifest", "Response row kind or deviation is invalid")
            _text(item.get("response_text"), "item.response_text")
            _text(item.get("deviation_note"), "item.deviation_note")
            evidence = _sequence(item.get("evidence", []), "item.evidence")
            if response_kind == "evidence" and not evidence:
                _fail("invalid_export_manifest", "Evidence response row has no evidence")
            if response_kind == "commitment" and evidence:
                _fail("invalid_export_manifest", "Commitment response row cannot carry evidence")
            for raw in evidence:
                evidence_item = _mapping(raw, "item.evidence[]")
                evidence_id = _text(evidence_item.get("id"), "evidence.id")
                if evidence_id in evidence_bindings:
                    _fail("invalid_export_manifest", "Evidence IDs cannot be repeated across rows")
                _text(evidence_item.get("confirmed_by"), "evidence.confirmed_by")
                _text(evidence_item.get("confirmed_at"), "evidence.confirmed_at")
                material_kind = evidence_item.get("material_kind")
                attachment_ordinal = evidence_item.get("attachment_ordinal")
                if material_kind == "user_supplied_pdf_page" or material_kind in IMAGE_MATERIALS:
                    attachment_ordinal = _integer(
                        attachment_ordinal, "evidence.attachment_ordinal", minimum=1
                    )
                    if material_kind in IMAGE_MATERIALS:
                        image_evidence.add(evidence_id)
                elif material_kind == "declaration":
                    _text(evidence_item.get("title"), "evidence.title")
                    if attachment_ordinal is not None:
                        _fail(
                            "invalid_export_manifest",
                            "Declaration evidence cannot reference an attachment",
                        )
                    attachment_ordinal = None
                else:
                    _fail("invalid_export_manifest", "Evidence material kind is unsupported")
                evidence_bindings[evidence_id] = (str(item["requirement_id"]), attachment_ordinal)
        elif kind == "comply_only":
            _text(item.get("disposition_by"), "item.disposition_by")
            _text(item.get("disposition_at"), "item.disposition_at")
        else:
            gaps += 1
            reasons = _sequence(item.get("gap_reasons"), "item.gap_reasons")
            if not reasons or any(reason not in GAP_LABELS for reason in reasons):
                _fail("invalid_export_manifest", "Gap reasons are missing or unsupported")
            if any(
                item.get(key) is not None
                for key in ("response_text", "deviation", "deviation_note")
            ):
                _fail("invalid_export_manifest", "Gap cannot contain candidate response content")
    if mode == "final_section" and gaps:
        _fail("export_gaps_present", "A final section cannot contain gaps")
    attachments = [
        _mapping(item, f"attachments[{index}]")
        for index, item in enumerate(_sequence(manifest.get("attachments", []), "attachments"))
    ]
    if len(attachments) > limits.max_attachments:
        _fail("export_resource_limit", "Attachment count exceeds the render profile")
    attachment_ordinals = [
        _integer(item.get("ordinal"), "attachment.ordinal", minimum=1) for item in attachments
    ]
    if attachment_ordinals != list(range(1, len(attachments) + 1)):
        _fail(
            "invalid_export_manifest", "Attachment ordinals must be unique, contiguous, and ordered"
        )
    seen_labels: set[str] = set()
    for index, attachment in enumerate(attachments):
        label = _text(attachment.get("label"), f"attachments[{index}].label")
        if label != f"E{index + 1:03d}" or label in seen_labels:
            _fail("invalid_export_manifest", "Attachment labels must be stable E001 ordinals")
        seen_labels.add(label)
        is_image = attachment.get("kind") == "image"
        if is_image:
            _text(attachment.get("rendition_id"), "attachment.rendition_id")
        elif "kind" in attachment:
            _fail("invalid_export_manifest", "Attachment kind is unsupported")
        else:
            _sha(attachment.get("original_sha256"), "attachment.original_sha256")
            _integer(attachment.get("page"), "attachment.page", minimum=1)
            _text(attachment.get("title"), "attachment.title")
        _sha(attachment.get("png_sha256"), "attachment.png_sha256")
        size = _integer(attachment.get("size_bytes"), "attachment.size_bytes", minimum=1)
        width = _integer(attachment.get("width"), "attachment.width", minimum=1)
        height = _integer(attachment.get("height"), "attachment.height", minimum=1)
        if (
            size > limits.max_page_bytes
            or width > limits.max_image_side
            or height > limits.max_image_side
            or width * height > limits.max_page_pixels
        ):
            _fail("export_resource_limit", "Attachment exceeds the render profile")
        linked_requirements = _sequence(
            attachment.get("requirement_ids"), "attachment.requirement_ids"
        )
        if not linked_requirements or not set(linked_requirements).issubset(set(requirement_ids)):
            _fail("invalid_export_manifest", "Attachment requirements do not match export items")
        evidence_ids = _sequence(attachment.get("evidence_ids"), "attachment.evidence_ids")
        if not evidence_ids:
            _fail("invalid_export_manifest", "Attachment must retain evidence references")
        expected_evidence = {
            evidence_id
            for evidence_id, (_, ordinal) in evidence_bindings.items()
            if ordinal == index + 1
        }
        if set(evidence_ids) != expected_evidence or any(
            (evidence_id in image_evidence) != is_image for evidence_id in evidence_ids
        ):
            _fail(
                "invalid_export_manifest",
                "Attachment evidence IDs do not match row evidence references",
            )
        expected_requirements = {
            requirement_id
            for _, (requirement_id, ordinal) in evidence_bindings.items()
            if ordinal == index + 1
        }
        if set(linked_requirements) != expected_requirements:
            _fail(
                "invalid_export_manifest",
                "Attachment requirement IDs do not match row evidence references",
            )
    referenced_ordinals = {
        ordinal for _, ordinal in evidence_bindings.values() if ordinal is not None
    }
    if referenced_ordinals != set(attachment_ordinals):
        _fail("invalid_export_manifest", "Every page evidence must resolve to one attachment")
    return template, items, attachments


def _fill_confidential(
    manifest: Mapping[str, object], items: list[Mapping[str, Any]], values: Mapping[str, str]
) -> list[Mapping[str, Any]]:
    fields = [
        _mapping(entry, "confidential[]")
        for entry in _sequence(manifest.get("confidential", []), "confidential")
    ]
    labels = {str(entry["key"]): str(entry["label"]) for entry in fields}
    filled = {str(entry["key"]) for entry in fields if entry.get("value_id") is not None}
    if set(values) != filled:
        _fail("invalid_export_manifest", "Confidential values do not match the fixed fields")

    def fill(value: object) -> object:
        return redaction.fill_secrets(value, values, labels) if isinstance(value, str) else value

    def evidence(entry: Mapping[str, Any]) -> Mapping[str, Any]:
        quoted = entry.get("input")
        if not isinstance(quoted, Mapping) or "quote" not in quoted:
            return entry
        return {**entry, "input": {**quoted, "quote": fill(quoted["quote"])}}

    return [
        {
            **item,
            **{
                name: fill(item[name])
                for name in ("response_text", "deviation_note")
                if name in item
            },
            **(
                {
                    "evidence": [
                        evidence(_mapping(entry, "item.evidence[]")) for entry in item["evidence"]
                    ]
                }
                if isinstance(item.get("evidence"), list)
                else {}
            ),
        }
        for item in items
    ]


def _find_anchor_paragraphs(document: DocumentObject) -> dict[str, tuple[Any, int]]:
    found: dict[str, tuple[Any, int]] = {}
    section_index = 0
    for paragraph in document.paragraphs:
        for section, marker in ANCHORS.items():
            if paragraph.text == marker:
                if section in found:
                    _fail("unsupported_template", "Template anchor was duplicated during render")
                found[section] = (paragraph, section_index)
        if paragraph._p.xpath("./w:pPr/w:sectPr"):
            section_index += 1
    if tuple(found) != SECTIONS:
        _fail("unsupported_template", "Template anchors changed before rendering")
    return found


def _set_paragraph_style_id(paragraph: Any, style_id: str) -> None:
    properties = paragraph._p.get_or_add_pPr()
    style = properties.find(qn("w:pStyle"))
    if style is None:
        style = OxmlElement("w:pStyle")
        properties.insert(0, style)
    style.set(qn("w:val"), style_id)


def _set_table_style_id(table: Any, style_id: str) -> None:
    properties = table._tbl.tblPr
    style = properties.find(qn("w:tblStyle"))
    if style is None:
        style = OxmlElement("w:tblStyle")
        properties.insert(0, style)
    style.set(qn("w:val"), style_id)


def _before_paragraph(document: DocumentObject, anchor: Any, text: str, style_id: str) -> Any:
    paragraph = document.add_paragraph(text)
    _set_paragraph_style_id(paragraph, style_id)
    for run in paragraph.runs:
        run.font.hidden = False
        run.font.color.rgb = RGBColor(0, 0, 0)
    anchor._p.addprevious(paragraph._p)
    return paragraph


def _before_table(
    document: DocumentObject, anchor: Any, rows: int, columns: int, style_id: str
) -> Any:
    table = document.add_table(rows=rows, cols=columns)
    _set_table_style_id(table, style_id)
    anchor._p.addprevious(table._tbl)
    return table


def _repeat_header(row: Any) -> None:
    properties = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    properties.append(header)


def _set_cell_text(cell: Any, value: str, *, bold: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(value)
    run.bold = bold
    run.font.hidden = False
    run.font.color.rgb = RGBColor(0, 0, 0)
    paragraph.paragraph_format.keep_together = True


def _set_fixed_widths(table: Any, widths: Sequence[Decimal], available_width: int) -> None:
    table.autofit = False
    properties = table._tbl.tblPr
    layout = properties.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        properties.append(layout)
    layout.set(qn("w:type"), "fixed")
    calculated = [int(available_width * width / Decimal(100)) for width in widths]
    calculated[-1] += available_width - sum(calculated)
    grid = table._tbl.tblGrid
    for grid_column, width in zip(grid.gridCol_lst, calculated, strict=True):
        grid_column.set(qn("w:w"), str(max(1, width // 635)))
    for row in table.rows:
        for cell, width in zip(row.cells, calculated, strict=True):
            cell.width = Emu(width)
            cell_properties = cell._tc.get_or_add_tcPr()
            tc_width = cell_properties.find(qn("w:tcW"))
            if tc_width is None:
                tc_width = OxmlElement("w:tcW")
                cell_properties.append(tc_width)
            tc_width.set(qn("w:type"), "dxa")
            tc_width.set(qn("w:w"), str(max(1, width // 635)))


def _requirement_text(item: Mapping[str, Any]) -> str:
    source = _mapping(item.get("source"), "item.source")
    quote = _text(source.get("quote"), "item.source.quote")
    # The tender's own ★/▲ stays as written; a starred requirement always shows one.
    if item.get("starred") is True and not quote.lstrip().startswith(("★", "▲")):
        return "★" + quote
    return quote


def _section_of(item: Mapping[str, Any]) -> str:
    kind = item.get("kind")
    return (
        str(item["table"]) if kind == "row" else "comply_only" if kind == "comply_only" else "gaps"
    )


def table_numbers(items: Sequence[Mapping[str, Any]]) -> dict[str, tuple[str, int]]:
    """Each requirement's section and its number within that section, in manifest order."""
    counters: dict[str, int] = {}
    numbers: dict[str, tuple[str, int]] = {}
    for item in items:
        section = _section_of(item)
        counters[section] = counters.get(section, 0) + 1
        numbers[str(item["requirement_id"])] = (section, counters[section])
    return numbers


def declaration_labels(items: Sequence[Mapping[str, Any]]) -> dict[str, str]:
    labels: dict[str, str] = {}
    for item in items:
        for raw in _sequence(item.get("evidence", []), "item.evidence"):
            evidence = _mapping(raw, "item.evidence[]")
            if evidence.get("material_kind") == "declaration":
                labels[str(evidence["id"])] = f"D{len(labels) + 1:03d}"
    return labels


def _evidence_refs(item: Mapping[str, Any], declarations: Mapping[str, str]) -> list[str]:
    result: list[str] = []
    for raw in _sequence(item.get("evidence", []), "item.evidence"):
        evidence = _mapping(raw, "item.evidence[]")
        ordinal = evidence.get("attachment_ordinal")
        if ordinal is not None:
            label = f"E{_integer(ordinal, 'evidence.attachment_ordinal', minimum=1):03d}"
        else:
            label = declarations[str(evidence["id"])]
        if label not in result:
            result.append(label)
    return result


def _reference_name(label: str) -> str:
    return ("附件 " if label.startswith("E") else "声明 ") + label


def _compliance(item: Mapping[str, Any], section: str) -> str:
    deviation = str(item["deviation"])
    if deviation == "none":
        return "响应且无负偏离" if section == "substantive" else "响应"
    return f"{DEVIATION_LABELS[deviation]}：{item['deviation_note']}"


def _clause_references(
    requirement_ids: Sequence[str], numbers: Mapping[str, tuple[str, int]]
) -> str:
    return "、".join(
        f"{SECTION_TITLES[numbers[rid][0]]}第 {numbers[rid][1]} 条"
        for rid in requirement_ids
        if rid in numbers
    )


def _add_internal_link(paragraph: Any, label: str, text_value: str | None = None) -> None:
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), f"bid_{label}")
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.extend((color, underline))
    text = OxmlElement("w:t")
    text.text = text_value or label
    run.extend((properties, text))
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _render_response_table(
    document: DocumentObject,
    anchor: Any,
    binding: Mapping[str, object],
    section: str,
    items: Sequence[Mapping[str, Any]],
    declarations: Mapping[str, str],
    available_width: int,
) -> None:
    columns = [
        _mapping(value, "binding.columns[]")
        for value in _sequence(binding["columns"], "binding.columns")
    ]
    keys = [str(column["key"]) for column in columns]
    if sorted(keys) != sorted(COLUMNS):
        _fail("export_binding_outdated", "Binding columns predate the current table layout")
    table = _before_table(
        document, anchor, max(2, len(items) + 1), len(keys), str(binding["table_style_id"])
    )
    _repeat_header(table.rows[0])
    for cell, key in zip(table.rows[0].cells, keys, strict=True):
        _set_cell_text(cell, COLUMN_TITLES[key], bold=True)
    widths = [Decimal(str(column["width_percent"])) for column in columns]
    _set_fixed_widths(table, widths, available_width)
    if not items:
        merged = table.rows[1].cells[0].merge(table.rows[1].cells[-1])
        _set_cell_text(merged, "本节无响应条目")
        return
    for number, item in enumerate(items, 1):
        refs = _evidence_refs(item, declarations)
        values = {
            "ordinal": str(number),
            "requirement": _requirement_text(item),
            "compliance": _compliance(item, section),
        }
        for cell, key in zip(table.rows[number].cells, keys, strict=True):
            if key == "response":
                _set_cell_text(cell, str(item["response_text"]))
                if refs:
                    paragraph = cell.paragraphs[-1]
                    paragraph.add_run("（见")
                    for index, label in enumerate(refs):
                        if index:
                            paragraph.add_run("、")
                        _add_internal_link(paragraph, label, _reference_name(label))
                    paragraph.add_run("）")
            else:
                _set_cell_text(
                    cell, values[key], bold=key == "compliance" and item["deviation"] == "negative"
                )


def _render_simple_table(
    document: DocumentObject,
    anchor: Any,
    binding: Mapping[str, object],
    columns: Sequence[tuple[str, int]],
    rows: Sequence[Sequence[str]],
    empty_text: str,
    available_width: int,
    bookmarks: Sequence[str | None] = (),
) -> None:
    headings = [heading for heading, _ in columns]
    table = _before_table(
        document, anchor, max(2, len(rows) + 1), len(headings), str(binding["table_style_id"])
    )
    _repeat_header(table.rows[0])
    for cell, heading in zip(table.rows[0].cells, headings, strict=True):
        _set_cell_text(cell, heading, bold=True)
    if not rows:
        merged = table.rows[1].cells[0].merge(table.rows[1].cells[-1])
        _set_cell_text(merged, empty_text)
    else:
        for row_index, values in enumerate(rows, 1):
            for cell, value in zip(table.rows[row_index].cells, values, strict=True):
                _set_cell_text(cell, value)
            name = bookmarks[row_index - 1] if row_index - 1 < len(bookmarks) else None
            if name is not None:
                # Response cells link here: the row of a declaration that has no page.
                _bookmark(table.rows[row_index].cells[0].paragraphs[0], name, 100_000 + row_index)
    _set_fixed_widths(table, [Decimal(width) for _, width in columns], available_width)


def _bookmark(paragraph: Any, name: str, identifier: int) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(identifier))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(identifier))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def _render_evidence_index(
    document: DocumentObject,
    anchor: Any,
    binding: Mapping[str, object],
    items: Sequence[Mapping[str, Any]],
    attachments: Sequence[Mapping[str, Any]],
    declarations: Mapping[str, str],
    numbers: Mapping[str, tuple[str, int]],
    available_width: int,
) -> None:
    observations: dict[int, list[str]] = {}
    excerpts: dict[int, list[str]] = {}
    declared: list[tuple[str, Mapping[str, Any], str]] = []
    for item in items:
        for raw in _sequence(item.get("evidence", []), "item.evidence"):
            evidence = _mapping(raw, "item.evidence[]")
            ordinal = evidence.get("attachment_ordinal")
            quote = str(_mapping(evidence.get("input", {}), "evidence.input").get("quote", ""))
            if ordinal is None:
                declared.append(
                    (declarations[str(evidence["id"])], evidence, str(item["requirement_id"]))
                )
                continue
            seen = (
                observations if evidence.get("material_kind") in IMAGE_MATERIALS else excerpts
            ).setdefault(int(ordinal), [])
            value = (
                str(evidence.get("visual_observation", ""))
                if evidence.get("material_kind") in IMAGE_MATERIALS
                else quote
            )
            if value and value not in seen:
                seen.append(value)
    rows: list[list[str]] = []
    bookmarks: list[str | None] = []
    for attachment in attachments:
        ordinal = int(attachment["ordinal"])
        if attachment.get("kind") == "image":
            material = "证据图片"
            content = "所见：" + "；".join(observations.get(ordinal, []))
        else:
            material = str(attachment["title"])
            content = f"原件第 {attachment['page']} 页"
            if excerpts.get(ordinal):
                content += "；摘录：" + "；".join(excerpts[ordinal])
        rows.append(
            [
                str(attachment["label"]),
                material,
                content,
                _clause_references([str(x) for x in attachment["requirement_ids"]], numbers),
            ]
        )
        bookmarks.append(None)
    for label, evidence, requirement_id in declared:
        quote = str(_mapping(evidence.get("input", {}), "evidence.input").get("quote", ""))
        rows.append(
            [
                label,
                f"{evidence['title']}（声明）",
                f"摘录：{quote}",
                _clause_references([requirement_id], numbers),
            ]
        )
        bookmarks.append(f"bid_{label}")
    _render_simple_table(
        document,
        anchor,
        binding,
        INDEX_COLUMNS,
        rows,
        "本次响应无证据附件",
        available_width,
        bookmarks,
    )


def _add_review_marker(document: DocumentObject, has_gaps: bool) -> None:
    document.settings.odd_and_even_pages_header_footer = False
    marker = "审阅件·存在缺口·不得提交" if has_gaps else "审阅件·不得提交"
    seen_parts: set[str] = set()
    for section in document.sections:
        section.different_first_page_header_footer = False
        footer = section.footer
        part_name = str(footer.part.partname)
        if part_name in seen_parts:
            continue
        seen_parts.add(part_name)
        paragraph = footer.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(marker)
        run.bold = True
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(192, 0, 0)


def _replace_metadata(document: DocumentObject, task: Mapping[str, Any]) -> None:
    replacements = {
        "{{bid.task_name}}": task.get("name"),
        "{{bid.tender_number}}": task.get("tender_number"),
    }
    for paragraph in document.paragraphs:
        if paragraph.text not in replacements:
            continue
        value = replacements[paragraph.text]
        if not isinstance(value, str) or not value.strip():
            _fail("missing_template_metadata", f"Task metadata is missing for {paragraph.text}")
        if not paragraph.runs:
            paragraph.add_run(value)
        else:
            paragraph.runs[0].text = value
            for run in paragraph.runs[1:]:
                run.text = ""


def _section_geometry(document: DocumentObject, section_index: int) -> tuple[int, int]:
    if section_index >= len(document.sections):
        _fail("unsupported_template", "Template anchor section is missing")
    section = document.sections[section_index]
    dimensions = (section.page_width, section.left_margin, section.right_margin)
    vertical = (section.page_height, section.top_margin, section.bottom_margin)
    if any(value is None for value in (*dimensions, *vertical)):
        _fail("unsupported_template", "Template section dimensions are incomplete")
    page_width, left_margin, right_margin, page_height, top_margin, bottom_margin = (
        int(value) for value in (*dimensions, *vertical) if value is not None
    )
    available_width = page_width - left_margin - right_margin
    available_height = page_height - top_margin - bottom_margin
    if available_width <= 0 or available_height <= CAPTION_RESERVE:
        _fail("unsupported_template", "Template section has no usable page area")
    return available_width, available_height


def _render_attachments(
    document: DocumentObject,
    anchor: Any,
    binding: Mapping[str, object],
    attachments: Sequence[Mapping[str, Any]],
    page_paths: Mapping[int, Path],
    limits: RenderLimits,
    available_width: int,
    available_height: int,
) -> None:
    image_height = available_height - CAPTION_RESERVE
    for attachment in attachments:
        ordinal = int(attachment["ordinal"])
        label = str(attachment["label"])
        content = _read_private(
            Path(page_paths[ordinal]), limits.max_page_bytes, "attachment_integrity"
        )
        if (
            len(content) != int(attachment["size_bytes"])
            or hashlib.sha256(content).hexdigest() != attachment["png_sha256"]
        ):
            _fail("attachment_integrity", "Evidence page bytes do not match the manifest")
        _validate_png(content, int(attachment["width"]), int(attachment["height"]))
        title = document.add_paragraph()
        _set_paragraph_style_id(title, str(binding["heading_style_id"]))
        title.paragraph_format.page_break_before = True
        # The caption never ends a page apart from its image.
        title.paragraph_format.keep_with_next = True
        title.add_run(
            f"附件 {label}　证据图片"
            if attachment.get("kind") == "image"
            else f"附件 {label}　{attachment['title']}　原件第 {attachment['page']} 页"
        )
        for run in title.runs:
            run.font.hidden = False
            run.font.color.rgb = RGBColor(0, 0, 0)
        _bookmark(title, f"bid_{label}", ordinal)
        anchor._p.addprevious(title._p)
        image_paragraph = document.add_paragraph()
        image_paragraph.paragraph_format.space_before = Pt(0)
        image_paragraph.paragraph_format.space_after = Pt(0)
        width = int(attachment["width"])
        height = int(attachment["height"])
        scale = min(available_width / width, image_height / height)
        run = image_paragraph.add_run()
        run.add_picture(
            io.BytesIO(content),
            width=Emu(max(1, int(width * scale))),
            height=Emu(max(1, int(height * scale))),
        )
        anchor._p.addprevious(image_paragraph._p)


def set_core_properties(document: DocumentObject) -> None:
    properties = document.core_properties
    fixed = datetime(2000, 1, 1, tzinfo=UTC)
    properties.author = "AI Bid Tool"
    properties.last_modified_by = "AI Bid Tool"
    properties.created = fixed
    properties.modified = fixed
    properties.revision = 1
    properties.comments = ""
    properties.category = ""
    properties.content_status = ""
    properties.identifier = ""
    properties.keywords = ""
    properties.language = "zh-CN"
    properties.subject = ""
    properties.title = "响应章节"
    properties.version = "1"


def _sort_attributes(element: etree._Element) -> None:
    for node in element.iter():
        if len(node.attrib) > 1:
            ordered = sorted(node.attrib.items())
            node.attrib.clear()
            for key, value in ordered:
                node.set(key, value)


def _normalize_xml(data: bytes, part: str) -> bytes:
    root = _xml(data, part)
    _sort_attributes(root)
    return etree.tostring(root, encoding="UTF-8", xml_declaration=True, standalone=True)


def _remove_personal_package_parts(files: dict[str, bytes]) -> None:
    """Remove inert template metadata that may identify the source workstation."""

    removed = {
        name for name in files if name.startswith("customXml/") or name == "docProps/thumbnail.jpeg"
    }
    for name in removed:
        del files[name]
    for rels_name in ("_rels/.rels", "word/_rels/document.xml.rels"):
        if rels_name not in files:
            continue
        root = _xml(files[rels_name], rels_name)
        for relation in list(root):
            rel_type = relation.get("Type", "")
            if rel_type.endswith("/customXml") or rel_type.endswith("/thumbnail"):
                root.remove(relation)
        files[rels_name] = etree.tostring(
            root, encoding="UTF-8", xml_declaration=True, standalone=True
        )
    types = _xml(files["[Content_Types].xml"], "[Content_Types].xml")
    for entry in list(types):
        part_name = entry.get("PartName", "").lstrip("/")
        extension = entry.get("Extension", "").casefold()
        if part_name in removed or (
            extension in {"jpeg", "jpg"}
            and not any(name.casefold().endswith((".jpeg", ".jpg")) for name in files)
        ):
            types.remove(entry)
    files["[Content_Types].xml"] = etree.tostring(
        types, encoding="UTF-8", xml_declaration=True, standalone=True
    )
    files["docProps/app.xml"] = (
        b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        b'<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties"'
        b' xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
        b"<Application>AI Bid Tool</Application><DocSecurity>0</DocSecurity>"
        b"<ScaleCrop>false</ScaleCrop><Company></Company><LinksUpToDate>false</LinksUpToDate>"
        b"<SharedDoc>false</SharedDoc><HyperlinksChanged>false</HyperlinksChanged>"
        b"<AppVersion>1.0</AppVersion></Properties>"
    )


def deterministic_package(content: bytes, limits: RenderLimits) -> bytes:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as source:
            if sum(entry.file_size for entry in source.infolist()) > limits.max_expanded_bytes:
                _fail(
                    "export_resource_limit",
                    "Rendered DOCX expanded size exceeds the render profile",
                )
            files = {entry.filename: source.read(entry) for entry in source.infolist()}
    except zipfile.BadZipFile as exc:
        raise ExportRenderError("render_failed", "Renderer produced an invalid DOCX") from exc
    _remove_personal_package_parts(files)
    for name in list(files):
        if name.endswith(".xml") or name.endswith(".rels"):
            files[name] = _normalize_xml(files[name], name)
    if sum(len(value) for value in files.values()) > limits.max_expanded_bytes:
        _fail(
            "export_resource_limit",
            "Normalized DOCX expanded size exceeds the render profile",
        )
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as archive:
        for name in sorted(files):
            info = zipfile.ZipInfo(name, date_time=_ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.create_system = 3
            info.external_attr = 0o600 << 16
            info.flag_bits = 0
            archive.writestr(info, files[name], compress_type=zipfile.ZIP_DEFLATED, compresslevel=9)
    result = output.getvalue()
    if len(result) > limits.max_output_bytes:
        _fail("export_resource_limit", "Rendered DOCX exceeds the output byte limit")
    return result


def render_export_docx(
    manifest: Mapping[str, object],
    inputs: AuthorizedExportInputs,
    *,
    limits: RenderLimits = DEFAULT_LIMITS,
    confidential: Mapping[str, str] | None = None,
) -> RenderCandidate:
    """Synchronously render one fixed export manifest for an isolated child process.

    `confidential` maps field keys to the values of the manifest's fixed value rows;
    they replace `{{secret.key}}` in response text and never enter the manifest hash."""

    template_manifest = _mapping(manifest.get("template"), "template")
    sections = [
        _mapping(value, f"template.sections[{index}]")
        for index, value in enumerate(
            _sequence(template_manifest.get("sections"), "template.sections")
        )
    ]
    headings = [
        _text(value, f"template.registered_headings[{index}]") or ""
        for index, value in enumerate(
            _sequence(template_manifest.get("registered_headings"), "template.registered_headings")
        )
    ]
    template_bytes = _read_private(
        inputs.template_path, limits.max_template_bytes, "template_integrity"
    )
    inspection = validate_template(template_bytes, sections, headings, limits=limits)
    template, items, attachments = _validate_manifest(manifest, inspection, limits)
    items = _fill_confidential(manifest, items, confidential or {})
    if set(inputs.evidence_page_paths) != {int(value["ordinal"]) for value in attachments}:
        _fail(
            "attachment_integrity", "Authorized attachment inputs do not exactly match the manifest"
        )
    try:
        document = Document(io.BytesIO(template_bytes))
        anchors = _find_anchor_paragraphs(document)
        _replace_metadata(document, _mapping(manifest.get("task"), "task"))
        by_section = _validated_sections(sections, _style_types(_safe_zip(template_bytes, limits)))
        numbers = table_numbers(items)
        declarations = declaration_labels(items)
        for section in TABLE_SECTIONS:
            anchor, section_index = anchors[section]
            available_width, _ = _section_geometry(document, section_index)
            section_items = [
                item for item in items if item.get("kind") == "row" and item.get("table") == section
            ]
            _before_paragraph(
                document,
                anchor,
                SECTION_TITLES[section],
                str(by_section[section]["heading_style_id"]),
            )
            _render_response_table(
                document,
                anchor,
                by_section[section],
                section,
                section_items,
                declarations,
                available_width,
            )
        comply = [item for item in items if item.get("kind") == "comply_only"]
        comply_anchor, comply_section_index = anchors["comply_only"]
        comply_width, _ = _section_geometry(document, comply_section_index)
        _before_paragraph(
            document,
            comply_anchor,
            SECTION_TITLES["comply_only"],
            str(by_section["comply_only"]["heading_style_id"]),
        )
        _render_simple_table(
            document,
            comply_anchor,
            by_section["comply_only"],
            COMPLY_COLUMNS,
            [
                [str(number), _requirement_text(item), "遵守"]
                for number, item in enumerate(comply, 1)
            ],
            "本次抽取范围内无须遵守条款",
            comply_width,
        )
        gaps = [item for item in items if item.get("kind") == "gap"]
        gaps_anchor, gaps_section_index = anchors["gaps"]
        gaps_width, _ = _section_geometry(document, gaps_section_index)
        _before_paragraph(
            document,
            gaps_anchor,
            SECTION_TITLES["gaps"],
            str(by_section["gaps"]["heading_style_id"]),
        )
        _render_simple_table(
            document,
            gaps_anchor,
            by_section["gaps"],
            GAP_COLUMNS,
            [
                [
                    str(number),
                    _requirement_text(item),
                    "；".join(GAP_LABELS[str(reason)] for reason in item["gap_reasons"]),
                ]
                for number, item in enumerate(gaps, 1)
            ],
            "本次抽取范围内无缺口",
            gaps_width,
        )
        evidence_anchor, evidence_section_index = anchors["evidence_appendix"]
        evidence_width, evidence_height = _section_geometry(document, evidence_section_index)
        _before_paragraph(
            document,
            evidence_anchor,
            SECTION_TITLES["evidence_appendix"],
            str(by_section["evidence_appendix"]["heading_style_id"]),
        )
        _render_evidence_index(
            document,
            evidence_anchor,
            by_section["evidence_appendix"],
            items,
            attachments,
            declarations,
            numbers,
            evidence_width,
        )
        _render_attachments(
            document,
            evidence_anchor,
            by_section["evidence_appendix"],
            attachments,
            inputs.evidence_page_paths,
            limits,
            evidence_width,
            evidence_height,
        )
        if manifest["mode"] == "review_copy":
            _add_review_marker(document, bool(gaps))
        for paragraph, _ in anchors.values():
            paragraph._element.getparent().remove(paragraph._element)
        set_core_properties(document)
        buffer = io.BytesIO()
        document.save(buffer)
        content = deterministic_package(buffer.getvalue(), limits)
    except ExportRenderError:
        raise
    except Exception as exc:
        raise ExportRenderError("render_failed", "DOCX rendering failed") from exc
    digest = hashlib.sha256(content).hexdigest()
    return RenderCandidate(
        content=content,
        sha256=digest,
        size_bytes=len(content),
        media_type=_DOCX_MEDIA_TYPE,
        renderer_profile=RENDERER_PROFILE,
        manifest_hash=manifest_sha256(manifest),
    )
